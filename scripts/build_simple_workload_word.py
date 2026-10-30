#!/usr/bin/env python3
"""Plain-language workload report, using saved experiment results only."""
from pathlib import Path
import csv
import json
import shutil
import hashlib
import re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/workload_simple_20261008'
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
NAME = '工作量_简单易懂版_20261008'
BLUE, ORANGE = '#3574A8', '#BD713E'
inputs = []

def rows(relative):
    p = ROOT / relative
    inputs.append(p)
    with p.open(encoding='utf-8') as f:
        return list(csv.DictReader(f))

random = rows('artifacts/random_stratification_ablation_20261006/tables/random_summary.csv')
mixture = rows('artifacts/controlled_risk_distribution_20261005/tables/panel_summary.csv')
scale = rows('artifacts/scale_only_confirmation_20261007/tables/summary.csv')
conditions = rows('artifacts/scale_only_confirmation_20261007/tables/area_condition_summary.csv')
ablation = rows('artifacts/random_stratification_ablation_20261006/tables/ablation_summary.csv')
paper = ROOT / 'docs/论文_最新.md'
inputs.append(paper)
case = re.search(r'\| ResNet-50，背向dog \| 本地 \| (\d+) \| ([\d.]+) \| ([\d.]+) \|', paper.read_text())
assert case
case_n, case_all, case_back = int(case[1]), float(case[2]), float(case[3])
scores = {}
for pool in ['B', 'C', 'D']:
    for source in ['local', 'online']:
        p = ROOT / f'artifacts/random_stratification_ablation_20261006/ablation/{pool}_{source}.npz'
        inputs.append(p)
        with np.load(p) as z:
            f1 = z['f1'][0]
            assert f1.shape == (24, 5)
            scores[pool, source] = f1.mean(axis=0) * 100
        r = next(r for r in random if r['pool'] == pool and r['source'] == source)
        assert abs(scores[pool, source][4] - scores[pool, source][0] - 100 * float(r['observed_delta'])) < 1e-8
assert all(int(r['repeats']) == 10000 and int(r['extreme_count']) == 0 for r in random)
assert all(float(r['mean_delta_auc']) < 0 for r in ablation)
mix_b = [r for r in mixture if r['pool'] == 'B' and r['source'] == 'local' and r['total'] == '120']
assert len(mix_b) == 21
mix_start = float(next(r for r in mix_b if r['high_risk_percent'] == '0')['macro_f1_mean']) * 100
mix_end = float(next(r for r in mix_b if r['high_risk_percent'] == '100')['macro_f1_mean']) * 100
condition_counts = {}
for pool in ['B', 'C', 'D']:
    rr = [r for r in conditions if r['pool'] == pool and r['evaluable'] == 'True']
    condition_counts[pool] = [len(rr), sum(float(r['area_adjusted_recall_gap']) < 0 for r in rr)]
assert condition_counts == {'B': [119, 109], 'C': [4, 2], 'D': [122, 111]}

font_manager.fontManager.addfont(str(ROOT / 'scripts/assets/NotoSansSC-Regular.ttf'))
plt.rcParams.update({'font.family': 'Noto Sans SC', 'font.size': 13,
                     'svg.fonttype': 'none', 'axes.spines.top': False,
                     'axes.spines.right': False, 'axes.linewidth': .7,
                     'xtick.major.size': 0, 'ytick.major.size': 0})
chart_values = []
def save(fig, stem):
    fig.savefig(FIG / f'{stem}.svg', bbox_inches='tight', facecolor='white')
    fig.savefig(FIG / f'{stem}.png', dpi=210, bbox_inches='tight', facecolor='white')
    plt.close(fig)

fig, ax = plt.subplots(figsize=(9.2, 3.1), layout='constrained')
x = np.arange(3)
low = [scores[p, 'local'][0] for p in ['B', 'C', 'D']]
high = [scores[p, 'local'][4] for p in ['B', 'C', 'D']]
for dx, vals, label, color in [(-.18, low, '较容易组', BLUE), (.18, high, '较困难组', ORANGE)]:
    bars = ax.bar(x + dx, vals, width=.32, label=label, color=color)
    ax.bar_label(bars, labels=[f'{v:.1f}' for v in vals], padding=4, fontsize=14)
ax.set(xticks=x, xticklabels=['COCO\n2400个目标', 'VOC2007\n600个目标', 'Objects365\n2400个目标'],
       ylim=(0, 112), yticks=[0, 25, 50, 75, 100], ylabel='F1综合评分（满分100）')
ax.legend(loc='upper center', bbox_to_anchor=(.5, 1.18), ncol=2, frameon=False)
save(fig, '01_three_sources')
for p, a, b in zip(['B', 'C', 'D'], low, high):
    chart_values.extend([['three_sources', p, 'local', 'Q1', a], ['three_sources', p, 'local', 'Q5', b]])

def two_bars(values, labels, ylabel, stem, suffix=''):
    fig, ax = plt.subplots(figsize=(8.6, 2.75), layout='constrained')
    bars = ax.bar([0, 1], values, width=.42, color=[BLUE, ORANGE])
    ax.bar_label(bars, labels=[f'{v:.1f}{suffix}' for v in values], padding=5, fontsize=17)
    ax.set(xticks=[0, 1], xticklabels=labels, xlim=(-.7, 1.7), ylim=(0, 110),
           yticks=[0, 25, 50, 75, 100], ylabel=ylabel)
    save(fig, stem)

two_bars([mix_start, mix_end], ['120个都来自前四组', '120个都来自较困难组'],
         'F1综合评分（满分100）', '02_same_number')
chart_values.extend([['mixture', 'B', 'local', '0% Q5', mix_start], ['mixture', 'B', 'local', '100% Q5', mix_end]])
two_bars([case_all, case_back], ['全部狗样本', '其中的背向狗样本'],
         '真实狗被认对的比例（%）', '03_back_facing_dog', '%')
chart_values.extend([['dog_recall', 'B', 'local', 'overall', case_all], ['dog_recall', 'B', 'local', 'back', case_back]])
with (OUT / 'chart_values.csv').open('w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['figure', 'pool', 'label_source', 'condition', 'value_100_scale'])
    w.writerows(chart_values)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
sec.header_distance = sec.footer_distance = Cm(.65)
normal = doc.styles['Normal']
normal.font.name = 'Noto Sans CJK SC'
normal.font.size = Pt(11.5)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
normal.paragraph_format.line_spacing = 1.18
normal.paragraph_format.space_after = Pt(5)
md = []

def runfont(run, size=11.5, bold=False, color='263442'):
    run.font.name = 'Noto Sans CJK SC'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def p(text, size=11.5, bold=False, color='263442'):
    pp = doc.add_paragraph()
    runfont(pp.add_run(text), size, bold, color)
    md.append(text + '\n')
    return pp

def heading(text):
    pp = p(text, 13, True, '24577E')
    pp.paragraph_format.keep_with_next = True
    pp.paragraph_format.space_before = Pt(6)
    md[-1] = '## ' + text + '\n'

def page(title, intro):
    if doc.paragraphs:
        doc.add_page_break()
    pp = p(title, 20, True, '24577E')
    pp.paragraph_format.keep_with_next = True
    md[-1] = '# ' + title + '\n'
    p(intro, 10, color='627180')

def table(headers, body, widths, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for j, value in enumerate(headers):
        t.rows[0].cells[j].text = str(value)
    for r in body:
        for c, v in zip(t.add_row().cells, r):
            c.text = str(v)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for j, c in enumerate(row.cells):
            c.width = Cm(widths[j])
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(4)
                pp.paragraph_format.space_before = Pt(4)
                pp.paragraph_format.line_spacing = 1.05
                for r in pp.runs:
                    runfont(r, size, i == 0)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EAF1F7')
                c._tc.get_or_add_tcPr().append(shade)
    md.append('| ' + ' | '.join(headers) + ' |\n|' + '|'.join(['---']*len(headers)) + '|\n')
    md.extend('| ' + ' | '.join(map(str, r)) + ' |\n' for r in body)
    md.append('\n')

def picture(stem, caption):
    pp = doc.add_paragraph()
    pp.paragraph_format.space_after = Pt(0)
    pp.paragraph_format.keep_with_next = True
    pp.add_run().add_picture(str(FIG / f'{stem}.png'), width=Cm(16.6))
    p(caption, 9, color='627180')
    md.append(f'![{caption}]({FIG / (stem + ".png")})\n')

header = sec.header.paragraphs[0]
runfont(header.add_run('研究工作回顾  ·  简单易懂版'), 8, color='627180')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
runfont(footer.add_run('2026.10.08  |  '), 8, color='627180')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

page('我到底做了哪些工作？', '依据你上传的《工作量》和截至10月7日的结果整理。以下按容易讲清楚的顺序排列。')
heading('先记住这一句话')
p('老师建议我把测试图片分组。我完成分组以后，又检查了分组有没有用、换一批图片还成不成立，以及能不能找出模型具体在哪些情况下容易认错。', 13, True)
heading('已经完成的工作清单')
work = [
    ['1', '整理图片，描述图片状态', '准备六类目标，并记录大小、朝向、遮挡等信息。'],
    ['2', '把图片分成五组', '同一类别内部，按规则从预计较容易到较困难排列。'],
    ['3', '测试，再换三批图片核验', '用相同模型比较各组表现；后续核验共5400个目标。'],
    ['4', '随机分组作比较', '每批随机分10000次，检查差别是不是随便分都能出现。'],
    ['5', '每次去掉一项规则再试', '依次去掉大小、可见部分、遮挡、画面切掉、朝向。'],
    ['6', '固定数量，换困难图片的比例', '每次120个目标，测试21种比例，每种抽样1000次。'],
    ['7', '补一个更简单的方法作比较', '只按目标大小分组，看是否也能发现明显差别。'],
    ['8', '找具体容易出错的情况', '例如背向的狗；再换图片、在大小相近组内检查。'],
    ['9', '整理论文、图表和计算材料', '更新结果说明，并核对已有预测能否算回相同结果。'],
]
table(['序号', '我完成了什么', '大白话说明'], work, [.8, 6, 10.4])
p('这些工作是在老师提出分层建议之后逐步完成的。上面的抽样次数用于反复核对已有数据，不等于新增了同样数量的图片。', 9, color='627180')
heading('这阶段最值得汇报的进展')
p('现在我不只知道模型总分，还能说清楚：哪些图片组分数低、哪些具体情况值得继续测，以及哪些结论还不能说得太满。')

page('1  先把“分层”这件事做出来', '“分层”就理解成分组：把不同条件的图片放到不同组里，分别看模型表现。')
heading('第一步：准备模型要认的目标')
p('研究对象是狗、猫、马、羊、人和车。每个测试样本对应图片里的一个目标；一张原图可能提供多个目标。模型的任务是判断它属于这六类中的哪一类。')
p('图片的真实类别来自原数据集。我用两套AI帮助描述图片状态：一套在本地运行，一套在线运行。它们负责描述状态，例如是否背向、是否被挡住；目标大小还由程序根据标注框计算。')
heading('第二步：规定分组时看哪些信息')
table(['看什么', '怎样理解'], [
    ['目标大小', '目标占画面的面积小不小。'],
    ['可见部分', '目标是基本完整，还是只露出一部分。'],
    ['遮挡', '有没有被别的东西挡住。'],
    ['被画面切掉', '有没有超出画面，身体被图片边缘切掉。'],
    ['朝向', '是常见视角，还是背向、俯视等视角。'],
], [4.4, 12.8], 11)
p('我按事先定好的规则给这些条件打分，在每个类别内部排队，再均分为五组。第一组是规则预计较容易的一组，第五组是预计较困难的一组。分组时不看模型已经答对还是答错。')
p('狗和狗比、猫和猫比，各组类别比例相同。这样可以减少“这一组只是某个类别更多，所以分数低”的影响。')
heading('第三步：让同一批模型分别做这五组“考题”')
p('我使用原来已经训练好的8种模型，每种有3份训练结果，共24份模型。后面的比较沿用这些模型；工作重点是组织测试图片、计算结果和做核验。')
p('注意：“较困难组”是按规则划出来的，不代表里面每一张图片都一定比其他组难。中间几组的分数也不一定逐组下降。', 10, color='627180')
heading('可以这样跟老师说')
p('“我先落实了您的分层建议：每类目标分别分成五组，再用相同的模型测试，比较规则预计较容易和较困难的图片。”')

page('2  分好组以后，我换了三批图片再测', '目的：检查结果是否只在原来那批图片上成立。')
heading('我具体做了什么？')
p('我保留原来的分组规则和模型，又整理了三批用于核验的目标：COCO的2400个、VOC2007的600个、Objects365的2400个，共5400个。')
p('COCO这里是换新样本；后两批还增加了不同图片来源。它们都包含相同的六类目标。')
heading('看图时，只看蓝柱和橙柱谁高')
picture('01_three_sources', '图1：蓝柱是较容易组，橙柱是较困难组。展示本地AI状态描述下，24份模型的平均结果；在线AI描述下也看到相同方向。')
p('这里的F1是分类表现的综合评分，换算成满分100，越高越好。它不是“100张图片中认对多少张”的正确率。', 10, color='627180')
table(['图片来源', '本地AI描述下的分差', '在线AI描述下的分差'], [
    [name, f'{-100*float(next(r for r in random if r["pool"]==pool and r["source"]=="local")["observed_delta"]):.2f}分',
     f'{-100*float(next(r for r in random if r["pool"]==pool and r["source"]=="online")["observed_delta"]):.2f}分']
    for pool, name in [('B','COCO'),('C','VOC2007'),('D','Objects365')]
], [4.4, 6.4, 6.4])
heading('这个结果说明什么？')
p('在这三批图片中，较困难组的平均分都更低。两套AI描述下，组间差值约为14到42分。这支持分组规则与模型表现有关。')
p('可以这样汇报：“分组有了差别以后，我又换了三批图片检查，差别仍然出现，所以结果并不只来自最初那批图片。”')

page('3  我又检查：是不是碰巧分出了差别？', '不只看一次结果，还做了几种简单的比较。')
heading('检查一：不按规则，随机分五组会怎样？')
p('我把同一类别的目标随机打乱，再分成五组；每批做10000次。各组大小和类别比例保持不变。')
p('结果：随机分组的分差大多在零附近波动，这些随机结果没有出现真实规则分组那么大的差值。这支持图片条件提供了有用的分组信息。')
heading('检查二：少用一项规则，结果还在吗？')
p('我分别去掉大小、可见部分、遮挡、被画面切掉和朝向，每次重新分组、重新比较。')
p('结果：去掉任意一项后，较困难组平均分更低的方向仍在。不过还不能据此确定每一项单独起了多大作用。')
heading('检查三：图片数量一样，只换困难图片的比例')
p('每次固定120个目标，六类各20个。我逐步增加其中较困难组的比例，共测21种比例，每种抽样1000次。')
picture('02_same_number', '图2：以COCO、本地AI描述结果为例。左边从前四组混合选120个，右边从第五组选120个；显示24份模型和重复抽样的平均F1。')
p('数量一样，考题变难以后，平均F1从90.2分降到69.4分。这里变化的是测试图片，模型没有重新训练。')
heading('可以这样跟老师说')
p('“我补做了随机分组、少用一项规则和固定测试数量的检查。尤其是数量不变时，换一批更困难的图片，总分仍会明显改变。汇报总分时需要同时说明测了什么图片。”')

page('4  最新补的工作：简单方法对比和具体问题', '不只证明复杂规则能分组，还检查简单方法能做到什么。')
heading('新增比较：只看目标大小，也能分组吗？')
p('我只按目标在图片里占的面积分五组，跟原来的五项规则作比较。结果发现，只看大小也能分出明显的表现差别。')
table(['图片来源', '五项规则的分差', '只看大小的分差'], [
    [name, f'{100*float(next(r for r in scale if r["pool"]==pool and r["source"]=="local")["full_gap"]):.2f}分',
     f'{100*float(next(r for r in scale if r["pool"]==pool and r["source"]=="local")["area_gap"]):.2f}分']
    for pool, name in [('B','COCO'),('C','VOC2007'),('D','Objects365')]
], [4.4, 6.4, 6.4])
p('上表用本地AI描述结果举例。完整比较还包含在线结果：五项规则没有在所有来源中都比只看大小更强。这里的分差大小也不等于识别模型总分高低。', 9, color='627180')
heading('继续找：具体哪些情况值得重点测试？')
p('我进一步整理“某个模型在某种图片条件下的表现”。例如，在COCO中，ResNet-50对全部狗样本的识别率约82%，但对背向狗只有约57%。')
picture('03_back_facing_dog', f'图3：真实狗被正确认成狗的比例，正式名称是Recall。背向狗子组有{case_n}个目标，数值是三份训练结果的平均；该例帮助解释问题，不代表所有模型。')
p('因此，我可以提出更具体的测试方向：“继续测背向狗”，而不仅说“这个模型分数不高”。之前发现的薄弱条件也被放到新图片上核验；有些成立，有些没有足够样本。')
p('我还把每类目标按大小分组，在大小相近的组内比较其他状态。COCO和Objects365中，多数能比较的条件仍表现较差；VOC很多条件样本不足，暂时不好判断。')
heading('可以这样跟老师说')
p('“最新我补了只看大小的对照。它也很强，所以不能说五项规则一定更好。状态描述更直接的用途是告诉我：模型具体在哪些情况下值得继续测。”')

page('5  你可以直接照着讲', '把工作按“先分组 → 换图片核验 → 做比较 → 找具体问题”讲清楚即可。')
heading('约3分钟的汇报稿')
speech = [
    '老师，上次您建议我把测试数据分层。我先把这个建议落实了：针对狗、猫、马、羊、人和车，按照目标大小、可见部分、遮挡、被画面切掉和朝向，把每类目标分成五组，再用原来已经训练好的模型比较各组表现。',
    '发现较困难组分数低以后，我又换了三批图片检查，共5400个目标。三批里都出现了相同方向，较容易组和较困难组的综合评分大约相差14到42分。',
    '然后我检查这个差别是不是偶然。每批随机分组10000次，并且分别去掉一项分组规则再试。我还把每次测试数量固定为120个，只增加困难图片的比例。结果说明：同一个模型换一套条件不同的考题，测出来的总分会明显改变。',
    '最新我补了只看目标大小的简单对照，发现它也能分出很大的差别，所以还不能说五项规则一定更好。我继续把结果细化到具体情况，例如背向狗比全部狗样本更容易认错，并在新图片、大小相近的分组中检查这些情况。',
    '最后，我把结果、案例和图表整理进了论文，并准备了计算材料。现在比较明确的成果，是能把一个总体分数展开成具体的测试情况，帮助决定下一步重点测什么。',
]
for text in speech:
    p(text)
heading('你需要分清的两件事')
p('① 工作量：你整理了数据、设计了分组、做了多种比较、核验了结果，还整理了论文材料。不能只说“我画了几张图”。')
p('② 研究价值：你更具体地描述了模型在不同条件下的表现，也查出了方法的局限。分数降低本身并不是成果，能解释这种差别才有用。')
heading('汇报时保留这三句话')
p('• 图片状态由AI自动描述，目前没有完整的人工核对，不能保证每项都正确。\n• 找到的是容易出错的相关情况，还没有证明它就是错误原因，也没有完成模型修复。\n• 五项规则没有全面超过只看大小的方法，论文主要强调分组观察和具体条件报告。', 10)
p('工作材料也已核对：利用保存的预测重新计算，8张新增结果表和60项历史比较与原结果一致。这是计算结果核对，不是重新训练全部模型。', 9, color='627180')
p('核查来源：论文_最新.md；random_stratification_ablation_20261006；controlled_risk_distribution_20261005；scale_only_confirmation_20261007；submission_closeout_20261007/portable_replay_audit.json。以上实验目录均在当前工作目录的artifacts下。', 8, color='627180')

doc.core_properties.title = '工作量：简单易懂版'
doc.core_properties.subject = '老师分层建议之后，已完成工作的通俗解释'
doc.core_properties.author = 'QE研究工作整理'
doc.save(OUT / f'{NAME}.docx')
(OUT / f'{NAME}.md').write_text('\n'.join(md), encoding='utf-8')
for ext in ['docx', 'md']:
    shutil.copy2(OUT / f'{NAME}.{ext}', ROOT / 'docs' / f'{NAME}.{ext}')
shutil.copy2(Path(__file__), OUT / Path(__file__).name)
(OUT / 'audit.json').write_text(json.dumps({
    'input_sha256': {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
    'planned_pages': 6, 'figures': 3, 'work_items': 9,
    'source_numbers_checked': True, 'case_support': case_n,
    'condition_counts': condition_counts,
    'new_experiments_run': False,
}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'word': str(ROOT / 'docs' / f'{NAME}.docx'), 'figures': str(FIG),
                  'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)}, ensure_ascii=False))
