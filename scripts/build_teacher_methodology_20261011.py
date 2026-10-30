#!/usr/bin/env python3
"""Export the teacher edition from one Markdown source; audit its numeric evidence."""
from pathlib import Path
import csv
import hashlib
import json
import re
import statistics
import sys
import zipfile
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
TASK = ROOT / 'artifacts/teacher_methodology_20261011'
sys.path.insert(0, str(TASK / 'vendor'))

from PIL import Image as PILImage, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, NextPageTemplate, PageBreak, PageTemplate,
    Paragraph, Spacer, Table, TableStyle,
)
import pymupdf

STEM = '视觉状态模型测试方法_老师审阅版_20261011'
SOURCE = ROOT / 'docs' / (STEM + '.md')
FIGURE = ROOT / 'output/imagegen/qe_method_final_stage_image2_20261010.png'
FONT = ROOT / 'scripts/assets/NotoSansSC-Regular.ttf'
BLUE = '#173653'
TEAL = '#176C75'
GRAY = '#55616D'
LINK = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')


def read_rows(path):
    with (ROOT / path).open(encoding='utf-8', newline='') as handle:
        return list(csv.DictReader(handle))


def evidence_audit():
    risk_path = 'artifacts/random_stratification_ablation_20261006/tables/random_summary.csv'
    mix_path = 'artifacts/controlled_risk_distribution_20261005/tables/panel_summary.csv'
    case_path = 'artifacts/sixclass_independent_confirmation_20261004/transfer/tables/state_validation.csv'
    area_path = 'artifacts/scale_only_confirmation_20261007/tables/summary.csv'
    risks = read_rows(risk_path)
    mixtures = read_rows(mix_path)
    cases = read_rows(case_path)
    areas = read_rows(area_path)
    expected = {
        ('B', 'local'): (-25.50, -20.78), ('B', 'online'): (-21.30, -16.62),
        ('C', 'local'): (-16.57, -15.92), ('C', 'online'): (-14.44, -12.29),
        ('D', 'local'): (-41.80, -34.77), ('D', 'online'): (-41.25, -31.84),
    }
    checked = []
    for key, (risk_expected, mix_expected) in expected.items():
        pool, source = key
        risk = next(r for r in risks if (r['pool'], r['source']) == key)
        endpoints = {
            int(r['high_risk_percent']): float(r['macro_f1_mean'])
            for r in mixtures if (r['pool'], r['source']) == key
            and r['total'] == '120' and r['high_risk_percent'] in ('0', '100')
        }
        risk_pp = float(risk['observed_delta']) * 100
        mix_pp = (endpoints[100] - endpoints[0]) * 100
        assert round(risk_pp, 2) == risk_expected, (key, risk_pp)
        assert round(mix_pp, 2) == mix_expected, (key, mix_pp)
        assert risk['repeats'] == '10000' and risk['extreme_count'] == '0'
        assert abs(float(risk['holm_six_p']) - 0.0005999400059994) < 1e-12
        checked.append({'pool': pool, 'source': source, 'risk_gap_pp': risk_pp,
                        'mixture_change_pp': mix_pp, 'corrected_p': float(risk['holm_six_p'])})
    case = [r for r in cases if r['source'] == 'local' and r['model'] == 'resnet50'
            and r['class_name'] == 'dog' and r['feature'] == 'body_orientation'
            and r['value'] == 'back']
    assert len(case) == 3 and {r['B_support'] for r in case} == {'46'}
    class_recall = statistics.mean(float(r['B_class_recall']) for r in case) * 100
    state_recall = statistics.mean(float(r['B_recall']) for r in case) * 100
    gap = statistics.mean(float(r['B_gap']) for r in case) * 100
    assert (round(class_recall, 2), round(state_recall, 2), round(gap, 2)) == (81.92, 57.25, -24.67)
    assert all(float(r['increment_mean']) < 0 for r in areas if r['pool'] == 'D')
    assert all(float(r['holm_six_p']) > .05 for r in areas if float(r['increment_mean']) > 0)
    return {
        'numeric_checks': checked,
        'condition_example': {'model': 'resnet50', 'class': 'dog', 'state': 'back',
                              'support': 46, 'seeds': 3, 'class_recall_percent': class_recall,
                              'state_recall_percent': state_recall, 'difference_pp': gap},
        'inputs': {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest()
                   for p in (risk_path, mix_path, case_path, area_path)},
    }


def parse_source():
    pages = [{'orientation': 'portrait', 'blocks': []}]
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        marker = re.fullmatch(r'<!-- page: (portrait|landscape) -->', line)
        if marker:
            pages.append({'orientation': marker[1], 'blocks': []})
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            pages[-1]['blocks'].append(('table', rows))
            continue
        if line.startswith('!['):
            match = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
            path = (SOURCE.parent / match[2]).resolve()
            assert path == FIGURE
            pages[-1]['blocks'].append(('image', str(path)))
        elif line.startswith('# '):
            pages[-1]['blocks'].append(('title', line[2:]))
        elif line.startswith('## '):
            pages[-1]['blocks'].append(('h1', line[3:]))
        elif line.startswith('### '):
            pages[-1]['blocks'].append(('h2', line[4:]))
        elif line.startswith('- '):
            pages[-1]['blocks'].append(('bullet', line[2:]))
        elif line.startswith(('图1', '图示说明：')):
            pages[-1]['blocks'].append(('caption', line))
        elif line.startswith('老师审阅版｜'):
            pages[-1]['blocks'].append(('subtitle', line))
        else:
            pages[-1]['blocks'].append(('paragraph', line))
        i += 1
    assert len(pages) == 6
    return pages


def set_font(run, size=None, bold=False, color=None):
    run.font.name = 'Microsoft YaHei'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.strip('#'))


def hyperlink(paragraph, label, url, size):
    node = OxmlElement('w:hyperlink')
    node.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Microsoft YaHei')
    fonts.set(qn('w:eastAsia'), '微软雅黑')
    props.append(fonts)
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '176C75')
    props.append(c)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(round(size * 2)))
    props.append(sz)
    run.append(props)
    text_node = OxmlElement('w:t')
    text_node.text = label
    run.append(text_node)
    node.append(run)
    paragraph._p.append(node)


def add_inline(paragraph, value, size=10.5, bold=False, color=None):
    pos = 0
    for match in LINK.finditer(value):
        set_font(paragraph.add_run(value[pos:match.start()]), size, bold, color)
        hyperlink(paragraph, match[1], match[2], size)
        pos = match.end()
    set_font(paragraph.add_run(value[pos:]), size, bold, color)


def layout_section(section, orientation):
    section.orientation = WD_ORIENT.LANDSCAPE if orientation == 'landscape' else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if orientation == 'landscape' else 21)
    section.page_height = Cm(21 if orientation == 'landscape' else 29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.55)
    section.left_margin = section.right_margin = Cm(1.65)
    section.header_distance = section.footer_distance = Cm(.7)


def table_ratios(rows):
    if rows[0][0] == '编号':
        return [.075, .50, .425]
    if rows[0][0] == '确认池／状态来源':
        return [.30, .31, .39]
    if len(rows[0]) == 2:
        return [.29, .71]
    return [.20, .50, .30]


def build_docx(pages):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [('Title', 23), ('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLUE.strip('#'))
        style.paragraph_format.space_before = Pt(9 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(5)
    header = doc.sections[0].header.paragraphs[0]
    set_font(header.add_run('QE研究方法说明 · 老师审阅版'), 8.5, color=GRAY)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run('2026-10-11  |  '), 8, color=GRAY)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for index, page in enumerate(pages):
        section = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        layout_section(section, page['orientation'])
        width = (section.page_width - section.left_margin - section.right_margin) / Cm(1)
        for kind, value in page['blocks']:
            if kind == 'image':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(4)
                p.add_run().add_picture(value, width=Cm(width))
            elif kind == 'table':
                rows = value
                ratios = table_ratios(rows)
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                for col, ratio in zip(table.columns, ratios):
                    col.width = Cm(width * ratio)
                for row_index, (row, cells) in enumerate(zip(table.rows, rows)):
                    props = row._tr.get_or_add_trPr()
                    props.append(OxmlElement('w:cantSplit'))
                    if row_index == 0:
                        props.append(OxmlElement('w:tblHeader'))
                    for cell, text_value, ratio in zip(row.cells, cells, ratios):
                        cell.width = Cm(width * ratio)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        tcpr = cell._tc.get_or_add_tcPr()
                        fill = OxmlElement('w:shd')
                        fill.set(qn('w:fill'), '173653' if row_index == 0 else ('EDF3F6' if row_index % 2 else 'F7F9FA'))
                        tcpr.append(fill)
                        margins = OxmlElement('w:tcMar')
                        for side in ('top', 'bottom', 'left', 'right'):
                            item = OxmlElement('w:' + side)
                            item.set(qn('w:w'), '65')
                            item.set(qn('w:type'), 'dxa')
                            margins.append(item)
                        tcpr.append(margins)
                        p = cell.paragraphs[0]
                        p.paragraph_format.line_spacing = 1.08
                        p.paragraph_format.space_after = Pt(1)
                        add_inline(p, text_value, 9.0, row_index == 0, '#FFFFFF' if row_index == 0 else None)
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
            else:
                style = {'title': 'Title', 'h1': 'Heading 1', 'h2': 'Heading 2'}.get(kind)
                p = doc.add_paragraph(style=style)
                if kind in ('title', 'h1', 'h2'):
                    size = {'title': 23, 'h1': 16, 'h2': 12}[kind]
                    add_inline(p, value, size, True, BLUE)
                else:
                    size = 9 if kind in ('caption', 'subtitle') else 10.5
                    if kind == 'bullet':
                        p.paragraph_format.left_indent = Cm(.35)
                        p.paragraph_format.first_line_indent = Cm(-.28)
                        value = '• ' + value
                    add_inline(p, value, size, color=GRAY if kind in ('caption', 'subtitle') else None)
    doc.core_properties.title = '视觉状态驱动的多类别模型测试方法'
    doc.core_properties.subject = '当前研究方法、总体框架与实证依据（老师审阅版）'
    doc.core_properties.author = 'QE项目'
    path = SOURCE.with_suffix('.docx')
    doc.save(path)
    return path


def pdf_text(value):
    chunks, pos = [], 0
    for match in LINK.finditer(value):
        chunks.append(escape(value[pos:match.start()]))
        chunks.append(f'<link href="{escape(match[2])}" color="{TEAL}">{escape(match[1])}</link>')
        pos = match.end()
    chunks.append(escape(value[pos:]))
    return ''.join(chunks)


def build_pdf(pages):
    pdfmetrics.registerFont(TTFont('NotoSC', str(FONT)))
    pdfmetrics.registerFontFamily('NotoSC', normal='NotoSC', bold='NotoSC', italic='NotoSC', boldItalic='NotoSC')
    styles = {}
    specs = {
        'title': (23, 30, 12), 'h1': (16, 23, 9), 'h2': (12, 18, 5),
        'paragraph': (10.5, 16.0, 7), 'bullet': (10.5, 16, 5),
        'caption': (9, 13.5, 7), 'subtitle': (9, 14, 12),
        'cell': (9, 12.7, 0), 'header': (9, 12.7, 0),
    }
    for key, (size, leading, after) in specs.items():
        styles[key] = ParagraphStyle(
            key, fontName='NotoSC', fontSize=size, leading=leading, spaceAfter=after,
            spaceBefore=8 if key == 'h2' else 0,
            textColor=colors.HexColor(BLUE if key in ('title', 'h1', 'h2') else GRAY if key in ('caption', 'subtitle') else '#18232D'),
            alignment=TA_LEFT, wordWrap='CJK', allowWidows=0, allowOrphans=0,
            keepWithNext=key in ('title', 'h1', 'h2'),
            leftIndent=10 if key == 'bullet' else 0,
            firstLineIndent=-8 if key == 'bullet' else 0,
        )
    styles['header'].textColor = colors.white
    path = SOURCE.with_suffix('.pdf')
    margin = 46.8
    top, bottom = 43, 44
    def on_page(canvas, doc):
        w, h = canvas._pagesize
        canvas.setFont('NotoSC', 8)
        canvas.setFillColor(colors.HexColor(GRAY))
        canvas.drawString(margin, h - 24, 'QE研究方法说明 · 老师审阅版')
        canvas.setStrokeColor(colors.HexColor('#DDE4E9'))
        canvas.line(margin, h - 29, w - margin, h - 29)
        canvas.drawRightString(w - margin, 23, f'2026-10-11  |  {doc.page}')
    doc = BaseDocTemplate(str(path), pagesize=A4, title='视觉状态驱动的多类别模型测试方法', author='QE项目')
    for orientation, dimensions in [('portrait', A4), ('landscape', landscape(A4))]:
        w, h = dimensions
        frame = Frame(margin, bottom, w - 2 * margin, h - top - bottom,
                      leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
                      id=orientation)
        doc.addPageTemplates(PageTemplate(id=orientation, frames=frame, onPage=on_page, pagesize=dimensions))
    story = []
    for index, page in enumerate(pages):
        if index:
            story.extend([NextPageTemplate(page['orientation']), PageBreak()])
        w = (landscape(A4) if page['orientation'] == 'landscape' else A4)[0] - 2 * margin
        for kind, value in page['blocks']:
            if kind == 'image':
                with PILImage.open(value) as im:
                    ratio = im.height / im.width
                story.extend([Image(value, width=w, height=w * ratio), Spacer(1, 6)])
            elif kind == 'table':
                ratios = table_ratios(value)
                contents = [[Paragraph(pdf_text(cell), styles['header' if i == 0 else 'cell'])
                             for cell in row] for i, row in enumerate(value)]
                table = Table(contents, colWidths=[w * r for r in ratios], repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(BLUE)),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#EDF3F6'), colors.HexColor('#F7F9FA')]),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ]))
                story.extend([table, Spacer(1, 9)])
            else:
                text_value = ('• ' if kind == 'bullet' else '') + value
                story.append(Paragraph(pdf_text(text_value), styles[kind]))
    doc.build(story)
    return path


def validate_and_render(docx_path, pdf_path, pages, audit):
    with zipfile.ZipFile(docx_path) as archive:
        media = [p for p in archive.namelist() if p.startswith('word/media/')]
        assert len(media) == 1
        assert hashlib.sha256(archive.read(media[0])).digest() == hashlib.sha256(FIGURE.read_bytes()).digest()
        xml = archive.read('word/document.xml').decode()
        assert xml.count('<w:sectPr') == 6
        assert '/home/hcf' not in xml
        rels = archive.read('word/_rels/document.xml.rels').decode()
        assert 'TargetMode="External"' in rels and 'arxiv.org/abs/2203.14960' in rels
    pdf = pymupdf.open(pdf_path)
    assert len(pdf) == len(pages), f'Unexpected PDF pagination: {len(pdf)} instead of {len(pages)}'
    rendered = TASK / 'preview'
    rendered.mkdir(parents=True, exist_ok=True)
    audit['pdf_pages'] = []
    thumbs = []
    for number, (page, logical) in enumerate(zip(pdf, pages), 1):
        w, h = page.rect.width, page.rect.height
        assert (w > h) == (logical['orientation'] == 'landscape'), (number, w, h)
        title = next(value for kind, value in logical['blocks'] if kind in ('h1', 'title'))
        extracted = page.get_text()
        assert title.replace(' ', '') in extracted.replace('\n', '').replace(' ', ''), (number, title)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(1.7, 1.7), alpha=False)
        path = rendered / f'page_{number:02d}.png'
        pix.save(path)
        im = PILImage.open(path).convert('RGB')
        im.thumbnail((500, 705))
        tile = PILImage.new('RGB', (530, 740), '#DEE5EA')
        tile.paste(im, ((530 - im.width) // 2, 25))
        ImageDraw.Draw(tile).text((12, 8), f'Page {number}', fill='#173653')
        thumbs.append(tile)
        audit['pdf_pages'].append({'number': number, 'orientation': logical['orientation'],
                                   'text_characters': len(extracted), 'external_links': len(page.get_links())})
    montage = PILImage.new('RGB', (1060, 2220), 'white')
    for index, thumb in enumerate(thumbs):
        montage.paste(thumb, ((index % 2) * 530, (index // 2) * 740))
    montage.save(rendered / 'contact_sheet.png')
    text = '\n'.join(p.get_text() for p in pdf)
    for value in ('81.92', '57.25', '24.67', '41.80', '34.77', '3368089.3409754'):
        # DOI targets are carried by clickable annotations rather than printed URLs.
        if value != '3368089.3409754':
            assert value in text, value
    assert len(pdf[1].get_images()) >= 1
    audit['figure_embedded'] = True
    audit['exports'] = {
        'docx': str(docx_path.relative_to(ROOT)), 'pdf': str(pdf_path.relative_to(ROOT)),
        'source': str(SOURCE.relative_to(ROOT)), 'figure': str(FIGURE.relative_to(ROOT)),
    }
    audit['pdf_generation'] = 'Same source as Word; PDF generated with ReportLab, not a render of Word.'
    (TASK / 'document_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'files': audit['exports'], 'pages': len(pdf), 'embedded_figure': True,
                      'validated_numeric_rows': len(audit['numeric_checks']),
                      'preview': str((rendered / 'contact_sheet.png').relative_to(ROOT))}, ensure_ascii=False, indent=2))


def main():
    TASK.mkdir(parents=True, exist_ok=True)
    audit = evidence_audit()
    pages = parse_source()
    docx = build_docx(pages)
    pdf = build_pdf(pages)
    validate_and_render(docx, pdf, pages, audit)


if __name__ == '__main__':
    main()
