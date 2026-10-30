#!/usr/bin/env python3
"""Revise the uploaded Methods document while retaining its styles and figure."""
from copy import deepcopy
from pathlib import Path
import hashlib
import json
import re
import zipfile

import export_paper_methods_20261011 as export

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
INPUT = Path('/home/hcf/.codex/attachments/7b71627d-a446-45c0-805b-be3460e9ab58/论文方法章节_视觉状态驱动模型测试_核验修订版_20261012.docx')
SOURCE = ROOT / 'docs/论文方法章节_视觉状态驱动模型测试_公式规范版_20261012.md'
OUT = ROOT / 'artifacts/methods_formula_format_20261012'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}


def structure_signature(element):
    return (element.tag, tuple(sorted(element.attrib.items())), element.text,
            tuple(structure_signature(child) for child in element))


def indicator(condition):
    one = export.mr('1')
    properties = OxmlElement('m:rPr')
    style = OxmlElement('m:sty')
    style.set(qn('m:val'), 'b')
    properties.append(style)
    one.insert(0, properties)
    return export.flat(one, export.delim(condition, '{', '}'))


original_inline = export.inline_math
original_display = export.display_math


def inline_math(tex):
    if tex == 'T':
        return export.flat('T')
    if tex in ('z_i^{(r)}', 'H_i^{(r)}'):
        return export.flat(export.script(tex[0], 'i', '(r)'))
    if tex == 'n_{c,s}^{(r)}(S)':
        return export.support()
    if tex in ('C_c^{(r)}(S)', 'C^{(r)}(S)'):
        subscript = 'c' if tex.startswith('C_c') else None
        return export.flat(export.script('C', subscript, '(r)'), export.delim('S'))
    if tex == r'g_c^{(r)}(i\mid S)':
        return export.flat(export.script('g', 'c', '(r)'), export.delim('i|S'))
    if tex == r'Q_{c,\ell}^{(r)}':
        return export.flat(export.script('Q', 'c,ℓ', '(r)'))
    if tex == r'\hat y_i':
        return export.flat(export.script(export.hat('y'), 'i'))
    if tex == r'\ell':
        return export.flat('ℓ')
    if tex == r'\ell=1,\ldots,5':
        return export.flat('ℓ=1,…,5')
    if tex == r'S=\varnothing':
        return export.flat('S=∅')
    if tex == r'\log(n+2)-\log(n+1)':
        return export.flat(export.mr('log', True), export.delim('n+2'), '−',
                           export.mr('log', True), export.delim('n+1'))
    if tex == r'\rho\in[0,1]':
        return export.flat('ρ∈', export.delim('0,1', '[', ']'))
    if tex == r'(1-\rho)m':
        return export.flat(export.delim('1−ρ'), 'm')
    if tex == r'\mathbf{1}\{\cdot\}':
        return indicator('·')
    return original_inline(tex)


def display_math(number):
    if number == 9:
        expectation = export.mr('E')
        properties = OxmlElement('m:rPr')
        script = OxmlElement('m:scr')
        script.set(qn('m:val'), 'double-struck')
        properties.append(script)
        expectation.insert(0, properties)
        return export.flat(expectation,
                           export.delim(indicator(export.mr('type observed', True)), '[', ']'),
                           '=1−', export.frac(export.binom('N−e', 'q'), export.binom('N', 'q')), '.')
    if number == 3:
        nodes = original_display(number)
        cardinality = nodes[3].find('m:den', NS)
        for child in list(cardinality):
            cardinality.remove(child)
        cardinality.extend(export.flat(export.delim(export.script('U', 'c'), '|', '|')))
        return nodes
    return original_display(number)


export.indicator = indicator
export.inline_math = inline_math
export.display_math = display_math


def rewrite(paragraph, text):
    properties = paragraph._p.find(qn('w:pPr'))
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)
    export.word_inline(paragraph, text)


def append_paragraph_after(document, anchor, text, size=11, caption=False):
    paragraph = document.add_paragraph()
    paragraph.style = 'Normal'
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.first_line_indent = Cm(0 if caption else .75)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = caption
    export.word_inline(paragraph, text, size)
    anchor.addnext(paragraph._p)
    return paragraph


def normalize_math(root):
    for delimiter in root.findall('.//m:d', NS):
        properties = delimiter.find('m:dPr', NS)
        if properties is None:
            properties = OxmlElement('m:dPr')
            delimiter.insert(0, properties)
        for tag, default in [('begChr', '('), ('endChr', ')')]:
            if properties.find('m:' + tag, NS) is None:
                node = OxmlElement('m:' + tag)
                node.set(qn('m:val'), default)
                properties.insert(0 if tag == 'begChr' else 1, node)
    for value in root.findall('.//m:type', NS):
        if value.get(qn('m:val')) == 'nobar':
            value.set(qn('m:val'), 'noBar')
    for math in root.findall('.//m:oMath', NS):
        for run in math.findall('.//m:r', NS):
            properties = run.find('w:rPr', NS)
            if properties is None:
                properties = OxmlElement('w:rPr')
                run.insert(1 if run.find('m:rPr', NS) is not None else 0, properties)
            fonts = properties.find('w:rFonts', NS)
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                properties.insert(0, fonts)
            fonts.set(qn('w:ascii'), 'Cambria Math')
            fonts.set(qn('w:hAnsi'), 'Cambria Math')


def write_document(blocks):
    document = Document(INPUT)
    original_styles = deepcopy(document.styles._element)
    original_section = deepcopy(document.sections[0]._sectPr)
    original_paragraphs = list(document.paragraphs)
    algorithm_table = document.tables[0]
    new_body = [text for kind, text, number in blocks if kind == 'body']
    changed_starts = [
        ('对同一候选池', '对同一候选池'),
        ('已有支持越少', '每个未饱和状态'),
        ('Sample表示无放回抽样', 'Sample表示无放回抽样'),
        ('对选样结果分别统计失败状态广度', '对选样结果分别统计失败状态广度'),
        ('各维度映射如表1所示', '各维度映射如表2所示'),
        ('风险检验包括', '风险检验包括'),
    ]
    for old_start, new_start in changed_starts:
        paragraph = next(p for p in original_paragraphs if p.text.startswith(old_start))
        text = next(t for t in new_body if t.startswith(new_start))
        rewrite(paragraph, text)
    for paragraph in original_paragraphs:
        if paragraph.text.startswith('表1　预定义风险映射'):
            rewrite(paragraph, next(text for kind, text, number in blocks
                                    if kind == 'caption' and text.startswith('表2')))
            for run in paragraph.runs:
                run.font.size = Pt(9)
        elif paragraph._p.find('.//m:oMathPara', NS) is None:
            for math in paragraph._p.findall('.//m:oMath', NS):
                if ''.join(math.itertext()) == '1·':
                    for child in list(math):
                        math.remove(child)
                    math.extend(indicator('·'))

    formulas = document._element.findall('.//m:oMathPara', NS)
    assert len(formulas) == 9
    for number, formula in enumerate(formulas, 1):
        for child in list(formula):
            formula.remove(child)
        math = OxmlElement('m:oMath')
        math.extend(display_math(number))
        formula.append(math)

    anchor = next(p for p in original_paragraphs if p.text.startswith('记被测模型为'))._p
    intro = append_paragraph_after(document, anchor, '主要符号及含义见表1，后文沿用同一记号。')
    caption = append_paragraph_after(document, intro._p, '表1　主要符号及含义。', 9, True)
    symbol_rows = next(text for kind, text, number in blocks if kind == 'table' and text[0][0] == '符号')
    table = document.add_table(rows=len(symbol_rows), cols=2)
    table.style = algorithm_table.style
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(6.2), Cm(10.8)
    for index, (row, values) in enumerate(zip(table.rows, symbol_rows)):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, value in zip(row.cells, values):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.paragraph_format.space_after = Pt(3)
            export.word_inline(paragraph, value, 9.5)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True
    caption._p.addnext(table._tbl)

    algorithm = next(text for kind, text, number in blocks if kind == 'algorithm')
    previous = algorithm_table._tbl
    for index, line in enumerate(algorithm):
        paragraph = append_paragraph_after(document, previous, line.strip(), 10)
        match = re.match(r'^(\d+) ( +)', line)
        indent = len(match[2]) if match else 0
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(indent * .12)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = index < len(algorithm) - 1
        previous = paragraph._p
    append_paragraph_after(document, previous,
                           '状态计数在风险层循环外初始化，并在各层之间累计。')
    algorithm_table._tbl.getparent().remove(algorithm_table._tbl)
    normalize_math(document._element)
    assert structure_signature(original_styles) == structure_signature(document.styles._element)
    assert structure_signature(original_section) == structure_signature(document.sections[0]._sectPr)
    path = SOURCE.with_suffix('.docx')
    document.save(path)
    return path


def audit(docx_path):
    with zipfile.ZipFile(INPUT) as original, zipfile.ZipFile(docx_path) as revised:
        tree = etree.fromstring(revised.read('word/document.xml'))
        original_tree = etree.fromstring(original.read('word/document.xml'))
        formulas = tree.findall('.//m:oMathPara', NS)
        assert len(formulas) == 9
        assert len(tree.findall('.//m:fPr/m:type[@m:val="noBar"]', NS)) == 2
        assert not tree.findall('.//m:type[@m:val="nobar"]', NS)
        assert all(d.find('m:dPr/m:begChr', NS) is not None
                   and d.find('m:dPr/m:endChr', NS) is not None
                   for d in tree.findall('.//m:d', NS))
        assert len(tree.findall('.//m:oMath', NS)) > 64
        style_ids = {paragraph.style.style_id for paragraph in Document(INPUT).paragraphs
                     if paragraph.style.name in ('Title', 'Heading 1', 'Heading 2')}
        def headings(root):
            return [structure_signature(paragraph) for paragraph in root.findall('.//w:body/w:p', NS)
                    if paragraph.find('w:pPr/w:pStyle', NS) is not None
                    and paragraph.find('w:pPr/w:pStyle', NS).get(qn('w:val')) in style_ids]
        assert headings(tree) == headings(original_tree)
        media = [name for name in revised.namelist() if name.startswith('word/media/')]
        old_media = [name for name in original.namelist()
                     if name.startswith('word/media/') and not name.endswith('/')]
        assert len(media) == len(old_media) == 1
        assert revised.read(media[0]) == original.read(old_media[0])
        assert len(tree.findall('.//w:tbl', NS)) == 2
        visible = ''.join(tree.itertext())
        assert '人工' not in visible and '\\binom' not in visible and '$' not in visible
    audit = {'input': str(INPUT), 'input_sha256': hashlib.sha256(INPUT.read_bytes()).hexdigest(),
             'output': str(docx_path), 'heading_runs_and_format_preserved': True,
             'styles_and_section_preserved': True, 'figure_preserved': True,
             'display_equations': 9, 'math_format': 'Native editable Office Math (OMML)',
             'combination_type': 'noBar', 'explicit_delimiters': True,
             'symbol_table_added': True, 'algorithm_is_paragraph_pseudocode': True,
             'new_experiments': False,
             'preview_scope': 'PDF is generated from the same source, not rendered from Word.',
             'math_schema_reference': 'https://learn.microsoft.com/en-us/dotnet/api/documentformat.openxml.math.fractiontypevalues?view=openxml-3.0.1'}
    (OUT / 'uploaded_document_revision_audit.json').write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def main():
    export.SOURCE = SOURCE
    export.OUT = OUT
    (OUT / 'equations').mkdir(parents=True, exist_ok=True)
    blocks = export.parse()
    word = write_document(blocks)
    pdf = export.pdf_export(blocks)
    with zipfile.ZipFile(INPUT) as package:
        media = next(name for name in package.namelist()
                     if name.startswith('word/media/') and not name.endswith('/'))
        figure_bytes = package.read(media)
    export.audit_and_preview(word, pdf, blocks, figure_bytes)
    audit(word)


if __name__ == '__main__':
    main()
