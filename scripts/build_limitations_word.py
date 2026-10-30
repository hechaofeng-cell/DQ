#!/usr/bin/env python3
"""Export the companion limitations report and verify its evidence scope."""
from pathlib import Path
from zipfile import ZipFile
import json
import xml.etree.ElementTree as ET

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from export_empirical_paper_docx import export, set_run_font

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/局限性.md'
TARGET = SOURCE.with_suffix('.docx')
OUT = ROOT / 'artifacts/limitations_report_20261006'


def main():
    text = SOURCE.read_text()
    paper = (ROOT / 'docs/visual_state_testing_paper_controls_20261006.md').read_text()
    official_rows = [
        '| 遮挡 | 本地 | 2289/2289 | 61.16 | 56.21 | 82.43 | 66.84 | 41.93 | 0.238 |',
        '| 遮挡 | 线上 | 2284/2289 | 66.07 | 60.50 | 82.10 | 69.67 | 51.58 | 0.331 |',
        '| 截断 | 本地 | 2289/2289 | 67.50 | 52.07 | 88.12 | 65.46 | 56.41 | 0.384 |',
        '| 截断 | 线上 | 2287/2289 | 79.62 | 66.73 | 83.25 | 74.08 | 77.67 | 0.576 |',
    ]
    for row in official_rows:
        assert row in paper
        p = [x.strip() for x in row.strip('|').split('|')]
        assert '| ' + ' | '.join([p[0], p[1], p[3], p[6], p[8]]) + ' |' in text
    table = pd.read_csv(ROOT / 'artifacts/random_stratification_ablation_20261006/tables/ablation_summary.csv')
    assert (table[table.variant != 'full'].holm_twelve_p == .09375).all()
    assert '0.09375' in text
    assert all(f'局限性{i}：' in text for i in range(1, 11))
    export(SOURCE, TARGET)
    doc = Document(TARGET)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(1.9)
    header = section.header.paragraphs[0]
    set_run_font(header.add_run('视觉状态与风险分层测试 / 结论边界'), 'Calibri', '微软雅黑', 9)
    for p in doc.paragraphs:
        heading = p.style.name.startswith('Heading')
        for r in p.runs:
            set_run_font(r, 'Calibri', '微软雅黑', 14 if heading else 20 if p.style.name == 'Title' else 11,
                         True if heading or p.style.name == 'Title' else r.bold)
            if heading:
                r.font.color.rgb = RGBColor.from_string('2166AC')
        if heading:
            p.paragraph_format.keep_with_next = True
        elif p.style.name == 'Normal':
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.first_line_indent = Cm(0)
    for t in doc.tables:
        for i, row in enumerate(t.rows):
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if i == 0:
                props.append(OxmlElement('w:tblHeader'))
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    for r in p.runs:
                        set_run_font(r, 'Calibri', '微软雅黑', 9.5, i == 0)
    doc.core_properties.title = '局限性：视觉状态与风险分层测试的结论边界'
    doc.core_properties.author = 'QE研究项目'
    doc.save(TARGET)
    with ZipFile(TARGET) as z:
        e = ET.fromstring(z.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert len(e.findall('.//w:tbl', ns)) == 3
        word_text = '\n'.join(n.text or '' for n in e.findall('.//w:t', ns))
        assert all(f'局限性{i}：' in word_text for i in range(1, 11))
        assert all(v in word_text for v in ('66.84', '0.238', '0.09375'))
    OUT.mkdir(parents=True, exist_ok=True)
    result = {'status': 'pass', 'word': str(TARGET), 'markdown': str(SOURCE),
              'limitations': 10, 'tables': 3, 'official_reference_numeric_cells_verified': 12,
              'ablation_holm_verified': True, 'new_experiments': False,
              'source': str(ROOT / 'docs/visual_state_testing_paper_controls_20261006.md'),
              'rendering': 'OOXML structure verified; pagination not rendered'}
    (OUT / 'audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
