#!/usr/bin/env python3
"""Revise the group-meeting deck with the agreed total-dataset workflow."""
from copy import deepcopy
from pathlib import Path
import sys

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE


SOURCE = Path(
    "/home/hcf/.codex/attachments/131f36dc-19fc-427e-8c9f-ccbe774470c7/"
    "QE_group_meeting_visual_test_adequacy_20260915_v2.pptx"
)
OUTPUT = Path("/home/hcf/project/QE/docs/QE_group_meeting_visual_test_adequacy_20260915_v4.pptx")


def set_text(shape, value: str, font_size=None):
    """Replace text while retaining the first run's visual style."""
    if not hasattr(shape, "text_frame"):
        return
    frame = shape.text_frame
    template = None
    if frame.paragraphs and frame.paragraphs[0].runs:
        run = frame.paragraphs[0].runs[0]
        template = {
            "name": run.font.name,
            "size": run.font.size,
            "bold": run.font.bold,
            "italic": run.font.italic,
            "color": run.font.color.rgb if run.font.color and run.font.color.type else None,
        }
    frame.clear()
    if not value:
        return
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    run = paragraph.runs[0]
    if template:
        run.font.name = template["name"]
        run.font.size = font_size or template["size"]
        run.font.bold = template["bold"]
        run.font.italic = template["italic"]
        if template["color"] is not None:
            run.font.color.rgb = template["color"]
    elif font_size:
        run.font.size = font_size
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def replace_by_index(slide, replacements):
    for index, value in replacements.items():
        set_text(slide.shapes[index], value)


def remove_slide(prs, index):
    """Remove a slide and its relationship without rebuilding the deck."""
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def main():
    prs = Presentation(SOURCE)

    # Keep the existing review and motivation section intact. The former next-step
    # slide becomes the concrete method proposed for the next research phase.
    method = prs.slides[12]
    replace_by_index(
        method,
        {
            0: "06 · RESEARCH METHOD",
            1: "一条主线：总数据集 → dog 子集 → 补图 → 回归总集",
            2: "先建立总数据集基线，再用 dog 子集的视觉状态缺口驱动补图，最后回到同一任务评价模型变化",
            7: "总数据集 D0",
            8: "按类别整理，固定训练 / 验证 / 测试划分",
            9: "所有类别先用同一模型建立基线指标 M_before。",
            13: "dog 子集 Ddog",
            14: "固定 COCO 目标，运行 dog_v3.1",
            15: "输出 23 项视觉状态，并计算补图前覆盖率 C_before。",
            16: "",
            20: "缺口驱动补图",
            21: "得到 Ddog+ 并清洁、去重、留痕",
            22: "只在训练候选中选择补充 dog 图片；测试集不参与选图。",
            23: "",
            27: "回归总集 D1",
            28: "将 Ddog+ 合并回总数据集",
            29: "用完全相同的模型配置重跑，得到 M_after 和 C_after。",
            30: "",
            32: "主线：D0 → Ddog → Ddog+ → D1",
            33: "两条结果线分别报告：覆盖变化 ΔC = C_after − C_before；模型变化 ΔM = M_after − M_before。",
            34: "当前已有 C_before = 78.64%、C_after = 84.17%；模型性能指标待固定测试集实验后填入。",
            35: "QE · MAIN",
        },
    )

    # Keep the conclusion section, but make the pending model validation explicit.
    conclusion = prs.slides[13]
    replace_by_index(
        conclusion,
        {
            0: "",
            1: "",
            2: "",
            11: "07 · EVALUATION & CONCLUSION",
            12: "覆盖率改善是中间证据，模型前后指标决定结论能否升级",
            13: "同一总数据集任务、同一固定测试集、同一模型配置，分别比较覆盖变化与性能变化",
            14: "当前已完成的证据",
            16: "dog 子集覆盖分析与补图流程可执行",
            17: "900/900 严格协议通过；所有标注、校验、选图和清洁步骤可回溯",
            18: "已得到的覆盖结果",
            19: "C_before 78.64% → C_after 84.17%",
            20: "ΔC = +5.53 个百分点（限当前候选池）",
            21: "待完成的性能验证",
            22: "在固定测试集上比较 M_before 与 M_after",
            23: "报告总体 Macro-F1、dog Recall/F1 及困难状态 Recall；当前没有填入模型结果",
            24: "研究方法落地顺序",
            25: "阶段 1 · 基线",
            26: "冻结 D0 划分；在总数据集上运行基线模型，记录 M_before",
            27: "阶段 2 · 补图",
            28: "构建 Ddog+ 并合并为 D1；保持模型配置不变，记录 M_after",
            29: "判定规则",
            30: "只有同时报告 ΔC 与 ΔM，才能判断覆盖改善是否伴随模型收益",
            31: "结论强度等于证据强度：目前支持流程可行性与覆盖改善；识别性能收益待独立测试集验证。",
            32: "QE · MAIN",
        },
    )

    # The main talk does not need the three detailed related-work slides or the
    # three backup slides. Keeping the summary related-work page and the focused
    # conclusion makes the agreed workflow visible without repeating context.
    for index in (16, 15, 14, 6, 5, 4):
        remove_slide(prs, index)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
