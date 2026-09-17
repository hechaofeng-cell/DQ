#!/usr/bin/env python3

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/hcf/project/QE")
SOURCE = ROOT / "docs/QE_group_meeting_visual_test_adequacy_20260916_final_with_summary.pptx"
OUTPUT = ROOT / "docs/QE_group_meeting_visual_test_adequacy_20260916_final_integrated_motivation.pptx"

LEFT_IMAGES = [
    (ROOT / "artifacts/dog594_release_v1_20260913/marked/000000143783.jpg", "正面 · 单犬"),
    (ROOT / "artifacts/dog594_release_v1_20260913/marked/000000530154.jpg", "清晰 · 无遮挡"),
    (ROOT / "artifacts/dog594_release_v1_20260913/marked/000000001688.jpg", "近景 · 轮廓完整"),
    (ROOT / "artifacts/dog594_release_v1_20260913/marked/000000285965.jpg", "常见视角 · 易识别"),
]

RIGHT_IMAGES = [
    (ROOT / "artifacts/dog_candidate400_v2_1_inputs_20260912/marked/000000247068.jpg", "多犬与人物"),
    (ROOT / "artifacts/dog_candidate400_v2_1_inputs_20260912/marked/000000438744.jpg", "背向与雪地"),
    (ROOT / "artifacts/dog_candidate400_v2_1_inputs_20260912/marked/000000495454.jpg", "极小目标"),
    (ROOT / "artifacts/dog_candidate400_v2_1_inputs_20260912/marked/000000280709.jpg", "跳跃与交互"),
]

COLORS = {
    "bg": "F7F9FB",
    "white": "FFFFFF",
    "ink": "17212B",
    "muted": "5F6B76",
    "line": "D8E0E6",
    "navy": "153B5B",
    "blue": "3183C5",
    "blue_soft": "E4F0FA",
    "teal": "188A87",
    "teal_soft": "DDF3EF",
    "orange": "E87524",
    "image_bg": "EEF2F5",
}


def rgb(name):
    return RGBColor.from_string(COLORS.get(name, name))


def add_text(slide, text, x, y, w, h, *, size=18, bold=False, fill="ink",
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font="Microsoft YaHei"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(fill)
    return box


def add_round_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(1)
    return shape


def delete_all_shapes(slide):
    shape_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        shape_tree.remove(shape._element)


def add_picture_contain(slide, image_path, x, y, w, h):
    add_round_rect(slide, x, y, w, h, "image_bg", "line", radius=True)
    with Image.open(image_path) as image:
        image_w, image_h = image.size
    scale = min(w / image_w, h / image_h)
    draw_w = image_w * scale
    draw_h = image_h * scale
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    slide.shapes.add_picture(
        str(image_path), Inches(draw_x), Inches(draw_y), Inches(draw_w), Inches(draw_h)
    )


def add_panel(slide, x, title, subtitle, accent, soft, images):
    add_round_rect(slide, x, 1.82, 5.84, 4.53, "white", "line", radius=True)
    add_round_rect(slide, x + 1.63, 2.02, 2.58, 0.36, soft, soft, radius=True)
    add_text(
        slide, title, x + 1.68, 2.03, 2.48, 0.31,
        size=13.5, bold=True, fill=accent, align=PP_ALIGN.CENTER
    )
    add_text(
        slide, subtitle, x + 0.35, 2.42, 5.14, 0.25,
        size=11.5, fill="muted", align=PP_ALIGN.CENTER
    )

    image_w = 2.48
    image_h = 1.17
    positions = [
        (x + 0.25, 2.78), (x + 3.11, 2.78),
        (x + 0.25, 4.47), (x + 3.11, 4.47),
    ]
    for (image_path, label), (image_x, image_y) in zip(images, positions):
        add_picture_contain(slide, image_path, image_x, image_y, image_w, image_h)
        add_text(
            slide, label, image_x, image_y + 1.23, image_w, 0.26,
            size=11.5, bold=True, fill=accent, align=PP_ALIGN.CENTER
        )


def main():
    for path, _ in LEFT_IMAGES + RIGHT_IMAGES:
        if not path.exists():
            raise FileNotFoundError(path)

    presentation = Presentation(SOURCE)
    if len(presentation.slides) != 9:
        raise RuntimeError(f"Expected 9 slides, found {len(presentation.slides)}")

    slide = presentation.slides[1]
    delete_all_shapes(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("bg")

    add_text(slide, "01 · MOTIVATION", 0.62, 0.26, 3.2, 0.24,
             size=10, bold=True, fill="teal", font="Aptos")
    add_text(slide, "同一类别、相同数量，不代表相同测试证据", 0.62, 0.58, 12.0, 0.53,
             size=28, bold=True, fill="navy")
    add_text(
        slide,
        "左侧是初始集中常见的正面、清晰、无遮挡状态；右侧是补图候选中的差异状态",
        0.64, 1.15, 11.95, 0.30, size=13, fill="muted"
    )
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.58), Inches(12.05), Inches(0.012)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = rgb("line")
    rule.line.color.rgb = rgb("line")

    add_panel(
        slide, 0.62, "A：常见条件集中", "正面 · 清晰 · 无遮挡",
        "blue", "blue_soft", LEFT_IMAGES
    )
    add_panel(
        slide, 6.87, "B：差异条件出现", "数量 · 朝向 · 尺度 · 动作与场景",
        "teal", "teal_soft", RIGHT_IMAGES
    )

    badge = add_round_rect(slide, 6.43, 3.77, 0.62, 0.62, "white", "line", radius=True)
    badge.line.width = Pt(1.2)
    add_text(slide, "≠", 6.43, 3.76, 0.62, 0.57, size=22, bold=True,
             fill="orange", align=PP_ALIGN.CENTER, font="Aptos")

    add_round_rect(slide, 0.86, 6.51, 11.61, 0.55, "navy", "navy", radius=True)
    add_text(
        slide,
        "只看 dog 的图片数量，会掩盖类别内部的视角、尺度、遮挡、动作和场景差异。",
        1.10, 6.59, 11.13, 0.34, size=17, bold=True, fill="white", align=PP_ALIGN.CENTER
    )
    add_text(
        slide, "真实图像示例来自 COCO 2017；红框为程序固定的 A 目标。",
        8.56, 7.10, 4.15, 0.15, size=7.5, fill="muted", align=PP_ALIGN.RIGHT
    )
    add_text(slide, "QE · MAIN", 0.62, 7.24, 1.0, 0.13,
             size=8, bold=True, fill="muted", font="Aptos")
    add_text(slide, "2", 12.53, 7.24, 0.14, 0.13,
             size=8, fill="muted", align=PP_ALIGN.RIGHT, font="Aptos")

    slide.notes_slide.notes_text_frame.text = (
        "这一页分成左右两类来看。左边四张都来自初始500张数据，主要是单犬、正面、清晰、无遮挡，"
        "属于比较常见、也比较容易识别的条件。如果大量图片集中在这类条件里，dog的总数即使很多，"
        "提供的测试证据仍然比较单一。右边仍然是dog类别，但出现了多犬和人物共现、背向和雪地、"
        "极小目标、跳跃与人物交互等差异条件。这里的八张图只是用于直观说明，并不代表完整分布统计。"
        "我想强调的是：同一类别、相同数量，不等于覆盖了相同的视觉条件。"
        "因此后面不再只统计dog有多少张，而是进一步测量类别内部的重要视觉状态是否出现、样本是否足够。"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
