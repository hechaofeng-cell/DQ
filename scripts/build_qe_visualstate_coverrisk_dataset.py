"""Build a traceable, cross-source visual-state dataset from existing annotations."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
from PIL import Image
from docx import Document

from run_category_five_dimension_pilot import score_sample, validate_input_binding, validate_label

ROOT = Path(__file__).resolve().parents[1]
NAME = 'QE_VisualState_CoverRisk_v1_20261013'
OUT = ROOT / 'artifacts' / NAME
DATA = ROOT / 'data' / NAME
CLASSES = ('dog', 'cat', 'horse', 'sheep', 'person', 'car')
POOLS = {
    'A_OpenImages': ROOT/'artifacts/sixclass_dual_vlm_replaced_20261004',
    'B_COCO2017': ROOT/'artifacts/sixclass_independent_confirmation_20261004',
    'C_VOC2007': ROOT/'artifacts/sixclass_voc_external_confirmation_20261005',
    'D_Objects365': ROOT/'artifacts/objects365_independent_confirmation_20261005',
}
FRESH = ROOT/'artifacts/fresh_vlm_feature_scoring_20261012'
CORE = {'target_scale','body_visibility','occlusion_level','truncation_level','body_orientation'}


def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def write_csv(path, rows):
    path=Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        columns=list(dict.fromkeys(k for r in rows for k in r))
        writer=csv.DictWriter(f,fieldnames=columns);writer.writeheader();writer.writerows(rows)


def write_json(path, obj):
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    Path(path).write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')


def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()


def note(stage, **kwargs):
    row=dict(stage=stage,time=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),**kwargs)
    write_json(OUT/'progress.json',row)
    print(json.dumps(row,ensure_ascii=False),flush=True)
    with (OUT/'research_log.md').open('a',encoding='utf-8') as f:f.write('\n'+json.dumps(row,ensure_ascii=False)+'\n')


def scale(area):
    return 'tiny' if area<.01 else 'small' if area<.03 else 'medium' if area<.1 else 'large'


def prepare():
    OUT.mkdir(parents=True,exist_ok=True);(OUT/'prepared').mkdir(exist_ok=True)
    base=next(iter(POOLS.values()))
    shared=json.loads((base/'configs/shared_object_schema.json').read_text())
    states=[('target_scale',v) for v in ('tiny','small','medium','large')]
    states += [(f['feature_id'],v) for f in shared['universal_features'] for v in f['possible_values'] if v!='unknown']
    assert len(states)==71 and len(set(states))==71
    config=json.loads((ROOT/'configs/category_five_dimension_pilot_v1.json').read_text())
    registration=dict(dataset=NAME,created_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),
        source_pools={k:str(p) for k,p in POOLS.items()},primary_annotation='local',sensitivity_annotation='online',
        coverage_dictionary='14 shared native fields plus original-image target scale; all 71 non-unknown states per category fixed before optimization',
        presence_denominator=426,support_requirement=20,
        risk_score='equal_global; five dimensions each 20 points; unknown lower/upper retained',
        optimization_objective='minimum cover for reachable states, then maximum lower-bound mean risk at fixed balanced budget',
        high_risk='top 20% per category in merged pool by midpoint score, ties sample_id; relative Q5, not empirical error probability',
        comparison=['coverage_greedy','risk_only_top','combined_coverage_risk','random'],random_repeats=1000,
        stopping='one predeclared finite sweep; minimum covers and risk optimizations time bounded; report solver bounds and gaps',
        new_annotations=False,classifier_feedback=False,
        fine_feature_policy='60 newer records attached separately to matching targets, not mixed into main coverage/risk protocol')
    write_json(OUT/'protocol.json',registration)
    write_json(OUT/'coverage_dictionary.json',dict(shared_schema=shared,states=states))
    write_json(OUT/'risk_protocol.json',config)
    (OUT/'research_log.md').write_text('# Dataset construction and optimization log\n\nClaims: availability ceiling under frozen dictionaries; achievable coverage/risk of selected subsets. No claim about model failure probability.\n',encoding='utf-8')
    frozen={str(Path(__file__).resolve()):sha(__file__),str(ROOT/'configs/category_five_dimension_pilot_v1.json'):sha(ROOT/'configs/category_five_dimension_pilot_v1.json')}
    records=[];all_hashes=set();crop_hashes=set();sample_ids=set();view_audits=[]
    for source,pool in POOLS.items():
        rows=read_csv(pool/'manifests/candidate_pool.csv')
        input_rows={r['image_id']:r for d in ('old','car') for r in read_csv(pool/'inputs'/f'{d}.csv')}
        schemas={d:json.loads((pool/'configs'/f'{d}_annotation_schema.json').read_text()) for d in ('old','car')}
        assert json.loads((pool/'configs/shared_object_schema.json').read_text())==shared
        for path in [pool/'manifests/candidate_pool.csv',pool/'inputs/old.csv',pool/'inputs/car.csv',*list((pool/'configs').glob('*.json'))]:frozen[str(path)]=sha(path)
        for row in rows:
            sid=row['sample_id'];category=row['class_name'];domain='car' if category=='car' else 'old'
            assert sid not in sample_ids;sample_ids.add(sid)
            assert int(row['label'])==CLASSES.index(category)
            with Image.open(row['image_path']) as image:
                width,height=image.size
                binding=validate_input_binding(row,input_rows[sid],image.size)
            xyxy=np.asarray(json.loads(row['bbox_norm_xyxy']),dtype=float)
            assert xyxy.shape==(4,) and np.isfinite(xyxy).all()
            assert np.all(xyxy[:2]>=-.00001) and np.all(xyxy[2:]<=1.00001) and np.all(xyxy[2:]>xyxy[:2])
            geometric_area=float(np.prod(xyxy[2:]-xyxy[:2]));area=float(row['bbox_area_ratio'])
            # COCO stored normalized coordinates are rounded; verify scale invariance.
            assert abs(geometric_area-area)<1e-6 and scale(geometric_area)==scale(area)
            digests={key:sha(row[key]) for key in ('image_path','crop_path','marked_image_path','feature_crop_path')}
            assert digests['image_path']==row['image_sha256'] and digests['crop_path']==row['crop_sha256']
            assert digests['image_path'] not in all_hashes and digests['crop_path'] not in crop_hashes
            all_hashes.add(digests['image_path']);crop_hashes.add(digests['crop_path'])
            for key,digest in digests.items():frozen[row[key]]=digest
            values={}
            for backend in ('local','online'):
                path=pool/'annotations'/backend/domain/'parsed'/f'{sid}.json'
                value=json.loads(path.read_text())
                features=validate_label(value,binding,schemas[domain])
                assert set(features)>=set(f['feature_id'] for f in shared['universal_features'])
                values[backend]=dict(native=value,risk=score_sample(area,features,category,'equal_global',config))
                frozen[str(path)]=sha(path)
            record=dict(row,source_pool=source,image_width=str(width),image_height=str(height),
                        annotation_target_instance_id=binding['target_instance_id'],annotations=values)
            records.append(record)
            view_audits.append(dict(sample_id=sid,source_pool=source,bbox_binding=True,
                                   image_hash=digests['image_path'],crop_hash=digests['crop_path']))
        note('source_validated',source=source,targets=len(rows),total=len(records))
    assert len(records)==7800 and Counter(r['class_name'] for r in records)==Counter(dict.fromkeys(CLASSES,1300))
    reference_hashes=set();reference_image_hashes=set();reference=[];missing_reference_originals=0
    for kind in ('train','eval'):
        path=base/'manifests'/f'{kind}.csv';frozen[str(path)]=sha(path)
        for row in read_csv(path):
            digest=sha(row['crop_path']);reference_hashes.add(digest)
            original=Path(row.get('image_path',''))
            if row.get('image_path') and original.is_file():reference_image_hashes.add(sha(original))
            else:missing_reference_originals+=1
            reference.append(dict(split=kind,sample_id=row.get('sample_id',''),crop_sha256=digest))
    assert not crop_hashes&reference_hashes and not all_hashes&reference_image_hashes
    write_csv(OUT/'audit/reference_targets.csv',reference)
    write_csv(OUT/'audit/source_identity_audit.csv',view_audits)
    with (OUT/'prepared/records.jsonl').open('w',encoding='utf-8') as f:
        for row in records:f.write(json.dumps(row,ensure_ascii=False,allow_nan=False)+'\n')
    gaps=[];ceiling=[];pool_manifest=[]
    for category in CLASSES:
        subset=[r for r in records if r['class_name']==category]
        matrices={}
        for backend in ('local','online'):
            matrix=np.zeros((len(subset),len(states)),dtype=np.uint8)
            for i,row in enumerate(subset):
                features=row['annotations'][backend]['native']['features']
                for j,(field,state) in enumerate(states):
                    value=scale(float(row['bbox_area_ratio'])) if field=='target_scale' else features[field]
                    matrix[i,j]=state in value if isinstance(value,list) else state==value
            matrices[backend]=matrix
            support=matrix.sum(axis=0);core=np.array([f in CORE for f,v in states])
            ceiling.append(dict(category=category,backend=backend,targets=len(subset),required_states=len(states),
                reached_states=int((support>0).sum()),presence_coverage=float((support>0).mean()),
                support_completion=float(np.minimum(support/20,1).mean()),
                core_required_states=int(core.sum()),core_reached_states=int((support[core]>0).sum()),
                core_presence_coverage=float((support[core]>0).mean())))
            for (field,state),count in zip(states,support):
                gaps.append(dict(category=category,backend=backend,field=field,state=state,
                                 pool_support=int(count),support_requirement=20,
                                 missing=count==0,shortfall=max(0,20-int(count))))
        risks={key:np.array([r['annotations']['local']['risk'][key] for r in subset]) for key in ('score_lower','score_midpoint','score_upper')}
        q5=np.zeros(len(subset),dtype=np.uint8)
        order=sorted(range(len(subset)),key=lambda i:(risks['score_midpoint'][i],subset[i]['sample_id']))
        q5[order[-len(subset)//5:]]=1
        np.savez_compressed(OUT/'prepared'/f'{category}.npz',matrix=matrices['local'],online_matrix=matrices['online'],
                            sample_ids=np.array([r['sample_id'] for r in subset]),
                            risk_lower=risks['score_lower'],risk_midpoint=risks['score_midpoint'],risk_upper=risks['score_upper'],high_risk=q5)
        for i,row in enumerate(subset):
            value={k:v for k,v in row.items() if k!='annotations'}
            value.update(risk_lower=float(risks['score_lower'][i]),risk_midpoint=float(risks['score_midpoint'][i]),
                         risk_upper=float(risks['score_upper'][i]),high_risk_q5=int(q5[i]),annotation_backend='local')
            pool_manifest.append(value)
    write_csv(OUT/'merged_manifest.csv',pool_manifest)
    write_csv(OUT/'pool_state_support.csv',gaps);write_csv(OUT/'pool_ceiling_by_category.csv',ceiling)
    write_json(OUT/'input_freeze.sha256.json',frozen)
    write_json(OUT/'preparation_audit.json',dict(status='pass',targets=len(records),per_class=1300,
        local_annotations=len(records),online_annotations=len(records),reference_targets=len(reference),
        exact_duplicate_targets=0,exact_duplicate_originals=0,exact_duplicate_crops=0,
        exact_train_eval_crop_overlap=0,reference_originals_available=len(reference)-missing_reference_originals,
        reference_originals_unavailable=missing_reference_originals,
        exact_available_reference_original_overlap=0,semantic_annotations_not_fully_human_validated=True,
        pooled_cross_source_near_duplicates_not_newly_checked=True,frozen_paths=len(frozen)))
    note('prepared',targets=len(records),dictionary_states=426)


def measures(z,indices):
    index=np.asarray(indices,dtype=int)
    support=z['matrix'][index].sum(axis=0);possible=z['matrix'].sum(axis=0)>0
    online=z['online_matrix'][index].sum(axis=0)
    return dict(n=len(index),presence_coverage=float((support>0).mean()),
                reachable_presence_coverage=float((support[possible]>0).mean()),
                support_completion=float(np.minimum(support/20,1).mean()),
                reached_states=int((support>0).sum()),required_states=len(support),
                risk_lower=float(z['risk_lower'][index].mean()),risk_midpoint=float(z['risk_midpoint'][index].mean()),
                risk_upper=float(z['risk_upper'][index].mean()),high_risk_q5_fraction=float(z['high_risk'][index].mean()),
                online_presence_coverage=float((online>0).mean()),
                online_support_completion=float(np.minimum(online/20,1).mean()))


def coverage_greedy(z,budget):
    matrix=z['matrix'];counts=np.zeros(matrix.shape[1],dtype=int);remaining=set(range(len(matrix)));selected=[]
    for _ in range(budget):
        def key(i):
            active=matrix[i].astype(bool)&(counts<20)
            gain=float((np.log(counts[active]+2)-np.log(counts[active]+1)).sum())
            return (-gain,str(z['sample_ids'][i]))
        chosen=min(remaining,key=key);selected.append(chosen);remaining.remove(chosen);counts+=matrix[chosen]
    return selected


def export():
    solution=json.loads((OUT/'optimization/solutions.json').read_text())
    records=[json.loads(s) for s in (OUT/'prepared/records.jsonl').read_text().splitlines()]
    lookup={r['sample_id']:r for r in records}
    arrays={c:dict(np.load(OUT/'prepared'/f'{c}.npz')) for c in CLASSES}
    budget=int(solution['compact_budget_per_class'])
    selections={}
    for mode in ('coverage_greedy','risk_only_top','combined_coverage_risk'):
        selections[mode]={}
        for category,z in arrays.items():
            if mode=='coverage_greedy':values=coverage_greedy(z,budget)
            elif mode=='risk_only_top':values=sorted(range(len(z['matrix'])),key=lambda i:(-z['risk_lower'][i],str(z['sample_ids'][i])))[:budget]
            else:values=solution['budgets'][str(budget)][category]['selected_indices']
            selections[mode][category]=values
    if solution.get('highest_q5_under_full_coverage'):
        selections['combined_highrisk_fraction']={c:solution['highest_q5_under_full_coverage'][c]['selected_indices'] for c in CLASSES}
    for value,per_class in solution['budgets'].items():
        if int(value)!=budget:selections[f'combined_n{int(value)*6}']={c:per_class[c]['selected_indices'] for c in CLASSES}
    metrics=[];suites={};rng=np.random.default_rng(20261013)
    for mode,selection in selections.items():
        suites[mode]=[]
        for category,index in selection.items():
            z=arrays[category];m=measures(z,index);metrics.append(dict(strategy=mode,category=category,**m))
            suites[mode]+=[str(z['sample_ids'][i]) for i in index]
    random_metrics=[]
    for repeat in range(1000):
        values=[measures(arrays[c],rng.choice(len(arrays[c]['matrix']),size=budget,replace=False)) for c in CLASSES]
        random_metrics.append(dict(repeat=repeat,total_targets=budget*6,
                                   **{k:float(np.mean([r[k] for r in values])) for k in values[0] if k not in ('n','reached_states','required_states')}))
    write_csv(OUT/'random_comparison.csv',random_metrics)
    write_csv(OUT/'selected_metrics_by_category.csv',metrics)
    summary=[]
    for mode in selections:
        values=[r for r in metrics if r['strategy']==mode]
        summary.append(dict(strategy=mode,total_targets=sum(r['n'] for r in values),per_class=values[0]['n'],
                            **{k:float(np.mean([r[k] for r in values])) for k in values[0] if k not in ('strategy','category','n','reached_states','required_states')}))
    random_summary=dict(strategy='random_mean_1000',total_targets=budget*6,per_class=budget)
    for key in random_metrics[0]:
        if key not in ('repeat','total_targets'):
            vals=np.array([r[key] for r in random_metrics]);random_summary[key]=float(vals.mean())
            random_summary[key+'_p025']=float(np.quantile(vals,.025));random_summary[key+'_p975']=float(np.quantile(vals,.975))
    summary.append(random_summary);write_csv(OUT/'strategy_comparison.csv',summary)
    DATA.mkdir(parents=True,exist_ok=True)
    for name in ('protocol.json','coverage_dictionary.json','risk_protocol.json'):shutil.copy2(OUT/name,DATA/name)
    (DATA/'annotations').mkdir(exist_ok=True)
    destination_rows=[];storage=Counter()
    for number,row in enumerate(records,1):
        sid=row['sample_id'];category=row['class_name']
        dest={k:v for k,v in row.items() if k!='annotations'}
        for key,folder in (('image_path','images'),('crop_path','crops'),('marked_image_path','marked'),('feature_crop_path','feature_crops')):
            source=Path(row[key]);target=DATA/folder/category/(sid+source.suffix.lower())
            target.parent.mkdir(parents=True,exist_ok=True)
            if not target.exists():
                try:os.link(source,target);storage['hardlinks']+=1
                except OSError:shutil.copy2(source,target);storage['copied_files']+=1
            dest['source_'+key]=str(source);dest[key]=str(target.relative_to(DATA))
        fresh=FRESH/'annotations'/f'{sid}.json'
        payload=dict(sample_id=sid,class_name=category,official_label=int(row['label']),
                     source_pool=row['source_pool'],bbox_norm_xyxy=json.loads(row['bbox_norm_xyxy']),
                     local=row['annotations']['local'],online=row['annotations']['online'])
        if fresh.exists():payload['fine_feature_supplement']=json.loads(fresh.read_text());storage['fine_supplements']+=1
        write_json(DATA/'annotations'/f'{sid}.json',payload)
        dest.update({k:row['annotations']['local']['risk'][k] for k in ('score_lower','score_midpoint','score_upper','fully_known','evaluable')})
        z=arrays[category];local_idx=np.flatnonzero(z['sample_ids']==sid);assert len(local_idx)==1
        dest['high_risk_q5']=int(z['high_risk'][local_idx[0]])
        dest['annotation_path']=f'annotations/{sid}.json'
        destination_rows.append(dest)
        if number%1300==0:note('materializing_dataset',completed=number,total=len(records))
    write_csv(DATA/'manifest.csv',destination_rows)
    destination_lookup={r['sample_id']:r for r in destination_rows}
    for mode,ids in suites.items():
        assert len(ids)==len(set(ids))
        write_csv(DATA/'subsets'/f'{mode}.csv',[destination_lookup[i] for i in ids])
    # Source labels remain intact; supplement is counted separately.
    assert storage['fine_supplements']==60
    ceilings=read_csv(OUT/'pool_ceiling_by_category.csv');local=[r for r in ceilings if r['backend']=='local']
    gaps=[r for r in read_csv(OUT/'pool_state_support.csv') if r['backend']=='local' and r['missing']=='True']
    primary=next(r for r in summary if r['strategy']=='combined_coverage_risk')
    lines=['# QE视觉状态覆盖与风险测试数据集', '',f'数据集名：{NAME}。位置：{DATA}。', '',
        '实际合并OpenImages、COCO2017、VOC2007、Objects365四个已标注目标池：7800个唯一目标，六类各1300个。保留原图、分类裁剪、框标记图、特征裁剪、官方标签和双VLM原生标注。最新60个细粒度标注作为匹配目标的补充单独保存。', '',
        '主覆盖清单在计算前固定：14项共享字段加目标尺度，每类71个非unknown状态，共426个。毛发、头发、车轮等原生/补充细节完整保留，但尚未全量标注的字段不混入主覆盖分母。', '',
        f"local标注下，完整候选池覆盖{sum(int(r['reached_states']) for r in local)}/426个状态，出现覆盖上限{np.mean([float(r['presence_coverage']) for r in local])*100:.4f}%；每状态20目标的支持完成度上限{np.mean([float(r['support_completion']) for r in local])*100:.4f}%。这些上限由全候选池计数给出，不依赖选样算法。", '',
        f'主推荐精简子集为combined_coverage_risk，共{budget*6}个目标、每类{budget}个。在全部可达状态至少一例的约束下，使用整数规划最大化风险下界。最小规模及风险最优性以optimization/solutions.json的solver状态和gap为准；未证明的结果只称已找到可行解。', '',
        '风险按既有equal_global规则计算，尺度、整体可见性、遮挡、截断、身体朝向各20分，范围0—100。不是最新60例的头部/躯干/肢体评分，也不是实测分类错误概率。未知维度保留下界、中点和上界；优化下界避免用未知值制造高风险。高风险占比为合并池各类别中点排序的相对Q5占比。', '',
        '## 相同紧凑预算对照及预算梯度', '',
        '| 策略 | 目标数 | 字典出现覆盖 | 可达状态覆盖 | 支持完成度（20） | 平均风险下界 | 平均风险中点 | Q5占比 |',
        '|---|---:|---:|---:|---:|---:|---:|---:|']
    for r in summary:
        lines.append(f"| {r['strategy']} | {r['total_targets']} | {100*r['presence_coverage']:.2f}% | {100*r['reachable_presence_coverage']:.2f}% | {100*r['support_completion']:.2f}% | {r['risk_lower']:.2f} | {r['risk_midpoint']:.2f} | {100*r['high_risk_q5_fraction']:.2f}% |")
    lines+=['','覆盖贪心对照按饱和阈值20的凹证据增益选择，不是字典出现覆盖的精确最优化。组合策略优先满足全部可达状态，再最大化风险；其全局最优结论只限于冻结字典、固定预算和风险下界目标。risk_only_top按风险下界排序，给出同预算无覆盖约束的最高平均风险，其他策略的风险均不得超过它。普通随机对照重复1000次；更细的共同风险配额随机对照尚未在本轮执行。','',
        '## 全池上限及缺失状态', '',
        '| 类别 | 已覆盖／要求 | 出现覆盖上限 | 20例支持完成度上限 | 五维核心覆盖上限 |', '|---|---:|---:|---:|---:|']
    for r in local:lines.append(f"| {r['category']} | {r['reached_states']}/{r['required_states']} | {100*float(r['presence_coverage']):.2f}% | {100*float(r['support_completion']):.2f}% | {100*float(r['core_presence_coverage']):.2f}% |")
    lines+=['', '完整字典缺失状态（保留在分母，未为制造100%删除）：', '']
    for r in gaps:lines.append(f"- {r['category']} / {r['field']} / {r['state']}")
    lines+=['', '## 使用与验证边界', '',
        'manifest.csv中的视图路径相对数据集根目录，subsets/中的清单引用同一套实体文件。同一个目标可以进入多个测试子集。停止原因区分需求达标与预算完成：组合解必须满足可达状态需求及预算；不可达状态始终留作缺口，不称全字典100%。', '',
        '已验证目标框与输入绑定、字典取值、原图/裁剪SHA256、精确去重及与现有1800训练目标、390验证/基准测试目标的全部裁剪哈希无交集；可用参考原图也无精确交集，缺失参考原图数量见preparation_audit.json。未重新进行全合并池跨源近重复审计；自动标注未全量人工复核；无法排除预训练来源重叠。', '',
        '没有重新训练、执行分类器或生成新图像，本轮风险和覆盖均来自已有自动特征标注。报告中的最好覆盖是该字典与已有标签下的可达上限，不能解释为视觉输入域完备。', '',
        '## 可复现文件', '',
        '主程序scripts/build_qe_visualstate_coverrisk_dataset.py；优化程序scripts/optimize_qe_coverrisk.py。input_freeze.sha256.json冻结来源、视图和标注，optimization/保留各类解与界。pool_state_support.csv记录所有状态及缺口。strategy_comparison.csv记录主结果，random_comparison.csv记录1000次随机对照。']
    text='\n'.join(lines)+'\n';(OUT/'report.md').write_text(text,encoding='utf-8');(DATA/'README.md').write_text(text,encoding='utf-8')
    doc=Document();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            cells=[[x.strip() for x in row.strip('|').split('|')] for row in block]
            table=doc.add_table(rows=1,cols=len(cells[0]));table.style='Table Grid'
            for cell,value in zip(table.rows[0].cells,cells[0]):cell.text=value
            for row in cells[2:]:
                for cell,value in zip(table.add_row().cells,row):cell.text=value
            continue
        if line.startswith('#'):doc.add_heading(line.lstrip('#').strip(),level=min(len(line)-len(line.lstrip('#')),3)-1)
        elif line:doc.add_paragraph(line)
        i+=1
    doc.save(OUT/'report.docx')
    write_json(OUT/'claim_evidence.json',dict(
        maximum_coverage=dict(evidence='pool_state_support.csv',definition='all 426 pre-frozen states; monotone support counts give exact upper bound'),
        selected_metrics=dict(evidence='selected_metrics_by_category.csv and data subsets manifests',definition='direct recomputation from frozen label matrices'),
        risk_optimality=dict(evidence='optimization/solutions.json',definition='solver status and dual bound/gap, no unverified optimum claim'),
        no_model_accuracy_claim=True,full_fine_feature_coverage_not_claimed=True))
    write_json(OUT/'export_audit.json',dict(status='pass',dataset=str(DATA),targets=len(records),subsets={k:len(v) for k,v in suites.items()},storage=dict(storage)))
    note('exported',dataset=str(DATA),targets=len(records),primary_subset=budget*6)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--phase',choices=('prepare','export'),required=True)
    args=p.parse_args()
    if args.phase=='prepare':prepare()
    else:export()
