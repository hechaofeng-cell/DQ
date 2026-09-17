#!/usr/bin/env python3
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
MD=ROOT/'docs/dog_dataset_paper_revised_zh.md'
OUT=ROOT/'docs/dog_dataset_paper_final_zh.docx'
FIG=ROOT/'output/imagegen/dog_v3_1_workflow_zh_polished.png'

def clean(s):
    s=re.sub(r'!\[([^]]*)\]\([^)]*\)',r'\1',s)
    s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1（\2）',s)
    return re.sub(r'[`*_]', '', s).strip()

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_font(run, name='宋体', size=10.5, bold=False):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:eastAsia'),name); run.font.size=Pt(size); run.bold=bold

def add_text(p,text,bold=False):
    r=p.add_run(clean(text));set_font(r,bold=bold);return r

def main():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    styles=doc.styles; styles['Normal'].font.name='宋体';styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');styles['Normal'].font.size=Pt(10.5)
    for level,size in [(1,16),(2,13),(3,11.5)]:
        st=styles[f'Heading {level}'];st.font.name='黑体';st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');st.font.size=Pt(size);st.font.bold=True
    lines=MD.read_text(encoding='utf-8').splitlines();i=0; figure_added=False
    while i<len(lines):
        line=lines[i]
        if not line.strip(): i+=1;continue
        if line.startswith('# '):
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(clean(line[2:]));set_font(r,'黑体',18,True);i+=1;continue
        if line.startswith('## '): doc.add_heading(clean(line[3:]),level=1);i+=1;continue
        if line.startswith('### '): doc.add_heading(clean(line[4:]),level=2);i+=1;continue
        if line.startswith('!['):
            m=re.search(r'!\[([^]]*)\]\(([^)]+)\)', line)
            alt=m.group(1) if m else '图'
            image_path=Path(m.group(2)) if m else FIG
            if not image_path.is_absolute():
                candidates=[MD.parent/image_path, ROOT/image_path]
                image_path=next((p for p in candidates if p.exists()), candidates[0])
            if not image_path.exists() and FIG.exists(): image_path=FIG
            if image_path.exists():
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(image_path),width=Inches(7.0 if 'workflow' in image_path.name else 6.5));cap=doc.add_paragraph(clean(alt));cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                figure_added=True
            i+=1;continue
        if line.startswith('|') and i+1<len(lines) and lines[i+1].startswith('|'):
            table_lines=[]
            while i<len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|\s*:?-+',lines[i]): table_lines.append([clean(x) for x in lines[i].strip('|').split('|')])
                i+=1
            if table_lines:
                t=doc.add_table(rows=len(table_lines),cols=len(table_lines[0]));t.style='Table Grid'
                for ri,row in enumerate(table_lines):
                    for ci,val in enumerate(row):
                        cell=t.cell(ri,ci);cell.text=val
                        for p in cell.paragraphs:
                            for r in p.runs:set_font(r,'宋体',9.2,ri==0)
                        if ri==0:shade(cell,'D9EAF2')
                doc.add_paragraph();
            continue
        if re.match(r'^\d+\.\s',line) or line.startswith('- '):
            p=doc.add_paragraph(style='List Number' if re.match(r'^\d+\.',line) else 'List Bullet');add_text(p,re.sub(r'^\d+\.\s|^-\s','',line));i+=1;continue
        if line.startswith('```'):
            code=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            i+=1;p=doc.add_paragraph();r=p.add_run('\n'.join(code));set_font(r,'等线',9);continue
        p=doc.add_paragraph();add_text(p,line);i+=1
    doc.add_heading('严格验证与投稿前检查',level=1)
    checks=[
        ('结果可追溯','已核对500张初始、400张候选、100张选择和596张清洁集的结果目录；论文中的数量、覆盖指数和成功率均对应本地 report.json 或脚本输出。'),
        ('指标边界','严格特征成功率仅表示JSON、字段、枚举和依赖规则通过；覆盖指数不是识别准确率；论文未把覆盖改善表述为模型性能提升。'),
        ('协议冻结','dog_v3.1 schema、提示词、输入三视图和候选选择结果单独保存；v2和v3.1历史结果不覆盖。'),
        ('人工验证','当前并非596张全部字段人工真值；论文明确说明人工抽查范围有限，后续应报告双人一致率和分歧裁决。'),
        ('比较实验','当前尚未完成随机、分层随机和贪心选择对照；论文不得声称缺口驱动选择普遍优于其他策略。'),
        ('识别实验','新增100张未运行独立分类VLM；论文不报告596张整体识别率，也不声称补图提升分类性能。'),
        ('泛化边界','证据仅来自COCO dog案例，不能推断跨类别泛化。'),
        ('数据许可','正式投稿前核对COCO原始图片的使用和再分发条款。'),
        ('引用核验','相关工作引用已链接到Visual Genome、Active Covering、TALISMAN和SELECT的论文或正式出版页面；投稿前按目标期刊格式补齐BibTeX和访问日期。'),
    ]
    t=doc.add_table(rows=1,cols=2);t.style='Table Grid';t.cell(0,0).text='检查项';t.cell(0,1).text='当前结论'
    for c in t.rows[0].cells:shade(c,'D9EAF2');
    for a,b in checks:
        cells=t.add_row().cells;cells[0].text=a;cells[1].text=b
        for c in cells:
            for p in c.paragraphs:
                for r in p.runs:set_font(r,'宋体',9.2)
    doc.add_paragraph('文档状态：阶段性论文初稿，适合内部评审和实验记录归档；完成对照实验、人工一致性和独立分类验证后再形成投稿终稿。')
    doc.save(OUT);print(OUT)
if __name__=='__main__':main()
