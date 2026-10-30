"""Organize the frozen 60-target annotations without changing annotations or scores."""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'artifacts/fresh_vlm_feature_scoring_20261012'
OUT = ROOT / 'artifacts/category_feature_dimension_review_20261013'
CLASSES = {'dog': '犬', 'cat': '猫', 'horse': '马', 'sheep': '羊', 'person': '人', 'car': '车'}
DIMS = ['尺度', '可见性', '遮挡', '截断', '朝向', '关联条件']
NAMES = {
    'body_visibility': '身体整体可见性', 'body_orientation': '身体朝向',
    'head_orientation': '头部朝向', 'head_visibility': '头部可见性',
    'pose_action': '姿态动作', 'leg_visibility': '腿部可见状态',
    'occlusion_level': '遮挡程度', 'occluder_type': '遮挡物类型',
    'truncation_level': '原图截断程度', 'truncation_parts': '截断部位',
    'group_context': '目标群体关系', 'interaction_context': '交互关系',
    'scene_environment': '场景环境', 'background_complexity': '背景复杂度',
    'blur_level': '模糊程度', 'illumination': '光照状态',
    'target_background_contrast': '目标与背景对比度',
    'target_recognizability': '目标可识别性', 'visual_representation': '视觉呈现类型',
    'car_color': '车身颜色', 'wheel_visibility': '可见车轮数量',
    'visible_vehicle_parts': '可见车辆部件集合', 'reflective_surface': '车身反光程度',
}


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))


def write_csv(name, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def grouping(field):
    if field.startswith('occlusion') or field == 'occluder_type':
        return '遮挡', '遮挡程度与原因'
    if field.startswith('truncation'):
        return '截断', '原图边界与部位'
    if 'orientation' in field:
        return '朝向', '身体及头部朝向'
    if field in ('tail_posture', 'ear_posture', 'pose_action', 'posture_action_state'):
        return '朝向', '关联姿态（非朝向评分项）'
    if 'visibility' in field or field in ('visible_vehicle_parts', 'hair_presence', 'eye_state', 'mouth_state'):
        if field in ('eye_state', 'mouth_state', 'hair_presence'):
            return '可见性', '部位状态附属描述（非缺失程度）'
        if any(part in field for part in ('torso', 'limb', 'leg', 'tail')):
            return '可见性', '躯干与肢体部位'
        if any(part in field for part in ('head', 'muzzle', 'mouth', 'ear', 'nose', 'face', 'eye', 'hair', 'tongue')):
            return '可见性', '头部及面部部位'
        if field == 'body_visibility':
            return '可见性', '整体呈现'
        return '可见性', '车辆部件'
    if field.startswith(('coat_', 'hair_')) or field in ('ear_shape', 'body_shape', 'car_color', 'reflective_surface', 'headwear', 'eyewear'):
        return '关联条件', '外观与配饰'
    if field in ('scene', 'scene_environment', 'group_context', 'interaction_context', 'background_complexity'):
        return '关联条件', '场景与交互'
    return '关联条件', '成像质量与整体判断'


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    annotations = read_csv(SOURCE / 'feature_annotations_with_evidence.csv')
    manifest = read_csv(SOURCE / 'manifest.csv')
    original_map = {(r['category'], r['field']): r for r in read_csv(SOURCE / 'feature_dimension_assignment.csv')}
    counts = Counter(r['class_name'] for r in manifest)
    by_field = {}
    for row in annotations:
        by_field.setdefault((row['class_name'], row['field']), []).append(json.loads(row['value']))
    mapping, distribution, count_rows = [], [], []
    schemas = {}
    for cls, zh in CLASSES.items():
        schema = json.loads((SOURCE / 'schemas' / f'{cls}.json').read_text())
        schemas[cls] = schema
        fields = schema['universal_features'] + schema['category_specific_features']
        for field in fields:
            fid = field['feature_id']
            dim, subgroup = grouping(fid)
            label = field.get('feature_name_zh') or NAMES.get(fid, fid)
            role = '参与评分' if original_map[(cls, fid)]['role'] == 'scored' else '已标注，不计分'
            mapping.append(dict(类别=zh, category=cls, 维度=dim, 子组=subgroup, 特征=label,
                                field=fid, 当前用途=role, 来源='VLM标注',
                                可选状态=' / '.join(field['possible_values'])))
            values = by_field[(cls, fid)]
            assert len(values) == counts[cls]
            state_counts = Counter(v for value in values for v in (value if isinstance(value, list) else [value]))
            assert set(state_counts) <= set(field['possible_values'])
            for state in field['possible_values']:
                distribution.append(dict(类别=zh, category=cls, 维度=dim, 特征=label, field=fid,
                                         状态=state, 目标数=state_counts[state], 类别目标数=counts[cls],
                                         占比=state_counts[state] / counts[cls],
                                         多选字段=field.get('value_type') == 'set'))
        mapping.append(dict(类别=zh, category=cls, 维度='尺度', 子组='原图目标比例',
                            特征='目标框面积占原图比例', field='bbox_area_ratio', 当前用途='参与评分',
                            来源='官方目标框计算', 可选状态='<1% / 1%–<3% / 3%–<10% / ≥10%'))
        scale_counts = Counter()
        for row in manifest:
            if row['class_name'] == cls:
                area = float(row['bbox_area_ratio'])
                scale_counts['<1%' if area < .01 else '1%–<3%' if area < .03 else '3%–<10%' if area < .1 else '≥10%'] += 1
        for state in ['<1%', '1%–<3%', '3%–<10%', '≥10%']:
            distribution.append(dict(类别=zh, category=cls, 维度='尺度', 特征='目标框面积占原图比例',
                                     field='bbox_area_ratio', 状态=state, 目标数=scale_counts[state],
                                     类别目标数=counts[cls], 占比=scale_counts[state] / counts[cls], 多选字段=False))
        for dim in DIMS:
            selected = [r for r in mapping if r['category'] == cls and r['维度'] == dim]
            count_rows.append(dict(类别=zh, 维度=dim, 字段数=len(selected),
                                   参与评分字段数=sum(r['当前用途'] == '参与评分' for r in selected)))
    assert len(annotations) == 1820 and len(mapping) == 188
    assert sum(r['字段数'] for r in count_rows) == len(mapping)
    write_csv('feature_dimension_catalog.csv', mapping)
    write_csv('feature_state_distribution.csv', distribution)
    write_csv('field_counts_by_dimension.csv', count_rows)
    lines = ['# 六类别细粒度特征的五维归类与实测分布', '',
             '本表依据2026-10-12实际完成的60个目标标注整理：每类10个，1820条特征记录。未重新运行VLM，未修改原始标注或评分。', '',
             '五维为尺度、可见性、遮挡、截断、朝向。特征后标★表示当前直接参与评分；其他条目属于已标注的细节或关联描述。', '',
             '头发可见性归入可见性；发色、长度、质地属于外观条件。耳朵姿态、尾巴姿态、动作放在朝向维度的关联姿态层，不等同于身体朝向，也不参与当前朝向评分。', '',
             '毛色、发色等可以作为“可见性×外观”的联合覆盖条件，但不能仅因它们属于某个部位就宣称是五维风险分量。关联条件完整保留，不强行赋分。', '',
             '尺度由官方目标框计算，不是VLM主观标注。犬猫的遮挡/截断状态为none、partial、heavy；其他类别为none、light、moderate、heavy，不能直接当作相同标尺。', '',
             '## 每类字段如何分配', '', '| 类别 | 尺度 | 可见性 | 遮挡 | 截断 | 朝向及关联姿态 | 关联条件 | VLM字段总数 |',
             '|---|---:|---:|---:|---:|---:|---:|---:|']
    for cls, zh in CLASSES.items():
        nums = [next(r['字段数'] for r in count_rows if r['类别'] == zh and r['维度'] == dim) for dim in DIMS]
        lines.append('| ' + ' | '.join([zh] + [str(n) for n in nums] + [str(sum(nums)-1)]) + ' |')
    lines += ['', '尺度列每类多出1项程序计算值，不计入VLM字段总数。可见性条目多是部位记录；当前犬猫马羊人仅汇总头部、躯干、肢体三项，车辆汇总八类适用部件。各维度预算均为20分，字段多不代表权重大。', '']
    for cls, zh in CLASSES.items():
        lines += [f'## {zh}：逐项归类', '', '| 维度 | 已输出特征（★参与当前评分） |', '|---|---|']
        for dim in DIMS:
            selected = [r for r in mapping if r['category'] == cls and r['维度'] == dim]
            names = [r['特征'] + ('★' if r['当前用途'] == '参与评分' else '') for r in selected]
            lines.append(f'| {dim} | ' + '；'.join(names) + ' |')
        lines += ['', '每个字段的完整可选值及是否计分见feature_dimension_catalog.csv。', '']
    lines += ['## 实际标注状态分布（每类10个）', '',
              '这里展示关键字段；完整CSV包含每个字段的全部可选状态，含未出现的0计数。多选部位可共存，各状态比例不必合计100%。这些是VLM输出分布，尚非人工确认真值，也不能代表整个候选池。', '',
              '| 特征 | 犬 | 猫 | 马 | 羊 | 人 | 车 |', '|---|---|---|---|---|---|---|']
    key_fields = [('bbox_area_ratio','尺度区间'), ('head_visibility','头部可见性'),
                  ('torso_visibility','躯干可见性'), ('limb_visibility','肢体可见性'),
                  ('occlusion_level','遮挡程度'), ('truncation_level','截断程度'),
                  ('body_orientation','身体朝向'), ('head_orientation','头部朝向')]
    for fid, label in key_fields:
        cells = []
        for cls in CLASSES:
            selected = [r for r in distribution if r['category'] == cls and r['field'] == fid and r['目标数']]
            cells.append('；'.join(f"{r['状态']}：{r['目标数']}" for r in selected) or '不适用')
        lines.append('| ' + ' | '.join([label] + cells) + ' |')
    lines += ['', '## 类别细节的实际分布', '']
    detail_fields = {'dog':['muzzle_visibility','mouth_visibility','eye_state','coat_texture'],
                     'cat':['muzzle_visibility','mouth_visibility','eye_state','coat_texture'],
                     'horse':['ear_visibility','tail_visibility','coat_texture'],
                     'sheep':['ear_visibility','tail_visibility','coat_texture'],
                     'person':['face_visibility','eye_visibility','hair_visibility','hair_presence','hair_color','hair_length','hair_texture','headwear','eyewear'],
                     'car':['front_visibility','rear_visibility','side_visibility','roof_visibility','windshield_visibility','wheels_visibility','windows_visibility','lights_visibility']}
    for cls, zh in CLASSES.items():
        lines += [f'### {zh}', '', '| 特征 | 实际状态及目标数 |', '|---|---|']
        for fid in detail_fields[cls]:
            selected = [r for r in distribution if r['category'] == cls and r['field'] == fid and r['目标数']]
            lines.append('| ' + selected[0]['特征'] + ' | ' + '；'.join(f"{r['状态']}：{r['目标数']}" for r in selected) + ' |')
        lines += ['']
    lines += ['## 尚未实现的细节', '',
              '马和羊目前共用27项字段，未标注马鬃、蹄部、羊角等类别专属字段。人物与车辆没有逐部位遮挡/截断字段；马羊人的截断部位采用共享字典，当前包含tail_or_rear这一兼容性条目。后续若细化，须更新字典并重新标注，不能把这些未测属性写成已输出结果。', '',
              '本表只整理字段归属。把面部、眼睛、头发等进一步纳入评分，需要另行定义层级权重并验证，避免头部及其子部位被重复计分。原始图像抽查已发现人物眼睛与车轮可见性有待复核；1例羊朝向未知，仍按原报告标为不可评价。']
    text = '\n'.join(lines) + '\n'
    (OUT / 'report.md').write_text(text, encoding='utf-8')
    doc = Document()
    doc.styles['Normal'].font.size = Pt(10)
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith('|'):
            block = []
            while index < len(lines) and lines[index].startswith('|'):
                block.append(lines[index]); index += 1
            cells = [[v.strip() for v in row.strip('|').split('|')] for row in block]
            table = doc.add_table(rows=1, cols=len(cells[0])); table.style = 'Table Grid'
            for cell, value in zip(table.rows[0].cells, cells[0]): cell.text = value
            for row in cells[2:]:
                for cell, value in zip(table.add_row().cells, row): cell.text = value
            continue
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            doc.add_heading(line.lstrip('#').strip(), level=min(level,3)-1)
        elif line:
            doc.add_paragraph(line)
        index += 1
    doc.save(OUT / 'report.docx')
    inputs = [Path(__file__), SOURCE/'manifest.csv', SOURCE/'feature_annotations_with_evidence.csv', SOURCE/'feature_dimension_assignment.csv']
    inputs += [SOURCE/'schemas'/f'{cls}.json' for cls in CLASSES]
    (OUT/'claim_evidence.json').write_text(json.dumps(dict(
        scope='60 frozen VLM-annotated targets, 10 per category; organization and descriptive frequencies only',
        annotation_records=1820, catalog_fields_including_computed_scale=188,
        input_sha256={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
        source_annotations_changed=False, scoring_rules_changed=False,
        frequency_definition='number of targets with each state / 10; multi-select states may co-occur',
        semantic_validation='partial; preserved pending review flags from source report'), ensure_ascii=False, indent=2),encoding='utf-8')
    print(json.dumps({'output':str(OUT),'catalog_fields':len(mapping),'state_rows':len(distribution),
                      'field_counts':count_rows},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
