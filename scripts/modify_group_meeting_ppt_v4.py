#!/usr/bin/env python3
"""Polish the supplied group-meeting deck and complete the two paper summaries."""
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


SOURCE = Path(
    "/home/hcf/.codex/attachments/4df473b5-3715-417e-bb00-8d6674a7de68/"
    "QE_group_meeting_visual_test_adequacy_20260915_v3.pptx"
)
OUTPUT = Path("/home/hcf/project/QE/docs/QE_group_meeting_visual_test_adequacy_20260915_v4.pptx")


def set_text(shape, value, font_size=None):
    if not hasattr(shape, "text_frame"):
        return
    frame = shape.text_frame
    template = None
    if frame.paragraphs and frame.paragraphs[0].runs:
        run = frame.paragraphs[0].runs[0]
        color = None
        try:
            if run.font.color.type is not None:
                color = run.font.color.rgb
        except (AttributeError, ValueError):
            pass
        template = {
            "name": run.font.name,
            "size": run.font.size,
            "bold": run.font.bold,
            "italic": run.font.italic,
            "color": color,
        }
    frame.clear()
    if not value:
        return
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    run = paragraph.runs[0]
    if template:
        run.font.name = template["name"]
        run.font.size = font_size or template["size"]
        run.font.bold = template["bold"]
        run.font.italic = template["italic"]
        if template["color"] is not None:
            run.font.color.rgb = template["color"]
    elif font_size:
        run.font.size = font_size
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def resize(shape, left=None, top=None, width=None, height=None):
    if left is not None:
        shape.left = Inches(left)
    if top is not None:
        shape.top = Inches(top)
    if width is not None:
        shape.width = Inches(width)
    if height is not None:
        shape.height = Inches(height)


def replace(slide, mapping):
    for index, value in mapping.items():
        set_text(slide.shapes[index], value)


def main():
    prs = Presentation(SOURCE)

    # Slide 8: make the two key papers directly comparable: problem, innovation,
    # and the precise lesson used by this project.
    slide = prs.slides[7]
    replace(
        slide,
        {
            1: "两篇关键论文：分别回答“能否信任测试结果”和“如何比较策展策略”",
            2: "TEASMA连接充分性与故障检出；SELECT把数据策展放进统一预算和统一评价",
            6: "TEASMA · IEEE TSE 2024",
            7: "主要问题",
            9: "测试集即使准确率很高，是否真的具备足够的故障检出能力？仅看准确率或单一覆盖分数，难以判断测试结果是否值得信任。",
            11: "主要创新",
            14: "基于训练集构建DNN特定的故障检出率（FDR）预测模型；比较DSC、LSC、IDC和Mutation Score，并用回归与相关性检验充分性。",
            17: "SELECT · NeurIPS 2024",
            19: "主要问题",
            18: "数据策展策略常在不同数据规模、来源和成本下被比较，难以判断新增数据或复杂选择规则是否真正提高数据效用。",
            13: "",
            21: "",
            23: "",
            25: "",
            26: "主要创新",
            27: "构建ImageNet++与SELECT基准，用相同模型和多种下游、鲁棒性指标，系统比较随机、检索、合成等策展策略。",
            29: "本项目借鉴：统一预算、同一任务、同一评价；边界：SELECT衡量训练数据效用，本项目衡量测试证据覆盖。",
        },
    )
    resize(slide.shapes[9], top=3.22, height=0.72)
    resize(slide.shapes[14], top=4.35, height=0.82)
    resize(slide.shapes[18], left=7.10, top=3.22, width=5.12, height=0.72)
    resize(slide.shapes[27], top=4.35, height=0.82)
    set_text(slide.shapes[9], slide.shapes[9].text, Pt(12))
    set_text(slide.shapes[14], slide.shapes[14].text, Pt(12))
    set_text(slide.shapes[18], slide.shapes[18].text, Pt(12))
    set_text(slide.shapes[27], slide.shapes[27].text, Pt(12))

    # Slide 9: reduce the four-stage chain and the three unresolved questions to
    # one visual grammar.
    slide = prs.slides[8]
    replace(
        slide,
        {
            1: "研究缺口：测量、目标、收益没有连成闭环",
            2: "现有方法分别覆盖局部环节，仍缺少面向视觉测试数据的可审计主线",
            5: "数据",
            6: "原始测试样本",
            8: "状态",
            9: "视角 / 尺度 / 遮挡 / 场景",
            11: "缺口",
            12: "缺失 / 稀疏 / 组合",
            14: "收益",
            15: "独立性能 / 错误",
            20: "问题 1",
            21: "状态测量可靠吗？",
            26: "问题 2",
            27: "缺口相对什么目标？",
            32: "问题 3",
            33: "补图是否改善评价？",
            35: "研究机会",
            36: "把可靠测量、目标约束、预算补图与独立验证组织成一条可审计流程。",
        },
    )

    # Slide 11: the workflow image remains the visual anchor; the footer now
    # explains the input/output boundary instead of repeating the title.
    slide = prs.slides[10]
    replace(
        slide,
        {
            1: "当前验证载体：犬类视觉状态覆盖流程",
            2: "固定目标、三视图输入、23项枚举特征、严格校验与缺口驱动选图",
            5: "输入：原图 + A框标记图 + A框裁剪图；输出：23项视觉特征、状态分布与缺口清单。",
            6: "边界：COCO程序固定目标框；VLM只负责属性测量；识别模型性能另行评估。",
        },
    )
    resize(slide.shapes[9], top=1.78, height=3.85)
    resize(slide.shapes[4], top=5.76, height=0.64)
    resize(slide.shapes[5], top=5.87, height=0.28)

    # Slide 16: make each backup answer a compact Q/A card with a consistent
    # evidence-first wording.
    slide = prs.slides[15]
    replace(
        slide,
        {
            1: "关键边界回答",
            2: "遇到质询时：先说当前证据，再说明不能推出什么，以及还需要什么实验",
            5: "Q1  阈值30从何而来？",
            7: "阶段性冻结的状态计数阈值，用来衡量最低证据目标；不宣称是普适的充分性阈值。",
            9: "Q2  覆盖上升是否必然带来性能上升？",
            11: "不必然。覆盖指数只说明状态分布更完整；模型收益必须在独立测试集上验证。",
            13: "Q3  VLM标签是否等于真值？",
            15: "不等于。100%只表示结构和枚举合法；还需要人工双标、一致性和标签扰动分析。",
            17: "Q4  为什么先做dog？",
            19: "COCO提供固定目标框，dog类内又有视角、姿态、遮挡和场景变化，适合先验证流程。",
            21: "Q5  核心创新是什么？",
            23: "受噪测量 → 缺口识别 → 预算补图 → 独立效用验证；贡献是可审计闭环，不是单一VLM或覆盖分数。",
            25: "回答原则：结论强度不超过证据强度。",
        },
    )
    for index in (7, 11, 15, 19, 23):
        set_text(slide.shapes[index], slide.shapes[index].text, Pt(12))

    # Slide 17: align the three dataset-count cards and make the denominator of
    # every reported result explicit.
    slide = prs.slides[16]
    replace(
        slide,
        {
            1: "结果口径与证据位置",
            2: "三种数量分别对应运行、合并和清洁口径；覆盖率只在清洁集上作最终报告",
            7: "900张标注运行",
            8: "500初始 + 400候选；确认v3.1结构化输出全部通过",
            10: "600张合并",
            11: "500初始 + 100补图；用于补图后状态分布计算",
            13: "596张清洁",
            14: "496张初始真实犬 + 100张补图；排除4张非真实犬",
            16: "主指标：覆盖指数 78.64% → 84.17%；低频状态 45 → 33。",
            19: "目标达成率按预冻结数量目标计算：63.41% → 74.80%。",
        },
    )
    for index in (8, 11, 14, 16, 19):
        set_text(slide.shapes[index], slide.shapes[index].text, Pt(12))
    resize(slide.shapes[5], left=1.02, top=2.13, width=6.48, height=3.78)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
