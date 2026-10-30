#!/usr/bin/env python3
"""Direct eight-model comparison charts, without explanatory report sections."""
from pathlib import Path
from collections import defaultdict
import csv
import json
import hashlib
import shutil
import zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/model_comparison_figures_20261008'
FIG=OUT/'figures';FIG.mkdir(parents=True,exist_ok=True)
NAME='八种模型_结果对比图_20261008'
SOURCE=ROOT/'artifacts/model_results_report_20261008'
MODELS=['ResNet-18','ResNet-50','MobileNetV3-Small','MobileNetV3-Large','DenseNet-121','EfficientNet-B0','ConvNeXt-Tiny','Swin-T']
POOLS={'B':'COCO','C':'VOC2007','D':'Objects365'}
BLUE,ORANGE='#3574A8','#C47538'
COLORS=['#1B6CA8','#E38D26','#409773','#BD5858','#8763AB','#87674E','#CD70AA','#4C8791']
MARKERS=['o','s','^','D','v','P','X','>']
paths=[SOURCE/'model_full_pool_scores.csv',SOURCE/'model_tier_scores.csv']
def rows(p):
    with p.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
overall={(r['pool'],r['model']):r for r in rows(paths[0])}
raw=rows(paths[1]);assert len(raw)==720
run_values=defaultdict(dict)
for r in raw:
    key=r['pool'],r['label_source'],r['model'],int(r['tier'])
    assert r['seed'] not in run_values[key]
    run_values[key][r['seed']]=float(r['macro_f1_100'])
assert len(run_values)==240 and all(len(v)==3 for v in run_values.values())
data={k:np.array(list(v.values())) for k,v in run_values.items()}
assert len(overall)==24
font_manager.fontManager.addfont(str(ROOT/'scripts/assets/NotoSansSC-Regular.ttf'))
plt.rcParams.update({'font.family':'Noto Sans SC','font.size':14,'svg.fonttype':'none','pdf.fonttype':42,
    'axes.spines.top':False,'axes.spines.right':False,'axes.linewidth':.7,'xtick.major.size':0,'ytick.major.size':0})
figures=[]
def save(fig,stem,title,caption):
    for ext in ['svg','png','pdf']:
        fig.savefig(FIG/f'{stem}.{ext}',dpi=250,bbox_inches='tight',facecolor='white')
    plt.close(fig);figures.append({'stem':stem,'title':title,'caption':caption})

fig,axes=plt.subplots(1,3,figsize=(15,6.2),layout='constrained',sharey=True)
for ax,(pool,title) in zip(axes,POOLS.items()):
    means=np.array([float(overall[pool,m]['mean_f1_100']) for m in MODELS])
    lo=np.array([float(overall[pool,m]['minimum_of_3_runs']) for m in MODELS])
    hi=np.array([float(overall[pool,m]['maximum_of_3_runs']) for m in MODELS])
    ax.barh(range(8),means,color=BLUE,height=.55)
    ax.errorbar(means,range(8),xerr=[means-lo,hi-means],fmt='none',ecolor='#203E58',capsize=3,lw=1)
    for i,v in enumerate(means):ax.text(max(v,hi[i])+1.2,i,f'{v:.2f}',va='center',fontsize=13)
    ax.set(title=title,yticks=range(8),yticklabels=MODELS,ylim=(7.65,-.65),xlim=(0,106),
           xticks=[0,25,50,75,100],xlabel='总体Macro-F1（满分100）')
save(fig,'01_overall_comparison','8种模型的总体成绩对比',
     '每种模型3份训练结果取平均；黑色短线为3份结果的最小—最大范围。总体F1使用该来源的全部目标计算。')

for source,source_name in [('local','本地AI描述'),('online','在线AI描述')]:
    fig,axes=plt.subplots(1,3,figsize=(15,6.25),layout='constrained',sharey=True)
    for ax,(pool,title) in zip(axes,POOLS.items()):
        for i,m in enumerate(MODELS):
            easy=data[pool,source,m,1];hard=data[pool,source,m,5]
            a,b=easy.mean(),hard.mean()
            ax.plot([b,a],[i,i],color='#C9D3DC',lw=2,zorder=1)
            for vals,c,label in [(easy,BLUE,'第1组：较容易'),(hard,ORANGE,'第5组：较困难')]:
                mean=vals.mean()
                ax.errorbar(mean,i,xerr=np.array([[mean-vals.min()],[vals.max()-mean]]),fmt='o',color=c,
                            markersize=6,capsize=2.5,lw=1,label=label if i==0 else None,zorder=3)
            ax.text(easy.max()+1.1,i,f'{a:.1f}',ha='left',va='center',color=BLUE,fontsize=12)
            ax.text(hard.min()-1.1,i,f'{b:.1f}',ha='right',va='center',color=ORANGE,fontsize=12)
        ax.set(title=title,yticks=range(8),yticklabels=MODELS,ylim=(7.65,-.65),xlim=(20,111),
               xticks=[25,50,75,100],xlabel='Macro-F1（满分100）')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='outside lower center',ncol=2,frameon=False,fontsize=14)
    save(fig,f'02_easy_hard_{source}',f'8种模型：较容易组与较困难组对比（{source_name}）',
         '蓝点=第1组，橙点=第5组；同一行对应同一模型。点为3份训练结果均值，短线为最小—最大范围；横轴显示20—111分。')

for source,source_name in [('local','本地AI描述'),('online','在线AI描述')]:
    fig,axes=plt.subplots(1,3,figsize=(15,6.25),layout='constrained',sharey=True)
    for ax,(pool,title) in zip(axes,POOLS.items()):
        for i,m in enumerate(MODELS):
            vals=[data[pool,source,m,t].mean() for t in range(1,6)]
            ax.plot(range(1,6),vals,color=COLORS[i],marker=MARKERS[i],markersize=5,lw=1.7,label=m)
        ax.set(title=title,xticks=range(1,6),xlim=(.85,5.15),ylim=(0,105),yticks=[0,25,50,75,100],
               xlabel='组别：1 → 5')
    axes[0].set_ylabel('Macro-F1（满分100）')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='outside lower center',ncol=4,frameon=False,fontsize=13)
    save(fig,f'03_five_groups_{source}',f'8种模型的五组成绩变化（{source_name}）',
         '每条线代表一种模型，点为3份训练结果的平均；组别1—5按事先规则从较容易到较困难划分。')

doc=Document();sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=Cm(29.7),Cm(21)
sec.left_margin=sec.right_margin=Cm(1.2)
sec.top_margin=sec.bottom_margin=Cm(1.15)
normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK SC'
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
normal.font.size=Pt(9.5)
normal.paragraph_format.space_after=Pt(4)
for i,record in enumerate(figures):
    if i:doc.add_page_break()
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
    run=p.add_run(record['title']);run.font.size=Pt(17);run.font.bold=True;run.font.color.rgb=RGBColor.from_string('24577E')
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/f'{record["stem"]}.png'),width=Cm(26.9))
    doc.add_paragraph(record['caption'])
doc.core_properties.title='8种模型结果对比图'
doc.core_properties.subject='总体F1、第一与第五组、五组变化；仅图表与简短图注'
doc.save(OUT/f'{NAME}.docx');shutil.copy2(OUT/f'{NAME}.docx',ROOT/'docs'/f'{NAME}.docx')
assert len(doc.inline_shapes)==5
for p in paths:shutil.copy2(p,OUT/p.name)
shutil.copy2(Path(__file__),OUT/Path(__file__).name)
audit={'figures':figures,'models':MODELS,'training_results_per_model':3,'source_sha256':{
    str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},
    'tier_rows':720,'new_training_or_inference':False,'planned_word_pages':5,'render_verified':False}
(OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
(OUT/'说明.txt').write_text('5张直接比较图：总体成绩、两套分组下第1与第5组、两套分组下五组曲线。每种模型用3份训练结果计算均值。PNG适合直接插入Word或PPT；SVG文字可编辑；PDF是单图矢量文件。两套AI负责描述和分组，不是被比较的8种分类模型。未新增训练或推理。\n',encoding='utf-8')
with zipfile.ZipFile(ROOT/'docs/八种模型_对比图与数据_20261008.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in FIG.glob('*'):z.write(p,p.relative_to(OUT))
    for name in ['model_full_pool_scores.csv','model_tier_scores.csv','audit.json','说明.txt',Path(__file__).name]:z.write(OUT/name,name)
    for p in paths:z.write(p,Path('source_data')/p.relative_to(ROOT))
    z.write(Path(__file__),'source_data/scripts/'+Path(__file__).name)
    z.write(ROOT/'scripts/assets/NotoSansSC-Regular.ttf','source_data/scripts/assets/NotoSansSC-Regular.ttf')
print(json.dumps({'word':str(ROOT/'docs'/f'{NAME}.docx'),'figures':5,'source_rows':len(raw)},ensure_ascii=False))
