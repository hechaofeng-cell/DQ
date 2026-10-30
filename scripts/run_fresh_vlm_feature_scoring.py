#!/usr/bin/env python3
"""New per-target VLM annotations, followed by fixed-rule visual risk scoring."""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
import math
import os
from pathlib import Path
import time

import numpy as np
from PIL import Image

import run_category_five_dimension_pilot as pilot
from run_dog_local_pipeline import ask, feature_schema

ROOT = Path(__file__).resolve().parents[1]
POOL = ROOT / "artifacts/sixclass_dual_vlm_replaced_20261004"
OUTPUT = ROOT / "artifacts/fresh_vlm_feature_scoring_20261012"
CLASSES = ("dog", "cat", "horse", "sheep", "person", "car")
PART_VALUES = ["clear", "partial", "not_visible", "not_in_view", "unknown"]


def now():
    from datetime import datetime
    from zoneinfo import ZoneInfo
    return datetime.now(ZoneInfo("Asia/Shanghai")).isoformat()


def atomic_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    temporary.replace(path)


def feature(fid, name, values, group, kind="enum"):
    f = dict(feature_id=fid, feature_name_zh=name, possible_values=values,
             group=group, value_type=kind, allow_unknown=True)
    if kind == "set":
        f["exclusive_values"] = [x for x in ("none", "unknown", "not_in_view") if x in values]
    return f


def build_schema(category):
    common = json.loads((ROOT / "configs/multiclass_object_state_schema_v1_0.json").read_text())
    if category in {"dog", "cat"}:
        relative = "configs/dog_feature_schema_v3_1.json" if category == "dog" else "configs/cat_feature_schema_v1_0.json"
        original = json.loads((ROOT / relative).read_text())
        fields = deepcopy(original["universal_features"])
        existing = {f["feature_id"] for f in fields}
        fields += [deepcopy(f) for f in common["universal_features"]
                   if f["feature_id"] not in existing and f["feature_id"] not in {"head_orientation", "leg_visibility", "pose_action", "scene_environment", "truncation_parts"}]
        extra = [feature("torso_visibility", "躯干可见性", PART_VALUES, "visibility"),
                 feature("limb_visibility", "适用肢体可见性", PART_VALUES, "visibility")]
        retained = [f["feature_id"] for f in original["universal_features"]]
    elif category == "car":
        native = json.loads((POOL / "configs/car_annotation_schema.json").read_text())
        fields = deepcopy(native["universal_features"] + native["category_specific_features"])
        extra = [feature(f"{part}_visibility", f"{name}可见性", PART_VALUES, "visibility")
                 for part, name in (("front", "车头"), ("rear", "车尾"), ("side", "车身侧面"),
                                    ("roof", "车顶"), ("windshield", "前风挡"), ("wheels", "车轮区域"),
                                    ("windows", "车窗"), ("lights", "灯具"))]
        retained = []
    else:
        fields = deepcopy(common["universal_features"])
        extra = [feature("torso_visibility", "躯干可见性", PART_VALUES, "visibility"),
                 feature("limb_visibility", "适用肢体可见性", PART_VALUES, "visibility"),
                 feature("muzzle_visibility" if category != "person" else "nose_visibility",
                         "口鼻部可见性" if category != "person" else "鼻部可见性", PART_VALUES, "visibility"),
                 feature("mouth_visibility", "嘴部可见性", PART_VALUES, "visibility")]
        if category == "person":
            extra += [feature("face_visibility", "面部可见性", PART_VALUES, "visibility"),
                      feature("eye_visibility", "眼睛可见性", ["both", "left_only", "right_only", "none", "not_in_view", "unknown"], "visibility"),
                      feature("hair_visibility", "头发可见性", PART_VALUES, "appearance"),
                      feature("hair_presence", "头发存在状态", ["present", "bald_or_shaved", "covered", "unknown"], "appearance"),
                      feature("hair_color", "头发颜色", ["black", "brown", "blond", "red", "gray", "white", "mixed", "other", "not_visible", "unknown"], "appearance"),
                      feature("hair_length", "头发长度", ["shaved", "short", "medium", "long", "not_visible", "unknown"], "appearance"),
                      feature("hair_texture", "头发质地", ["straight", "wavy", "curly", "coily", "braided", "covered", "not_visible", "unknown"], "appearance"),
                      feature("headwear", "头部配饰", ["none", "hat", "cap", "helmet", "headscarf", "hood", "other", "unknown"], "appearance"),
                      feature("eyewear", "眼部配饰", ["none", "glasses", "sunglasses", "other", "unknown"], "appearance")]
        else:
            extra += [feature("ear_visibility", "耳部可见性", PART_VALUES, "visibility"),
                      feature("tail_visibility", "尾部可见性", PART_VALUES, "visibility"),
                      feature("coat_primary_color", "被毛主色", ["black", "white", "brown", "gray", "tan", "cream", "mixed", "other", "unknown"], "appearance"),
                      feature("coat_texture", "被毛质地", ["smooth", "wavy", "curly", "wiry", "fluffy", "other", "unknown"], "appearance")]
        retained = []
    fields += extra
    if len({f["feature_id"] for f in fields}) != len(fields):
        raise ValueError("Duplicate schema feature")
    return dict(schema_version=f"fresh_{category}_fine_v1", category=category,
                universal_features=fields, category_specific_features=[],
                preserved_early_field_ids=retained)


def scoring_protocol():
    base = json.loads(pilot.DEFAULT_CONFIG.read_text())
    return dict(version="fresh_five_dimensions_v1", dimension_budget=20,
                maps=base["maps"], scale_edges=base["scale_edges"],
                visibility_regions=dict(noncar=["head_visibility", "torso_visibility", "limb_visibility"],
                                        car=[f"{p}_visibility" for p in ("front", "rear", "side", "roof", "windshield", "wheels", "windows", "lights")]),
                unknown_values=["unknown", "not_observable"],
                not_applicable_values=["not_in_view"], minimum_known_weight=.5,
                nominal_fields_policy="preserve descriptions, never assign an unvalidated risk ordering",
                class_prediction_policy="official labels and frozen classifier predictions; VLM does not classify or score")


PROMPT = """你是固定目标的细粒度视觉特征标注器。输入顺序为原图、带目标框图、目标裁剪放大图。只描述指定类别的A框目标target_1，不更换对象。
逐字段观察，并严格按附带字典输出全部features及一一对应的evidence，每项证据用简短中文（通常不超过20字）说明可见事实或无法判断原因。只输出JSON，不输出评分、分类预测、置信度或额外字段。
原图判断原始边界、背景及遮挡来源；标记图确认目标；裁剪图看局部。紧贴目标框的裁剪不代表原图截断。occlusion是被其他可见物体遮蔽；truncation是目标超出原图边界，两者分开。none和unknown不能与其他set值共存，set不为空。
可见性字段中clear表示适用区域在该视角下清楚可观察，partial表示仅部分可观察，not_visible表示适用区域因遮挡、截断等缺失而不可见。只有字典允许not_in_view时，才可用它表示正常视角不呈现的部位，例如正面车图的车尾；不要把遮挡、低清或无法判断标为not_in_view。unknown表示证据不足，不能猜测。
torso_visibility表示躯干整体，limb_visibility表示当前视角通常应呈现的肢体区域完整程度，不以固定应见四条腿判定所有类别。轮胎数量按可辨认数量填写，但正常侧视不因只有两轮就判定重风险。
头部朝向与身体朝向分别判断，左右采用目标自身方位。表面姿态动作只描述直接可见姿态，静态图不能推断尾巴正在摆动或车正在行驶。犬猫原协议中的not_observable表示无法观察；尾巴不可观察时tail_posture必须unknown。嘴部not_visible时mouth_state、tongue_visibility（若字段存在）必须unknown。
hair_color、hair_length、hair_texture仅描述可见头发，头发不可观察时用对应字典中的not_visible或unknown，不猜测。不得推断年龄、性别、种族、身份、品种、健康、情绪等。眼睛和嘴部的状态以可见事实为准，不推测被遮挡部分。
保留原字典全部枚举名，不自由造词。若某字段的字典没有not_in_view，无法可靠描述时用unknown。实际schema_version必须使用附带版本。
"""


def prepare(output, per_class, seed):
    output.mkdir(parents=True, exist_ok=True)
    config_path = output / "runtime.json"
    configuration = dict(model="qwen3-vl:30b-a3b-instruct", url="http://localhost:11434/api/chat",
                         num_ctx=16384, num_predict=4000, temperature=0, seed=seed,
                         timeout_seconds=900, max_retries=0, retry_wait_seconds=2,
                         per_class=per_class, selection="fixed class-stratified random IDs, no classifier feedback")
    if config_path.exists() and json.loads(config_path.read_text()) != configuration:
        raise ValueError("Existing run configuration differs; use another output directory")
    atomic_json(config_path, configuration)
    rows = pilot.read_csv(POOL / "manifests/candidate_pool.csv")
    rng = np.random.default_rng(seed)
    chosen = {}
    for c in CLASSES:
        indices = [i for i, r in enumerate(rows) if r["class_name"] == c]
        if not 1 <= per_class <= len(indices):
            raise ValueError("Invalid per-class size")
        chosen[c] = [rows[i] for i in rng.choice(indices, per_class, replace=False)]
    # Interleave categories so the first six real calls check all category schemas.
    selected = [chosen[c][j] for j in range(per_class) for c in CLASSES]
    native_inputs = {}
    for domain in ("old", "car"):
        native_inputs.update({r["image_id"]: r for r in pilot.read_csv(POOL / "inputs" / f"{domain}.csv")})
    manifest = []
    for row in selected:
        native = native_inputs[row["sample_id"]]
        with Image.open(row["image_path"]) as image:
            pilot.validate_input_binding(row, native, image.size)
        manifest.append(dict(row, annotation_target_instance_id=native["target_instance_id"],
                             annotation_bbox=native["bbox"]))
    manifest_path = output / "manifest.csv"
    if manifest_path.exists() and pilot.read_csv(manifest_path) != manifest:
        raise ValueError("Selected target manifest changed")
    pilot.write_csv(manifest_path, manifest)
    paths = [config_path, manifest_path, Path(__file__)]
    for c in CLASSES:
        schema = build_schema(c)
        path = output / "schemas" / f"{c}.json"
        if path.exists() and json.loads(path.read_text()) != schema:
            raise ValueError("Frozen schema changed")
        atomic_json(path, schema)
        paths.append(path)
    prompt_path = output / "prompt.txt"
    prompt_path.write_text(PROMPT, encoding="utf-8")
    protocol_path = output / "scoring_protocol.json"
    protocol = scoring_protocol()
    if protocol_path.exists() and json.loads(protocol_path.read_text()) != protocol:
        raise ValueError("Frozen scoring protocol changed")
    atomic_json(protocol_path, protocol)
    paths += [prompt_path, protocol_path]
    for row in manifest:
        for field in ("image_path", "marked_image_path", "feature_crop_path"):
            paths.append(Path(row[field]))
    freeze = {str(p.resolve()): pilot.sha256(p) for p in paths}
    freeze_path = output / "input_freeze.sha256.json"
    if freeze_path.exists() and json.loads(freeze_path.read_text()) != freeze:
        raise ValueError("Inputs changed since freeze")
    atomic_json(freeze_path, freeze)
    log = output / "research_log.md"
    if not log.exists():
        log.write_text("# 新VLM逐目标标注后评分\n\n"
                       "目标：实际重新标注类别专用细粒度字段，再用冻结规则评分。"
                       "本轮先验证标注、证据、可见性分量与评分链路；样本较少，不确认方法优越性。"
                       "选样不读取历史分类错误，新标注不复用旧标签。犬猫保留原23项字段及值域。"
                       "五维各20分，部件not_in_view不算缺失，unknown保留区间，外观名义属性不任意排序。\n\n"
                       "主比较在同一目标上使用旧整体标签评分、旧局部标签评分、新细粒度标签评分；"
                       "新旧方案字段与规则并不完全相同，不作标注来源的单因素因果推断。"
                       "先每类一目标真实预检，再完成固定清单。失败会写明并要求完整JSON重试，"
                       "有效标注数量不够时不输出complete。\n", encoding="utf-8")
    return manifest, configuration


def validate_annotation(value, row, schema):
    if value.get("schema_version") != schema["schema_version"]:
        raise ValueError("Schema version mismatch")
    pilot.validate_label(value, dict(sample_id=row["sample_id"], target_instance_id=row["annotation_target_instance_id"]), schema)
    evidence = value.get("evidence")
    features = value["features"]
    if not isinstance(evidence, dict) or set(evidence) != set(features):
        raise ValueError("Missing field-level evidence")
    if any(not isinstance(x, str) or not x.strip() for x in evidence.values()):
        raise ValueError("Empty field evidence")
    if features.get("mouth_visibility") == "not_visible":
        for key in ("mouth_state", "tongue_visibility"):
            if key in features and features[key] != "unknown":
                raise ValueError("Invisible mouth dependency")
    if features.get("tail_visibility") == "not_observable" and features.get("tail_posture") != "unknown":
        raise ValueError("Unobservable tail dependency")
    return value


def annotate(manifest, configuration, output, limit=0):
    started = time.monotonic()
    errors, records = [], []
    work = manifest[:limit] if limit else manifest
    for position, row in enumerate(work, 1):
        path = output / "annotations" / f"{row['sample_id']}.json"
        schema_path = output / "schemas" / f"{row['class_name']}.json"
        schema = json.loads(schema_path.read_text())
        state = dict(status="running", stage="annotation", planned=len(manifest), current_batch=len(work),
                     processed=position - 1, successful=len(records), failed=len(errors),
                     current_sample=row["sample_id"], current_class=row["class_name"],
                     updated_at=now(), pid=os.getpid())
        atomic_json(output / "progress.json", state)
        if path.exists():
            value = validate_annotation(json.loads(path.read_text()), row, schema)
            records.append(value)
            print(f"[resume] {position}/{len(work)} {row['class_name']} {row['sample_id']}", flush=True)
            continue
        compact = dict(schema_version=schema["schema_version"], category=row["class_name"],
                       fields=schema["universal_features"])
        prompt = PROMPT + "\n冻结特征字典：\n" + json.dumps(compact, ensure_ascii=False)
        native = dict(image_id=row["sample_id"])
        last_error = None
        # Each attempt has a distinct raw response, including semantic retries.
        for attempt in range(1, 4):
            try:
                correction = "" if last_error is None else f"\n上一份输出校验错误：{last_error}。重新输出全部字段和证据。"
                candidate = ask(configuration, prompt + correction, native,
                                [Path(row["image_path"]), Path(row["marked_image_path"]), Path(row["feature_crop_path"])],
                                output / "raw" / f"{row['sample_id']}_semantic{attempt}.json",
                                feature_schema(schema_path))
                validate_annotation(candidate, row, schema)
                candidate["annotation_model"] = configuration["model"]
                candidate["annotated_at"] = now()
                candidate["schema_sha256"] = pilot.sha256(schema_path)
                candidate["prompt_sha256"] = pilot.sha256(output / "prompt.txt")
                atomic_json(path, candidate)
                records.append(candidate)
                last_error = None
                break
            except Exception as exc:
                last_error = str(exc)
                errors.append(dict(sample_id=row["sample_id"], category=row["class_name"], attempt=attempt, error=last_error))
                print(f"[retry] {row['sample_id']} attempt={attempt}: {last_error}", flush=True)
        if last_error is not None:
            pilot.write_csv(output / "annotation_attempt_errors.csv", errors)
            atomic_json(output / "progress.json", dict(state, status="failed", failure=last_error,
                                                         successful=len(records), updated_at=now()))
            raise ValueError(f"Annotation failed; incomplete run: {row['sample_id']}: {last_error}")
        if errors:
            pilot.write_csv(output / "annotation_attempt_errors.csv", errors)
        print(f"[annotated] {position}/{len(work)} {row['class_name']} fields={len(candidate['features'])} "
              f"elapsed={candidate['elapsed_seconds']:.1f}s", flush=True)
    atomic_json(output / "annotation_audit.json", dict(status="pass", expected=len(work), valid=len(records),
                                                        fresh_or_valid_resumed=len(records),
                                                        elapsed_seconds=time.monotonic() - started,
                                                        attempt_errors=len(errors), complete_manifest=len(work) == len(manifest)))
    return len(records)


def region_score(fields, maps):
    applicable = [key for key, value in fields.items() if value != "not_in_view"]
    if not applicable:
        return pilot.aggregate_components([(None, 1.0)])
    values = []
    for key in applicable:
        value = fields[key]
        if value in {"unknown", "not_observable"}:
            risk = None
        elif value in maps:
            risk = maps[value]
        else:
            raise ValueError(f"Unmapped visibility state: {value}")
        values.append((risk, 1 / len(applicable)))
    return pilot.aggregate_components(values)


def fresh_score(row, value, protocol):
    f, c = value["features"], row["class_name"]
    maps = protocol["maps"]
    old = c in {"dog", "cat"}
    regions = protocol["visibility_regions"]["car" if c == "car" else "noncar"]
    components = {"visibility": region_score({k: f[k] for k in regions}, maps["part_visibility"])}
    config = dict(scale_edges=protocol["scale_edges"])
    components["scale"] = pilot.aggregate_components([(pilot.scale_risk(float(row["bbox_area_ratio"]), config), 1)])
    for dimension, key in (("occlusion", "occlusion_level"), ("truncation", "truncation_level")):
        mapping = maps[f"legacy_{key}" if old else key]
        risk = None if f[key] == "unknown" else mapping[f[key]]
        components[dimension] = pilot.aggregate_components([(risk, 1)])
    direction_fields = ["body_orientation"] if c == "car" else ["body_orientation", "head_orientation"]
    direction = []
    for key in direction_fields:
        mapping = maps[f"legacy_{key}" if old else key]
        risk = None if f[key] == "unknown" else mapping[f[key]]
        direction.append((risk, 1 / len(direction_fields)))
    components["orientation"] = pilot.aggregate_components(direction)
    lower = 20 * sum(v["lower"] for v in components.values())
    upper = 20 * sum(v["upper"] for v in components.values())
    result = dict(score_lower=lower, score_upper=upper, score_midpoint=(lower + upper) / 2,
                  fully_known=all(math.isclose(v["known_weight"], 1) for v in components.values()),
                  evaluable=all(v["known_weight"] >= protocol["minimum_known_weight"] - 1e-12 for v in components.values()),
                  unknown_dimensions=sum(v["unknown_components"] > 0 for v in components.values()),
                  applicable_visibility_regions=sum(f[k] != "not_in_view" for k in regions))
    for dimension, values in components.items():
        for key, v in values.items():
            result[f"{dimension}_{key}"] = v
    if not -1e-8 <= lower <= result["score_midpoint"] <= upper <= 100 + 1e-8:
        raise ValueError("Score outside fixed budget")
    return result


def score_and_report(manifest, output, limit=0):
    work = manifest[:limit] if limit else manifest
    protocol = json.loads((output / "scoring_protocol.json").read_text())
    rows, features_long = [], []
    for row in work:
        schema = json.loads((output / "schemas" / f"{row['class_name']}.json").read_text())
        path = output / "annotations" / f"{row['sample_id']}.json"
        value = validate_annotation(json.loads(path.read_text()), row, schema)
        score = fresh_score(row, value, protocol)
        rows.append(dict(sample_id=row["sample_id"], class_name=row["class_name"], schema_version=schema["schema_version"],
                         **score, **{f"feature_{k}": json.dumps(v, ensure_ascii=False) if isinstance(v, list) else v for k, v in value["features"].items()}))
        for f in schema["universal_features"]:
            key = f["feature_id"]
            features_long.append(dict(sample_id=row["sample_id"], class_name=row["class_name"], field=key,
                                      feature_name=f.get("feature_name_zh", key),
                                      value=json.dumps(value["features"][key], ensure_ascii=False), evidence=value["evidence"][key]))
    pilot.write_csv(output / "scores.csv", rows)
    pilot.write_csv(output / "feature_annotations_with_evidence.csv", features_long)
    # Prediction loading only begins after all requested annotations have passed validation.
    original = pilot.read_csv(POOL / "manifests/candidate_pool.csv")
    index = {r["sample_id"]: i for i, r in enumerate(original)}
    old_config = json.loads(pilot.DEFAULT_CONFIG.read_text())
    classifiers = []
    for path in sorted((POOL / "benchmark/runs").glob("*/*/candidate_outputs.npz")):
        with np.load(path) as cache:
            if cache["sample_ids"].tolist() != [r["sample_id"] for r in original]:
                raise ValueError("Prediction identity mismatch")
            classifiers.append((path.parent.parent.name, path.parent.name, cache["logits"].argmax(1)))
    comparable = []
    for category in CLASSES:
        selected = [r for r in work if r["class_name"] == category]
        if not selected:
            continue
        current = {r["sample_id"]: r for r in rows}
        for source in ("local", "online"):
            scores = {name: [] for name in ("old_global", "old_parts", "fresh_parts")}
            for row in selected:
                domain = "car" if category == "car" else "old"
                old_value = json.loads((POOL / "annotations" / source / domain / "parsed" / f"{row['sample_id']}.json").read_text())
                for scheme, name in (("equal_global", "old_global"), ("equal_parts", "old_parts")):
                    scores[name].append(pilot.score_sample(float(row["bbox_area_ratio"]), old_value["features"], category, scheme, old_config)["score_midpoint"])
                scores["fresh_parts"].append(current[row["sample_id"]]["score_midpoint"])
            for model, seed, predictions in classifiers:
                errors = np.asarray([predictions[index[row["sample_id"]]] != CLASSES.index(category) for row in selected], dtype=int)
                for scheme, risk in scores.items():
                    comparable.append(dict(class_name=category, old_source=source, scheme=scheme, model=model, seed=seed,
                                           n=len(selected), errors=int(errors.sum()), error_auc=pilot.error_auc(errors, np.asarray(risk))))
    pilot.write_csv(output / "classifier_comparison.csv", comparable)
    summary = []
    for category in CLASSES:
        group = [r for r in rows if r["class_name"] == category]
        if group:
            summary.append(dict(class_name=category, n=len(group), fields_per_target=len(build_schema(category)["universal_features"]),
                                score_mean=float(np.mean([r["score_midpoint"] for r in group])),
                                fully_known=sum(r["fully_known"] for r in group), evaluable=sum(r["evaluable"] for r in group),
                                uncertainty_mean=float(np.mean([r["score_upper"] - r["score_lower"] for r in group]))))
    pilot.write_csv(output / "summary_by_class.csv", summary)
    audit = dict(status="pass", completed_at=now(), targets=len(rows), expected=len(work),
                 complete_manifest=len(work) == len(manifest), feature_evidence_records=len(features_long),
                 model="qwen3-vl:30b-a3b-instruct", llm_direct_risk_scoring=False,
                 annotation_before_scoring=True, prediction_runs=len(classifiers), summary=summary)
    atomic_json(output / "audit.json", audit)
    atomic_json(output / "progress.json", dict(status="complete" if audit["complete_manifest"] else "smoke_complete",
                                                stage="scored", successful=len(rows), planned=len(manifest), updated_at=now()))
    report = ["# 新VLM标注后五维评分", "",
              "本次实际调用本地Qwen3-VL，对固定目标逐字段重新标注，并为每字段保存视觉证据。模型未输出风险分；评分由冻结的程序规则完成。", "",
              "犬猫保留原23项字段及值域，并增加躯干、肢体可见性等字段；人类加入头发、面部和配饰，车辆加入八种部件可见性。", "",
              f"本次有效目标{len(rows)}，字段与证据记录{len(features_long)}。", "",
              "| 类别 | 目标数 | 每目标字段数 | 平均分 | 完整评分证据目标 | 可评价目标 |", "|---|---:|---:|---:|---:|---:|"]
    for v in summary:
        report.append(f"| {v['class_name']} | {v['n']} | {v['fields_per_target']} | {v['score_mean']:.2f} | {v['fully_known']} | {v['evaluable']} |")
    report += ["", "五维各20分。犬猫、马羊和人的可见性汇总头部、躯干、肢体三个区域；车辆汇总该视角适用的部件。正常视角下不呈现的not_in_view部件不作为缺失计分；unknown保留分数区间。身体与头部朝向分别记录，毛色、发色等没有天然风险顺序的属性保留作条件字段。", "",
               "犬猫原协议的partial遮挡／截断本轮各映射为0.5，该版本独立定义，不等同于旧共享协议的moderate。新旧评分改变了标签和部分字段规则，不能把差异全部归因于LLM重标。", "",
               "classifier_comparison.csv在相同目标上检查旧整体、旧局部和新局部评分。样本较少、某些类别或模型无分类错误时AUROC不可评价，不能据此宣称新评分优于原方案。", "",
               "可追溯输出：manifest.csv为固定目标；schemas/及prompt.txt为字典和提示；raw/保留模型原始响应；annotations/保留逐目标特征及证据；feature_annotations_with_evidence.csv为长表；scores.csv为五维及总分；audit.json检查数量与执行状态。", "",
               "复现：`.venv/bin/python scripts/run_fresh_vlm_feature_scoring.py --phase all`。已存在且通过校验的标注可续用，不重复请求；修改字典或评分规则须用新输出目录。"]
    (output / "report.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    from docx import Document
    document = Document()
    table = None
    for line in report:
        if line.startswith("|---"):
            continue
        if line.startswith("|"):
            cells = [v.strip() for v in line.strip("|").split("|")]
            if table is None:
                table = document.add_table(rows=0, cols=len(cells));table.style = "Table Grid"
            for cell, text in zip(table.add_row().cells, cells):
                cell.text = text
        else:
            table = None
            if line.startswith("# "):
                document.add_heading(line[2:], 0)
            elif line:
                document.add_paragraph(line)
    document.save(output / "report.docx")
    (output / "claim_evidence.md").write_text("# 主张与证据\n\n支持：真实模型逐目标新标注后按规则评分，见raw/、annotations/、scores.csv及audit.json。\n\n"
                                            "未支持：新评分普遍优于原评分、标注全部正确、所有名义外观属性具有风险顺序。\n", encoding="utf-8")
    with (output / "research_log.md").open("a", encoding="utf-8") as log:
        log.write(f"\n{now()}：完成{len(rows)}个目标的新标注及评分；本轮完整清单={audit['complete_manifest']}。\n")
    print(json.dumps(audit, ensure_ascii=False, indent=2), flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--per-class", type=int, default=10)
    parser.add_argument("--seed", type=int, default=20261012)
    parser.add_argument("--phase", choices=("prepare", "annotate", "score", "all"), default="all")
    parser.add_argument("--limit", type=int, default=0, help="First N targets, for a real smoke run")
    args = parser.parse_args()
    manifest, config = prepare(args.output, args.per_class, args.seed)
    if args.phase in {"annotate", "all"}:
        annotate(manifest, config, args.output, args.limit)
    if args.phase in {"score", "all"}:
        score_and_report(manifest, args.output, args.limit)
    if args.phase == "prepare":
        print(json.dumps(dict(targets=len(manifest), per_class=args.per_class,
                              schemas={c: len(build_schema(c)["universal_features"]) for c in CLASSES}), ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
