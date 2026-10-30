#!/usr/bin/env python3
"""Build a three-page, teacher-facing result brief from verified report assets."""
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys
import zipfile

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw

import export_current_method_conclusions_20261012 as report
import recompile_methods_latex_20261012 as methods

ROOT = Path(__file__).resolve().parents[1]
NAME = '视觉状态驱动模型测试_教师汇报结果_20261012'
SOURCE = ROOT / 'docs' / (NAME + '.md')
OUT = ROOT / 'artifacts/teacher_results_20261012'
LATEX = OUT / 'latex'
TABLES = ROOT / 'artifacts/current_method_conclusions_20261012/tables'
PAGEBREAK = '<!-- PAGEBREAK -->'


def source_blocks():
    report.SOURCE = SOURCE
    return report.parse()


def add_text(paragraph, text):
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else part)
        run.bold = part.startswith('**')


def set_fonts(style, size, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn('w:eastAsia'), '黑体' if bold else '宋体')


def word_export(blocks, table_widths=None):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.9)
    set_fonts(doc.styles['Normal'], 10.5)
    set_fonts(doc.styles['Title'], 18, True)
    set_fonts(doc.styles['Heading 1'], 13, True)
    set_fonts(doc.styles['Heading 2'], 11, True)
    set_fonts(doc.styles['Caption'], 9)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = 1.15
    normal.space_after = Pt(5)
    for heading in ('Heading 1', 'Heading 2'):
        doc.styles[heading].paragraph_format.space_before = Pt(9)
        doc.styles[heading].paragraph_format.space_after = Pt(5)
    for kind, value in blocks:
        if kind == 'body' and value == PAGEBREAK:
            doc.add_page_break()
        elif kind == 'heading1':
            p = doc.add_paragraph(value, 'Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind in ('heading2', 'heading3'):
            doc.add_paragraph(value, 'Heading 1' if kind == 'heading2' else 'Heading 2')
        elif kind == 'image':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(value[1]), width=Cm(17))
            caption = doc.add_paragraph(value[0], 'Caption')
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_next = True
        elif kind == 'table':
            table = doc.add_table(rows=len(value), cols=len(value[0]))
            table.style = 'Light Shading Accent 1'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = ((table_widths or {}).get(len(value[0]))
                      or ([4.4, 2.4, 5.1, 5.1] if len(value[0]) == 4 else [4.4, 6.3, 6.3]))
            for row_index, (row, cells) in enumerate(zip(table.rows, value)):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if row_index == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for col_index, (cell, text) in enumerate(zip(row.cells, cells)):
                    cell.width = Cm(widths[col_index])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(text)
                    run.font.size = Pt(9)
                    run.bold = row_index == 0
        else:
            p = doc.add_paragraph()
            add_text(p, value)
            if value.startswith(('图', '数据依据：', '截至')):
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(9)
            if value.startswith('截至'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    output = SOURCE.with_suffix('.docx')
    doc.save(output)
    return output


def pdf_export(blocks, table_weights=None):
    (LATEX / 'fonts').mkdir(parents=True, exist_ok=True)
    fonts = ROOT / 'artifacts/methods_explanations_overlay_20261012/latex/fonts'
    for path in fonts.iterdir():
        if path.is_file():
            shutil.copy2(path, LATEX / 'fonts' / path.name)
    preamble = methods.PREAMBLE.replace('[11pt,a4paper]', '[10pt,a4paper]')
    preamble = preamble.replace('top=2.1cm,bottom=2cm', 'top=1.9cm,bottom=1.9cm')
    preamble = preamble.replace(r'\setstretch{1.22}', r'\setstretch{1.17}')
    parts = [preamble]
    figure_index = 0
    for kind, value in blocks:
        if kind == 'body' and value == PAGEBREAK:
            parts.append(r'\clearpage')
        elif kind == 'heading1':
            parts.append(r'\begin{center}{\Large\sffamily\bfseries '
                         + report.inline_tex(value) + r'\par}\end{center}')
        elif kind in ('heading2', 'heading3'):
            command = r'\section*' if kind == 'heading2' else r'\subsection*'
            parts.append(command + '{' + report.inline_tex(value) + '}')
        elif kind == 'image':
            figure_index += 1
            label, png = value
            target = LATEX / f'figure_{figure_index}.pdf'
            shutil.copy2(png.with_suffix('.pdf'), target)
            parts.append(r'\begin{center}\includegraphics[width=\linewidth]{'
                         + target.name + r'}\par\vspace{2pt}{\small '
                         + report.inline_tex(label) + r'\par}\end{center}')
        elif kind == 'table':
            n = len(value[0])
            weights = ((table_weights or {}).get(n)
                       or ([.26, .14, .30, .30] if n == 4 else [.26, .37, .37]))
            columns = '@{}' + ''.join(
                r'>{\raggedright\arraybackslash}p{\dimexpr ' + str(w)
                + r'\linewidth-' + str((n - 1) * 2 / n) + r'\tabcolsep\relax}'
                for w in weights) + '@{}'
            parts.append(r'\begin{center}{\small\renewcommand{\arraystretch}{1.25}'
                         + r'\begin{tabular}{' + columns + r'}\toprule')
            for row_index, row in enumerate(value):
                cells = [r'\textbf{' + report.inline_tex(cell) + '}' if row_index == 0
                         else report.inline_tex(cell) for cell in row]
                parts.append(' & '.join(cells) + r'\\' + (r'\midrule' if row_index == 0 else ''))
            parts.append(r'\bottomrule\end{tabular}}\end{center}')
        elif value.startswith(('图', '数据依据：', '截至')):
            parts.append(r'{\small\noindent ' + report.inline_tex(value) + r'\par}')
        else:
            parts.append(report.inline_tex(value) + '\n')
    parts.append(r'\end{document}')
    source = LATEX / 'main.tex'
    source.write_text('\n\n'.join(parts) + '\n', encoding='utf-8')
    methods.OUT, methods.SOURCE = OUT, SOURCE
    return methods.compile_latex(source)


def check_results():
    text = SOURCE.read_text(encoding='utf-8')
    coverage = pd.read_csv(TABLES / 'coverage_curves.csv')
    endpoint = coverage.loc[coverage.groupby(['pool', 'source']).total_budget.idxmax()]
    for value in endpoint.delta_pp:
        assert f'+{value:.2f}个百分点' in text
    mixture = pd.read_csv(TABLES / 'risk_mixture_endpoints.csv')
    for column in ('low_f1_percent', 'high_f1_percent'):
        for value in mixture[column]:
            assert f'{value:.2f}%' in text
    support = pd.read_csv(TABLES / 'condition_support.csv')
    for _, row in support.groupby('pool')[['eligible', 'negative_mean', 'insufficient']].sum().iterrows():
        assert '| ' + ' | '.join(str(v) for v in row) + ' |' in text
    return {'coverage_maxima_rows': len(endpoint), 'mixture_endpoint_values': 12,
            'condition_aggregates': support.groupby('pool')[['eligible', 'negative_mean', 'insufficient']]
                .sum().to_dict(orient='index')}


def audit(word, pdf, evidence):
    sys.path.insert(0, str(ROOT / 'artifacts/teacher_methodology_20261011/vendor'))
    import pymupdf
    preview = OUT / 'preview'
    preview.mkdir(exist_ok=True)
    document = pymupdf.open(pdf)
    assert len(document) == 3, len(document)
    expected = ['结果一', '结果二', '结果三']
    thumbs = []
    for index, page in enumerate(document):
        content = page.get_text()
        assert expected[index] in content, (index, content)
        assert '\ufffd' not in content
        path = preview / f'page_{index + 1:02d}.png'
        page.get_pixmap(matrix=pymupdf.Matrix(1.7, 1.7)).save(path)
        im = Image.open(path).convert('RGB')
        im.thumbnail((510, 720))
        thumbs.append(im)
    sheet = Image.new('RGB', (1590, 765), '#eeeeee')
    draw = ImageDraw.Draw(sheet)
    for index, im in enumerate(thumbs):
        x = 10 + 530 * index
        sheet.paste(im, (x, 30))
        draw.text((x, 10), f'Page {index + 1}', fill='black')
    sheet.save(preview / 'contact_sheet.png')
    log = (OUT / 'latex_compilation.log').read_text(encoding='utf-8')
    assert not re.search(r'Missing character|Undefined control sequence|Overfull', log)
    with zipfile.ZipFile(word) as package:
        from lxml import etree
        xml = etree.fromstring(package.read('word/document.xml'))
        assert len(xml.findall('.//w:drawing', methods.NS)) == 3
        assert len(xml.findall('.//w:tbl', methods.NS)) == 3
        assert len(xml.findall('.//w:br[@w:type="page"]', methods.NS)) == 2
    result = {'source': str(SOURCE), 'word': str(word), 'pdf': str(pdf),
              'pdf_pages': 3, 'figures': 3, 'tables': 3, 'word_rendered': False,
              'uses_verified_existing_figures': True, 'evidence_checks': evidence,
              'sha256': {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                         for p in (SOURCE, word, pdf)}}
    (OUT / 'document_audit.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    evidence = check_results()
    blocks = source_blocks()
    word = word_export(blocks)
    pdf = pdf_export(blocks)
    audit(word, pdf, evidence)


if __name__ == '__main__':
    main()
