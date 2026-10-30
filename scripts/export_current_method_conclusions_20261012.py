#!/usr/bin/env python3
"""Export a same-source conclusion report to Word/PDF and bundle editable figures."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from PIL import Image

import recompile_methods_latex_20261012 as methods

ROOT = Path(__file__).resolve().parents[1]
NAME = '当前方法实证结论与结果图_20261012'
SOURCE = ROOT / 'docs' / (NAME + '.md')
OUT = ROOT / 'artifacts/current_method_conclusions_20261012'
LATEX = OUT / 'latex'
REFERENCE = ROOT / 'docs/论文方法章节_视觉状态驱动模型测试_说明完善版_20261012.docx'


def inline_tex(text):
    text = re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', text)
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    result = []
    for part in parts:
        if not part:
            continue
        value = part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part
        escaped = methods.escape_text(value).replace(r'\_', r'\_\allowbreak{}')
        # File identifiers in the evidence table may wrap at slash boundaries.
        escaped = escaped.replace('/', r'/\allowbreak{}')
        if part.startswith('**'):
            escaped = r'\textbf{' + escaped + '}'
        result.append(escaped)
    return ''.join(result)


def parse():
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    blocks, index = [], 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells = [cell.strip() for cell in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                index += 1
            blocks.append(('table', rows))
            continue
        image = re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)', line)
        heading = re.fullmatch(r'(#{1,3})\s+(.+)', line)
        if image:
            blocks.append(('image', (image.group(1), (SOURCE.parent / image.group(2)).resolve())))
        elif heading:
            blocks.append(('heading' + str(len(heading.group(1))), heading.group(2)))
        elif re.match(r'^\d+\. ', line):
            blocks.append(('item', re.sub(r'^\d+\. ', '', line)))
        else:
            blocks.append(('body', line))
        index += 1
    return blocks


def word_export():
    output = SOURCE.with_suffix('.docx')
    process = subprocess.run([str(methods.PANDOC), str(SOURCE), '--from=markdown',
                              '--to=docx', '--resource-path=' + str(SOURCE.parent),
                              '--reference-doc=' + str(REFERENCE), '--output=' + str(output)],
                             text=True, capture_output=True, check=True)
    (OUT / 'word_export.log').write_text(process.stderr, encoding='utf-8')
    assert not process.stderr.strip(), process.stderr
    document = Document(output)
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2.1)
    for name, size in [('Normal', 11), ('Title', 18), ('Heading 1', 15),
                       ('Heading 2', 13), ('Heading 3', 11.5), ('Caption', 9)]:
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name in ('Title', 'Heading 1', 'Heading 2', 'Heading 3')
        style._element.get_or_add_rPr().get_or_add_rFonts().set(
            qn('w:eastAsia'), '黑体' if style.font.bold else '宋体')
    for paragraph in document.paragraphs:
        if paragraph.text.startswith('视觉状态驱动模型测试：'):
            paragraph.style = document.styles['Title']
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        if paragraph.style.name in ('Normal', 'Body Text', 'First Paragraph'):
            paragraph.paragraph_format.line_spacing = 1.25
        if paragraph._p.findall('.//' + qn('w:drawing')):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.startswith('图') and '说明：' in paragraph.text:
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.15
            for run in paragraph.runs:
                run.font.size = Pt(9)
    for image in document.inline_shapes:
        ratio = image.height / image.width
        image.width, image.height = Cm(17), int(Cm(17) * ratio)
    for table in document.tables:
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                        run.bold = row_index == 0
    document.save(output)
    return output


def latex_export(blocks):
    (LATEX / 'fonts').mkdir(parents=True, exist_ok=True)
    # Same portable font files as the latest methods chapter.
    fonts = ROOT / 'artifacts/methods_explanations_overlay_20261012/latex/fonts'
    for path in fonts.iterdir():
        if path.is_file():
            shutil.copy2(path, LATEX / 'fonts' / path.name)
    parts = [methods.PREAMBLE.replace(r'\begin{document}',
                                      r'\usepackage{enumitem}' + '\n' + r'\begin{document}')]
    in_items = False
    figure_number = 0
    for index, (kind, value) in enumerate(blocks):
        if in_items and kind != 'item':
            parts.append(r'\end{enumerate}')
            in_items = False
        if kind == 'heading1':
            parts.append(r'{\Large\sffamily\bfseries ' + inline_tex(value) + r'\par}\vspace{8pt}')
        elif kind in ('heading2', 'heading3'):
            command = r'\section*' if kind == 'heading2' else r'\subsection*'
            space = r'4\baselineskip'
            following = blocks[index + 1:index + 4]
            # Keep a section title with its immediately introduced figure/table.
            if following and following[0][0] == 'table':
                space = r'8\baselineskip'
            for distance, (next_kind, next_value) in enumerate(following):
                if next_kind == 'image':
                    with Image.open(next_value[1]) as im:
                        width, height = im.size
                    space = str(round(483.78 * height / width + 100 + distance * 40)) + 'pt'
                    break
                if next_kind not in ('heading2', 'heading3', 'body'):
                    break
            parts.append(r'\needspace{' + space + '}' + command + '{' + inline_tex(value) + '}')
        elif kind == 'item':
            if not in_items:
                parts.append(r'\begin{enumerate}[leftmargin=2em,itemsep=2pt,topsep=3pt]')
                in_items = True
            parts.append(r'\item ' + inline_tex(value))
        elif kind == 'image':
            figure_number += 1
            label, png = value
            assert png.is_file() and png.with_suffix('.pdf').is_file()
            target = LATEX / f'figure_{figure_number}.pdf'
            shutil.copy2(png.with_suffix('.pdf'), target)
            parts.append(r'\needspace{9\baselineskip}\begin{center}\begin{minipage}{\linewidth}\centering '
                         + r'\includegraphics[width=\linewidth,height=.48\textheight,keepaspectratio]{'
                         + target.name + r'}\par\vspace{3pt}{\small ' + inline_tex(label)
                         + r'\par}\end{minipage}\end{center}')
        elif kind == 'table':
            n = len(value[0])
            if n == 2:
                weights = [.30, .70]
            elif n == 3:
                weights = [.20, .33, .47]
            elif n == 5:
                weights = [.19, .17, .23, .23, .18]
            elif n == 6:
                weights = [.19, .13, .17, .18, .15, .18]
            elif n == 4:
                weights = [.22, .25, .25, .28]
            else:
                raise ValueError(n)
            columns = '@{}' + ''.join(
                r'>{\raggedright\arraybackslash}p{\dimexpr ' + str(w)
                + r'\linewidth-' + str((n - 1) * 2 / n) + r'\tabcolsep\relax}'
                for w in weights) + '@{}'
            parts.append(r'{\footnotesize\renewcommand{\arraystretch}{1.25}'
                         + r'\begin{longtable}{' + columns + r'}\toprule')
            header = ' & '.join(r'\textbf{' + inline_tex(cell) + '}' for cell in value[0])
            parts.append(header + r'\\\midrule\endfirsthead')
            parts.append(header + r'\\\midrule\endhead')
            for row in value[1:]:
                assert len(row) == n
                parts.append(' & '.join(inline_tex(cell) for cell in row) + r'\\')
            parts.append(r'\bottomrule\end{longtable}}')
        else:
            if value.startswith('图') and '说明：' in value:
                parts.append(r'{\small\noindent ' + inline_tex(value) + r'\par}')
            else:
                parts.append(inline_tex(value) + '\n')
    if in_items:
        parts.append(r'\end{enumerate}')
    parts.append(r'\end{document}')
    source = LATEX / 'main.tex'
    source.write_text('\n\n'.join(parts) + '\n', encoding='utf-8')
    methods.SOURCE, methods.OUT, methods.NAME = SOURCE, OUT, NAME
    return methods.compile_latex(source)


def preview_and_package(word, pdf):
    sys.path.insert(0, str(ROOT / 'artifacts/teacher_methodology_20261011/vendor'))
    import pymupdf
    from PIL import Image, ImageDraw
    preview = OUT / 'preview'
    preview.mkdir(exist_ok=True)
    document = pymupdf.open(pdf)
    images = []
    for index, page in enumerate(document):
        text = page.get_text()
        assert '\ufffd' not in text
        image_path = preview / f'page_{index + 1:02d}.png'
        page.get_pixmap(matrix=pymupdf.Matrix(1.6, 1.6)).save(image_path)
        thumb = Image.open(image_path).convert('RGB')
        thumb.thumbnail((390, 560))
        images.append(thumb)
    contact = Image.new('RGB', (4 * 410, ((len(images) + 3) // 4) * 600), '#eeeeee')
    draw = ImageDraw.Draw(contact)
    for index, im in enumerate(images):
        x, y = (index % 4) * 410 + 10, (index // 4) * 600 + 25
        contact.paste(im, (x, y))
        draw.text((x, y - 18), f'Page {index + 1}', fill='black')
    contact.save(preview / 'contact_sheet.png')
    log = (OUT / 'latex_compilation.log').read_text(encoding='utf-8')
    assert not re.search(r'Missing character|Undefined control sequence|Overfull', log)
    with zipfile.ZipFile(word) as package:
        root = etree.fromstring(package.read('word/document.xml'))
        assert len(root.findall('.//w:drawing', methods.NS)) == 7
        assert len(root.findall('.//w:tbl', methods.NS)) == 7
    archive_path = ROOT / 'docs' / (NAME + '_图表与数据.zip')
    with zipfile.ZipFile(archive_path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for folder in (OUT / 'figures', OUT / 'tables'):
            for file in sorted(folder.iterdir()):
                if file.is_file() and file.suffix in ('.svg', '.png', '.pdf', '.csv', '.json', '.txt'):
                    archive.write(file, file.relative_to(OUT))
        for file in (ROOT / 'scripts/plot_current_method_conclusions_20261012.py',
                     ROOT / 'scripts/plot_current_method_recent_extension_20261012.py',
                     OUT / 'result_audit.json', OUT / 'README.md', word, pdf):
            archive.write(file, ('scripts/' + file.name) if file.suffix == '.py' else file.name)
        portable_text = SOURCE.read_text(encoding='utf-8').replace(
            '../artifacts/current_method_conclusions_20261012/', '').replace('../scripts/', 'scripts/')
        archive.writestr(SOURCE.name, portable_text)
    audit = {'source': str(SOURCE), 'word': str(word), 'pdf': str(pdf),
             'pages': len(document), 'figures': 7, 'tables': 7,
             'pdf_engine': 'Tectonic 0.17.0', 'no_missing_glyphs_or_overfull': True,
             'word_rendered': False, 'source_and_result_scope_reviewed': True,
             'figure_and_data_archive': str(archive_path),
             'sha256': {str(file): hashlib.sha256(file.read_bytes()).hexdigest()
                        for file in (SOURCE, word, pdf, archive_path)}}
    (OUT / 'document_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(audit, ensure_ascii=False, indent=2), flush=True)


def main():
    word = word_export()
    pdf = latex_export(parse())
    preview_and_package(word, pdf)


if __name__ == '__main__':
    main()
