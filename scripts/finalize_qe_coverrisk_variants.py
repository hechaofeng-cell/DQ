"""Export independently optimized label-source variants and global ratio results."""
import json
from pathlib import Path

import numpy as np
from docx import Document

from build_qe_visualstate_coverrisk_dataset import (
    OUT, DATA, CLASSES, read_csv, write_csv, write_json, sha, measures,
)


def main():
    rows=read_csv(DATA/'manifest.csv');master={r['sample_id']:r for r in rows}
    variants=[];perclass=[];certificates=[]
    for source,root in [('local',OUT),('online',OUT/'online_analysis')]:
        arrays={c:dict(np.load(root/'prepared'/f'{c}.npz')) for c in CLASSES}
        fixed=json.loads((root/'optimization/solutions.json').read_text())
        global_results=json.loads((root/'global_ratio/solutions.json').read_text())
        selections={}
        if source=='online':
            k=fixed['compact_budget_per_class']
            for budget,items in fixed['budgets'].items():
                selections[f'online_cover_risk_n{int(budget)*6}']={c:items[c]['selected_indices'] for c in CLASSES}
            selections[f'online_cover_highrisk_n{k*6}']={c:fixed['highest_q5_under_full_coverage'][c]['selected_indices'] for c in CLASSES}
        for objective,result in global_results.items():
            if objective not in ('risk_mean','high_risk_fraction'):continue
            selections[f'{source}_global_{objective}']=result['selected_indices']
            certificates.append(dict(annotation_source=source,objective=objective,
                global_optimum_proved=result['global_optimum_proved'],
                per_class=result['budget_per_class'],result_path=str(root/'global_ratio/solutions.json')))
        for name,selection in selections.items():
            chosen=[];values=[]
            for c,indices in selection.items():
                z=arrays[c];m=measures(z,indices);values.append(m)
                perclass.append(dict(strategy=name,annotation_source=source,category=c,**m))
                for i in indices:
                    sid=str(z['sample_ids'][i]);r=dict(master[sid])
                    for key in ('score_lower','score_midpoint','score_upper','high_risk_q5','fully_known','evaluable'):
                        r['master_local_'+key]=r[key]
                    r.update(annotation_source=source,score_lower=float(z['risk_lower'][i]),
                             score_midpoint=float(z['risk_midpoint'][i]),score_upper=float(z['risk_upper'][i]),
                             high_risk_q5=int(z['high_risk'][i]),
                             fully_known=bool(np.isclose(z['risk_lower'][i],z['risk_upper'][i],atol=1e-8,rtol=0)),
                             evaluable=bool(np.isclose(z['risk_lower'][i],z['risk_upper'][i],atol=1e-8,rtol=0)))
                    chosen.append(r)
            assert len({r['sample_id'] for r in chosen})==len(chosen)
            write_csv(DATA/'variants'/f'{name}.csv',chosen)
            variants.append(dict(strategy=name,annotation_source=source,total_targets=len(chosen),per_class=values[0]['n'],
                                  **{key:float(np.mean([v[key] for v in values])) for key in values[0] if key not in ('n','required_states','reached_states')},
                                  manifest=str(DATA/'variants'/f'{name}.csv')))
    write_csv(OUT/'variant_metrics_by_category.csv',perclass)
    write_csv(OUT/'variant_comparison.csv',variants)
    core=read_csv(OUT/'strategy_comparison.csv')
    merged=[dict(r,annotation_source='local',manifest=str(DATA/'subsets'/f"{r['strategy']}.csv") if r['strategy']!='random_mean_1000' else '') for r in core]+variants
    write_csv(OUT/'merged_strategy_comparison.csv',merged)
    pool=read_csv(OUT/'pool_ceiling_by_category.csv')
    gaps=read_csv(OUT/'pool_state_support.csv')
    appendix=['','## 双标注来源与全预算最高风险核验','',
        '两种标注来源分别评价，不把两种自动标注的状态并集当作真实覆盖。来源差异本身说明最高覆盖受标注器影响，不能据此判断线上标注更准确。','',
        '| 标注来源 | 固定字典出现覆盖上限 | 每状态20例支持完成度上限 | 五维计分核心状态出现覆盖上限 |',
        '|---|---:|---:|---:|']
    for source in ('local','online'):
        selected=[r for r in pool if r['backend']==source]
        appendix.append(f"| {source} | {100*np.mean([float(r['presence_coverage']) for r in selected]):.4f}% | {100*np.mean([float(r['support_completion']) for r in selected]):.4f}% | {sum(int(r['core_reached_states']) for r in selected)}/{sum(int(r['core_required_states']) for r in selected)} |")
    appendix+=['','五维计分核心清单只含目标尺度、整体可见性、遮挡程度、截断程度、身体朝向，共每类22状态；它小于包含关联条件的完整71状态清单。','',
        'online仍缺失两个状态：'+ '；'.join(f"{r['category']}/{r['field']}/{r['state']}" for r in gaps if r['backend']=='online' and r['missing']=='True')+'。','',
        '全预算优化同时要求六类数量相同、各来源全部可达状态至少一例；每类数量允许从已证明最小共同预算12变化到1300。整数分式规划分别优化平均风险下界与Q5占比，保留各迭代MILP状态、目标界与最终精确残差；这两个目标对应不同最优测试集。','',
        '| 来源／目标 | 目标数 | 固定字典覆盖 | 平均风险下界 | Q5占比 | 全局最优已证明 |',
        '|---|---:|---:|---:|---:|---|']
    for cert in certificates:
        name=f"{cert['annotation_source']}_global_{cert['objective']}"
        r=next(v for v in variants if v['strategy']==name)
        appendix.append(f"| {name} | {r['total_targets']} | {100*r['presence_coverage']:.4f}% | {r['risk_lower']:.4f}/100 | {100*r['high_risk_q5_fraction']:.4f}% | {cert['global_optimum_proved']} |")
    appendix+=['','最高Q5比例是相对风险层的样本构成，非分类错误概率。最大平均风险下界也只是既定视觉状态评分的最优值，不表示真实风险已校准。最高Q5目标没有继续优化平均风险，同分最优解的平均风险可不同。上述全局性只对当前字典、标注、均衡约束和可达状态至少1例成立，不对所有可能图像、未标注细节或每状态20例约束成立。','',
        'variants/保存online与全预算变体清单。变体明确annotation_source，其score与Q5列使用对应来源；master_local_*列保留主清单原值。全部变体复用同一套实体图像。主比较仍使用local，最新60个细粒度记录仍是辅助标注，不混入主覆盖优化。']
    path=OUT/'report.md';original=path.read_text()
    marker='\n## 双标注来源与全预算最高风险核验'
    if marker in original:original=original.split(marker)[0]
    text=original.rstrip()+'\n'+'\n'.join(appendix)+'\n';path.write_text(text)
    (DATA/'README.md').write_text(text)
    doc=Document();lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            cells=[[v.strip() for v in row.strip('|').split('|')] for row in block]
            table=doc.add_table(rows=1,cols=len(cells[0]));table.style='Table Grid'
            for cell,value in zip(table.rows[0].cells,cells[0]):cell.text=value
            for row in cells[2:]:
                for cell,value in zip(table.add_row().cells,row):cell.text=value
            continue
        if line.startswith('#'):doc.add_heading(line.lstrip('#').strip(),level=min(len(line)-len(line.lstrip('#')),3)-1)
        elif line:doc.add_paragraph(line)
        i+=1
    doc.save(OUT/'report.docx')
    write_json(OUT/'global_optimization_certificates.json',certificates)
    paths=[Path(__file__),OUT/'report.md',OUT/'report.docx',OUT/'merged_strategy_comparison.csv',OUT/'variant_metrics_by_category.csv',OUT/'variant_comparison.csv',
           ROOT_SCRIPT_DIR/'optimize_qe_coverrisk.py',ROOT_SCRIPT_DIR/'optimize_qe_coverrisk_global_ratio.py']
    paths += [p for source in (OUT,OUT/'online_analysis') for p in [source/'optimization/solutions.json',source/'global_ratio/solutions.json']]
    paths += list((DATA/'variants').glob('*.csv'))
    write_json(OUT/'final_result_freeze.sha256.json',{str(p):sha(p) for p in paths})
    print(json.dumps(dict(status='exported_variants',variants=len(variants),global_certificates=certificates),ensure_ascii=False,indent=2))


ROOT_SCRIPT_DIR=Path(__file__).resolve().parent
if __name__=='__main__':main()
