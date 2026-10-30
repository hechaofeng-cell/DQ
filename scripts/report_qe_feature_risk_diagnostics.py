"""Export usable stratified manifests and explain the new pool's scoring pipeline."""
from pathlib import Path
import csv
import hashlib
import json
from collections import Counter

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data/QE_VisualState_CoverRisk_v1_20261013'
FULL = ROOT / 'artifacts/QE_ThreeStrategy_Full_20261013'
OUT = ROOT / 'artifacts/QE_FeatureRisk_Diagnostics_20261013'
CLASSES = ('dog', 'cat', 'horse', 'sheep', 'person', 'car')
DIMS = ('scale', 'visibility', 'occlusion', 'truncation', 'orientation')
CN = dict(zip(DIMS, ('尺度', '可见性', '遮挡', '截断', '朝向')))
NAMES = {'coverage': '覆盖选样', 'risk': '风险分层随机', 'risk_coverage': '风险＋覆盖'}


def read(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def write(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def table(headers, rows):
    return ['| ' + ' | '.join(map(str, headers)) + ' |',
            '| ' + ' | '.join('---' for _ in headers) + ' |'] + [
            '| ' + ' | '.join(map(str, r)) + ' |' for r in rows]


def main():
    scores = read(OUT / 'target_scores.csv')
    lookup = {(r['annotation_backend'], r['sample_id']): r for r in scores}
    manifest = read(DATA / 'manifest.csv')
    master = {r['sample_id']: r for r in manifest}
    assert len(master) == 7800 and len(lookup) == 15600
    native = {}
    for r in manifest:
        ann = json.loads((DATA / r['annotation_path']).read_text())
        for backend in ('local', 'online'):
            native[backend, r['sample_id']] = ann[backend]['native']['features']

    def enriched(backend, sid):
        original = master[sid]
        score = lookup[backend, sid]
        row = {k: original[k] for k in ('sample_id', 'class_name', 'label', 'source_pool', 'bbox_area_ratio', 'bbox_norm_xyxy')}
        for k in ('image_path', 'crop_path', 'marked_image_path', 'feature_crop_path', 'annotation_path'):
            row[k] = str(DATA / original[k])
        row.update(score)
        # The master manifest holds local scores; replace legacy score aliases for this backend.
        for k in ('lower', 'midpoint', 'upper'):
            row['score_' + k] = score['global_risk_' + k]
        row['high_risk_q5'] = int(score['risk_tier'] == 'Q5')
        for k, value in native[backend, sid].items():
            row['feature_' + k] = json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value
        return row

    exports = []
    for backend in ('local', 'online'):
        for tier in range(1, 6):
            ids = [r['sample_id'] for r in scores if r['annotation_backend'] == backend and r['risk_tier_integer'] == str(tier)]
            assert len(ids) == len(set(ids)) == 1560
            assert Counter(master[s]['class_name'] for s in ids) == Counter(dict.fromkeys(CLASSES, 260))
            path = OUT / 'stratified_manifests' / f'{backend}_Q{tier}.csv'
            write(path, [enriched(backend, sid) for sid in ids])
            exports.append(dict(path=str(path), backend=backend, kind='whole_pool_risk_layer', risk_layer=tier, targets=len(ids), random_repetition=''))
        arrays = {c: np.load((ROOT / 'artifacts/QE_VisualState_CoverRisk_v1_20261013' / ('prepared' if backend == 'local' else 'online_analysis/prepared') / f'{c}.npz'), allow_pickle=False) for c in CLASSES}
        selection = np.load(FULL / 'selection' / f'{backend}_n1200.npz', allow_pickle=False)
        for method in NAMES:
            indices = selection[method][0] if selection[method].ndim == 3 else selection[method]
            ids = [str(arrays[c]['sample_ids'][i]) for ci, c in enumerate(CLASSES) for i in indices[ci]]
            assert len(ids) == len(set(ids)) == 1200
            path = OUT / 'selected_scored_manifests' / f'{backend}_{method}_n1200.csv'
            write(path, [enriched(backend, sid) for sid in ids])
            exports.append(dict(path=str(path), backend=backend, kind='selected_scored_example', method=method, targets=1200, random_repetition='0' if method == 'risk' else ''))
    write(OUT / 'manifest_index.csv', exports)

    summary = read(FULL / 'summary.csv')
    dimensions = read(OUT / 'dimension_score_summary.csv')
    strata = read(OUT / 'risk_strata_summary.csv')
    contributions = read(OUT / 'selection_dimension_score_totals.csv')
    errors = read(OUT / 'risk_tier_error_relationship.csv')
    parts = read(OUT / 'equal_parts_sensitivity_summary.csv')
    inventory = json.loads((OUT / 'inventory_audit.json').read_text())
    lines = ['# QE新数据集：特征标注、评分、分层与三策略实验结果', '',
             '对象是QE合并数据集的7800个目标，犬、猫、马、羊、人、车各1300。每个目标分别保留local、online原生标注；三种方法实际在该池内选样并评价24组冻结分类预测。本次补充计算与导出逐目标、逐特征、五维与风险层结果，没有新增VLM标注、模型推理、训练或微调。', '',
             '## 1. 每个目标是怎样得到评分的', '',
             'VLM输出可见状态和文字证据；尺度由官方目标框面积占原图比例计算。程序按冻结映射将有序状态转换成0—100的风险代理分，五维各乘0.2，总风险R＝0.2×(尺度＋可见性＋遮挡＋截断＋朝向)，每维最多贡献20分。该分数不是模型失败概率。', '']
    lines += table(['维度', '当前主评分字段', '状态到风险分的映射'], [
        ['尺度', 'bbox_area_ratio', '≥10%：0；3%—10%：33.33；1%—3%：66.67；<1%：100（左闭右开）'],
        ['可见性', 'body_visibility', 'full/most：0；partial：66.67；minimal：100'],
        ['遮挡', 'occlusion_level', 'none/light/moderate/heavy：0/33.33/66.67/100'],
        ['截断', 'truncation_level', 'none/light/moderate/heavy：0/25/50/100'],
        ['朝向', 'body_orientation', '正面/左右侧/three_quarter：0；背面：50；俯视或特殊：100']])
    lines += ['', '未知状态保留该维[0,100]区间，中点50仅用于排序，总分同时保存下界、中点和上界。它没有被当成确定标签或期望真实风险。', '']
    example = next(r for r in scores if r['annotation_backend'] == 'local' and r['sample_id'] == manifest[0]['sample_id'])
    lines += [f"真实示例 {example['sample_id']}：五维分为" + '、'.join(f"{CN[d]}{float(example[d+'_risk_midpoint_0_100']):.2f}" for d in DIMS) + f"；加权总分{float(example['global_risk_midpoint']):.2f}，归入{example['risk_tier']}。头部可见性等其他标注见逐特征表，不会自动重复加到主总分。", '',
              '## 2. 7800目标整体评分与风险层', '']
    pool_rows = []
    for backend in ('local', 'online'):
        row = next(r for r in summary if r['source'] == backend and r['total_targets'] == '7800' and r['method'] == 'coverage')
        known = sum(r['fully_known'] == 'True' for r in scores if r['annotation_backend'] == backend)
        means = [np.mean([float(r['mean']) for r in dimensions if r['annotation_backend'] == backend and r['dimension'] == d]) for d in DIMS]
        pool_rows.append([backend] + [f'{v:.2f}' for v in means] + [f"{float(row['risk_midpoint']):.2f}", f"[{float(row['risk_lower']):.2f}, {float(row['risk_upper']):.2f}]", f'{known}/7800', f"{float(row['presence_coverage']):.2%}", f"{float(row['support_completion']):.2%}"])
    lines += table(['标注源'] + [CN[d] + '平均分' for d in DIMS] + ['总风险中点', '平均风险区间', '五维信息全已知', '状态覆盖', '20例支持完成度'], pool_rows)
    ties = sum(r['upper_boundary_tied_next_tier'] == 'True' for r in strata if r['risk_tier'] != 'Q5')
    lines += ['', f'每个标注源、每个类别按总分中点再按目标ID排序，分成Q1—Q5，每层260个目标；六类合计每层1560。48条相邻层边界中有{ties}条存在同分值跨层，所以Q1—Q5是类别内相对等人数层，不是0—20、20—40等绝对分段，也不能将Q5占比称为失败概率。', '',
              'stratified_manifests/提供十份可直接读取图像、评分与标签的分层清单，路径指向新数据集真实文件。两个来源的层清单使用各自风险评分。', '',
              '## 3. 三种方法选出来的数据有什么变化', '',
              '全部10预算点均保留。下面以1200目标、每类200的实测结果展示差异；风险随机组是1000次抽样均值，单个随机样本清单仅导出第0次并明确标识。', '']
    result_rows = []
    for backend in ('local', 'online'):
        for method in NAMES:
            r = next(r for r in summary if r['source'] == backend and r['total_targets'] == '1200' and r['method'] == method)
            result_rows.append([backend, NAMES[method], f"{float(r['presence_coverage']):.2%}", f"{float(r['support_completion']):.2%}", f"{float(r['risk_midpoint']):.2f}", f"[{float(r['risk_lower']):.2f}, {float(r['risk_upper']):.2f}]", f"{float(r['q5_fraction']):.2%}", f"{float(r['risk_unknown_fraction']):.2%}", f"{float(r['macro_f1']):.2%}"])
    lines += table(['标注源', '方法', '状态覆盖', '支持完成度', '总风险中点', '风险区间均值', 'Q5占比', '目标评分有未知比例', 'Macro-F1'], result_rows)
    contribution_rows = []
    for backend in ('local', 'online'):
        for method in NAMES:
            contribution_rows.append([backend, NAMES[method]] + [f"{float(next(r for r in contributions if r['backend'] == backend and r['total_targets'] == '1200' and r['method'] == method and r['dimension'] == d)['midpoint_points_mean']):.3f}" for d in DIMS])
    lines += ['', '对应五维对总风险的平均贡献，每列0—20：', ''] + table(['标注源', '方法'] + [CN[d] for d in DIMS], contribution_rows)
    lines += ['', '风险随机与组合组每层均240目标，总风险均值与未知比例精确匹配，但各维贡献仍可能不同；这不是逐维匹配。覆盖组没有风险配额，因而会改变风险构成。完整各类别、各层、各字段支持数见selection_*表。全池7800终点三种方法全部数据及结果一致，选样优势体现在子集预算。', '',
              '## 4. 风险层是否对应更高的分类错误', '']
    error_rows = []
    for backend in ('local', 'online'):
        error_rows.append([backend] + [f"{float(next(r for r in errors if r['backend'] == backend and r['category'] == 'ALL_CLASSES' and r['risk_layer'] == str(q))['mean_error_rate']):.2%}" for q in range(1, 6)])
    lines += table(['标注源'] + [f'Q{q}平均错误率' for q in range(1, 6)], error_rows)
    lines += ['', '上述错误率由官方类别真值和24组冻结预测独立计算，没有参与风险分定义。整体呈上升趋势，单个类别不保证单调；各层跨数据源组成不同。这支持在本池中用该风险代理组织条件分析，不构成校准概率或因果结论。扩展三策略18项检验受8架构双侧检验分辨率限制，未得到校正显著性；详细统计保留在全量实验报告。', '',
              '## 5. 头部、腿部、毛发、头发等细项到底有多少', '',
              '全池非车目标各有19个原生字段，车各18个；另有官方几何尺度。头部可见性、腿部可见性、头部朝向等是非车原生字段，可按既有映射导出评分；本轮主equal_global仅用整体可见性和身体朝向，部位分不会自动进入主风险。', '',
              '另已计算equal_parts敏感性：非车可见性由头部与腿部各占一半、朝向由身体与头部各占一半；车保留整体方案。该列使用原生字段，未重做三种选样，不能据此宣称部位方案更好。', '']
    lines += table(['标注源', '类别', '整体总风险均值', '部位方案均值', '分数变化目标数'], [[r['annotation_backend'], r['class_name'], f"{float(r['global_mean']):.2f}", f"{float(r['equal_parts_mean']):.2f}", r['changed_targets']] for r in parts])
    fine_scores = read(OUT / 'fine_score_summary.csv')
    lines += ['', '新增丰富细项的独立标注只有60目标、每类10个。完整字段数量与细项评分见fine_feature_inventory.csv、fine_feature_annotations.csv及fine_score_summary.csv。头发、毛色、毛纹、嘴部等新增字段没有全池标注；现有数值不能代表7800目标的全量细粒度评分。颜色、毛纹、发型等名义属性保留状态，不凭空赋予高低风险。部位不可适用、不可观察与真正被遮挡应按对应原协议处理。', '',
              '下表保留原60例细项方案的评分；均值包含保留区间的中点，不能与全池整体方案直接解释为模型改善或风险降低。', '']
    lines += table(['类别', '每目标细项数', '已标目标/该类全池', '可评价目标数', '10例总风险中点均值', '可见性平均贡献0—20'], [[r['class_name'], r['fine_fields_per_target'], f"{r['targets']}/{r['pool_targets']}", r['evaluable_targets'], f"{float(r['mean_score_midpoint']):.2f}", f"{float(r['visibility_mean_contribution_points']):.2f}"] for r in fine_scores])
    lines += ['', f'库存核查与标注真实性记录：inventory_audit.json；细评分汇总共{len(fine_scores)}行。早期596犬样本与新池的匹配范围见该审计，未跨数据集冒接细标签。', '',
              '## 6. 数据文件与复现', '',
              'target_scores.csv：每个目标双来源总分、五维状态/区间/贡献与相对层。feature_scores_long.csv：309400条原生字段＋尺度值、文字证据、映射、是否参与主评分；无有效有序映射的字段分数留空。risk_strata_summary.csv：每类五层实际分数范围、同分边界与未知数。feature_state_counts.csv：固定状态缺口及原生字段计数。selection_feature_states.csv：三种方法选出的状态支持、零支持概率与覆盖。selected_scored_manifests/：1200目标三方法的实际目标、图片路径、原生特征、正确标注来源评分；risk清单为第0次抽样。', '',
              '评分和特征标签是已有自动标注，未做全量人工语义核验；序数映射及五维权重是冻结方案，尚不能作为已验证的最优特征权重。完整图像或多维联合特征空间的覆盖不由426个单字段状态覆盖替代。']
    (OUT / 'report.md').write_text('\n'.join(lines) + '\n')
    fig, axes = plt.subplots(1, 2, figsize=(12, 4), layout='constrained')
    colors = {'local': '#2166ac', 'online': '#e08214'}
    for backend in ('local', 'online'):
        e = [next(r for r in errors if r['backend'] == backend and r['category'] == 'ALL_CLASSES' and r['risk_layer'] == str(q)) for q in range(1, 6)]
        axes[0].plot(range(1, 6), [100 * float(r['mean_error_rate']) for r in e], marker='o', label=backend, color=colors[backend])
    axes[0].set(xlabel='Relative risk layer', ylabel='Mean classifier error rate (%)', xticks=range(1, 6))
    for offset, method in enumerate(NAMES):
        rows = [r for r in contributions if r['backend'] == 'local' and r['total_targets'] == '1200' and r['method'] == method]
        values = [float(next(r for r in rows if r['dimension'] == d)['midpoint_points_mean']) for d in DIMS]
        axes[1].bar(np.arange(5) + (offset - 1) * .25, values, .25, label=method)
    axes[1].set(xticks=np.arange(5), xticklabels=DIMS, ylabel='Mean contribution (0–20)', title='Local labels, 1,200 selected targets')
    for ax in axes:
        ax.legend(fontsize=8); ax.grid(axis='y', alpha=.2); ax.spines[['top', 'right']].set_visible(False)
    (OUT / 'figures').mkdir(exist_ok=True)
    for ext in ('png', 'svg'):
        fig.savefig(OUT / 'figures' / f'scoring_and_risk_layers.{ext}', dpi=180)
    plt.close(fig)
    doc = Document()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                block.append(lines[i]); i += 1
            cells = [[v.strip() for v in x.strip('|').split('|')] for x in block]
            t = doc.add_table(rows=1, cols=len(cells[0])); t.style = 'Table Grid'
            for c, val in zip(t.rows[0].cells, cells[0]): c.text = val
            for row in cells[2:]:
                for c, val in zip(t.add_row().cells, row): c.text = val
            continue
        if line.startswith('#'):
            doc.add_heading(line.lstrip('#').strip(), level=min(3, len(line) - len(line.lstrip('#'))) - 1)
        elif line:
            doc.add_paragraph(line)
        i += 1
    doc.add_picture(str(OUT / 'figures/scoring_and_risk_layers.png'), width=Inches(6.4))
    doc.save(OUT / 'report.docx')
    files = [Path(__file__).resolve(), OUT / 'report.md', OUT / 'report.docx', OUT / 'manifest_index.csv'] + [Path(r['path']) for r in exports] + list((OUT / 'figures').glob('*'))
    dependencies = [OUT / n for n in ('target_scores.csv', 'dimension_score_summary.csv', 'risk_strata_summary.csv', 'selection_dimension_score_totals.csv', 'risk_tier_error_relationship.csv', 'equal_parts_sensitivity_summary.csv', 'inventory_audit.json', 'fine_score_summary.csv')] + [FULL / 'summary.csv', DATA / 'manifest.csv']
    (OUT / 'report_freeze.sha256.json').write_text(json.dumps({str(p): digest(p) for p in files + dependencies}, ensure_ascii=False, indent=2) + '\n')
    audit = dict(status='pass', scope='7800 target pool, reused labels and frozen classifier predictions', stratified_manifests=10, targets_per_layer=1560, selected_scored_examples=6, example_targets=1200, boundary_ties=ties, report_tables=sum(x.startswith('| ---') for x in lines), unknown_intervals_retained=True, source_correct_score_fields=True, inventory_reference=inventory)
    (OUT / 'report_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: v for k, v in audit.items() if k != 'inventory_reference'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
