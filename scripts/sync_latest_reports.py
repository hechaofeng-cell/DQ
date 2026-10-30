#!/usr/bin/env python3
"""Synchronize the bounded evidence narrative, companion reports and package."""
import json
from pathlib import Path
import posixpath
import re
import shutil
from zipfile import ZipFile

from docx import Document
import pandas as pd

import build_journal_revision as base
import build_submission_closeout as manuscript
import package_submission_closeout as package

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
OUT=ROOT/'artifacts/submission_closeout_20261007'
POSITIVES=DOCS/'优点_20261007.md'
LIMITS=DOCS/'局限性_20261007.md'


def verify_reports():
    positive=base.table_blocks(POSITIVES.read_text())
    limits=base.table_blocks(LIMITS.read_text())
    assert len(positive)==4 and len(limits)==4
    random=pd.read_csv(ROOT/'artifacts/random_stratification_ablation_20261006/tables/random_summary.csv')
    panel=pd.read_csv(ROOT/'artifacts/controlled_risk_distribution_20261005/tables/panel_summary.csv')
    contrasts=pd.read_csv(ROOT/'artifacts/controlled_risk_distribution_20261005/tables/distribution_contrasts.csv')
    scale=pd.read_csv(manuscript.EXP/'tables/summary.csv')
    cond=pd.read_csv(manuscript.EXP/'tables/area_condition_summary.csv')
    checks=0
    def close(value,expected,decimals=2):
        nonlocal checks
        base.close(value,expected,decimals);checks+=1
    def interval(value,lo,hi):
        for actual,expected in zip(value.strip('[]').split(','),[lo,hi]):close(actual,100*expected)
    for row,r in zip(base.cells(positive[1])[1:],random.itertuples()):
        assert row[0].split('/')[0]==r.pool
        close(row[2],-100*r.observed_delta)
        interval(row[3],r.random_sampling_low,r.random_sampling_high)
        close(row[4],r.holm_six_p,5)
    for row,r in zip(base.cells(positive[2])[1:],scale.itertuples()):
        assert row[0]==r.pool
        p=panel[(panel.pool==r.pool)&(panel.source==r.source)&(panel.total==120)].set_index('high_risk_percent')
        c=contrasts[(contrasts.pool==r.pool)&(contrasts.source==r.source)&(contrasts.total==120)&(contrasts.model=='panel_mean')].iloc[0]
        close(row[2],100*p.loc[0].macro_f1_mean);close(row[3],100*p.loc[100].macro_f1_mean)
        close(row[4],100*c.macro_f1_delta)
        interval(row[5],c.macro_f1_sampling_low,c.macro_f1_sampling_high)
    for row,r in zip(base.cells(positive[3])[1:],scale.itertuples()):
        d=cond[(cond.pool==r.pool)&(cond.source==r.source)];valid=d[d.evaluable]
        close(row[2],len(d),0);close(row[3],len(valid),0)
        close(row[4],(valid.area_adjusted_recall_gap<0).sum(),0)
        close(row[5],100*valid.area_adjusted_recall_gap.mean())
    for row,r in zip(base.cells(limits[1])[1:],scale.itertuples()):
        close(row[2],100*r.full_gap);close(row[3],100*r.area_gap)
        close(row[4],100*r.increment_mean);close(row[5],r.holm_six_p,5)
    paper_tables=base.table_blocks(manuscript.ZH.read_text())
    for report,official in zip(base.cells(limits[2])[1:],base.cells(paper_tables[11])[1:]):
        assert report[:2]==official[:2]
        for value,expected,decimals in zip(report[2:],official[3:],[2,2,3]):close(value,base.number(expected),decimals)
    assert '119个可评价组合有109个负差' in POSITIVES.read_text()
    assert '122个有111个负差' in POSITIVES.read_text()
    replay=json.loads((OUT/'portable_replay_audit.json').read_text())
    assert replay['status']=='pass' and len(replay['all_scale_tables_match'])==8
    assert replay['historical_selection_contrasts_checked']==60
    return checks


def main():
    checked=verify_reports()
    df,cond,_,_=manuscript.scale_tables()
    manuscript.supplement(df,cond)
    manuscript_records=manuscript.render_documents()
    report_records=[]
    for source,stable in [(POSITIVES,'优点'),(LIMITS,'局限性')]:
        record=base.word(source,False)
        report_records.append(record)
        doc=Document(source.with_suffix('.docx'))
        expected_images=len(re.findall(r'^!\[',source.read_text(),re.M))
        assert len(doc.tables)==4 and len(doc.inline_shapes)==expected_images
        assert '修订版v2' in '\n'.join(p.text for p in doc.paragraphs)
        for suffix in ['.md','.docx']:
            shutil.copy2(source.with_suffix(suffix),DOCS/(stable+suffix))
    # A convenient latest alias does not replace the date-versioned manuscript.
    for suffix in ['.md','.docx']:
        shutil.copy2(manuscript.ZH.with_suffix(suffix),DOCS/('论文_最新'+suffix))
    audit=package.audit()
    result=dict(status='pass',revision='2026-10-07 v2',new_experiments=False,
                report_numeric_cells_checked=checked,report_documents=report_records,
                manuscript_documents=manuscript_records,
                manuscript_numeric_cells_checked=audit['old_numeric_cells']+audit['new_numeric_cells'],
                bilingual_numeric_tables_match=True,
                reused_empirical_figures=True,original_inputs_unchanged=audit['original_inputs_unchanged'],
                portable_extraction_replay_passed=True,
                backup_directory='docs/archive/20261007_before_report_sync',
                author_facts_pending=True,full_page_rendering_verified=False)
    (OUT/'report_sync_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    package.base.word(DOCS/'投稿收口与尺度对照结果_20261007.md',False)
    package.packages(OUT/'portable')
    with ZipFile(DOCS/'visual_state_testing_submission_materials_20261007.zip') as z:
        members=set(z.namelist());linked=0
        for member in members:
            if not (member.startswith('docs/') and member.endswith('.md')):continue
            text=z.read(member).decode('utf-8')
            for target in re.findall(r'!\[.*?\]\((.*?)\)',text):
                resolved=posixpath.normpath(posixpath.join(posixpath.dirname(member),target))
                assert resolved in members,(member,target)
                linked+=1
    assert linked==20,linked
    result['packaged_markdown_image_links_verified']=linked
    for original,alias in [(POSITIVES,DOCS/'优点.md'),(LIMITS,DOCS/'局限性.md'),
                           (manuscript.ZH,DOCS/'论文_最新.md')]:
        assert original.read_bytes()==alias.read_bytes()
        assert original.with_suffix('.docx').read_bytes()==alias.with_suffix('.docx').read_bytes()
    result['latest_aliases_match']=True
    (OUT/'report_sync_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    # Refresh only the archive's audit record after the archive checks.
    with ZipFile(DOCS/'visual_state_testing_submission_materials_20261007.zip','a') as z:
        z.writestr('review/package_link_audit.json',json.dumps({
            'status':'pass','markdown_image_links':linked,'latest_aliases_match':True},indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
