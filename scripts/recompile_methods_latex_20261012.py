#!/usr/bin/env python3
"""Build real LaTeX PDF and editable Word mathematics from one manuscript source."""
from copy import deepcopy
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/methods_latex_recompiled_20261012'
NAME = '论文方法章节_视觉状态驱动模型测试_LaTeX重编译版_20261012'
SOURCE = ROOT / 'docs' / (NAME + '.md')
REFERENCE = ROOT / 'docs/论文方法章节_视觉状态驱动模型测试_公式规范版_20261012.docx'
REFERENCE_SOURCE = REFERENCE.with_suffix('.md')
PANDOC = OUT / 'vendor/pypandoc/files/pandoc'
TECTONIC = ROOT / 'artifacts/journal_revision_20261006/vendor/tectonic'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
MATH_PATTERN = re.compile(r'(\$[^$]+\$)')


def blocks(path):
    lines = path.read_text(encoding='utf-8').splitlines()
    result, index = [], 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == '$$':
            content = []
            index += 1
            while lines[index].strip() != '$$':
                content.append(lines[index])
                index += 1
            result.append(('equation', '\n'.join(content)))
        elif line == '```algorithm':
            index += 1
            while lines[index].strip() != '```':
                result.append(('algorithm', lines[index].rstrip()))
                index += 1
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells = [s.strip() for s in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(':?-+:?', s) for s in cells):
                    rows.append(cells)
                index += 1
            result.append(('table', rows))
            continue
        elif line.startswith('!['):
            result.append(('image', re.fullmatch(r'!\[[^]]*\]\(([^)]+)\)', line)[1]))
        elif line.startswith('# '):
            result.append(('title', line[2:]))
        elif line.startswith('## '):
            result.append(('heading', line[3:]))
        elif line.startswith('### '):
            result.append(('subheading', line[4:]))
        elif re.match(r'^(?:图|表|算法)\d+(?:[ \u3000]|$)', line):
            result.append(('caption', line))
        else:
            result.append(('body', line))
        index += 1
    return result


def inline_tex(text):
    return MATH_PATTERN.findall(text)


def math_inventory(content):
    result = []
    for kind, text in content:
        if kind == 'equation':
            result.append(('display', text))
        elif kind == 'table':
            for row in text:
                for cell in row:
                    result.extend(('inline', s[1:-1]) for s in inline_tex(cell))
        elif kind != 'image':
            result.extend(('inline', s[1:-1]) for s in inline_tex(text))
    return result


def convert_math(inventory):
    # Pandoc parses actual TeX, rather than mapping equation numbers to handwritten XML.
    pool = '\n\n'.join('$$\n' + s + '\n$$' if kind == 'display' else '$' + s + '$'
                       for kind, s in inventory)
    temporary = OUT / 'all_math_pandoc.docx'
    command = [str(PANDOC), '--from=markdown+tex_math_dollars', '--to=docx',
               '--output=' + str(temporary)]
    process = subprocess.run(command, input=pool, text=True, capture_output=True, check=True)
    (OUT / 'pandoc_math_conversion.log').write_text(process.stderr, encoding='utf-8')
    if process.stderr.strip():
        raise RuntimeError('Unexpected Pandoc math-conversion warning: ' + process.stderr)
    with zipfile.ZipFile(temporary) as package:
        root = etree.fromstring(package.read('word/document.xml'))
    equations = root.findall('.//m:oMath', NS)
    assert len(equations) == len(inventory), (len(equations), len(inventory))
    converted = {}
    for (kind, source), element in zip(inventory, equations):
        key = (kind, source)
        if key in converted:
            assert etree.tostring(converted[key]) == etree.tostring(element)
        converted[key] = element
    return converted


def native_math(source, kind, converted, halfpoints=22):
    math = deepcopy(converted[(kind, source)])
    for run in math.findall('.//m:r', NS):
        properties = run.find('w:rPr', NS)
        if properties is None:
            properties = OxmlElement('w:rPr')
            run.insert(1 if run.find('m:rPr', NS) is not None else 0, properties)
        fonts = properties.find('w:rFonts', NS)
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            properties.append(fonts)
        for attribute in ('ascii', 'hAnsi'):
            fonts.set(qn('w:' + attribute), 'Cambria Math')
        size = properties.find('w:sz', NS)
        if size is None:
            size = OxmlElement('w:sz')
            properties.append(size)
        size.set(qn('w:val'), str(halfpoints))
    return math


def rewrite(paragraph, text, converted, halfpoints=22):
    sample = next((r._r.find(qn('w:rPr')) for r in paragraph.runs
                   if r._r.find(qn('w:rPr')) is not None), None)
    properties = paragraph._p.find(qn('w:pPr'))
    for child in list(paragraph._p):
        if child is not properties:
            paragraph._p.remove(child)
    for part in MATH_PATTERN.split(text):
        if part.startswith('$') and part.endswith('$'):
            paragraph._p.append(native_math(part[1:-1], 'inline', converted, halfpoints))
        elif part:
            run = paragraph.add_run(part)
            if sample is not None:
                run._r.insert(0, deepcopy(sample))


def plain(text):
    return MATH_PATTERN.sub('', text).strip()


def write_word(content, converted):
    before = blocks(REFERENCE_SOURCE)
    assert len(before) >= len(content)
    if len(content) < len(before):
        assert before[len(content)] == ('heading', '3.7 参数依据、可靠性检验与适用边界')
    document = Document(REFERENCE)
    elements = [e for e in document._element.body
                if e.tag == qn('w:tbl') or
                e.tag == qn('w:p') and
                (e.findall('.//w:t', NS) or e.findall('.//m:oMath', NS)
                 or e.findall('.//w:drawing', NS))]
    assert len(elements) == len(before), (len(elements), len(before))
    for element, (kind, old_text), (new_kind, text) in zip(elements, before, content):
        assert kind == new_kind
        if kind == 'table':
            table = Table(element, document)
            assert len(table.rows) == len(text)
            for row, new_cells, old_cells in zip(table.rows, text, old_text):
                for cell, value, old_value in zip(row.cells, new_cells, old_cells):
                    assert cell.paragraphs[0].text.strip() == plain(old_value)
                    rewrite(cell.paragraphs[0], value, converted, 19)
        elif kind == 'image':
            assert element.findall('.//w:drawing', NS)
        elif kind == 'equation':
            para = element.find('m:oMathPara', NS)
            assert para is not None
            properties = para.find('m:oMathParaPr', NS)
            for child in list(para):
                if child is not properties:
                    para.remove(child)
            para.append(native_math(text, 'display', converted))
        elif kind in ('title', 'heading', 'subheading'):
            assert Paragraph(element, document).text.strip() == text
        else:
            paragraph = Paragraph(element, document)
            assert paragraph.text.strip() == plain(old_text), (kind, paragraph.text, plain(old_text))
            rewrite(paragraph, text.strip(), converted)
    for element in elements[len(content):]:
        element.getparent().remove(element)
    path = SOURCE.with_suffix('.docx')
    document.save(path)
    return path


def escape_text(text):
    mapping = {'\\': r'\textbackslash{}', '&': r'\&', '%': r'\%', '$': r'\$',
               '#': r'\#', '_': r'\_', '{': r'\{', '}': r'\}',
               '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}
    return ''.join(mapping.get(c, c) for c in text)


def tex_inline(text):
    return ''.join(r'\(' + part[1:-1] + r'\)' if part.startswith('$') and part.endswith('$')
                   else escape_text(part) for part in MATH_PATTERN.split(text))


PREAMBLE = r'''\documentclass[11pt,a4paper]{article}
\usepackage[margin=2cm,top=2.1cm,bottom=2cm]{geometry}
\usepackage{fontspec,xeCJK}
\setmainfont[Path=fonts/,BoldFont=DejaVuSerif-Bold.ttf,ItalicFont=DejaVuSerif-Italic.ttf,BoldItalicFont=DejaVuSerif-BoldItalic.ttf]{DejaVuSerif.ttf}
\setCJKmainfont[Path=fonts/,BoldFont=NotoSansSC-Bold.ttf]{NotoSansSC-Regular.ttf}
\setCJKsansfont[Path=fonts/,BoldFont=NotoSansSC-Bold.ttf]{NotoSansSC-Regular.ttf}
\usepackage{amsmath,amssymb,graphicx,booktabs,longtable,array,needspace,setspace}
\usepackage{titlesec}
\setstretch{1.22}
\setlength{\parindent}{2em}
\setlength{\parskip}{3pt}
\setlength{\tabcolsep}{5pt}
\setlength{\emergencystretch}{2em}
\setlength{\abovedisplayskip}{9pt}
\setlength{\belowdisplayskip}{9pt}
\titleformat{\section}{\large\sffamily\bfseries}{}{0pt}{}
\titleformat{\subsection}{\normalsize\sffamily\bfseries}{}{0pt}{}
\titlespacing*{\section}{0pt}{13pt}{6pt}
\titlespacing*{\subsection}{0pt}{9pt}{4pt}
\raggedbottom
\begin{document}
'''


def write_latex(content):
    latex_dir = OUT / 'latex'
    fonts = latex_dir / 'fonts'
    fonts.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / 'scripts/assets/NotoSansSC-Regular.ttf', fonts)
    shutil.copy2(ROOT / 'output/research-figure-composer/fonts/NotoSansSC-Bold.ttf', fonts)
    for font in ('DejaVuSerif.ttf', 'DejaVuSerif-Bold.ttf', 'DejaVuSerif-Italic.ttf',
                 'DejaVuSerif-BoldItalic.ttf'):
        shutil.copy2(Path('/usr/share/fonts/truetype/dejavu') / font, fonts)
    shutil.copy2(ROOT / 'scripts/assets/NotoSansSC-LICENSE.txt', fonts)
    shutil.copy2('/usr/share/doc/fonts-dejavu-core/copyright', fonts / 'DejaVu-LICENSE.txt')
    # Preserve the exact figure embedded in the user's edited document.
    with zipfile.ZipFile(REFERENCE) as package:
        image = next(p for p in package.namelist()
                     if p.startswith('word/media/') and not p.endswith('/'))
        (latex_dir / 'figure_1.png').write_bytes(package.read(image))
    parts, equation_number, in_algorithm = [PREAMBLE], 0, False
    for index, (kind, value) in enumerate(content):
        if in_algorithm and kind != 'algorithm':
            parts.append(r'\end{minipage}\end{flushleft}')
            in_algorithm = False
        if kind == 'title':
            parts.append(r'{\Large\sffamily\bfseries ' + tex_inline(value) + r'\par}\vspace{7pt}')
        elif kind in ('heading', 'subheading'):
            parts.append(('\\section*{' if kind == 'heading' else '\\subsection*{')
                         + tex_inline(value) + '}')
        elif kind == 'image':
            parts.append(r'\begin{center}\includegraphics[width=\linewidth]{figure_1.png}\end{center}')
        elif kind == 'equation':
            equation_number += 1
            parts.append(r'\begin{equation}' + '\n' + value
                         + '\n' + r'\label{eq:' + str(equation_number) + r'}\end{equation}')
        elif kind == 'caption':
            parts.append(r'\needspace{3\baselineskip}\noindent{\small ' + tex_inline(value) + r'\par}')
        elif kind == 'table':
            width = '.36' if value[0][0] == '符号' else '.20'
            parts.append(r'{\small\renewcommand{\arraystretch}{1.25}'
                         + '\n' + r'\begin{longtable}{@{}>{\raggedright\arraybackslash}p{' + width
                         + r'\linewidth}>{\raggedright\arraybackslash}p{' + ('.60' if width == '.36' else '.76')
                         + r'\linewidth}@{}}\toprule')
            parts.append(' & '.join(r'\textbf{' + tex_inline(x) + '}' for x in value[0])
                         + r'\\\midrule\endfirsthead')
            parts.append(' & '.join(r'\textbf{' + tex_inline(x) + '}' for x in value[0])
                         + r'\\\midrule\endhead')
            for row in value[1:]:
                parts.append(' & '.join(tex_inline(x) for x in row) + r'\\')
            parts.append(r'\bottomrule\end{longtable}}')
        elif kind == 'algorithm':
            if not in_algorithm:
                parts.append(r'\begin{flushleft}\begin{minipage}{\linewidth}\small')
                in_algorithm = True
            match = re.match(r'^(\d+) ( +)', value)
            indentation = len(match[2]) if match else 0
            parts.append(r'\noindent\hspace*{' + str(indentation * .12)
                         + 'cm}' + tex_inline(value.strip()) + r'\par')
        else:
            parts.append(tex_inline(value) + '\n')
    if in_algorithm:
        parts.append(r'\end{minipage}\end{flushleft}')
    parts.append(r'\end{document}')
    target = latex_dir / 'main.tex'
    target.write_text('\n\n'.join(parts) + '\n', encoding='utf-8')
    return target


def compile_latex(source):
    command = [str(TECTONIC), '--keep-logs', '--keep-intermediates', '--print',
               '--outdir', str(source.parent), str(source)]
    print('Compiling complete LaTeX manuscript with Tectonic.', flush=True)
    with (OUT / 'latex_compilation.log').open('w', encoding='utf-8') as log:
        process = subprocess.run(command, cwd=source.parent, text=True,
                                 stdout=log, stderr=subprocess.STDOUT)
    output = (OUT / 'latex_compilation.log').read_text(encoding='utf-8')
    if process.returncode:
        print(output[-6000:], flush=True)
        raise RuntimeError('LaTeX compilation failed')
    print(output[-4000:], flush=True)
    pdf = SOURCE.with_suffix('.pdf')
    shutil.copy2(source.with_suffix('.pdf'), pdf)
    return pdf


def audit(docx, pdf, content, inventory):
    with zipfile.ZipFile(docx) as package, zipfile.ZipFile(REFERENCE) as before:
        root = etree.fromstring(package.read('word/document.xml'))
        old = etree.fromstring(before.read('word/document.xml'))
        assert package.read('word/styles.xml') == before.read('word/styles.xml')
        assert etree.tostring(root.find('.//w:sectPr', NS)) == etree.tostring(old.find('.//w:sectPr', NS))
        assert len(root.findall('.//m:oMathPara', NS)) == 9
        assert len(root.findall('.//m:oMath', NS)) == len(inventory)
        assert len(root.findall('.//m:fPr/m:type[@m:val="noBar"]', NS)) == 2
        assert not root.findall('.//m:type[@m:val="nobar"]', NS)
        assert root.findall('.//m:eqArr', NS), 'Multiline equations must remain native.'
        assert not any('$' in (t.text or '') for t in root.findall('.//w:t', NS))
        images = [s for s in package.namelist() if s.startswith('word/media/') and not s.endswith('/')]
        assert all(package.read(s) == before.read(s) for s in images)
    log = (OUT / 'latex_compilation.log').read_text(encoding='utf-8')
    assert not re.search(r'Missing character|Undefined control sequence|Overfull', log)
    record = {'source': str(SOURCE), 'pdf': str(pdf), 'word': str(docx),
              'pdf_engine': subprocess.check_output([str(TECTONIC), '--version'], text=True).strip(),
              'word_converter': subprocess.check_output([str(PANDOC), '--version'], text=True).splitlines()[0],
              'single_source_math': True, 'math_objects': len(inventory), 'display_equations': 9,
              'editable_omml': True, 'no_overfull_or_missing_glyphs': True,
              'word_styles_and_figure_preserved': True, 'new_experiments': False,
              'files_sha256': {str(s): hashlib.sha256(s.read_bytes()).hexdigest()
                               for s in (SOURCE, docx, pdf, OUT / 'latex/main.tex')}}
    (OUT / 'build_audit.json').write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False, indent=2), flush=True)


def package_and_preview(pdf):
    sys.path.insert(0, str(ROOT / 'artifacts/teacher_methodology_20261011/vendor'))
    import pymupdf
    from PIL import Image, ImageDraw
    target = OUT / 'preview'
    target.mkdir(exist_ok=True)
    document = pymupdf.open(pdf)
    thumbnails = []
    for index, page in enumerate(document):
        assert '\ufffd' not in page.get_text()
        page.get_pixmap(matrix=pymupdf.Matrix(1.7, 1.7)).save(target / f'page_{index + 1:02d}.png')
        thumb = Image.open(target / f'page_{index + 1:02d}.png').convert('RGB')
        thumb.thumbnail((400, 570))
        thumbnails.append(thumb)
    sheet = Image.new('RGB', (4 * 420, ((len(thumbnails) + 3) // 4) * 600), '#eeeeee')
    draw = ImageDraw.Draw(sheet)
    for index, thumb in enumerate(thumbnails):
        left, top = (index % 4) * 420 + 10, (index // 4) * 600 + 25
        sheet.paste(thumb, (left, top))
        draw.text((left, top - 18), f'Page {index + 1}', fill='black')
    sheet.save(target / 'contact_sheet.png')
    package = ROOT / 'docs' / (NAME + '_LaTeX源码.zip')
    with zipfile.ZipFile(package, 'w', zipfile.ZIP_DEFLATED) as archive:
        for file in sorted((OUT / 'latex').rglob('*')):
            if file.is_file() and (file.suffix in ('.tex', '.ttf', '.png', '.md', '.txt')):
                archive.write(file, file.relative_to(OUT / 'latex'))
    print(f'PDF pages: {len(document)}; LaTeX source package: {package}', flush=True)


def main():
    global SOURCE, OUT, NAME
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE)
    parser.add_argument('--output-dir', type=Path, default=OUT)
    args = parser.parse_args()
    SOURCE = args.source.resolve()
    OUT = args.output_dir.resolve()
    NAME = SOURCE.stem
    OUT.mkdir(parents=True, exist_ok=True)
    content = blocks(SOURCE)
    inventory = math_inventory(content)
    assert sum(kind == 'display' for kind, _ in inventory) == 9
    converted = convert_math(inventory)
    docx = write_word(content, converted)
    source = write_latex(content)
    pdf = compile_latex(source)
    audit(docx, pdf, content, inventory)
    package_and_preview(pdf)


if __name__ == '__main__':
    main()
