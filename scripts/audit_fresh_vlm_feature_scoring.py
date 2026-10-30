#!/usr/bin/env python3
"""Audit the real labeling pilot and compare only matched complete evidence."""
from pathlib import Path
import csv
import json
import math

import numpy as np

import run_fresh_vlm_feature_scoring as fresh
import run_category_five_dimension_pilot as pilot


def main():
    output = fresh.OUTPUT
    manifest = pilot.read_csv(output / "manifest.csv")
    scores = pilot.read_csv(output / "scores.csv")
    long = pilot.read_csv(output / "feature_annotations_with_evidence.csv")
    protocol = json.loads((output / "scoring_protocol.json").read_text())
    index = {r["sample_id"]: r for r in scores}
    frozen = json.loads((output / "input_freeze.sha256.json").read_text())
    for path, expected in frozen.items():
        if pilot.sha256(path) != expected:
            raise ValueError(f"Frozen input changed: {path}")
    assert len(manifest) == len(scores) == len(index) == 60
    assert len(long) == 1820
    for row in manifest:
        schema = json.loads((output / "schemas" / f"{row['class_name']}.json").read_text())
        value = json.loads((output / "annotations" / f"{row['sample_id']}.json").read_text())
        fresh.validate_annotation(value, row, schema)
        recalculated = fresh.fresh_score(row, value, protocol)
        record = index[row["sample_id"]]
        for key, number in recalculated.items():
            if isinstance(number, bool):
                assert record[key] == str(number)
            else:
                assert math.isclose(float(record[key]), number, rel_tol=0, abs_tol=1e-10)
        # Independently check the five fixed budgets and total interval.
        for key in ("lower", "upper", "midpoint"):
            total = 20 * sum(float(record[f"{d}_{key}"]) for d in pilot.DIMENSIONS)
            assert abs(total - float(record[f"score_{key}"])) < 1e-10
        if row["class_name"] in {"dog", "cat"}:
            original = json.loads((fresh.ROOT / "configs" / ("dog_feature_schema_v3_1.json" if row["class_name"] == "dog" else "cat_feature_schema_v1_0.json")).read_text())
            new_fields = {f["feature_id"]: f for f in schema["universal_features"]}
            assert len(original["universal_features"]) == 23
            for field in original["universal_features"]:
                assert new_fields[field["feature_id"]]["possible_values"] == field["possible_values"]
                assert f"feature_{field['feature_id']}" in record
    candidates = pilot.read_csv(fresh.POOL / "manifests/candidate_pool.csv")
    lookup = {r["sample_id"]: i for i, r in enumerate(candidates)}
    classifiers = []
    paths = []
    for path in sorted((fresh.POOL / "benchmark/runs").glob("*/*/candidate_outputs.npz")):
        paths.append(path)
        with np.load(path) as cache:
            assert cache["sample_ids"].tolist() == [r["sample_id"] for r in candidates]
            assert np.isfinite(cache["logits"]).all()
            classifiers.append((path.parent.parent.name, path.parent.name, cache["logits"].argmax(1)))
    config = json.loads(pilot.DEFAULT_CONFIG.read_text())
    comparisons, cohorts = [], []
    for source in ("local", "online"):
        for category in fresh.CLASSES:
            group, old_scores = [], {}
            for row in manifest:
                if row["class_name"] != category or index[row["sample_id"]]["fully_known"] != "True":
                    continue
                domain = "car" if category == "car" else "old"
                path = fresh.POOL / "annotations" / source / domain / "parsed" / f"{row['sample_id']}.json"
                paths.append(path)
                old = json.loads(path.read_text())["features"]
                records = {name: pilot.score_sample(float(row["bbox_area_ratio"]), old, category, scheme, config)
                           for name, scheme in (("old_global", "equal_global"), ("old_parts", "equal_parts"))}
                if all(v["fully_known"] for v in records.values()):
                    group.append(row)
                    old_scores[row["sample_id"]] = records
            cohorts.append(dict(source=source, class_name=category, eligible=len(group), selected=10))
            for model, seed, predictions in classifiers:
                errors = np.asarray([predictions[lookup[r["sample_id"]]] != fresh.CLASSES.index(category) for r in group], dtype=int)
                for scheme in ("old_global", "old_parts", "fresh_parts"):
                    risk = np.asarray([float(index[r["sample_id"]]["score_midpoint"]) if scheme == "fresh_parts" else
                                       old_scores[r["sample_id"]][scheme]["score_midpoint"] for r in group])
                    auc = pilot.error_auc(errors, risk) if len(group) else None
                    comparisons.append(dict(source=source, class_name=category, scheme=scheme,
                                            model=model, seed=seed, n=len(group), errors=int(errors.sum()), error_auc=auc))
    pilot.write_csv(output / "classifier_comparison_complete.csv", comparisons)
    pilot.write_csv(output / "comparison_cohorts.csv", cohorts)
    # Keep the initial all-midpoint table identifiable as sensitivity only.
    report_path = output / "report.md"
    report = report_path.read_text()
    note = ("\n## 最终核验与复核事项\n\n"
            "60个目标、1820条特征与证据记录均通过身份、字典、依赖及评分范围校验。"
            "59个目标评分分量全部已知；另1例羊的朝向缺少证据，评分区间为20—40分，"
            "30分仅是区间中点，状态为不可评价。\n\n"
            "classifier_comparison_complete.csv仅比较新旧三方案共同完整证据目标，样本支持见comparison_cohorts.csv。"
            "原classifier_comparison.csv使用全部目标的中点，保留为未知处理敏感性表，不能作为完整证据主比较。\n\n"
            "有限图像抽查发现一例人物双眼可见性待复核，以及一例车辆可能将视角导致的车轮不可见解释为缺失。"
            "原始标注保留，具体见visual_spot_review.csv；结构合法不表示语义已验证。"
            "本轮各类只有10个目标，不足以确认评分优越性。\n\n"
            "feature_dimension_assignment.csv逐项注明字段所属维度或条件及是否计分。"
            "完整标注不代表所有名义外观属性均获得风险分值。\n")
    # Derive this particular unknown interval from the output, rather than assume it.
    uncertain = [r for r in scores if r["evaluable"] != "True"]
    if len(uncertain) == 1:
        u = uncertain[0]
        note = note.replace("20—40分", f"{float(u['score_lower']):g}—{float(u['score_upper']):g}分")
        note = note.replace("30分仅", f"{float(u['score_midpoint']):g}分仅")
    if "## 最终核验与复核事项" not in report:
        report_path.write_text(report + note, encoding="utf-8")
        from docx import Document
        document = Document(output / "report.docx")
        document.add_heading("最终核验与复核事项", 1)
        for paragraph in note.split("\n\n")[1:]:
            document.add_paragraph(paragraph.strip())
        document.save(output / "report.docx")
    audit = dict(status="pass", targets=len(scores), feature_evidence_records=len(long),
                 frozen_inputs_verified=len(frozen), recomputed_scores=len(scores),
                 known_risk_targets=sum(r["fully_known"] == "True" for r in scores),
                 unevaluable_targets=len(uncertain), complete_comparison_rows=len(comparisons),
                 semantic_review_is_partial=True)
    fresh.atomic_json(output / "final_audit.json", audit)
    paths += [Path(__file__), pilot.DEFAULT_CONFIG, output / "scores.csv", output / "feature_annotations_with_evidence.csv"]
    fresh.atomic_json(output / "analysis_freeze.sha256.json", {str(p.resolve()): pilot.sha256(p) for p in paths})
    with (output / "research_log.md").open("a") as log:
        log.write(f"\n{fresh.now()}：最终核验60目标与1820条字段证据；核对原23项值域保留；"
                  "增补完整证据匹配对照，原全部中点对照保留为敏感性结果。命令："
                  "`.venv/bin/python scripts/audit_fresh_vlm_feature_scoring.py`。\n")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
