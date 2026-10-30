#!/usr/bin/env python3
"""Apply evidence-checked text revisions to the user's edited methods DOCX."""
from copy import deepcopy
from pathlib import Path
import difflib
import hashlib
import json
import re
import subprocess
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from lxml import etree

import recompile_methods_latex_20261012 as build

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path('/home/hcf/.codex/attachments/c625a115-ca4a-493c-b226-7b38760633f0/论文方法章节_视觉状态驱动模型测试_单行公式版_20261012.docx')
BASELINE = ROOT / 'docs/论文方法章节_视觉状态驱动模型测试_方法正文版_20261012.md'
SOURCE = ROOT / 'docs/论文方法章节_视觉状态驱动模型测试_说明完善版_20261012.md'
OUT = ROOT / 'artifacts/methods_explanations_overlay_20261012'
NS = build.NS


def content_elements(document):
    return [element for element in document._element.body
            if element.tag == qn('w:tbl') or element.tag == qn('w:p') and
            (element.findall('.//w:t', NS) or element.findall('.//m:oMath', NS)
             or element.findall('.//w:drawing', NS))]


def signature(element):
    """Compare XML content, independent of namespace-prefix serialization."""
    if element is None:
        return None
    attributes = dict(element.attrib)
    if element.tag == qn('m:type') and attributes.get(qn('m:val')) == 'nobar':
        attributes[qn('m:val')] = 'noBar'
    return (element.tag, tuple(sorted(attributes.items())), element.text,
            tuple(signature(child) for child in element))


def validate_text(document, elements, content):
    assert len(elements) == len(content), (len(elements), len(content))
    for index, (element, (kind, value)) in enumerate(zip(elements, content)):
        if kind == 'image':
            assert element.findall('.//w:drawing', NS), index
        elif kind == 'equation':
            assert element.find('m:oMathPara', NS) is not None, index
        elif kind == 'table':
            table = Table(element, document)
            assert len(table.rows) == len(value), index
            for row, cells in zip(table.rows, value):
                for cell, expected in zip(row.cells, cells):
                    actual = cell.paragraphs[0].text.strip().replace('\u200b', '')
                    assert actual == build.plain(expected), (index, actual, expected)
        else:
            actual = Paragraph(element, document).text.strip().replace('\u200b', '')
            assert actual == build.plain(value), (index, kind, actual, build.plain(value))


def copy_body_paragraph(template):
    paragraph = deepcopy(template)
    paragraph.attrib.pop(qn('w14:paraId'), None)
    paragraph.attrib.pop(qn('w14:textId'), None)
    return paragraph


def overlay(before, after, converted):
    document = Document(REFERENCE)
    elements = content_elements(document)
    validate_text(document, elements, before)
    unchanged = []
    revised = []
    operations = difflib.SequenceMatcher(
        None, [repr(item) for item in before], [repr(item) for item in after],
        autojunk=False).get_opcodes()
    for operation, a, b, c, d in operations:
        if operation == 'equal':
            unchanged.extend((c + offset, signature(element))
                             for offset, element in enumerate(elements[a:b]))
            continue
        assert operation in ('replace', 'insert'), operation
        assert all(kind == 'body' for kind, _ in after[c:d]), operation
        if operation == 'replace':
            assert all(kind == 'body' for kind, _ in before[a:b])
            anchor = elements[a]
            template = deepcopy(anchor)
            replacements = [anchor]
            for _ in after[c + 1:d]:
                added = copy_body_paragraph(template)
                replacements[-1].addnext(added)
                replacements.append(added)
            for element in elements[a + 1:b]:
                element.getparent().remove(element)
        else:
            assert a > 0
            anchor = elements[a - 1]
            assert before[a - 1][0] == 'body'
            template = deepcopy(anchor)
            replacements = []
            for _ in after[c:d]:
                added = copy_body_paragraph(template)
                anchor.addnext(added)
                anchor = added
                replacements.append(added)
        for element, (_, text) in zip(replacements, after[c:d]):
            build.rewrite(Paragraph(element, document), text, converted)
        revised.append({'operation': operation, 'old_blocks': [a, b], 'new_blocks': [c, d]})
    fixes = 0
    for element in document._element.findall('.//m:type[@m:val="nobar"]', NS):
        element.set(qn('m:val'), 'noBar')
        fixes += 1
    assert fixes == 2
    output = SOURCE.with_suffix('.docx')
    document.save(output)
    return output, unchanged, revised


def audit(docx, pdf, before_content, content, inventory, unchanged, revised):
    document = Document(docx)
    elements = content_elements(document)
    validate_text(document, elements, content)
    assert all(signature(elements[index]) == expected for index, expected in unchanged)
    with zipfile.ZipFile(docx) as current, zipfile.ZipFile(REFERENCE) as previous:
        root = etree.fromstring(current.read('word/document.xml'))
        old = etree.fromstring(previous.read('word/document.xml'))
        assert signature(etree.fromstring(current.read('word/styles.xml'))) == signature(
            etree.fromstring(previous.read('word/styles.xml')))
        assert signature(root.find('.//w:sectPr', NS)) == signature(old.find('.//w:sectPr', NS))
        assert all(current.read(name) == previous.read(name)
                   for name in current.namelist()
                   if name.startswith('word/media/') and not name.endswith('/'))
        for name in previous.namelist():
            if re.fullmatch(r'word/(?:header|footer)\d+\.xml', name):
                assert signature(etree.fromstring(current.read(name))) == signature(
                    etree.fromstring(previous.read(name)))
        assert len(root.findall('.//m:oMath', NS)) == len(inventory)
        assert len(root.findall('.//m:fPr/m:type[@m:val="noBar"]', NS)) == 2
        assert not root.findall('.//m:type[@m:val="nobar"]', NS)
        display_positions = [index for index, (kind, _) in enumerate(content) if kind == 'equation']
        assert len(display_positions) == 9
        # Four additional math paragraphs belong to symbol-table cells, not body equations.
        assert len(root.findall('.//m:oMathPara', NS)) == 13
        assert elements[display_positions[5]].findall('.//m:eqArr', NS) == []
        assert not any('$' in (node.text or '') for node in root.findall('.//w:t', NS))
        assert not any((node.text or '').startswith('3.7') for node in root.findall('.//w:t', NS))
    log = (OUT / 'latex_compilation.log').read_text(encoding='utf-8')
    assert not re.search(r'Missing character|Undefined control sequence|Overfull', log)
    record = {
        'uploaded_reference': str(REFERENCE), 'source': str(SOURCE),
        'word': str(docx), 'pdf': str(pdf),
        'changes': revised, 'unchanged_content_blocks': len(unchanged),
        'original_block_count': len(before_content), 'revised_block_count': len(content),
        'display_equations': 9, 'math_objects': len(inventory),
        'symbol_table_math_paragraphs_preserved': 4,
        'existing_display_equations_preserved': True, 'equation_6_single_line': True,
        'editable_native_math': True, 'combination_type_case_fixes': 2,
        'uploaded_headings_fonts_styles_page_settings_and_figure_preserved': True,
        'section_3_7_absent': True, 'new_experiments': False,
        'pdf_engine': subprocess.check_output([str(build.TECTONIC), '--version'], text=True).strip(),
        'word_conversion': 'Changed paragraphs only; LaTeX-to-OMML with Pandoc',
        'pdf_is_word_page_rendering': False,
        'no_overfull_or_missing_glyphs': True,
        'files_sha256': {str(path): hashlib.sha256(path.read_bytes()).hexdigest()
                         for path in (REFERENCE, SOURCE, docx, pdf, OUT / 'latex/main.tex')},
    }
    (OUT / 'build_audit.json').write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False, indent=2), flush=True)


def main():
    build.SOURCE, build.REFERENCE, build.OUT, build.NAME = SOURCE, REFERENCE, OUT, SOURCE.stem
    OUT.mkdir(parents=True, exist_ok=True)
    before = build.blocks(BASELINE)
    before = [(kind, '3.3 状态覆盖度量与覆盖引导选样'
               if kind == 'heading' and value.startswith('3.3 ') else
               '3.4 视觉风险分层与测试组织'
               if kind == 'heading' and value.startswith('3.4 ') else value)
              for kind, value in before]
    content = build.blocks(SOURCE)
    inventory = build.math_inventory(content)
    assert [value for kind, value in before if kind == 'equation'] == [
        value for kind, value in content if kind == 'equation']
    converted = build.convert_math(inventory)
    docx, unchanged, revised = overlay(before, content, converted)
    latex = build.write_latex(content)
    pdf = build.compile_latex(latex)
    audit(docx, pdf, before, content, inventory, unchanged, revised)
    build.package_and_preview(pdf)


if __name__ == '__main__':
    main()
