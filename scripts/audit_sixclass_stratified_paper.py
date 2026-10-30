#!/usr/bin/env python3
"""Check manuscript numerical tables against frozen outputs and Word contents."""
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile
import pandas as pd
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'artifacts/sixclass_difficulty_supplement_20261004'
MAIN=ROOT/'artifacts/sixclass_dual_vlm_replaced_20261004'
PAPER=ROOT/'docs/visual_state_coverage_difficulty_stratified_paper_20261004.md'

def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def main():
    raw=PAPER.read_text();lines=raw.splitlines();tables=[];index=0;checks=0
    while index<len(lines):
        if lines[index].startswith('|'):
            group=[]
            while index<len(lines) and lines[index].startswith('|'):
                group.append([v.strip() for v in lines[index].strip('|').split('|')]);index+=1
            tables.append(group)
        else:index+=1
    def read(name):return pd.read_csv(BASE/'tables'/f'{name}.csv')
    def pct(actual,expected):
        nonlocal checks
        value=float(actual.replace('%','').replace(' pp','').replace('−','-'))
        assert abs(value-100*float(expected))<=.00501,(actual,expected)
        checks+=1
    def num(actual,expected,tol=.000051):
        nonlocal checks
        assert abs(float(actual.replace('−','-'))-float(expected))<=tol,(actual,expected)
        checks+=1
    full=read('full_stratum_metrics');full=full[full.unknown_mode=='midpoint']
    means=full.groupby(['source','tier']).mean(numeric_only=True)
    for tier,row in enumerate(tables[2][2:],1):
        for j,source in enumerate(['local','online']):
            m=means.loc[source,tier]
            for offset,key in enumerate(['macro_f1','accuracy','macro_precision']):pct(row[1+j*3+offset],m[key])
    batch=read('selected_batch_metrics').groupby(['source','tier']).mean(numeric_only=True)
    cum=read('cumulative_metrics').groupby(['source','tier']).mean(numeric_only=True)
    for tier,row in enumerate(tables[3][2:],1):
        assert int(row[1])==120*tier
        for j,source in enumerate(['local','online']):
            for off,value in enumerate([batch.loc[source,tier].macro_f1,cum.loc[source,tier].macro_f1,cum.loc[source,tier].macro_frequency_coverage]):pct(row[2+j*3+off],value)
    m=pd.read_csv(MAIN/'analysis/tables/cross_fixed_metrics.csv');m=m[m.budget==100]
    r=pd.read_csv(MAIN/'analysis/tables/cross_random_summary.csv');r=r[r.budget==100]
    for row,source in zip(tables[4][2:4],['local','online']):
        x=m[(m.method=='coverage')&(m.selection_source==source)&(m.evaluation_source==source)].mean(numeric_only=True)
        for actual,key in zip(row[1:],['macro_frequency_coverage','accuracy','macro_f1','macro_error_rate']):pct(actual,x[key])
    x=r[r.evaluation_source=='local'].mean(numeric_only=True)
    for actual,value in zip(tables[4][4][2:],[x.accuracy_mean,x.macro_f1_mean,x.macro_error_rate_mean]):pct(actual,value)
    for row,method in zip(tables[4][5:],['deepgini','dsa','kcenter']):
        x=m[(m.method==method)&(m.evaluation_source=='local')].mean(numeric_only=True)
        for actual,key in zip(row[2:],['accuracy','macro_f1','macro_error_rate']):pct(actual,x[key])
    primary=pd.read_csv(MAIN/'analysis/tables/independent_comparisons.csv');primary=primary[primary.baseline=='random']
    for i,row in enumerate(tables[5][2:]):
        source=['local','online'][i//2];metric=['confusion','clusters'][i%2]
        value=primary[(primary.selection_source==source)&(primary.metric=='class_equal_error_'+metric)].iloc[0]
        num(row[2],value.mean_delta_auc);lo,hi=re.findall(r'[+−-]?\d+\.\d+',row[3]);num(lo,value.ci95_low);num(hi,value.ci95_high)
        num(row[5],value.holm_primary_four_p,tol=1e-9);assert int(row[4].split('/')[0])==value.positive_runs
    comp=read('matched_equal_error_comparisons')
    for row,(_,value) in zip(tables[6][2:],comp.iterrows()):
        num(row[2],value.mean_delta_auc);lo,hi=re.findall(r'[+−-]?\d+\.\d+',row[3]);num(lo,value.ci95_low);num(hi,value.ci95_high)
        num(row[5],value.holm_four_p,tol=1e-9);assert int(row[4].split('/')[0])==value.positive_runs
    names=['convnext_tiny','swin_t','resnet50','efficientnet_b0','densenet121','mobilenet_v3_large','resnet18','mobilenet_v3_small']
    rankings=read('model_robustness')
    for row,model in zip(tables[7][2:],names):
        for j,source in enumerate(['local','online']):
            x=rankings[(rankings.source==source)&(rankings.model==model)].iloc[0]
            for off,key in enumerate(['low_risk_f1','high_risk_f1','drop']):pct(row[1+j*3+off],x[key])
    for row in tables[8][2:]:
        for j,source in enumerate(['local','online']):
            vals=row[1+j*2].split('→');low=means.loc[source,1][row[0]+'_recall'];high=means.loc[source,5][row[0]+'_recall']
            pct(vals[0],low);pct(vals[1],high);pct(row[2+j*2],high-low)
    states=read('state_recall').groupby(['source','class_name','feature','value','support']).recall.mean()
    for row,(field,value) in zip(tables[9][2:],[('blur_level','severe'),('occlusion_level','heavy'),('target_scale','tiny'),('body_orientation','back')]):
        for j,source in enumerate(['local','online']):
            n,v=row[1+j].split('/');pct(v,states.loc[source,'person',field,value,int(n)])
    for figure in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',raw):assert (PAPER.parent/figure).is_file()
    for svg in (BASE/'figures').glob('*.svg'):
        parsed=ET.parse(svg);ns={'s':'http://www.w3.org/2000/svg'}
        assert parsed.findall('.//s:text',ns) and not parsed.findall('.//s:image',ns)
    doc=Document(PAPER.with_suffix('.docx'));assert len(doc.inline_shapes)==4 and len(doc.tables)==11
    with ZipFile(PAPER.with_suffix('.docx')) as z:
        body=z.read('word/document.xml').decode();assert body.count('<m:oMath>')==1 and '$$' not in body
    frozen=json.loads((BASE/'input_code_order_freeze.sha256.json').read_text())
    assert all(digest(path)==sha for path,sha in frozen.items())
    files=[PAPER,PAPER.with_suffix('.docx'),Path(__file__),ROOT/'scripts/export_sixclass_stratified_paper.py']
    files+=list((BASE/'tables').glob('*.csv'))+list((BASE/'figures').glob('*'))
    (BASE/'paper_artifact_manifest.sha256.json').write_text(json.dumps({str(p):digest(p) for p in files},indent=2)+'\n')
    result=dict(status='pass',numeric_cells_checked=checks,manuscript_characters=len(raw),word_tables=len(doc.tables),
                word_figures=len(doc.inline_shapes),word_native_equations=1,editable_svg_figures=4,source_freeze_unchanged=True)
    (BASE/'paper_audit.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result))

if __name__=='__main__':main()
