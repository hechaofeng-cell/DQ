#!/usr/bin/env python3
"""Build an evidence-linked Chinese progress report with figures and speaker notes."""
from pathlib import Path
import csv
import hashlib
import json
import shutil
import zipfile
import os
import re
import xml.etree.ElementTree as ET
from collections import defaultdict

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/progress_report_20261008'
FIG = OUT / 'figures'
TABLE = OUT / 'tables'
NAME = '老师分层建议之后的进度汇报_图文讲解版_20261008'
BLUE, ORANGE, GRAY = '#2166AC', '#B8642A', '#D5D9DE'
POOLS = {'B': 'COCO', 'C': 'VOC2007', 'D': 'Objects365'}
SOURCES = {'local': '本地VLM', 'online': '线上VLM'}
inputs, figures, evidence_rows = [], [], []


def csv_rows(relative):
    p = ROOT / relative
    inputs.append(p)
    with p.open(encoding='utf-8') as f:
        return list(csv.DictReader(f))


def save_chart(fig, stem):
    for suffix in ['.svg', '.png']:
        fig.savefig(FIG / (stem + suffix), dpi=230, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    figures.append(stem)
    return FIG / (stem + '.png')


def copy_chart(relative, stem):
    p = ROOT / relative
    translations = {
        'Official labels':'官方标注', 'Target boxes':'目标框', '+ categories':'与类别真值',
        'Dual VLM states':'双VLM状态', 'No classifier':'不读取分类器', 'outcomes':'预测结果',
        'Frozen classifiers':'冻结分类器', '24 prediction sets':'24组预测缓存',
        'Coverage + risk':'状态覆盖与风险', 'Test-suite':'测试集合', 'organization':'组织',
        'Conditional results':'条件结果', 'Performance':'分类表现', '+ diagnosis':'与状态诊断',
        'Discovery A: freeze conditions':'发现池A：冻结薄弱条件',
        'Confirm B / C / D: reuse frozen models':'确认池B/C/D：沿用冻结模型',
        'Back-facing dog':'背向dog', 'Small-area cat':'小尺度cat', 'Blurred sheep':'严重模糊sheep',
        'Medium-area dog':'中等尺度dog', 'Recall (%)':'Recall（%）',
        'Class recall':'类别总体Recall', 'Conditional recall':'状态Recall',
        'Q1 − Q5 Macro-F1 (pp)':'Q1 − Q5 Macro-F1（百分点）',
        'a  Endpoint separation':'a  两方案端点区分', 'Area only':'仅面积', 'Full risk':'完整风险',
        'Full − area-only gap (pp)':'完整评分 − 仅面积（百分点）', 'b  Paired difference':'b  配对增量',
        'Original + official box':'原图与官方目标框', 'Actual classifier input':'分类器实际输入',
        'Automatic state + outcome':'自动状态与预测', 'Area ratio: 0.0026':'面积比：0.0026',
        'Area ratio: 0.0266':'面积比：0.0266', 'Area ratio: 0.0383':'面积比：0.0383',
        'Area ratio: 0.0506':'面积比：0.0506', 'Predictions (3 seeds):':'三个种子的预测：',
        'GT: dog':'真实类别：dog', 'GT: cat':'真实类别：cat', 'GT: sheep':'真实类别：sheep',
        'Failure example':'错误案例', 'Correct example':'正确案例',
        'local: body_orientation':'本地：身体朝向', 'local: target_scale':'程序：目标尺度',
        'online: blur_level':'线上：模糊程度', '= back':'= 背向', '= small':'= 小', '= medium':'= 中等', '= severe':'= 严重',
    }
    for pool in POOLS:
        for source in SOURCES:
            translations[f'{pool} · {source}'] = f'{pool} · ' + ('本地' if source == 'local' else '线上')
    inputs.extend([p.with_suffix('.svg'),p.with_suffix('.png')])
    tree = ET.parse(p.with_suffix('.svg'))
    for e in tree.getroot().iter():
        if e.tag.endswith('}text') and e.text in translations:
            e.text = translations[e.text]
        if 'style' in e.attrib:
            e.attrib['style'] = re.sub(r'font-family:[^;]+', "font-family: 'Noto Sans SC'", e.attrib['style'])
    dest = FIG / (stem + '.svg')
    tree.write(dest,encoding='utf-8',xml_declaration=True)
    import cairosvg
    width = Image.open(p.with_suffix('.png')).width
    cairosvg.svg2png(url=str(dest),write_to=str(FIG/(stem+'.png')),output_width=width)
    figures.append(stem)
    return FIG / (stem + '.png')


def font(run, size=10.5, bold=False, color='263442'):
    run.font.name = 'Noto Sans CJK SC'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')


doc = Document()
md = []
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.5)
section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
section.header_distance, section.footer_distance = Cm(.65), Cm(.65)
normal = doc.styles['Normal']
normal.font.name = 'Noto Sans CJK SC'
normal.font.size = Pt(10.5)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
normal.paragraph_format.line_spacing = 1.2
normal.paragraph_format.space_after = Pt(5)
header = section.header.paragraphs[0]
font(header.add_run('QE研究进度  |  从分层建议到条件诊断'), 8, color='697685')
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(footer.add_run('2026.10.08  ·  '), 8, color='697685')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def para(text, size=10.5, bold=False, color='263442'):
    p = doc.add_paragraph()
    font(p.add_run(text), size, bold, color)
    md.append(text + '\n')
    return p


def label(text):
    p = para(text, 11, True, '2166AC')
    p.paragraph_format.keep_with_next = True
    return p


def page(title, subtitle=None):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    p = para(title, 17, True, '163F66')
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    md[-1] = '# ' + title + '\n'
    if subtitle:
        para(subtitle, 10, False, '697685')


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = str(v)
    for i, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, c in enumerate(row.cells):
            if widths:
                c.width = Cm(widths[j])
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8F0F7')
                c._tc.get_or_add_tcPr().append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                for r in p.runs:
                    font(r, 9, i == 0)
    md.append('| ' + ' | '.join(headers) + ' |\n| ' + ' | '.join(['---'] * len(headers)) + ' |\n')
    md.extend('| ' + ' | '.join(str(x) for x in r) + ' |\n' for r in rows)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def picture(path, caption, max_height=9):
    w, h = Image.open(path).size
    width = min(17.2, max_height * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width))
    cap = para(caption, 8.5, color='697685')
    cap.paragraph_format.space_after = Pt(7)
    md.append(f'![{caption}](../artifacts/progress_report_20261008/figures/{path.name})\n')


def speak(text):
    label('可以这样向老师讲')
    p = para(text)
    p.paragraph_format.left_indent = Cm(.3)
    p.paragraph_format.right_indent = Cm(.3)
    border = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    for k, v in [('val', 'single'), ('sz', '12'), ('space', '9'), ('color', '2166AC')]:
        left.set(qn('w:' + k), v)
    border.append(left)
    p._p.get_or_add_pPr().append(border)


def transition(text):
    para('接下来追问：' + text, 10, True, '965623')


def main():
    FIG.mkdir(parents=True, exist_ok=True)
    TABLE.mkdir(parents=True, exist_ok=True)
    config = OUT/'fontconfig.xml'
    config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd"><fontconfig><dir>'
                      +str(ROOT/'scripts/assets')+'</dir><dir>/usr/share/fonts</dir><cachedir>'
                      +str(OUT/'fontcache')+'</cachedir></fontconfig>',encoding='utf-8')
    os.environ['FONTCONFIG_FILE'] = str(config)
    fpath = ROOT / 'scripts/assets/NotoSansSC-Regular.ttf'
    font_manager.fontManager.addfont(str(fpath))
    family = font_manager.FontProperties(fname=str(fpath)).get_name()
    plt.rcParams.update({'font.family': family, 'font.size': 11, 'svg.fonttype': 'none',
                         'pdf.fonttype': 42, 'axes.spines.top': False, 'axes.spines.right': False,
                         'axes.unicode_minus': False, 'axes.linewidth': .8, 'legend.frameon': False})
    random = csv_rows('artifacts/random_stratification_ablation_20261006/tables/random_summary.csv')
    ablation = csv_rows('artifacts/random_stratification_ablation_20261006/tables/ablation_summary.csv')
    scale = csv_rows('artifacts/scale_only_confirmation_20261007/tables/summary.csv')
    conditions = csv_rows('artifacts/scale_only_confirmation_20261007/tables/area_condition_summary.csv')
    panel = csv_rows('artifacts/controlled_risk_distribution_20261005/tables/panel_summary.csv')
    assert len(random) == len(scale) == 6
    assert all(int(r['extreme_count']) == 0 for r in random)
    scores = {}
    for p in POOLS:
        for s in SOURCES:
            path = ROOT / f'artifacts/random_stratification_ablation_20261006/ablation/{p}_{s}.npz'
            inputs.append(path)
            with np.load(path) as z:
                assert str(z['variants'][0]) == 'full'
                values = z['f1'][0]
                assert values.shape == (24, 5)
                assert (values[:, 0] > values[:, 4]).all()
                scores[p, s] = values.mean(axis=0) * 100
            row = next(r for r in random if r['pool'] == p and r['source'] == s)
            assert abs(scores[p, s][4] - scores[p, s][0] - 100 * float(row['observed_delta'])) < 1e-9

    workflow = copy_chart('artifacts/submission_closeout_20261007/figures/01_workflow', '01_test_workflow')
    cases = copy_chart('artifacts/submission_closeout_20261007/figures/04_condition_cases', '03_condition_cases')
    randfig = copy_chart('artifacts/positive_results_report_20261006/figures/01_risk_gap_vs_random', '04_risk_vs_random')
    mixture = copy_chart('artifacts/positive_results_report_20261006/figures/04_risk_share_response_delta', '06_fixed_budget_response')
    scalefig = copy_chart('artifacts/submission_closeout_20261007/figures/05_scale_only_control', '07_scale_control')
    real = copy_chart('artifacts/submission_closeout_20261007/figures/06_real_target_cases', '09_real_inputs')

    fig, axes = plt.subplots(1, 3, figsize=(11.2, 4.9), layout='constrained', sharey=True)
    for ax, p in zip(axes, POOLS):
        for x, s in enumerate(SOURCES):
            lo, hi = scores[p, s][0], scores[p, s][4]
            ax.plot([x, x], [hi, lo], color='#B7C8D9', lw=3)
            ax.scatter(x, lo, s=85, marker='o', color=BLUE, label='Q1低风险' if x == 0 else None)
            ax.scatter(x, hi, s=85, marker='s', color=ORANGE, label='Q5高风险' if x == 0 else None)
            ax.text(x, lo + 3, f'{lo:.2f}%', ha='center', fontsize=10, color=BLUE)
            ax.text(x, hi - 6, f'{hi:.2f}%', ha='center', fontsize=10, color=ORANGE)
            ax.text(x + .14, (lo + hi) / 2, f'差{lo-hi:.2f}\n个百分点', fontsize=9, va='center')
        ax.set(title=POOLS[p], xticks=[0, 1], xticklabels=['本地标签', '线上标签'],
               xlim=(-.5, 1.75), ylim=(0, 108))
    axes[0].set_ylabel('24个冻结模型平均Macro-F1（%）')
    axes[0].legend(loc='lower left', fontsize=10)
    confirmfig = save_chart(fig, '02_confirmation_endpoints')

    variants = ['full', 'drop_size', 'drop_body_visibility', 'drop_occlusion_level', 'drop_truncation_level', 'drop_body_orientation', 'equal_normalized']
    actual_variants = {r['variant'] for r in ablation}
    for i, v in enumerate(variants):
        if v not in actual_variants and v == 'equal_weight':
            variants[i] = next(x for x in actual_variants if 'equal' in x)
    mat = np.array([[-100 * float(next(r for r in ablation if r['pool'] == p and r['source'] == s and r['variant'] == v)['mean_delta_auc'])
                     for p in POOLS for s in SOURCES] for v in variants])
    assert (mat > 0).all()
    fig, ax = plt.subplots(figsize=(10.5, 5.7), layout='constrained')
    im = ax.pcolormesh(np.arange(-.5,6.5),np.arange(-.5,7.5),mat,
                       cmap='Blues',vmin=0,vmax=46,shading='flat',rasterized=False)
    ax.set_ylim(6.5,-.5)
    ax.set(xticks=range(6), xticklabels=[f'{POOLS[p]}\n{SOURCES[s]}' for p in POOLS for s in SOURCES],
           yticks=range(7), yticklabels=['完整评分', '删除尺度', '删除可见性', '删除遮挡', '删除截断', '删除朝向', '等权归一化'])
    ax.tick_params(length=0)
    for i in range(7):
        for j in range(6):
            ax.text(j, i, f'{mat[i,j]:.2f}', ha='center', va='center', color='white' if mat[i,j]>26 else '#263442', fontsize=11)
    cbar = fig.colorbar(im, ax=ax, shrink=.85, pad=.025)
    cbar.solids.set_rasterized(False)
    cbar.set_label('Q1 − Q5 Macro-F1（百分点）')
    abfig = save_chart(fig, '05_dimension_ablation')

    count = []
    for p in POOLS:
        allrows = [r for r in conditions if r['pool'] == p]
        evalrows = [r for r in allrows if r['evaluable'] == 'True']
        negative = sum(float(r['area_adjusted_recall_gap']) < 0 for r in evalrows)
        count.append([POOLS[p], len(allrows), len(evalrows), negative, len(evalrows)-negative, len(allrows)-len(evalrows)])
    assert count[0][2:4] == [119, 109] and count[1][2:4] == [4, 2] and count[2][2:4] == [122, 111]
    with (TABLE/'area_condition_counts.csv').open('w', encoding='utf-8', newline='') as f:
        w=csv.writer(f);w.writerow(['pool','frozen','evaluable','negative','nonnegative','insufficient']);w.writerows(count)
    fig, ax = plt.subplots(figsize=(10.5, 4.2), layout='constrained')
    yy = np.arange(3)
    n, other, missing = (np.array([r[k] for r in count]) for k in [3,4,5])
    ax.barh(yy, n, color=BLUE, label='Recall负差')
    ax.barh(yy, other, left=n, color=ORANGE, hatch='//', label='非负差')
    ax.barh(yy, missing, left=n+other, color=GRAY, label='支持不足')
    for i,r in enumerate(count):
        if r[3]>10:ax.text(r[3]/2,i,str(r[3]),ha='center',va='center',color='white',fontweight='bold')
        if r[5]>10:ax.text(r[3]+r[4]+r[5]/2,i,str(r[5]),ha='center',va='center',color='#4D4D4D')
        ax.text(145,i,f'{r[3]} / {r[2]}可评价组合为负',va='center',fontsize=11)
    ax.set(yticks=yy, yticklabels=POOLS.values(), xlabel='冻结的非尺度状态—架构组合数量（两标注来源合计）', xlim=(0,207), ylim=(2.65,-.8))
    ax.legend(loc='upper center',bbox_to_anchor=(.48,1.13),ncol=3,fontsize=10)
    countfig = save_chart(fig,'08_area_adjusted_conditions')

    page('从老师的分层建议，到可解释的模型诊断', '研究进度汇报 · 图文讲解版 · 2026年10月8日')
    para('汇报起点：上次老师建议对测试数据进行分层。本轮先完成分层，再验证分层是否有效、能否复现、能指出哪些具体薄弱条件。',12)
    label('本次汇报的一条主线')
    para('把总体分数展开为风险层表现和具体条件报告；通过新样本、新来源和简单对照，明确这种测试组织的用途及边界。',12,True)
    table(['步骤','我做了什么','回答什么问题'],[
        ['1','五维风险、类别内五等分','怎样落实分层建议？'],
        ['2','新COCO样本、VOC、Objects365确认','换图片、换来源还成立吗？'],
        ['3','冻结薄弱状态，在新池确认','能指出下一步具体该测什么？'],
        ['4','随机划层与风险字段消融','差异是随机的吗？依赖单字段吗？'],
        ['5','固定120目标，改变Q5占比','测试构成如何影响测得能力？'],
        ['6','仅面积对照、面积层内条件诊断','状态信息超出了目标大小吗？'],
        ['7','论文、真实案例与复算材料','结果如何解释、复查和使用？'],
    ],[.9,7.8,8.7])
    speak('老师，上次您建议我做分层。我先把分层规则落实，再围绕结果做了几类验证：检查新样本与新来源是否复现，检查随机分组和风险字段的影响，进一步形成具体状态诊断，最后补了仅按目标大小分层的对照。')
    para('阅读方式：每一页对应一个汇报问题，并提供“怎么看图”“可以怎么讲”与后续追问。顺序服务于解释，部分实验实际执行时交叉进行。',9,color='697685')

    page('01  我怎样落实老师的分层建议？','先按视觉条件组织样本，再评价冻结分类器')
    para('六分类：dog、cat、horse、sheep、person、car。person是人体目标。八种架构各三个种子，共24份权重；训练和checkpoint选择完成后保持冻结。')
    picture(workflow,'图1  官方类别负责分类真值，VLM负责状态描述；分类器预测用于结果评价。',6.2)
    table(['风险维度','评分依据'],[
        ['目标尺度','官方框面积/原图面积：越小分数越高（0—3）'],
        ['身体可见性','完整/大部分0，部分2，极少3'],
        ['遮挡','无0、轻1、中2、重3'],
        ['截断','无0、轻0.5、中1、重2'],
        ['朝向','常规0、背向1、俯视/异常2'],
    ],[3.2,14.2])
    para('五项相加得到0—13分风险代理，每类内部排序并均分Q1—Q5。每类400目标时，每层每类80个、六类共480个。Q5是池内相对高风险，分层不读取分类错误或置信度。')
    speak('我用五项视觉信息，在每个类别内部划分五个等数量风险层。先组织条件、再看模型表现，并保持高低层类别比例相同。这样可以减少“某层只是包含更多难类别”的解释。')
    transition('在原样本上能区分表现，换一批图片是否仍然成立？')

    page('02  分层结果能在新样本、新来源上复现吗？','冻结规则后的三个确认池：5400个目标')
    para('初始池中，本地标签Q1/Q5平均F1为94.90%/82.67%，线上为93.68%/84.35%。之后建立新COCO样本池，并加入VOC2007、Objects365；沿用模型、评分和指标。')
    picture(confirmfig,'图2  每条线连接同一确认设置的Q1与Q5。点为24运行平均值；本图没有添加误差条。',8.0)
    table(['确认池','目标数','本地Q1−Q5','线上Q1−Q5'],[
        ['COCO','2400','25.50个百分点','21.30个百分点'],
        ['VOC2007','600','16.57个百分点','14.44个百分点'],
        ['Objects365','2400','41.80个百分点','41.25个百分点'],
    ],[4.2,2.2,5.5,5.5])
    label('怎么看图')
    para('同一设置中，蓝色低风险点均高于橙色高风险点。落差在多个已测来源重复出现；中间层不要求严格递减。新COCO只属于新样本确认，VOC/Objects365增加来源证据。')
    speak('分层有初步差异以后，我固定规则继续确认。三个确认池中，高风险层平均F1都较低，落差约十四到四十二个百分点。现象能够在多个来源重复出现，但不意味着每个具体状态都能迁移。')

    page('03  分层之外，我定位了哪些具体薄弱条件？','把“模型平均表现怎样”展开为“下一步具体该测什么”')
    para('在发现池A中冻结175条状态—架构组合：状态至少20个目标、Recall比类别总体低至少10个百分点；随后沿用同一清单到新池确认。')
    picture(cases,'图3  同一模型的类别总体Recall与具体状态Recall；第四例保留未复现的正向差。',8.2)
    label('重点讲一个例子')
    para('COCO中，ResNet-50的dog总体Recall为81.92%，背向dog的46个目标Recall为57.25%，低24.67个百分点。Recall表示真实dog被正确识别为dog的比例。')
    table(['确认池','175条清单中可评价组合','解释'],[
        ['COCO','153','支持充分时可检查既有薄弱条件'],
        ['VOC2007','4','多数状态支持不足，不概括为稳定迁移'],
        ['Objects365','160；其中142个为负方向','组合会重叠，不是142个独立缺陷'],
    ],[3,6,8.4])
    speak('接着我把分层往条件诊断推进。例如dog总体Recall超过百分之八十，但背向dog只有百分之五十七左右。报告就可以提出明确的后续测试需求。我提前冻结条件，再到新池确认，减少只挑当前失败案例的问题。')

    page('04  任意随机分成五组，会得到同样差异吗？','随机划层对照：每池10000次，层大小和类别比例不变')
    picture(randfig,'图4  蓝色菱形为实际风险分层Q5−Q1；灰色条为固定池中随机划层的中央95%范围。',9.2)
    label('怎么看图')
    para('先看零线：随机分组差异在零附近波动。再看蓝色菱形：实际风险落差明显更负。COCO随机范围为−3.24至+3.22个百分点，实际落差约−21至−25个百分点。')
    para('三个池各10000次随机划层，没有达到实际落差的绝对值；六项Holm校正p约0.00060。这个范围描述随机分组波动，不是部署性能置信区间，也不证明评分最优。')
    speak('为了检查分层是否只是分组偶然性，我在每个类别内随机打乱并等分五层，重复一万次。随机落差基本围绕零，而真实风险层落差明显更大。这支持风险规则提供了与模型表现有关的分组信息。')
    transition('这种差异是不是必须依赖某一个风险字段？')

    page('05  结果是否过度依赖某一项风险字段？','逐项删除字段，重新评分并重新划层')
    picture(abfig,'图5  颜色与数字均表示Q1−Q5 F1差；所有变体保持正差，即Q5得分较低。',8.7)
    label('怎么看图')
    para('先读完整评分一行，再向下看删除某项后的落差。Objects365删除尺度后，落差从41.80/41.25缩小到32.37/27.99个百分点；删除遮挡后仍保留明显差异。')
    para('各池消融检验的Holm校正p为0.09375，未达到显著标准。这里支持方向保留，不能把删维前后变化解释为单维独立因果贡献，也不能据此验证标签准确率。')
    speak('我逐项去掉五个风险字段重新划层，检查结果是否依赖某个字段存在。高风险层较低的方向仍保留，尺度在部分来源中的影响较大。不过字段可能相关，成员也改变，所以我保留消融的不确定性。')
    transition('累计测试中数量和条件同时改变，能不能把数量固定再看？')

    page('06  数量不变，风险构成会怎样影响评分？','每次120目标，六类各20个；21个真实比例点，每点1000次抽样')
    picture(mixture,'图6  相对0% Q5起点的平均F1变化；每条线连接21个实测比例点，没有插入新的观测。',7.5)
    table(['来源','本地100%−0% Q5','线上100%−0% Q5'],[
        ['COCO','−20.78个百分点','−16.62个百分点'],
        ['VOC2007','−15.92个百分点','−12.29个百分点'],
        ['Objects365','−34.77个百分点','−31.84个百分点'],
    ],[4.2,6.6,6.6])
    label('怎么看图')
    para('横轴是高风险Q5占比，纵轴是相对起点的F1变化。测试量和类别比例固定，改变的只有成员与风险构成。0%来自Q1—Q4混合，不是单独Q1。')
    speak('我每次保持一百二十个目标和六类比例，只提高其中高风险目标的比例。总体F1终点下降约十二到三十五个百分点，部分模型名次也改变。这说明模型总分需要结合测试构成解释。模型没有继续训练，变化的是测得能力。')
    para('覆盖描述测到了哪些状态及证据量；风险构成描述相对高风险目标占比。此实验不能解释为“覆盖率提高导致F1下降”。',9,bold=True)

    page('07  最新对照：只按目标大小分层是否就够了？','连续官方框面积对照，不使用VLM状态；与完整评分保持同等配额')
    picture(scalefig,'图7  B=COCO、C=VOC2007、D=Objects365。右图为完整减面积的差；区间按架构/种子重抽。',9.1)
    label('怎么看图')
    para('左图两方案都能形成较大落差；右图看零线两侧，完整评分的增量并非都为正。COCO本地+2.12、线上−2.08；VOC本地+5.62、线上+3.48；Objects365为−2.09/−2.64个百分点。')
    para('完整评分的正增量均未通过六项Holm校正；Objects365面积方案更强的两项校正p为0.04688。不能仅靠更大的高低层差证明VLM必要。')
    speak('最新我增加了只按目标大小分层的简单对照。它本身就很强，多维评分没有普遍更好的端点区分。因此我进一步把状态信息的用途落到条件描述和诊断：它能否在相近粗尺度构成下仍指出薄弱条件？')
    para('“删除尺度”检查其他字段是否仍能区分；“只用尺度”检查便宜元数据本身能做到什么。这是两个不同问题。',9,bold=True)

    page('08  在粗尺度调整后，还有条件诊断信息吗？','沿用冻结假设，排除纯尺度条件；每类按面积五分位进行比较')
    picture(countfig,'图8  每池141条冻结非尺度组合；蓝色为负差、橙色为非负差、灰色为支持不足。',7.4)
    table(['池','可评价','负差','非负差','支持不足'],[[r[0],r[2],r[3],r[4],r[5]] for r in count],[4.2,3.3,3.3,3.3,3.3])
    label('怎么看图')
    para('COCO和Objects365多数可评价组合仍为负差，VOC主要是灰色支持不足。支持不足不等于没有风险；负差计数是重叠组合的描述性结果，不是独立缺陷数量。')
    speak('我在每个类别的面积五分位层中，比较某状态与同字段的已知非状态，并按支持量加权。COCO119个可评价组合中109个为负，Objects365122个中111个为负。这支持许多状态在粗尺度调整后仍有诊断关联。VOC只有四个可评价，不能概括为稳定复现。')
    para('面积五分位不等于连续面积精确匹配，其他条件也可能共现。因此当前结果是关联证据，不能认定某状态导致分类错误。',9,bold=True)

    page('09  用真实案例把状态、输入和预测连起来','原图上下文、分类器实际裁剪输入与自动标签的用途不同')
    picture(real,'图9  四个既有条件的真实目标和实际输入；前三例按三种子均错、第四例按均对的规则选择。',16.5)
    para('这组案例用于解释测试流程，不估计总体失败比例，也不是人工语义复核。Resize与中心裁剪可能丢失细长目标的信息，结果描述的是分类器加预处理流程。',9)
    speak('我还补了实际输入案例，帮助理解模型到底看到了什么。状态标注可以看到原图上下文，但分类器使用固定裁剪输入。因此错误不能全部归因于架构或某个状态，预处理也属于测试对象的一部分。')

    page('10  现在可以向老师明确交代的成果和边界','研究定位：自动视觉状态测试与条件诊断的实证研究')
    table(['已经完成','带来的证据或用途'],[
        ['五层测试与三池确认','风险层能揭示总体指标掩盖的条件差异'],
        ['冻结薄弱条件确认','支持充分的新来源可提供后续测试需求'],
        ['随机划层、删维和风险混合','检查分组波动、字段依赖及测试构成响应'],
        ['仅面积与面积层内对照','限定多维评分优势，明确条件诊断用途'],
        ['中英文论文、图表与补充材料','最新结果和结论边界已经同步'],
        ['便携缓存复算','已有独立复算记录：8张新表、60项历史比较一致'],
    ],[5.5,11.9])
    label('汇报中必须保留的边界')
    para('1. 分类真值来自数据集；自动状态没有完整人工语义真值。双VLM一致不等于正确。\n2. 条件差是关联，不是因果根源；面积调整没有控制所有混杂。\n3. Coverage收益具有条件性，没有证明普遍超过强基线。\n4. 重复抽样不是新增数据，三个种子不是三个独立架构。\n5. 缓存复算不等于从零重跑训练、VLM调用和原始聚类。')
    speak('这一阶段我按照您的建议完成了分层，又通过新样本、新来源和多种对照明确了它的作用。最有用的成果是把总体表现展开为可复用的条件报告，帮助指定进一步测试。论文和复算材料已更新，接下来重点是审阅贡献表述和完善投稿材料。')
    label('建议带到组会的材料')
    para('主讲本报告01—08节及第10节；真实案例、标签敏感性和Coverage基线放备用材料。最新论文：docs/论文_最新.docx；完整材料包：docs/visual_state_testing_submission_materials_20261007.zip。',9)

    page('附录A  可直接讲的连贯汇报稿','约5—7分钟；不需要逐个念全部数字')
    fulltext = [
        '老师，上次您建议我对测试数据做分层。这一阶段我先落实了这个建议，再围绕分层结果做了一系列验证。',
        '首先，我使用目标大小、身体可见性、遮挡、截断和朝向五个维度，在每个类别内部划分五个等数量风险层。分层不读取分类器是否答错，分类真值来自数据集，状态由程序和两套VLM提供。模型训练完成后保持冻结，所以后续指标变化反映测试条件变化。',
        '初步结果发现，高风险层的F1低于低风险层。为检查这是不是原样本的特殊现象，我固定分层规则，在新的COCO样本、VOC和Objects365中确认。三个确认池共五千四百个目标，六类均衡，八种架构各三个种子。结果在这些来源中重复出现，高低风险层的平均F1相差约十四到四十二个百分点。',
        '接着，我把整体分层往具体状态诊断推进。比如ResNet-50的dog总体Recall约百分之八十二，但背向dog约百分之五十七。这样的报告能指出下一步该测什么。我在发现池冻结薄弱条件，再到新池确认。COCO和Objects365支持充分时，多数条件保留薄弱方向；VOC很多状态样本太少，具体迁移证据有限。',
        '为了检查分层是否只是分组偶然性，我每池做了一万次类别平衡随机划层。随机差围绕零波动，明显小于实际风险层落差。我也逐项删除风险字段重新划层，高风险层较低的方向仍保留，但没有把消融解释成各字段的独立因果贡献。',
        '然后我固定测试总量和类别比例，只改变高风险目标占比。每次一百二十个目标、六类各二十个，共二十一个比例点。即使数量不变，风险构成变化仍对应平均F1下降约十二到三十五个百分点，部分模型名次也改变。这说明模型总分需要附带测试条件解释，不能把下降理解成模型训练后退化。',
        '最新我增加了只按目标大小分层的对照。这个简单方法也能形成很强的区分，多维评分没有普遍更强。因此我进一步检查尺度之外的条件信息。在粗面积层内比较后，COCO一百一十九个可评价组合中一百零九个仍有Recall负差，Objects365一百二十二个中一百一十一个为负。这支持状态报告用于条件诊断，但不证明因果根源或全部自动标签正确。',
        '目前的成果是：完成了风险分层，并通过新样本、新来源、随机和尺度对照明确了它的用途；把总体表现展开为可复用的薄弱条件报告。最新论文和复算材料已同步整理，接下来重点是审阅贡献表述和完善投稿材料。',
    ]
    for text in fulltext: para(text,11)

    page('附录B  老师可能追问什么？','先回答具体问题，再说明证据范围')
    qa = [
        ('你比上次多做了什么？','分层落实了上次建议。后续新增新样本/来源确认、冻结状态确认、随机与删维检查、固定预算混合，以及最新面积单独对照和面积层内诊断。'),
        ('F1下降为什么算进展？','进展在于发现并确认条件能力差异。冻结模型面对不同图片时测得分数改变；F1下降本身不能证明测试更好。'),
        ('只按大小就能分层，VLM有什么用途？','仅要端点区分时，面积很强。状态提供更具体的条件描述和可复用测试需求；当前未证明其成本效益普遍更好。'),
        ('你找到了模型错误原因吗？','发现的是关联的薄弱条件。状态会共现，尺度调整不完整，裁剪也会影响输入；没有隔离因果根源。'),
        ('状态标签可靠吗？','已做官方部分属性参考、双VLM分歧和敏感性分析，没有完整人工真值。结论限定在自动协议下。'),
        ('覆盖和风险有什么不同？','覆盖表示测到了哪些状态及证据量；风险表示预定义困难条件的构成。更多错误也不自动等于更多不同失败。'),
    ]
    for q,a in qa:label(q);para(a,10)

    page('附录C  数值与图表的来源','本报告复用既有实验；没有新增训练、推理或VLM标注')
    table(['内容','本工作目录下的证据'],[
        ['初始五层','sixclass_difficulty_supplement_20261004/tables/full_stratum_metrics.csv'],
        ['三池确认与随机','random_stratification_ablation_20261006/tables/random_summary.csv；ablation/*.npz'],
        ['具体条件案例','论文_最新.md 表10—11；submission_closeout_20261007/figures/04_condition_cases.svg'],
        ['逐项删维','random_stratification_ablation_20261006/tables/ablation_summary.csv'],
        ['固定120目标风险混合','controlled_risk_distribution_20261005/tables/panel_summary.csv、distribution_contrasts.csv'],
        ['连续面积对照','scale_only_confirmation_20261007/tables/summary.csv'],
        ['粗尺度内条件','scale_only_confirmation_20261007/tables/area_condition_summary.csv'],
        ['已有独立复算记录','submission_closeout_20261007/portable_replay_audit.json'],
    ],[4,13.4])
    para('上表未另写前缀的实验路径位于artifacts/；论文位于docs/。汇报起点“老师建议分层”依据用户回忆，其余工作依据聊天与结果文件，不推定为老师逐项要求。',9)
    para('口径提醒：Macro-F1是六类F1等权平均；表中的pp/百分点不是相对百分比。Q1−Q5是完整层比较；100%−0%风险混合是另一实验，0%来自Q1—Q4。自动属性对齐F1也不能当作分类F1。',9)
    para('图表编辑：材料包保留SVG、PNG与生成脚本。图中的数值直接来自保存结果；统计范围与适用边界在每页图注中说明。',9)

    word = ROOT / 'docs' / (NAME + '.docx')
    doc.save(word)
    shutil.copy2(word, OUT / (NAME + '.docx'))
    (ROOT/'docs'/(NAME+'.md')).write_text('\n'.join(md),encoding='utf-8')
    shutil.copy2(Path(__file__), OUT/'build_teacher_progress_word.py')
    # Preserve small primary tables alongside the figures and source.
    for p in inputs:
        if p.suffix == '.csv':shutil.copy2(p,TABLE/(p.parent.parent.name+'_'+p.name))
    audit = {'status':'pass','document':str(word),'figures':len(figures),
             'explicit_pages':14,'tables':len(doc.tables),'speaker_note_blocks':11,
             'checks':['24-run endpoint differences agree with random_summary',
                       'all plotted ablation endpoint gaps are positive',
                       'area-conditioned counts reconstructed from evaluable rows'],
             'sources':[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in dict.fromkeys(inputs)],
             'full_word_rendering':'pending'}
    with zipfile.ZipFile(word) as z:
        media=[n for n in z.namelist() if n.startswith('word/media/')]
        assert len(media)==9
        audit['embedded_images']=len(media)
    (OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'word':str(word),'figures':len(figures),'tables':len(doc.tables),'embedded_images':len(media)},ensure_ascii=False))


if __name__=='__main__':
    main()
