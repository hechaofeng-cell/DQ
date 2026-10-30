#!/usr/bin/env python3
"""Show each architecture separately, using saved six-class predictions only."""
from pathlib import Path
import csv
import hashlib
import json
import shutil
import zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from sklearn.metrics import f1_score
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxcompose.composer import Composer

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/model_results_report_20261008'
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
NAME = '八种模型结果_逐模型讲解版_20261008'
FULL = '工作量_逐图解释版_含八种模型结果_20261008'
OLD = ROOT / 'docs/工作量_逐图解释版_20261008.docx'
MODELS = {'resnet18':'ResNet-18', 'resnet50':'ResNet-50',
          'mobilenet_v3_small':'MobileNetV3-Small', 'mobilenet_v3_large':'MobileNetV3-Large',
          'densenet121':'DenseNet-121', 'efficientnet_b0':'EfficientNet-B0',
          'convnext_tiny':'ConvNeXt-Tiny', 'swin_t':'Swin-T'}
POOLS = {'B':('COCO','sixclass_independent_confirmation_20261004'),
         'C':('VOC2007','sixclass_voc_external_confirmation_20261005'),
         'D':('Objects365','objects365_independent_confirmation_20261005')}
CLASSES = ['dog','cat','horse','sheep','person','car']
SOURCE_LABELS = {'local':'本地AI描述后分组', 'online':'在线AI描述后分组'}
COLORS = {'B':'#3574A8', 'C':'#788D59', 'D':'#BD713E'}
MARKERS = {'B':'o','C':'s','D':'^'}
inputs, data, overall, source_rows = [], {}, {}, []

def read_csv(path):
    with path.open(encoding='utf-8') as f: return list(csv.DictReader(f))

for pool, (name, directory) in POOLS.items():
    manifest = ROOT / 'artifacts' / directory / 'manifests/candidate_pool.csv'
    inputs.append(manifest)
    targets = read_csv(manifest)
    truth = np.array([CLASSES.index(r['class_name']) for r in targets])
    sample_ids = [r['sample_id'] for r in targets]
    prev_predictions = None
    for source in SOURCE_LABELS:
        path = ROOT / f'artifacts/random_stratification_ablation_20261006/ablation/{pool}_{source}.npz'
        inputs.append(path)
        with np.load(path) as saved:
            assert saved['sample_ids'].tolist() == sample_ids
            assert saved['variants'][0] == 'full'
            predictions = saved['predictions'].copy()
            models = saved['models'].copy()
            seeds = saved['seeds'].copy()
            scores = saved['f1'][0].copy() * 100
            orders = saved['orders'][0].copy()
        if prev_predictions is not None: assert np.array_equal(predictions, prev_predictions)
        prev_predictions = predictions
        assert sorted(orders.ravel().tolist()) == list(range(len(targets)))
        for category in range(6): assert (truth[orders[category].ravel()] == category).all()
        for i, (model,seed) in enumerate(zip(models,seeds)):
            for tier in range(5):
                idx = orders[:,tier,:].ravel()
                direct = 100*f1_score(truth[idx], predictions[i,idx], average='macro', labels=range(6))
                assert abs(direct-scores[i,tier]) < 1e-9
                source_rows.append([pool,name,source,MODELS[model],model,int(seed),tier+1,len(idx),direct])
        for model in MODELS:
            idx = np.flatnonzero(models == model)
            assert len(idx) == 3
            assert len(set(seeds[idx].tolist())) == 3
            data[pool,source,model] = scores[idx]
            assert (scores[idx,0] > scores[idx,4]).all()
            if source == 'local':
                overall[pool,model] = np.array([100*f1_score(truth, predictions[i], average='macro', labels=range(6)) for i in idx])

with (OUT/'model_tier_scores.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['pool','dataset','label_source','model','model_id','seed','tier','targets','macro_f1_100']);w.writerows(source_rows)
with (OUT/'model_full_pool_scores.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['pool','dataset','model','mean_f1_100','minimum_of_3_runs','maximum_of_3_runs'])
    for model in MODELS:
        for pool,(name,_) in POOLS.items():
            vals=overall[pool,model]
            w.writerow([pool,name,MODELS[model],vals.mean(),vals.min(),vals.max()])

font_manager.fontManager.addfont(str(ROOT/'scripts/assets/NotoSansSC-Regular.ttf'))
plt.rcParams.update({'font.family':'Noto Sans SC','font.size':14,'svg.fonttype':'none',
    'axes.spines.top':False,'axes.spines.right':False,'axes.linewidth':.7,
    'xtick.major.size':0,'ytick.major.size':0})
figure_names=[]
def save(fig, stem):
    fig.savefig(FIG/f'{stem}.svg',bbox_inches='tight',facecolor='white')
    fig.savefig(FIG/f'{stem}.png',dpi=220,bbox_inches='tight',facecolor='white')
    plt.close(fig);figure_names.append(stem)

matrix=np.array([[overall[p,m].mean() for p in POOLS] for m in MODELS])
fig,ax=plt.subplots(figsize=(8.8,5.25),layout='constrained')
ax.imshow(matrix,cmap='Blues',vmin=0,vmax=100,aspect='auto')
ax.set(xticks=range(3),xticklabels=[v[0] for v in POOLS.values()],yticks=range(8),yticklabels=MODELS.values())
ax.tick_params(pad=10)
for i in range(8):
    for j in range(3): ax.text(j,i,f'{matrix[i,j]:.2f}',ha='center',va='center',fontsize=16,color='white')
for spine in ax.spines.values():spine.set_visible(False)
save(fig,'01_eight_models_overall')

for model,label in MODELS.items():
    fig,axes=plt.subplots(1,2,figsize=(9,4.45),layout='constrained',sharey=True)
    for ax,source in zip(axes,SOURCE_LABELS):
        endpoints=[]
        for pool,(name,_) in POOLS.items():
            runs=data[pool,source,model];mean=runs.mean(axis=0)
            ax.fill_between(range(1,6),runs.min(axis=0),runs.max(axis=0),color=COLORS[pool],alpha=.12,linewidth=0)
            ax.plot(range(1,6),mean,color=COLORS[pool],marker=MARKERS[pool],markersize=5,lw=2,label=name)
            endpoints.append((mean[4],pool))
        last_y=-100
        for score,pool in sorted(endpoints):
            label_y=max(score,last_y+8)
            ax.annotate(f'{score:.1f}',xy=(5,score),xytext=(5.18,label_y),fontsize=12,
                        va='center',color=COLORS[pool],arrowprops={'arrowstyle':'-','color':COLORS[pool],'lw':.65})
            last_y=label_y
        ax.set(title=SOURCE_LABELS[source],xlim=(.85,5.85),ylim=(0,105),yticks=[0,25,50,75,100],xticks=range(1,6),xlabel='组别：1 → 5')
    axes[0].set_ylabel('F1综合评分（满分100）')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='outside lower center',ncol=3,frameon=False,fontsize=12)
    save(fig,f'02_{model}_five_groups')

doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.9)
sec.top_margin=sec.bottom_margin=Cm(1.6)
sec.header_distance=sec.footer_distance=Cm(.65)
normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(11)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
normal.paragraph_format.line_spacing=1.15
normal.paragraph_format.space_after=Pt(5)
md=[]
def runfont(run,size=11,bold=False,color='263442'):
    run.font.name='Noto Sans CJK SC';run.font.size=Pt(size);run.font.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
def para(text,size=11,bold=False,color='263442'):
    pp=doc.add_paragraph();runfont(pp.add_run(text),size,bold,color);md.append(text+'\n');return pp
def heading(text):
    pp=para(text,12.5,True,'24577E');pp.paragraph_format.keep_with_next=True
    md[-1]='## '+text+'\n'
def page(title,subtitle):
    if doc.paragraphs:doc.add_page_break()
    pp=para(title,18,True,'24577E');pp.paragraph_format.keep_with_next=True
    md[-1]='# '+title+'\n';para(subtitle,9.5,color='627180')
def picture(stem,caption):
    pp=doc.add_paragraph();pp.paragraph_format.keep_with_next=True;pp.paragraph_format.space_after=Pt(0)
    pp.add_run().add_picture(str(FIG/f'{stem}.png'),width=Cm(16.7))
    para(caption,9,color='627180');md.append(f'![{caption}]({FIG/f"{stem}.png"})\n')
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.space_before=Pt(3)
                for run in pp.runs:runfont(run,9.5,i==0)
            if i==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EAF1F7');c._tc.get_or_add_tcPr().append(shd)
    md.append('| '+' | '.join(headers)+' |\n|'+'|'.join(['---']*len(headers))+'|\n')
    md.extend('| '+' | '.join(map(str,r))+' |\n' for r in rows)

runfont(sec.header.paragraphs[0].add_run('研究工作回顾  ·  各模型成绩与图表解释'),8,color='627180')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runfont(footer.add_run('2026.10.08  |  '),8,color='627180')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

page('8种模型，分别看成绩', '补充到进度汇报中：每种模型单独展示，数据来自已经保存的最新核验结果。')
para('8个模型名字代表8种图片分类模型，每种都有3份已完成的训练结果，共24份。它们负责判断狗、猫、马、羊、人和车；两套AI负责描述状态，是另一项任务。')
picture('01_eight_models_overall','图M1  在每批全部目标上计算的F1，再对同一模型的3份训练结果取平均；满分100，越高越好。')
heading('先看这个总览，再看后面8页')
para('每行是一种模型，每列是一个图片来源。这里计算的是整批目标的总体分，不是把五组F1简单平均，也不受采用哪套AI分组影响。')
heading('当前结果可以怎样讲？')
best=max(MODELS,key=lambda m:overall['B',m].mean())
assert best=='convnext_tiny'
assert all(max(MODELS,key=lambda m:overall[p,m].mean())=='convnext_tiny' for p in POOLS)
assert all(max(MODELS,key=lambda m:data[p,s,m].mean(axis=0)[4])=='convnext_tiny' for p in POOLS for s in SOURCE_LABELS)
para('“我比较了8种模型。当前三批目标上，ConvNeXt-Tiny的平均总体F1最高。但总体高不代表各类图片都一样好，后面要看每种模型在五组中的具体变化。”')
para('这只是本任务和已测目标上的数值比较，没有据此证明某模型在所有任务都最好，也没有进行模型优劣的显著性检验。',9,color='627180')

for number,(model,label) in enumerate(MODELS.items(),2):
    page(f'{number-1}  {label} 的具体结果', '一个模型一页：展示三种图片来源、两套状态描述和3份训练结果的波动。')
    picture(f'02_{model}_five_groups',f'图M{number}  实线是该模型3份训练结果的平均；浅色带是这3份结果的最小到最大范围，不是置信区间。线尾数字为第5组平均分。')
    heading('图怎么看？')
    para('左图按本地AI描述分组，右图按在线AI描述分组；它们测的是同一模型、同一批目标，分组成员可能不同。蓝线=COCO，绿线=VOC2007，橙线=Objects365。横轴1到5表示规则预计从较容易到较困难，纵轴是F1综合分。',10.5)
    table(['图片来源','总体F1','本地\n第1组','本地\n第5组','在线\n第1组','在线\n第5组'],[
        [name,f'{overall[pool,model].mean():.2f}',
         f'{data[pool,"local",model].mean(axis=0)[0]:.2f}',f'{data[pool,"local",model].mean(axis=0)[4]:.2f}',
         f'{data[pool,"online",model].mean(axis=0)[0]:.2f}',f'{data[pool,"online",model].mean(axis=0)[4]:.2f}']
        for pool,(name,_) in POOLS.items()])
    heading('这个模型的结果说明什么？')
    b=data['B','local',model].mean(axis=0);d=data['D','local',model].mean(axis=0)
    para(f'以COCO、本地描述为例，第1组{b[0]:.1f}分，第5组{b[4]:.1f}分，相差{b[0]-b[4]:.1f}分。换到Objects365，第5组为{d[4]:.1f}分。两套分组下，三种来源的第5组都低于第1组；中间各组不要求逐步下降。',10.5)
    heading('可以这样向老师讲')
    para(f'“{label}在COCO全部目标上的平均F1是{overall["B",model].mean():.1f}分，但较困难组只有{b[4]:.1f}分。把它单独展开，就能看到总体成绩没有反映出来的条件差别。”',10.5)
    para('同一个分组中还可以比较不同模型；组间分差较小不自动表示模型更强，要同时看总体分和困难组的实际分数。',9,color='627180')

doc.core_properties.title='8种模型的结果与逐模型解释'
doc.core_properties.subject='六分类模型的总体F1、五组F1和三次训练波动'
doc.save(OUT/f'{NAME}.docx')
shutil.copy2(OUT/f'{NAME}.docx',ROOT/'docs'/f'{NAME}.docx')
(OUT/f'{NAME}.md').write_text('\n'.join(md),encoding='utf-8')
shutil.copy2(OUT/f'{NAME}.md',ROOT/'docs'/f'{NAME}.md')
assert len(doc.inline_shapes)==9
assert len(doc.element.xpath('//w:br[@w:type="page"]'))==8

# Start the complete report with the models, then retain all earlier explanations.
doc.add_page_break()
composer=Composer(doc)
composer.append(Document(OLD))
composer.doc.core_properties.title='工作量逐图解释：含8种模型结果'
composer.save(OUT/f'{FULL}.docx')
shutil.copy2(OUT/f'{FULL}.docx',ROOT/'docs'/f'{FULL}.docx')
full_md='\n'.join(md)+'\n\n'+OLD.with_suffix('.md').read_text()
(ROOT/'docs'/f'{FULL}.md').write_text(full_md,encoding='utf-8')
assert len(composer.doc.inline_shapes)==20
shutil.copy2(Path(__file__),OUT/Path(__file__).name)
audit={'models':MODELS,'source_files':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
       'direct_prediction_checks':720,'architectures':8,'runs_per_architecture':3,
       'model_figures':figure_names,'standalone_planned_pages':9,'full_planned_pages':21,
       'prediction_label_order_verified':True,'new_training_or_inference':False,'render_verified':False}
(OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
(OUT/'说明.md').write_text('图M1是完整候选池总体F1。其余8图每种模型一张，包含两套状态描述、三个来源、每组3份训练结果；浅色带是三份结果最小到最大范围。数据来自已有预测，不新增训练或推理。720个分组F1已由官方类别清单与保存预测逐格重新核对。生成脚本依赖numpy、matplotlib、scikit-learn、python-docx和docxcompose。完整Word前9页为模型结果，后12页保留原工作量逐图解释。\n',encoding='utf-8')
with zipfile.ZipFile(ROOT/'docs/八种模型结果_图表与数据_20261008.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in FIG.glob('*'):z.write(p,p.relative_to(OUT))
    for name in ['model_tier_scores.csv','model_full_pool_scores.csv','audit.json','说明.md',Path(__file__).name]:z.write(OUT/name,name)
    for p in inputs:z.write(p,Path('source_data')/p.relative_to(ROOT))
    z.write(ROOT/'scripts/assets/NotoSansSC-Regular.ttf','source_data/scripts/assets/NotoSansSC-Regular.ttf')
    z.write(OLD,Path('source_data')/OLD.relative_to(ROOT))
    z.write(OLD.with_suffix('.md'),Path('source_data')/OLD.with_suffix('.md').relative_to(ROOT))
    z.write(Path(__file__),'source_data/scripts/'+Path(__file__).name)
print(json.dumps({'model_word':str(ROOT/'docs'/f'{NAME}.docx'),'full_word':str(ROOT/'docs'/f'{FULL}.docx'),
                  'models':8,'new_figures':9,'checked_tier_f1_cells':720},ensure_ascii=False))
