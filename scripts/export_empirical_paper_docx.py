#!/usr/bin/env python3
"""Export the empirical-study Markdown draft to an editable Word document."""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


def set_run_font(run, western: str, east_asia: str, size: float, bold: bool | None = None) -> None:
    run.font.name = western
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("`", "")
    return text


def add_inline(paragraph, text: str, size: float = 10.5) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\[[^]]+\]\([^)]+\))", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else clean_inline(part)
        run = paragraph.add_run(value)
        set_run_font(run, "Consolas" if code else "Times New Roman", "宋体", size, bold)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document: Document, raw_rows: list[str]) -> None:
    parsed = [[clean_inline(cell.strip()) for cell in row.strip().strip("|").split("|")] for row in raw_rows]
    if len(parsed) > 1 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in parsed[1]):
        parsed.pop(1)
    if not parsed:
        return
    columns = max(len(row) for row in parsed)
    table = document.add_table(rows=len(parsed), cols=columns)
    table.style = "Table Grid"
    for row_index, values in enumerate(parsed):
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            cell.text = value
            if row_index == 0:
                shade_cell(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, "Times New Roman", "宋体", 8.5, row_index == 0)
    document.add_paragraph()


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def export(markdown_path: Path, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5), ("Heading 4", 10.5)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.8)
                paragraph.paragraph_format.space_after = Pt(6)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, "Consolas", "等线", 8.5)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line.strip())
        if image_match:
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            if image_path.is_file():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.2))
                caption = document.add_paragraph(image_match.group(1))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in caption.runs:
                    set_run_font(run, "Times New Roman", "宋体", 9)
            index += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1 and not document.paragraphs:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text)
            else:
                document.add_heading(text, level=level)
            index += 1
            continue
        if line.strip() == "---":
            document.add_section(WD_SECTION.NEW_PAGE)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue
        if line.strip() == "\\[":
            equation = []
            index += 1
            while index < len(lines) and lines[index].strip() != "\\]":
                equation.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(" ".join(equation))
            set_run_font(run, "Cambria Math", "宋体", 10)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r"^(#{1,4})\s|^\||^```|^- |^\d+\.\s|^!\[|^\\\[$", lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        paragraph.paragraph_format.line_spacing = 1.35
        paragraph.paragraph_format.space_after = Pt(4)
        add_inline(paragraph, " ".join(paragraph_lines))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    export(args.markdown.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
