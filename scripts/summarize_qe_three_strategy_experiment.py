"""Add interpretable conclusions and source composition to completed comparisons."""
import json
from pathlib import Path

import numpy as np
from docx import Document

from run_qe_three_strategy_experiment import SOURCE,OUT,DATA,CLASSES,BUDGETS,MODELS,load_data
from build_qe_visualstate_coverrisk_dataset import read_csv,write_csv,write_json,sha


def main():
    rows=read_csv(OUT/'summary.csv');runs=read_csv(OUT/'model_run_metrics.csv');tests=read_csv(OUT/'primary_tests.csv')
    master={r['sample_id']:r for r in read_csv(DATA/'manifest.csv')}
    pools=sorted({r['source_pool'] for r in master.values()});composition=[]
    for source in ('local','online'):
        arrays=load_data(source)
        lookup={c:np.array([master[str(s)]['source_pool'] for s in arrays[c]['sample_ids']]) for c in CLASSES}
        for budget in BUDGETS:
            selections=dict(np.load(OUT/'selection'/f'{source}_n{budget*6}.npz'))
            for method,index in selections.items():
                if index.ndim==2:index=index[None]
                for pool in pools:
                    count=sum((lookup[c][index[:,i]]==pool).sum(1) for i,c in enumerate(CLASSES))
                    composition.append(dict(source=source,budget_per_class=budget,method=method,source_pool=pool,
                                             mean_targets=float(count.mean()),mean_fraction=float(count.mean()/(6*budget)),
                                             target_p025=float(np.quantile(count,.025)),target_p975=float(np.quantile(count,.975))))
    write_csv(OUT/'source_composition.csv',composition)
    model_summary=[];recall=[]
    for source in ('local','online'):
        for method in ('coverage','risk','risk_coverage'):
            for model in MODELS:
                selected=[r for r in runs if r['source']==source and r['method']==method and r['model']==model and int(r['budget_per_class'])==40]
                assert len(selected)==3
                model_summary.append(dict(source=source,method=method,model=model,budget_per_class=40,
                                          **{key:float(np.mean([float(r[key]) for r in selected])) for key in ('accuracy','macro_f1','errors','confusion_type_count')}))
            selected=[r for r in rows if r['source']==source and r['method']==method]
            for row in selected:
                for c in CLASSES:recall.append(dict(source=source,method=method,budget_per_class=int(row['budget_per_class']),category=c,recall=float(row[f'{c}_recall'])))
    write_csv(OUT/'model_summary_primary_budget.csv',model_summary)
    write_csv(OUT/'category_recall_summary.csv',recall)
    deltas=[]
    for source in ('local','online'):
        for budget in BUDGETS:
            lookup_row={r['method']:r for r in rows if r['source']==source and int(r['budget_per_class'])==budget}
            left,right=lookup_row['risk_coverage'],lookup_row['risk']
            test=next(r for r in tests if r['source']==source and int(r['budget_per_class'])==budget)
            for key in ('risk_lower','risk_midpoint','risk_upper','risk_unknown_fraction'):
                assert np.isclose(float(left[key]),float(right[key]),atol=1e-10,rtol=0)
            deltas.append(dict(source=source,total_targets=budget*6,budget_per_class=budget,
                              coverage_delta_pp=100*(float(left['presence_coverage'])-float(right['presence_coverage'])),
                              support_delta_pp=100*(float(left['support_completion'])-float(right['support_completion'])),
                              other_label_support_delta_pp=100*(float(left['other_label_support_completion'])-float(right['other_label_support_completion'])),
                              macro_f1_delta_pp=100*(float(left['macro_f1'])-float(right['macro_f1'])),
                              additional_errors=float(left['errors'])-float(right['errors']),
                              additional_confusion_types=float(left['confusion_type_count'])-float(right['confusion_type_count']),
                              equal_error_breadth_delta_pp=100*float(test['mean_equal_error_delta']),
                              holm_p=float(test['holm_six_p'])))
    write_csv(OUT/'combo_vs_risk_deltas.csv',deltas)
    primary={r['source']:r for r in deltas if r['budget_per_class']==40}
    larger={r['source']:r for r in deltas if r['budget_per_class']==100}
    findings=['','## 本轮结果解释','',
              '在所有六个来源—预算设置下，风险＋覆盖组相对共同风险配额随机组提高了支持完成度，也发现更多错误与更多原始混淆类型。两组平均风险分区间和未知比例一致，因此这一差异不是由协议所测风险均值变化产生；未测语义难度及源数据集构成仍可不同。source_composition.csv保存各组四个来源占比，不宣称仅操纵覆盖率的因果效应。','',
              f"主预算240目标：local出现覆盖提高{primary['local']['coverage_delta_pp']:.2f}百分点，支持完成度提高{primary['local']['support_delta_pp']:.2f}百分点；online对应提高{primary['online']['coverage_delta_pp']:.2f}和{primary['online']['support_delta_pp']:.2f}百分点。相同错误数量后的混淆广度平均增益分别为{primary['local']['equal_error_breadth_delta_pp']:.2f}与{primary['online']['equal_error_breadth_delta_pp']:.2f}百分点，两来源均7/8架构方向为正，但六项校正后p={primary['local']['holm_p']:.4f}，未达到0.05。因此主预算只能报告正向趋势，不能宣称稳健优越。",'',
              f"600目标：local与online支持完成度分别提高{larger['local']['support_delta_pp']:.2f}与{larger['online']['support_delta_pp']:.2f}百分点，等错误数混淆广度提高{larger['local']['equal_error_breadth_delta_pp']:.2f}与{larger['online']['equal_error_breadth_delta_pp']:.2f}百分点，两来源均8/8架构方向为正，Holm p={larger['local']['holm_p']:.6f}。在这一已测预算与混淆类型代理指标上，结果支持组合选样具有额外诊断收益；120目标未显示可靠的等错误数广度增益。",'',
              '无风险约束的覆盖组达到更高的原始覆盖，但它的风险分和Q5比例也更高，不能把其F1下降单独归因于覆盖增加。风险配额随机组测得F1最高并不表示它是最有效的缺陷发现策略；组合组发现更多错误也不表示模型参数退化。','',
              '风险配额在本轮固定为五层各20%，尚未扫描其他风险混合比例；该实验不能替代此前不同风险比例的响应实验。没有对训练数据进行补充或微调，因此没有训练改善结论。']
    path=OUT/'report.md';text=path.read_text();marker='\n## 本轮结果解释'
    if marker in text:text=text.split(marker)[0]
    text=text.rstrip()+'\n'+'\n'.join(findings)+'\n';path.write_text(text)
    lines=text.splitlines();doc=Document();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            cells=[[v.strip() for v in row.strip('|').split('|')] for row in block]
            table=doc.add_table(rows=1,cols=len(cells[0]));table.style='Table Grid'
            for cell,v in zip(table.rows[0].cells,cells[0]):cell.text=v
            for row in cells[2:]:
                for cell,v in zip(table.add_row().cells,row):cell.text=v
            continue
        if line.startswith('#'):doc.add_heading(line.lstrip('#').strip(),level=min(len(line)-len(line.lstrip('#')),3)-1)
        elif line:doc.add_paragraph(line)
        i+=1
    doc.save(OUT/'report.docx')
    evidence=json.loads((OUT/'claim_evidence.json').read_text())
    evidence['matched_risk_verified']=dict(source='raw/ condition_risk*_risk_{lower,midpoint,upper,unknown_fraction}',status='all6settings identical')
    evidence['primary_budget_diagnostic_gain']=dict(source='primary_tests.csv',status='positive direction, Holm p .0625; not significant')
    evidence['larger_budget_diagnostic_gain']=dict(source='primary_tests.csv',status='600-targets only; both sources Holm p .046875; confusion-type proxy only')
    write_json(OUT/'claim_evidence.json',evidence)
    freeze={str(Path(__file__).resolve()):sha(__file__)}
    for name in ['summary.csv','primary_tests.csv','source_composition.csv','model_summary_primary_budget.csv','category_recall_summary.csv','combo_vs_risk_deltas.csv','report.md','report.docx','claim_evidence.json']:
        freeze[str(OUT/name)]=sha(OUT/name)
    for p in (OUT/'figures').glob('*'):freeze[str(p)]=sha(p)
    write_json(OUT/'final_result_freeze.sha256.json',freeze)
    print(json.dumps(dict(status='summarized',settings=len(deltas),source_composition_rows=len(composition),model_summary_rows=len(model_summary)),ensure_ascii=False))


if __name__=='__main__':main()
