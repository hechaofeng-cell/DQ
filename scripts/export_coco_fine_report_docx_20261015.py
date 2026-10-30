#!/usr/bin/env python3
"""Package the audited Chinese result report and its figures as a Word document."""
import argparse
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def clean(text):
    return text.replace('**','').replace('`','')


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,required=True)
    root=parser.parse_args().input.resolve()
    assert (root/'independent_result_audit.json').exists()
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(1.8)
    section.left_margin=section.right_margin=Cm(1.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.line_spacing=1.2
    normal.paragraph_format.space_after=Pt(5)
    for name in ('Title','Heading 1','Heading 2','Heading 3'):
        doc.styles[name].font.name='Calibri'
        doc.styles[name].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
    lines=(root/'report_zh.md').read_text().splitlines();i=0;figures=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[clean(c.strip()) for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells):block.append(cells)
                i+=1
            table=doc.add_table(rows=1,cols=len(block[0]));table.style='Table Grid'
            for cell,text in zip(table.rows[0].cells,block[0]):cell.text=text
            header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
            for row in block[1:]:
                for cell,text in zip(table.add_row().cells,row):cell.text=text
            for row in table.rows:
                no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:r.font.size=Pt(7.5 if len(block[0])>=6 else 9)
            doc.add_paragraph();continue
        image=re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if image:
            path=root/image.group(2);assert path.exists()
            doc.add_picture(str(path),width=Cm(17.2));figures+=1
        elif line.startswith('# '):doc.add_heading(clean(line[2:]),0)
        elif line.startswith('## '):doc.add_heading(clean(line[3:]),1)
        elif line.startswith('### '):doc.add_heading(clean(line[4:]),2)
        elif line.startswith('- '):doc.add_paragraph(clean(line[2:]),style='List Bullet')
        else:doc.add_paragraph(clean(line))
        i+=1
    footer=section.footer.paragraphs[0];footer.alignment=2
    footer.add_run('COCO 单源实验 · 2026-10-15　')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=root/'report_zh.docx';doc.save(path)
    check=Document(path)
    assert len(check.inline_shapes)==7 and figures==7
    assert len(check.tables)==3
    print(f'Word report: {path}; 7 embedded figures, 3 result tables.')


if __name__=='__main__':main()
