#!/usr/bin/env python3

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SOURCE = Path(
    "/home/hcf/.codex/attachments/597e8f55-1e62-454b-ae06-d453d48c10ce/"
    "QE_group_meeting_visual_test_adequacy_20260915555555.pptx"
)
OUTPUT = Path(
    "/home/hcf/project/QE/docs/"
    "QE_group_meeting_visual_test_adequacy_20260916_final_with_summary.pptx"
)


COLORS = {
    "bg": "F7F9FB",
    "white": "FFFFFF",
    "ink": "17212B",
    "muted": "5F6B76",
    "line": "D8E0E6",
    "navy": "153B5B",
    "blue": "2E6F9E",
    "blue_soft": "E7F0F7",
    "teal": "187D7B",
    "teal_soft": "E5F3F1",
    "orange": "E87524",
    "orange_soft": "FCECDD",
    "green": "2F7D5A",
    "green_soft": "E7F2EC",
    "red": "B84A3A",
}


def color(value):
    return RGBColor.from_string(COLORS.get(value, value))


def add_text(slide, text, x, y, w, h, *, size=18, bold=False, fill="ink",
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font="Microsoft YaHei"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(fill)
    return box


def add_card(slide, x, title, items, accent, soft, number):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.98), Inches(3.62), Inches(3.12)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = color("white")
    card.line.color.rgb = color("line")
    card.line.width = Pt(1)

    stripe = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.98), Inches(3.62), Inches(0.09)
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = color(accent)
    stripe.line.color.rgb = color(accent)

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + 0.25), Inches(2.28), Inches(0.43), Inches(0.43)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = color(accent)
    badge.line.color.rgb = color(accent)
    add_text(slide, number, x + 0.25, 2.29, 0.43, 0.39, size=14, bold=True,
             fill="white", align=PP_ALIGN.CENTER, font="Aptos")

    add_text(slide, title, x + 0.83, 2.25, 2.45, 0.48, size=20, bold=True, fill=accent)

    body = slide.shapes.add_textbox(
        Inches(x + 0.28), Inches(2.93), Inches(3.06), Inches(1.86)
    )
    frame = body.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.text = f"•  {item}"
        paragraph.space_after = Pt(12)
        paragraph.line_spacing = 1.08
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(15)
            run.font.bold = index == len(items) - 1
            run.font.color.rgb = color("ink")

    tag = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x + 0.28), Inches(4.61), Inches(3.06), Inches(0.31)
    )
    tag.fill.solid()
    tag.fill.fore_color.rgb = color(soft)
    tag.line.color.rgb = color(soft)


def main():
    presentation = Presentation(SOURCE)
    slide = presentation.slides.add_slide(presentation.slides[-1].slide_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color("bg")

    add_text(slide, "07 · SUMMARY", 0.62, 0.28, 3.1, 0.25, size=10, bold=True,
             fill="teal", font="Aptos")
    add_text(slide, "总结：从覆盖诊断走向可验证的数据行动", 0.62, 0.62, 12.0, 0.58,
             size=29, bold=True, fill="navy")
    add_text(
        slide,
        "当前完成了状态测量与缺口驱动补图；更强结论仍需要公平基线和独立效用验证",
        0.64, 1.20, 11.9, 0.34, size=13, fill="muted"
    )
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.64), Inches(12.05), Inches(0.012)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = color("line")
    rule.line.color.rgb = color("line")

    add_card(
        slide, 0.72, "已经完成",
        ["固定目标与三视图输入", "23项视觉状态和严格校验", "缺口驱动选择100张补图"],
        "blue", "blue_soft", "1"
    )
    add_card(
        slide, 4.85, "当前证据",
        ["覆盖指数：78.64% → 84.17%", "低频状态：45 → 33", "只支持状态覆盖改善"],
        "orange", "orange_soft", "2"
    )
    add_card(
        slide, 8.98, "下一步验证",
        ["人工一致性与标签有效性", "等预算随机/分层基线", "固定测试集上的 ΔM 与故障检出"],
        "green", "green_soft", "3"
    )

    takeaway = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.86), Inches(5.45), Inches(11.61), Inches(0.92)
    )
    takeaway.fill.solid()
    takeaway.fill.fore_color.rgb = color("navy")
    takeaway.line.color.rgb = color("navy")
    add_text(slide, "TAKE-HOME", 1.16, 5.61, 1.55, 0.24, size=11, bold=True,
             fill="teal_soft", font="Aptos")
    add_text(
        slide,
        "覆盖的价值不是替代模型性能，而是指出“哪些测试证据还没有”，并指导有限预算下的补样。",
        2.55, 5.57, 9.45, 0.38, size=19, bold=True, fill="white", align=PP_ALIGN.CENTER
    )
    add_text(
        slide,
        "当前结论边界：流程已运行、状态覆盖已改善；策略优越性和模型收益仍待验证。",
        1.05, 6.57, 10.95, 0.30, size=14.5, bold=True, fill="red", align=PP_ALIGN.CENTER
    )
    add_text(slide, "QE · SUMMARY", 0.62, 7.05, 1.4, 0.18, size=8.5, bold=True,
             fill="muted", font="Aptos")

    slide.notes_slide.notes_text_frame.text = (
        "总结时按三列讲。第一，当前已经跑通固定目标、23项状态提取、严格校验和缺口驱动补图。"
        "第二，已有证据是覆盖指数由78.64%提高到84.17%，低频状态由45降到33；这只支持状态覆盖改善。"
        "第三，下一步验证标签有效性、等预算基线、固定测试集上的模型变化和故障检出能力。"
        "最后读出Take-home message：覆盖不是模型性能的替代品，而是用于发现仍然缺少的测试证据。"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
