"""Compare coverage, risk-quota random, and quota-constrained coverage tests."""
from __future__ import annotations

import argparse
import itertools
import json
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from scipy.special import gammaln
from docx import Document

from build_qe_visualstate_coverrisk_dataset import read_csv, write_csv, write_json, sha

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data/QE_VisualState_CoverRisk_v1_20261013'
SOURCE=ROOT/'artifacts/QE_VisualState_CoverRisk_v1_20261013'
OUT=ROOT/'artifacts/QE_ThreeStrategy_Test_20261013'
CLASSES=('dog','cat','horse','sheep','person','car')
MODELS=('resnet18','resnet50','mobilenet_v3_small','mobilenet_v3_large','densenet121','efficientnet_b0','convnext_tiny','swin_t')
SEEDS=(20260916,20260917,20260918)
BUDGETS=(20,40,100)
REPEATS=1000
BASE_SEED=20261013


def note(stage, **extra):
    from datetime import datetime
    from zoneinfo import ZoneInfo
    row=dict(stage=stage,time=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),**extra)
    write_json(OUT/'progress.json',row)
    with (OUT/'research_log.md').open('a') as f:f.write('\n'+json.dumps(row,ensure_ascii=False)+'\n')
    print(json.dumps(row,ensure_ascii=False),flush=True)


def load_data(source):
    root=SOURCE if source=='local' else SOURCE/'online_analysis'
    return {c:dict(np.load(root/'prepared'/f'{c}.npz',allow_pickle=False)) for c in CLASSES}


def largest_remainder(cells,budget):
    total=sum(len(v) for v in cells.values())
    # Integer arithmetic preserves mathematical ties in the remainder ordering.
    numerator={k:budget*len(v) for k,v in cells.items()}
    quota={k:v//total for k,v in numerator.items()}
    remainder=budget-sum(quota.values())
    for k in sorted(cells,key=lambda k:(-(numerator[k]%total),k))[:remainder]:quota[k]+=1
    assert sum(quota.values())==budget and all(quota[k]<=len(cells[k]) for k in cells)
    return quota


def structure(z,budget):
    n=len(z['sample_ids']);assert n==1300 and budget%5==0
    order=sorted(range(n),key=lambda i:(z['risk_midpoint'][i],str(z['sample_ids'][i])))
    tiers=np.empty(n,dtype=int)
    for t,portion in enumerate(np.array_split(order,5),1):tiers[portion]=t
    unknown=np.rint((z['risk_upper']-z['risk_lower'])/20).astype(int)
    assert np.allclose(unknown*20,z['risk_upper']-z['risk_lower'])
    cells=defaultdict(list)
    for i in range(n):cells[(int(tiers[i]),float(z['risk_midpoint'][i]),int(unknown[i]))].append(i)
    quota={}
    for tier in range(1,6):
        groups={k:v for k,v in cells.items() if k[0]==tier}
        quota.update(largest_remainder(groups,budget//5))
    assert sum(quota.values())==budget
    group_by_idx={i:k for k,v in cells.items() for i in v}
    return dict(tiers=tiers,unknown=unknown,cells=dict(cells),quota=quota,group_by_idx=group_by_idx)


def greedy(z,budget,risk_structure=None):
    matrix=z['matrix'].astype(bool)
    counts=np.zeros(matrix.shape[1],dtype=int);available=np.ones(len(matrix),dtype=bool)
    selected=[];used=Counter();audit=[]
    for step in range(budget):
        feasible=np.flatnonzero(available)
        if risk_structure is not None:
            feasible=np.array([i for i in feasible if used[risk_structure['group_by_idx'][int(i)]]<risk_structure['quota'][risk_structure['group_by_idx'][int(i)]]])
        assert len(feasible)
        weights=np.where(counts<20,np.log(counts+2)-np.log(counts+1),0)
        gains=matrix[feasible]@weights
        best=float(gains.max())
        ties=feasible[np.isclose(gains,best,atol=1e-12,rtol=0)]
        chosen=int(min(ties,key=lambda i:str(z['sample_ids'][i])))
        before=np.minimum(counts/20,1).mean()
        selected.append(chosen);available[chosen]=False;counts+=matrix[chosen]
        if risk_structure is not None:used[risk_structure['group_by_idx'][chosen]]+=1
        audit.append(dict(rank=step+1,sample_id=str(z['sample_ids'][chosen]),concave_gain=best,
                          frequency_coverage_gain=float(np.minimum(counts/20,1).mean()-before),
                          presence_coverage=float((counts>0).mean()),support_completion=float(np.minimum(counts/20,1).mean())))
    assert len(set(selected))==budget
    if risk_structure is not None:assert used==Counter({k:v for k,v in risk_structure['quota'].items() if v})
    return np.array(selected,dtype=np.int16),audit


def selection():
    OUT.mkdir(parents=True,exist_ok=True);(OUT/'selection').mkdir(exist_ok=True)
    assert json.loads((SOURCE/'final_audit.json').read_text())['status']=='pass'
    protocol=dict(dataset=str(DATA),candidate_targets=7800,classes=list(CLASSES),primary_source='local',sensitivity_source='online',
        budgets_per_class=list(BUDGETS),primary_budget_per_class=40,risk_random_repeats=REPEATS,seed=BASE_SEED,
        methods=['coverage','risk','risk_coverage'],coverage_gain='capped concave sum log(n+2)-log(n+1) for n<20',
        coverage_dictionary='all71 nonunknown shared15 states per category,426fixed total; unreachable states remain',
        coverage_support_requirement=20,risk_score='existing equal_global five equal weights, [0,100], intervals retained',
        tiers='per-source pooled-class midpoint score then sample_id; five equal260-target tiers',
        risk_quota='each tier budget/5; within-tier midpoint-score/unknown-count joint cells proportional largest remainder using exact integer numerator/remainders, ties cell key',
        coverage='greedy full pool, same class budget, no risk constraints',
        risk='1000 unique random selections per class, obey exact common risk cell quotas, no coverage input',
        risk_coverage='greedy globally across still-feasible cells under same exact quotas as risk random',
        random_design='paired source-specific RNG; repetitions are selection variation, not independent trained models',
        frozen_models=dict(models=list(MODELS),seeds=list(SEEDS),retraining=False,cache_reuse=True),
        stopping='each method ends at common budget; risk quota equality asserted; zero gain recorded but no unequal-size early exit',
        primary_tests='risk_coverage vs risk: per-class equal-error expected confusion breadth, primary budget40; Holm6 source-budget tests',
        failure_types='wrong (official true class, predicted class) pairs among 30possible; no feature-cluster proxy',
        annotation_accuracy='existing automatic labels; no new human validation',
        claims='selection affects measured metrics; lower F1 or more errors alone does not show better tests')
    write_json(OUT/'protocol.json',protocol)
    (OUT/'research_log.md').write_text('# Three-strategy comparison\n\nFrozen factors: merged pool, state universe, score version, per-class budgets, classifier checkpoints and preprocessing. Primary question: whether quota-constrained coverage adds diagnostic breadth over same-quota random. All declared budgets and both annotation sources reported; no result-dependent replacement or new model search.\n')
    freeze={str(Path(__file__).resolve()):sha(__file__),str(DATA/'manifest.csv'):sha(DATA/'manifest.csv'),str(SOURCE/'coverage_dictionary.json'):sha(SOURCE/'coverage_dictionary.json'),str(SOURCE/'risk_protocol.json'):sha(SOURCE/'risk_protocol.json'),str(OUT/'protocol.json'):sha(OUT/'protocol.json')}
    quota_rows=[];tier_rows=[];traces=[]
    for source in ('local','online'):
        arrays=load_data(source)
        for budget in BUDGETS:
            deterministic={m:[] for m in ('coverage','risk_coverage')}
            random=np.empty((REPEATS,len(CLASSES),budget),dtype=np.int16)
            for c,category in enumerate(CLASSES):
                z=arrays[category];s=structure(z,budget)
                seed=BASE_SEED+(0 if source=='local' else 100000)+budget*100+c
                rng=np.random.default_rng(seed)
                for k in sorted(s['cells']):
                    quota_rows.append(dict(source=source,budget_per_class=budget,category=category,tier=k[0],risk_midpoint=k[1],unknown_count=k[2],candidates=len(s['cells'][k]),quota=s['quota'][k]))
                for repeat in range(REPEATS):
                    chosen=[]
                    for k in sorted(s['cells']):
                        q=s['quota'][k]
                        if q:chosen.extend(rng.choice(s['cells'][k],q,replace=False).tolist())
                    assert len(set(chosen))==budget
                    random[repeat,c]=chosen
                for method,constraint in [('coverage',None),('risk_coverage',s)]:
                    index,audit=greedy(z,budget,constraint);deterministic[method].append(index)
                    traces += [dict(source=source,budget_per_class=budget,category=category,method=method,**v) for v in audit]
                if budget==BUDGETS[0]:
                    tier_rows += [dict(source=source,category=category,sample_id=str(sid),tier=int(s['tiers'][i]),risk_lower=float(z['risk_lower'][i]),risk_midpoint=float(z['risk_midpoint'][i]),risk_upper=float(z['risk_upper'][i]),unknown_count=int(s['unknown'][i])) for i,sid in enumerate(z['sample_ids'])]
                assert np.all(s['tiers'][random[:,c]].sum(1)==3*budget)
            path=OUT/'selection'/f'{source}_n{budget*6}.npz'
            np.savez_compressed(path,coverage=np.stack(deterministic['coverage']),risk_coverage=np.stack(deterministic['risk_coverage']),risk=random)
            freeze[str(path)]=sha(path)
            note('selection_frozen',source=source,per_class=budget,repeats=REPEATS)
        prepared=SOURCE/'prepared' if source=='local' else SOURCE/'online_analysis/prepared'
        for category in CLASSES:freeze[str(prepared/f'{category}.npz')]=sha(prepared/f'{category}.npz')
    for name,rows in [('quotas.csv',quota_rows),('tiers.csv',tier_rows),('greedy_trace.csv',traces)]:
        path=OUT/'selection'/name;write_csv(path,rows);freeze[str(path)]=sha(path)
    write_json(OUT/'selection/input_freeze.sha256.json',freeze)
    write_json(OUT/'selection/audit.json',dict(status='pass',predictions_read=False,source_budget_settings=6,random_selections=6000,balanced_budgets=True,common_exact_risk_quotas=True))


def selected_indices(arrays,selection):
    # Orders align with the class-specific prepared files; samples need no new inference.
    id_lookup={sid:i for i,sid in enumerate([str(s) for c in CLASSES for s in arrays[c]['sample_ids']])}
    result=np.empty(selection.shape,dtype=np.int32)
    for c,category in enumerate(CLASSES):
        result[:,c]=np.array([id_lookup[str(s)] for s in arrays[category]['sample_ids']])[selection[:,c]]
    return result.reshape(selection.shape[0],-1)


def load_predictions():
    manifest=read_csv(DATA/'manifest.csv');master_ids=[r['sample_id'] for c in CLASSES for r in manifest if r['class_name']==c]
    protocol=json.loads((SOURCE/'protocol.json').read_text());pools=protocol['source_pools']
    all_predictions=[];freeze={};cache_ids=None
    for model,seed in itertools.product(MODELS,SEEDS):
        mapping={}
        for source,path in pools.items():
            pool=Path(path);rows=read_csv(pool/'manifests/candidate_pool.csv')
            npz=pool/'benchmark/runs'/model/str(seed)/'candidate_outputs.npz'
            with np.load(npz,allow_pickle=False) as z:
                ids=z['sample_ids'].tolist();assert ids==[r['sample_id'] for r in rows]
                logits=z['logits'];assert logits.shape==(len(rows),6) and np.isfinite(logits).all()
                preds=logits.argmax(1)
            for sid,pred in zip(ids,preds):
                assert sid not in mapping;mapping[sid]=int(pred)
            freeze[str(npz)]=sha(npz)
        assert set(mapping)==set(master_ids)
        all_predictions.append([mapping[s] for s in master_ids])
        note('prediction_cache_verified',model=model,seed=seed,targets=len(mapping))
    write_json(OUT/'prediction_input_freeze.sha256.json',freeze)
    predicted=np.array(all_predictions,dtype=np.int8)
    np.savez_compressed(OUT/'merged_predictions.npz',sample_ids=np.array(master_ids),predictions=predicted,
                        models=np.array([m for m,s in itertools.product(MODELS,SEEDS)]),seeds=np.array([s for m,s in itertools.product(MODELS,SEEDS)]))
    return predicted,master_ids


def state_metrics(arrays,selection):
    n=selection.shape[0];totals=defaultdict(lambda:np.zeros(n,dtype=float));tier_hist=np.zeros((n,5),dtype=int)
    for c,category in enumerate(CLASSES):
        z=arrays[category];indices=selection[:,c];support=z['matrix'][indices].sum(1);other=z['online_matrix'][indices].sum(1)
        tier=structure(z,selection.shape[-1])['tiers']
        for t in range(1,6):tier_hist[:,t-1]+=(tier[indices]==t).sum(1)
        totals['presence_coverage']+=(support>0).mean(1)/6
        totals['support_completion']+=np.minimum(support/20,1).mean(1)/6
        totals['other_label_presence_coverage']+=(other>0).mean(1)/6
        totals['other_label_support_completion']+=np.minimum(other/20,1).mean(1)/6
        for key in ('lower','midpoint','upper'):totals['risk_'+key]+=z['risk_'+key][indices].mean(1)/6
        totals['risk_unknown_fraction']+=(z['risk_upper'][indices]>z['risk_lower'][indices]+1e-8).mean(1)/6
    for t in range(5):totals[f'q{t+1}_fraction']=tier_hist[:,t]/(6*selection.shape[-1])
    return dict(totals)


def model_metrics(predicted,global_index,budget):
    # predictions gathered into [selection, class, target] arrays separately per run.
    n=len(global_index);conf=np.zeros((n,6,6),dtype=int)
    block=predicted[global_index].reshape(n,6,budget)
    for c in range(6):
        for label in range(6):conf[:,c,label]=(block[:,c]==label).sum(1)
    true=conf.sum(2);pred=conf.sum(1);correct=np.diagonal(conf,axis1=1,axis2=2)
    f1=np.divide(2*correct,true+pred,out=np.zeros_like(correct,dtype=float),where=(true+pred)>0)
    wrong=conf.copy();wrong[:,np.arange(6),np.arange(6)]=0
    return dict(accuracy=correct.sum(1)/(6*budget),macro_f1=f1.mean(1),
                errors=6*budget-correct.sum(1),confusion_type_count=(wrong>0).sum((1,2)),
                confusion_breadth=(wrong>0).sum((1,2))/30,
                **{f'{c}_recall':correct[:,i]/budget for i,c in enumerate(CLASSES)}),wrong


def rarefaction(counts,q):
    # Expected distinct wrong predicted classes among q uniformly sampled errors.
    n=counts.sum(-1)
    q=np.asarray(q);assert np.all(q>=0) and np.all(q<=n)
    remain=n[...,None]-counts;qq=q[...,None]
    log_den=gammaln(n+1)-gammaln(q+1)-gammaln(n-q+1)
    absent=np.ones(counts.shape,dtype=float)
    valid=(qq>0)&(counts>0)&(remain>=qq)
    log_num=gammaln(remain+1)-gammaln(qq+1)-gammaln(remain-qq+1)
    vals=log_num-log_den[...,None]
    absent[valid]=np.exp(np.minimum(0,vals[valid]))
    absent[(qq>0)&(counts>0)&(remain<qq)]=0
    return (1-absent).sum(-1)/5


def exact_signflip(values):
    observed=abs(float(np.mean(values)))
    signs=np.array(list(itertools.product((-1,1),repeat=len(values))))
    distribution=np.abs((signs*np.array(values)).mean(1))
    return float(np.mean(distribution>=observed-1e-12))


def holm(p):
    ids=np.argsort(p);out=np.zeros(len(p));running=0
    for rank,i in enumerate(ids):running=max(running,min(1,(len(p)-rank)*p[i]));out[i]=running
    return out


def evaluate():
    freeze=json.loads((OUT/'selection/input_freeze.sha256.json').read_text())
    assert all(sha(p)==v for p,v in freeze.items())
    assert json.loads((OUT/'selection/audit.json').read_text())['predictions_read'] is False
    predicted,master_ids=load_predictions();runs=list(itertools.product(MODELS,SEEDS))
    per_run=[];state_rows=[];paired_rows=[];summary=[];raw_files=[]
    manifest={r['sample_id']:r for r in read_csv(DATA/'manifest.csv')}
    for source in ('local','online'):
        arrays=load_data(source)
        assert [str(s) for c in CLASSES for s in arrays[c]['sample_ids']]==master_ids
        for budget in BUDGETS:
            path=OUT/'selection'/f'{source}_n{budget*6}.npz'
            orders=dict(np.load(path));orders={k:v if v.ndim==3 else v[None] for k,v in orders.items()}
            metric_arrays={};wrong_arrays={};global_indices={};condition={}
            for method,selection in orders.items():
                assert selection.shape[1:]==(6,budget)
                condition[method]=state_metrics(arrays,selection)
                global_indices[method]=selected_indices(arrays,selection)
                values=[];wrongs=[]
                for run,(model,seed) in enumerate(runs):
                    metrics,wrong=model_metrics(predicted[run],global_indices[method],budget)
                    values.append(metrics);wrongs.append(wrong)
                    per_run.append(dict(source=source,budget_per_class=budget,method=method,model=model,seed=seed,
                        repetitions=len(selection),**{k:float(v.mean()) for k,v in metrics.items()}))
                metric_arrays[method]={k:np.stack([v[k] for v in values]) for k in values[0]}
                wrong_arrays[method]=np.stack(wrongs)
                for repetition in range(len(selection)):
                    state_rows.append(dict(source=source,budget_per_class=budget,method=method,repetition=repetition,
                        **{k:float(v[repetition]) for k,v in condition[method].items()}))
                name=f'{source}_{method}_n{budget*6}'
                subset=[manifest[master_ids[i]] for i in global_indices[method][0]]
                write_csv(OUT/'test_sets'/f'{name}.csv',subset)
            pairs={}
            for left in ('coverage','risk_coverage'):
                fixed=wrong_arrays[left][:,0]
                random=wrong_arrays['risk']
                q=np.minimum(fixed.sum(-1)[:,None],random.sum(-1))
                fixed_breadth=rarefaction(np.broadcast_to(fixed[:,None],random.shape),q).mean(-1)
                random_breadth=rarefaction(random,q).mean(-1)
                delta=fixed_breadth-random_breadth
                evidence=(q>0).mean(-1)
                pairs[left]=dict(equal_error_confusion_delta=delta,evaluable_class_fraction=evidence,
                                 left_equal_error_breadth=fixed_breadth,right_equal_error_breadth=random_breadth)
                for run,(model,seed) in enumerate(runs):
                    paired_rows.append(dict(source=source,budget_per_class=budget,contrast=f'{left}_minus_risk',model=model,seed=seed,
                        equal_error_confusion_delta=float(delta[run].mean()),evaluable_class_fraction=float(evidence[run].mean()),
                        error_delta=float(metric_arrays[left]['errors'][run,0]-metric_arrays['risk']['errors'][run].mean()),
                        raw_confusion_breadth_delta=float(metric_arrays[left]['confusion_breadth'][run,0]-metric_arrays['risk']['confusion_breadth'][run].mean()),
                        macro_f1_delta=float(metric_arrays[left]['macro_f1'][run,0]-metric_arrays['risk']['macro_f1'][run].mean())))
            for method in ('coverage','risk','risk_coverage'):
                row=dict(source=source,budget_per_class=budget,total_targets=budget*6,method=method,repetitions=len(orders[method]),model_seed_runs=24)
                row.update({k:float(v.mean()) for k,v in condition[method].items()})
                row.update({k:float(v.mean()) for k,v in metric_arrays[method].items()})
                if method=='risk':
                    # Selection intervals average the same24 models before quantiles.
                    for key,vals in metric_arrays[method].items():
                        v=vals.mean(0);row[key+'_selection_p025']=float(np.quantile(v,.025));row[key+'_selection_p975']=float(np.quantile(v,.975))
                    for key,vals in condition[method].items():row[key+'_selection_p025']=float(np.quantile(vals,.025));row[key+'_selection_p975']=float(np.quantile(vals,.975))
                summary.append(row)
            raw={f'metrics_{m}_{k}':v for m,d in metric_arrays.items() for k,v in d.items()}
            raw.update({f'wrong_counts_{m}':v for m,v in wrong_arrays.items()})
            raw.update({f'condition_{m}_{k}':v for m,d in condition.items() for k,v in d.items()})
            raw.update({f'paired_{m}_{k}':v for m,d in pairs.items() for k,v in d.items()})
            output=OUT/'raw'/f'{source}_n{budget*6}.npz';output.parent.mkdir(exist_ok=True)
            np.savez_compressed(output,**raw);raw_files.append(output)
            note('evaluated',source=source,per_class=budget,model_runs=24,random_repeats=REPEATS)
    tests=[]
    for source in ('local','online'):
        for budget in BUDGETS:
            values=[r for r in paired_rows if r['source']==source and r['budget_per_class']==budget and r['contrast']=='risk_coverage_minus_risk']
            architecture=[np.mean([r['equal_error_confusion_delta'] for r in values if r['model']==m]) for m in MODELS]
            tests.append(dict(source=source,budget_per_class=budget,primary_budget=budget==40,
                mean_equal_error_delta=float(np.mean(architecture)),positive_architectures=int(np.sum(np.array(architecture)>0)),
                architectures=8,signflip_two_sided_p=exact_signflip(architecture)))
    adjusted=holm([t['signflip_two_sided_p'] for t in tests])
    for row,p in zip(tests,adjusted):row['holm_six_p']=float(p)
    for name,rows in [('summary.csv',summary),('model_run_metrics.csv',per_run),('condition_metrics.csv',state_rows),('paired_model_run_metrics.csv',paired_rows),('primary_tests.csv',tests)]:write_csv(OUT/name,rows)
    result_freeze={str(p):sha(p) for p in raw_files+[OUT/'selection/input_freeze.sha256.json',OUT/'prediction_input_freeze.sha256.json',OUT/'merged_predictions.npz']}
    for path in [OUT/'summary.csv',OUT/'model_run_metrics.csv',OUT/'condition_metrics.csv',OUT/'paired_model_run_metrics.csv',OUT/'primary_tests.csv']:result_freeze[str(path)]=sha(path)
    write_json(OUT/'analysis_result_freeze.sha256.json',result_freeze)
    report(summary,tests)
    note('complete',settings=6,strategies=3,classifier_runs=24)


def report(summary,tests):
    names={'coverage':'覆盖选样','risk':'风险配额随机选样','risk_coverage':'风险＋覆盖选样'}
    lines=['# QE合并数据集：三种测试选样策略实验', '',
        '在7800目标合并池上实际构造三种测试集合，复用经ID、标签、有限值及来源校验的24组冻结模型分类预测，计算每种集合的测试结果。没有重新训练或重新执行模型推理。覆盖与风险仅基于统一共享14字段＋尺度，最新60个细粒度标注未混入本实验。', '',
        '八种架构各三个训练种子；每类预算20、40、100，总量120、240、600。主预算为每类40。覆盖组不控制风险构成；风险组和组合组每类Q1—Q5各占20%，并进一步精确匹配风险中点和未知维数联合格配额。风险组无放回随机重复1000次，组合组确定性贪心。全部选到预设预算后停止，不按模型性能、覆盖率或零增益提前终止。', '',
        '同分风险按目标ID确定层界，Q1—Q5为各标注来源和类别内相对层。未知评分保留区间；自动标签未全量人工验证。共同风险格确保本协议下风险组与组合组风险均值及未知比例一致，但不能控制所有语义难度。', '',
        '出现覆盖固定分母426，支持完成度按每状态20个目标计算；贪心直接优化凹证据增益，不是出现覆盖或支持完成度的精确全局优化。对方标注来源下的覆盖也保存，避免只在选样标签上评价。', '',
        '## 主预算：240目标', '',
        '| 标注来源 | 方法 | 出现覆盖 | 支持完成度 | 平均风险分 | Q5占比 | Accuracy | Macro-F1 | 错误数 | 混淆类型数 |',
        '|---|---|---:|---:|---:|---:|---:|---:|---:|---:|']
    for row in summary:
        if row['budget_per_class']==40:
            lines.append(f"| {row['source']} | {names[row['method']]} | {100*row['presence_coverage']:.2f}% | {100*row['support_completion']:.2f}% | {row['risk_midpoint']:.2f} | {100*row['q5_fraction']:.2f}% | {100*row['accuracy']:.2f}% | {100*row['macro_f1']:.2f}% | {row['errors']:.2f} | {row['confusion_type_count']:.2f} |")
    lines+=['', '模型指标是24次运行平均；风险组另对1000次选样求平均。混淆类型指真实类别到错误预测类别的有向组合，最多30种；数量多不表示每种类型具有足够证据。', '',
            '## 风险＋覆盖相对风险配额随机的额外作用', '',
            '两组按真实类别匹配错误数，使用无放回解析稀释计算预期不同错误预测类别数，再除以每类最多5种并作六类平均。匹配错误数为0的类别贡献0且单独报告可评价类别比例。该指标用于区分错误数增加与错误类型丰富，不代表所有失败机制。统计以八架构为单位，先平均三个种子和1000随机重复，做256种精确符号翻转；六个来源×预算检验统一Holm校正。', '',
            '| 来源 | 总目标数 | 支持完成度增量 | 等错误数混淆广度增量 | 正向架构数 | Holm p |',
            '|---|---:|---:|---:|---:|---:|']
    for t in tests:
        left=next(r for r in summary if r['source']==t['source'] and r['budget_per_class']==t['budget_per_class'] and r['method']=='risk_coverage')
        right=next(r for r in summary if r['source']==t['source'] and r['budget_per_class']==t['budget_per_class'] and r['method']=='risk')
        lines.append(f"| {t['source']} | {t['budget_per_class']*6} | {100*(left['support_completion']-right['support_completion']):+.2f}百分点 | {100*t['mean_equal_error_delta']:+.2f}百分点 | {t['positive_architectures']}/8 | {t['holm_six_p']:.4f} |")
    lines+=['', '低F1、高错误数不能独立证明选样更有效。覆盖组与其他组的差异还包含风险构成变化；风险组与组合组共同风险配额的比较是更直接的覆盖附加作用对照。结果支持范围依据上表，无论正负均报告，不能借用前一轮MILP最大风险集合来替代本轮贪心策略。', '',
            '## 所有预算与模型结果', '',
            'summary.csv保存18个来源—预算—方法结果及风险随机组的95%选样区间，model_run_metrics.csv保存每个架构种子指标，paired_model_run_metrics.csv保存等错误数比较，condition_metrics.csv保存每次随机选样的条件构成，selection/保存可复用的目标索引和选样轨迹。test_sets/保存两种固定集合和风险组第1次抽样的实体manifest引用；1000次风险集合全部索引在NPZ中。', '',
            '图表仅连接实际预算点，不插值制造实验观测。prediction_input_freeze.sha256.json冻结96个来源模型缓存；selection/input_freeze.sha256.json证明读取分类预测前已确定选样。被测模型之间共享同一测试集。']
    (OUT/'report.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    doc=Document();i=0
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            cells=[[c.strip() for c in row.strip('|').split('|')] for row in block]
            table=doc.add_table(rows=1,cols=len(cells[0]));table.style='Table Grid'
            for cell,val in zip(table.rows[0].cells,cells[0]):cell.text=val
            for row in cells[2:]:
                for cell,val in zip(table.add_row().cells,row):cell.text=val
            continue
        if line.startswith('#'):doc.add_heading(line.lstrip('#').strip(),level=min(len(line)-len(line.lstrip('#')),3)-1)
        elif line:doc.add_paragraph(line)
        i+=1
    doc.save(OUT/'report.docx')
    (OUT/'figures').mkdir(exist_ok=True)
    fig,axes=plt.subplots(2,3,figsize=(13,7.4),layout='constrained')
    colors={'coverage':'#377eb8','risk':'#999999','risk_coverage':'#e58139'}
    for i,source in enumerate(('local','online')):
        for j,(key,label,scale_factor) in enumerate([('support_completion','State support completion (%)',100),('macro_f1','Macro-F1 (%)',100),('confusion_type_count','Wrong-class pair types',1)]):
            ax=axes[i,j]
            for method in names:
                rows=sorted((r for r in summary if r['source']==source and r['method']==method),key=lambda r:r['budget_per_class'])
                ax.plot([r['total_targets'] for r in rows],[r[key]*scale_factor for r in rows],marker='o',color=colors[method],label=method)
            ax.set_title(f'{source}: {label}');ax.set_xlabel('Targets');ax.grid(alpha=.2);ax.spines[['top','right']].set_visible(False)
            if i==0 and j==0:ax.legend(fontsize=8)
    for ext in ('png','svg'):fig.savefig(OUT/'figures'/f'three_strategy_comparison.{ext}',dpi=180)
    plt.close(fig)
    claims=dict(design=dict(source='protocol.json and selection/input_freeze.sha256.json',status='executed'),
                metrics=dict(source='raw/*.npz and merged_predictions.npz',status='computed'),
                primary_comparison=dict(source='primary_tests.csv',status='report exact sign and corrected p, no universal superiority assumed'),
                scope='merged annotated finite pool, frozen24classifier runs, coverage/risk proxy labels; no human semantic validation')
    write_json(OUT/'claim_evidence.json',claims)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--phase',choices=('select','evaluate','all'),default='all')
    args=parser.parse_args()
    if args.phase in ('select','all'):selection()
    if args.phase in ('evaluate','all'):evaluate()
