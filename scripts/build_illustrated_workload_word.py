#!/usr/bin/env python3
"""Plain-language workload report, using saved experiment results only."""
from pathlib import Path
import csv
import json
import shutil
import hashlib
import re
import zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/workload_illustrated_20261008'
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
NAME = '工作量_逐图解释版_20261008'
BLUE, ORANGE = '#3574A8', '#BD713E'
inputs = []

def rows(relative):
    p = ROOT / relative
    inputs.append(p)
    with p.open(encoding='utf-8') as f:
        return list(csv.DictReader(f))

random = rows('artifacts/random_stratification_ablation_20261006/tables/random_summary.csv')
mixture = rows('artifacts/controlled_risk_distribution_20261005/tables/panel_summary.csv')
scale = rows('artifacts/scale_only_confirmation_20261007/tables/summary.csv')
conditions = rows('artifacts/scale_only_confirmation_20261007/tables/area_condition_summary.csv')
ablation = rows('artifacts/random_stratification_ablation_20261006/tables/ablation_summary.csv')
paper = ROOT / 'docs/论文_最新.md'
inputs.append(paper)
case = re.search(r'\| ResNet-50，背向dog \| 本地 \| (\d+) \| ([\d.]+) \| ([\d.]+) \|', paper.read_text())
assert case
case_n, case_all, case_back = int(case[1]), float(case[2]), float(case[3])
scores = {}
for pool in ['B', 'C', 'D']:
    for source in ['local', 'online']:
        p = ROOT / f'artifacts/random_stratification_ablation_20261006/ablation/{pool}_{source}.npz'
        inputs.append(p)
        with np.load(p) as z:
            f1 = z['f1'][0]
            assert f1.shape == (24, 5)
            scores[pool, source] = f1.mean(axis=0) * 100
        r = next(r for r in random if r['pool'] == pool and r['source'] == source)
        assert abs(scores[pool, source][4] - scores[pool, source][0] - 100 * float(r['observed_delta'])) < 1e-8
assert all(int(r['repeats']) == 10000 and int(r['extreme_count']) == 0 for r in random)
assert all(float(r['mean_delta_auc']) < 0 for r in ablation)
mix_b = [r for r in mixture if r['pool'] == 'B' and r['source'] == 'local' and r['total'] == '120']
assert len(mix_b) == 21
mix_start = float(next(r for r in mix_b if r['high_risk_percent'] == '0')['macro_f1_mean']) * 100
mix_end = float(next(r for r in mix_b if r['high_risk_percent'] == '100')['macro_f1_mean']) * 100
condition_counts = {}
for pool in ['B', 'C', 'D']:
    rr = [r for r in conditions if r['pool'] == pool and r['evaluable'] == 'True']
    condition_counts[pool] = [len(rr), sum(float(r['area_adjusted_recall_gap']) < 0 for r in rr)]
assert condition_counts == {'B': [119, 109], 'C': [4, 2], 'D': [122, 111]}

font_manager.fontManager.addfont(str(ROOT / 'scripts/assets/NotoSansSC-Regular.ttf'))
plt.rcParams.update({'font.family': 'Noto Sans SC', 'font.size': 13,
                     'svg.fonttype': 'none', 'axes.spines.top': False,
                     'axes.spines.right': False, 'axes.linewidth': .7,
                     'xtick.major.size': 0, 'ytick.major.size': 0})
chart_values = []
def save(fig, stem):
    fig.savefig(FIG / f'{stem}.svg', bbox_inches='tight', facecolor='white')
    fig.savefig(FIG / f'{stem}.png', dpi=210, bbox_inches='tight', facecolor='white')
    plt.close(fig)

fig, ax = plt.subplots(figsize=(9.2, 3.1), layout='constrained')
x = np.arange(3)
low = [scores[p, 'local'][0] for p in ['B', 'C', 'D']]
high = [scores[p, 'local'][4] for p in ['B', 'C', 'D']]
for dx, vals, label, color in [(-.18, low, '较容易组', BLUE), (.18, high, '较困难组', ORANGE)]:
    bars = ax.bar(x + dx, vals, width=.32, label=label, color=color)
    ax.bar_label(bars, labels=[f'{v:.1f}' for v in vals], padding=4, fontsize=14)
ax.set(xticks=x, xticklabels=['COCO\n2400个目标', 'VOC2007\n600个目标', 'Objects365\n2400个目标'],
       ylim=(0, 112), yticks=[0, 25, 50, 75, 100], ylabel='F1综合评分（满分100）')
ax.legend(loc='upper center', bbox_to_anchor=(.5, 1.18), ncol=2, frameon=False)
save(fig, '01_three_sources')
for p, a, b in zip(['B', 'C', 'D'], low, high):
    chart_values.extend([['three_sources', p, 'local', 'Q1', a], ['three_sources', p, 'local', 'Q5', b]])

def two_bars(values, labels, ylabel, stem, suffix=''):
    fig, ax = plt.subplots(figsize=(8.6, 2.75), layout='constrained')
    bars = ax.bar([0, 1], values, width=.42, color=[BLUE, ORANGE])
    ax.bar_label(bars, labels=[f'{v:.1f}{suffix}' for v in values], padding=5, fontsize=17)
    ax.set(xticks=[0, 1], xticklabels=labels, xlim=(-.7, 1.7), ylim=(0, 110),
           yticks=[0, 25, 50, 75, 100], ylabel=ylabel)
    save(fig, stem)

two_bars([mix_start, mix_end], ['120个都来自前四组', '120个都来自较困难组'],
         'F1综合评分（满分100）', '02_same_number')
chart_values.extend([['mixture', 'B', 'local', '0% Q5', mix_start], ['mixture', 'B', 'local', '100% Q5', mix_end]])
two_bars([case_all, case_back], ['全部狗样本', '其中的背向狗样本'],
         '真实狗被认对的比例（%）', '03_back_facing_dog', '%')
chart_values.extend([['dog_recall', 'B', 'local', 'overall', case_all], ['dog_recall', 'B', 'local', 'back', case_back]])
with (OUT / 'chart_values.csv').open('w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['figure', 'pool', 'label_source', 'condition', 'value_100_scale'])
    w.writerows(chart_values)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
sec.header_distance = sec.footer_distance = Cm(.65)
normal = doc.styles['Normal']
normal.font.name = 'Noto Sans CJK SC'
normal.font.size = Pt(11.5)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
normal.paragraph_format.line_spacing = 1.18
normal.paragraph_format.space_after = Pt(5)
md = []

def runfont(run, size=11.5, bold=False, color='263442'):
    run.font.name = 'Noto Sans CJK SC'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def p(text, size=11.5, bold=False, color='263442'):
    pp = doc.add_paragraph()
    runfont(pp.add_run(text), size, bold, color)
    md.append(text + '\n')
    return pp

def heading(text):
    pp = p(text, 13, True, '24577E')
    pp.paragraph_format.keep_with_next = True
    pp.paragraph_format.space_before = Pt(6)
    md[-1] = '## ' + text + '\n'

def page(title, intro):
    if doc.paragraphs:
        doc.add_page_break()
    pp = p(title, 20, True, '24577E')
    pp.paragraph_format.keep_with_next = True
    md[-1] = '# ' + title + '\n'
    p(intro, 10, color='627180')

def table(headers, body, widths, size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for j, value in enumerate(headers):
        t.rows[0].cells[j].text = str(value)
    for r in body:
        for c, v in zip(t.add_row().cells, r):
            c.text = str(v)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for j, c in enumerate(row.cells):
            c.width = Cm(widths[j])
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(4)
                pp.paragraph_format.space_before = Pt(4)
                pp.paragraph_format.line_spacing = 1.05
                for r in pp.runs:
                    runfont(r, size, i == 0)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EAF1F7')
                c._tc.get_or_add_tcPr().append(shade)
    md.append('| ' + ' | '.join(headers) + ' |\n|' + '|'.join(['---']*len(headers)) + '|\n')
    md.extend('| ' + ' | '.join(map(str, r)) + ' |\n' for r in body)
    md.append('\n')

def picture(stem, caption):
    pp = doc.add_paragraph()
    pp.paragraph_format.space_after = Pt(0)
    pp.paragraph_format.keep_with_next = True
    pp.add_run().add_picture(str(FIG / f'{stem}.png'), width=Cm(16.6))
    p(caption, 9, color='627180')
    md.append(f'![{caption}]({FIG / (stem + ".png")})\n')

header = sec.header.paragraphs[0]
runfont(header.add_run('研究工作回顾  ·  每项工作逐图解释'), 8, color='627180')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
runfont(footer.add_run('2026.10.08  |  '), 8, color='627180')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

# Method diagrams preserve real editable text and a valid Draw.io source.
import xml.etree.ElementTree as ET
import html
import base64
import io
import os
import cairosvg
from PIL import Image, ImageDraw

fc = OUT / 'fontconfig.xml'
fc.write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{ROOT / "scripts/assets"}</dir><dir>/usr/share/fonts</dir><cachedir>{OUT / "fontcache"}</cachedir></fontconfig>')
os.environ['FONTCONFIG_FILE'] = str(fc)

def diagram(stem, width, height, nodes, edges, notes=()):
    ns = 'http://www.w3.org/2000/svg'
    ET.register_namespace('', ns)
    svg = ET.Element(f'{{{ns}}}svg', {'width': str(width), 'height': str(height), 'viewBox': f'0 0 {width} {height}'})
    def add(tag, attrs, content=None):
        e = ET.SubElement(svg, f'{{{ns}}}{tag}', {k: str(v) for k,v in attrs.items()})
        if content is not None: e.text = content
        return e
    add('rect', {'width':width,'height':height,'fill':'white'})
    defs = ET.SubElement(svg, f'{{{ns}}}defs')
    marker = ET.SubElement(defs, f'{{{ns}}}marker', {'id':'arrow','markerWidth':'8','markerHeight':'8','refX':'7','refY':'3','orient':'auto'})
    ET.SubElement(marker, f'{{{ns}}}path', {'d':'M0,0 L0,6 L7,3 Z','fill':'#778896'})
    mxfile=ET.Element('mxfile',{'host':'app.diagrams.net'})
    mxdiagram=ET.SubElement(mxfile,'diagram',{'name':'工作说明','id':stem})
    graph=ET.SubElement(mxdiagram,'mxGraphModel',{'page':'1','pageWidth':str(width),'pageHeight':str(height)})
    root=ET.SubElement(graph,'root')
    ET.SubElement(root,'mxCell',{'id':'0'})
    ET.SubElement(root,'mxCell',{'id':'1','parent':'0'})
    for i,(a,b) in enumerate(edges):
        na,nb=nodes[a],nodes[b]
        ax,ay,aw,ah=na[:4];bx,by,bw,bh=nb[:4]
        if abs(ay-by)<3:
            coords=(ax+aw,ay+ah/2,bx,by+bh/2)
        else: coords=(ax+aw/2,ay+ah,bx+bw/2,by)
        add('line',dict(zip(['x1','y1','x2','y2'],coords)) | {'stroke':'#778896','stroke-width':2,'marker-end':'url(#arrow)'})
        cell=ET.SubElement(root,'mxCell',{'id':f'e{i}','parent':'1','edge':'1','source':f'n{a}','target':f'n{b}','style':'edgeStyle=orthogonalEdgeStyle;endArrow=block;strokeColor=#778896;'})
        ET.SubElement(cell,'mxGeometry',{'relative':'1','as':'geometry'})
    for i,(x,y,w,h,lines,fill) in enumerate(nodes):
        add('rect',{'x':x,'y':y,'width':w,'height':h,'rx':3,'fill':fill,'stroke':'#8AA1B5','stroke-width':1.2})
        for j,line in enumerate(lines):
            add('text',{'x':x+w/2,'y':y+h/2-(len(lines)-1)*15+j*30+7,'text-anchor':'middle','font-family':'Noto Sans SC','font-size':23 if j==0 else 19,'fill':'#263442'},line)
        cell=ET.SubElement(root,'mxCell',{'id':f'n{i}','parent':'1','vertex':'1','value':'<br>'.join(html.escape(t) for t in lines),'style':f'rounded=1;whiteSpace=wrap;html=1;fillColor={fill};strokeColor=#8AA1B5;fontSize=20;fontFamily=Noto Sans SC;'})
        ET.SubElement(cell,'mxGeometry',{'x':str(x),'y':str(y),'width':str(w),'height':str(h),'as':'geometry'})
    for i,(x,y,text) in enumerate(notes):
        add('text',{'x':x,'y':y,'font-family':'Noto Sans SC','font-size':21,'fill':'#516575'},text)
        cell=ET.SubElement(root,'mxCell',{'id':f't{i}','parent':'1','vertex':'1','value':text,'style':'text;html=0;fontSize=20;fontFamily=Noto Sans SC;align=left;'})
        ET.SubElement(cell,'mxGeometry',{'x':str(x),'y':str(y-24),'width':str(width-x-5),'height':'30','as':'geometry'})
    ET.ElementTree(svg).write(FIG/f'{stem}.svg',encoding='utf-8',xml_declaration=True)
    ET.ElementTree(mxfile).write(FIG/f'{stem}.drawio',encoding='utf-8',xml_declaration=True)
    cairosvg.svg2png(url=str(FIG/f'{stem}.svg'),write_to=str(FIG/f'{stem}.png'),output_width=1800)

diagram('04_workflow',1000,310,[
    (20,25,285,90,['先整理图片','描述大小、朝向等状态'],'#EDF3F8'),
    (355,25,285,90,['按规则分五组','用相同模型分别测试'],'#EDF3F8'),
    (695,25,285,90,['换三批图片核验','检查能否重复看到差别'],'#EDF3F8'),
    (195,180,285,90,['补做几种比较','随机、删一项、只看大小'],'#F6EFE8'),
    (535,180,285,90,['找具体薄弱情况','整理案例、图表和论文'],'#EDF3F8'),
],[(0,1),(1,2),(1,3),(3,4)])
nodes=[(330,15,340,75,['COCO：某一类400个目标'],'#EDF3F8'),
       (245,125,510,75,['先按五项条件排序','不看模型已经答对还是答错'],'#EDF3F8')]
for i in range(5):
    nodes.append((15+i*198,250,178,90,[f'第{i+1}组','80个目标'],'#F6EFE8' if i==4 else '#EDF3F8'))
diagram('05_split_example',1000,410,nodes,[(0,1)]+[(1,i) for i in range(2,7)],
        [(15,380,'六类都这样分 → 每组6 × 80 = 480个目标；各组类别比例相同。')])

plt.rcParams['font.size']=15
fig, axes=plt.subplots(1,3,figsize=(9.2,3.4),layout='constrained',sharey=True)
for ax,pool,name in zip(axes,['B','C','D'],['COCO','VOC2007','Objects365']):
    vals=scores[pool,'local']
    ax.plot(range(1,6),vals,color=BLUE,marker='o',lw=2.2)
    for j,v in enumerate(vals): ax.text(j+1,v+3,f'{v:.1f}',ha='left' if j==0 else 'right' if j==4 else 'center',fontsize=13)
    ax.set(title=name,xticks=range(1,6),ylim=(0,110),yticks=[0,25,50,75,100],xlabel='组别：1 → 5')
axes[0].set_ylabel('F1综合评分（满分100）')
save(fig,'06_all_five_groups')

fig,ax=plt.subplots(figsize=(9.5,3.5),layout='constrained')
for i,(pool,name) in enumerate(zip(['B','C','D'],['COCO','VOC2007','Objects365'])):
    rr=next(r for r in random if r['pool']==pool and r['source']=='local')
    lo,hi=-100*float(rr['random_sampling_high']),-100*float(rr['random_sampling_low'])
    actual=-100*float(rr['observed_delta'])
    ax.plot([lo,hi],[i,i],color='#BFC6CE',lw=12,solid_capstyle='butt',label='随机分组的通常范围' if i==0 else None)
    ax.scatter(actual,i,color=ORANGE,s=100,marker='D',label='按五项规则分组' if i==0 else None,zorder=3)
    ax.text(actual+1.2,i,f'{actual:.1f}分',va='center',fontsize=13)
ax.axvline(0,color='#83909D',lw=1,ls='--')
ax.set(yticks=[0,1,2],yticklabels=['COCO','VOC2007','Objects365'],ylim=(2.5,-.5),
       xlim=(-9,50),xlabel='第1组分数 − 第5组分数（分）')
ax.legend(loc='upper center',bbox_to_anchor=(.5,1.2),ncol=2,frameon=False,fontsize=12)
save(fig,'07_random_check')

variants=['full','drop_size','drop_body_visibility','drop_occlusion_level','drop_truncation_level','drop_body_orientation']
labels=['五项全部使用','去掉大小','去掉可见部分','去掉遮挡','去掉被画面切掉','去掉朝向']
fig,axes=plt.subplots(1,3,figsize=(9.5,4.3),layout='constrained',sharey=True)
for ax,pool,name in zip(axes,['B','C','D'],['COCO','VOC2007','Objects365']):
    vals=[-100*float(next(r for r in ablation if r['pool']==pool and r['source']=='local' and r['variant']==v)['mean_delta_auc']) for v in variants]
    bars=ax.barh(np.arange(6),vals,color=[BLUE]+['#A8BFD2']*5,height=.6)
    ax.bar_label(bars,labels=[f'{v:.1f}' for v in vals],padding=3,fontsize=13)
    ax.set(title=name,yticks=range(6),yticklabels=labels,xlim=(0,53),ylim=(5.7,-.7),xticks=[0,20,40],xlabel='第1组比第5组高几分')
save(fig,'08_remove_one_rule')

fig,ax=plt.subplots(figsize=(9.5,3.6),layout='constrained')
for pool,name,color,marker in [('B','COCO',BLUE,'o'),('C','VOC2007','#778D64','s'),('D','Objects365',ORANGE,'^')]:
    rr=sorted([r for r in mixture if r['pool']==pool and r['source']=='local' and r['total']=='120'],key=lambda r:float(r['high_risk_percent']))
    assert len(rr)==21
    xx=[float(r['high_risk_percent']) for r in rr];yy=[100*float(r['macro_f1_mean']) for r in rr]
    ax.plot(xx,yy,label=name,color=color,marker=marker,markersize=4,lw=1.8)
    ax.text(101,yy[-1],f'{yy[-1]:.1f}',color=color,va='center')
ax.set(xlim=(-2,109),ylim=(0,105),yticks=[0,25,50,75,100],xticks=[0,25,50,75,100],
       xlabel='120个目标中，第5组目标的占比（%）',ylabel='F1综合评分（满分100）')
ax.legend(frameon=False,ncol=3,loc='lower left',fontsize=12)
save(fig,'09_fixed_120_curve')

fig,ax=plt.subplots(figsize=(9.5,3.3),layout='constrained')
for dx,key,label,color in [(-.18,'full_gap','五项规则',BLUE),(.18,'area_gap','只看目标大小',ORANGE)]:
    vals=[100*float(next(r for r in scale if r['pool']==pool and r['source']=='local')[key]) for pool in ['B','C','D']]
    bars=ax.bar(np.arange(3)+dx,vals,width=.32,color=color,label=label)
    ax.bar_label(bars,labels=[f'{v:.1f}' for v in vals],padding=4,fontsize=13)
ax.set(xticks=range(3),xticklabels=['COCO','VOC2007','Objects365'],ylim=(0,55),yticks=[0,10,20,30,40,50],
       ylabel='第1组比第5组高几分')
ax.legend(frameon=False,ncol=2,loc='upper center',bbox_to_anchor=(.5,1.18))
save(fig,'10_size_only_comparison')

fig,ax=plt.subplots(figsize=(9.8,3.6),layout='constrained')
neg=np.array([condition_counts[p][1] for p in ['B','C','D']])
other=np.array([condition_counts[p][0]-condition_counts[p][1] for p in ['B','C','D']])
missing=np.array([141-condition_counts[p][0] for p in ['B','C','D']])
for vals,left,label,color in [(neg,np.zeros(3),'仍较差',BLUE),(other,neg,'相同或更好',ORANGE),(missing,neg+other,'样本不足','#D5DBE1')]:
    ax.barh(range(3),vals,left=left,label=label,color=color,height=.58)
    for i,(v,l) in enumerate(zip(vals,left)):
        if v>=8: ax.text(l+v/2,i,str(v),ha='center',va='center',color='white' if color==BLUE else '#263442',fontsize=13)
for i,pool in enumerate(['B','C','D']):
    total,n=condition_counts[pool]
    ax.text(147,i,f'{n}/{total}个可比较项仍较差',va='center',fontsize=12)
ax.set(yticks=range(3),yticklabels=['COCO','VOC2007','Objects365'],ylim=(2.5,-.5),xlim=(0,213),
       xticks=[0,50,100,141],xlabel='事先保留的“模型＋具体条件”核验项（每批141项）')
ax.legend(frameon=False,ncol=3,loc='upper center',bbox_to_anchor=(.5,1.2),fontsize=12)
save(fig,'11_after_size_grouping')

# Reuse actual existing images, with shorter Chinese labels. No synthetic cases.
case_source=ROOT/'artifacts/submission_closeout_20261007/figures/06_real_target_cases.svg'
inputs.append(case_source)
tree=ET.parse(case_source)
imgs=[e.attrib['{http://www.w3.org/1999/xlink}href'] for e in tree.getroot().iter() if e.tag.endswith('}image')]
case_spec={'source':str(case_source.relative_to(ROOT)),'panels':['原图中标出的真实狗目标','分类器实际看到的裁剪输入'],'truth':'dog','predictions':['person','car','person'],'automatic_state':'back','not_human_verified':True}
(FIG/'12_real_case.json').write_text(json.dumps(case_spec,ensure_ascii=False,indent=2))
sx=['<svg xmlns="http://www.w3.org/2000/svg" width="1000" height="440" viewBox="0 0 1000 440"><rect width="1000" height="440" fill="white"/>']
for i,(uri,label) in enumerate(zip(imgs[:2],['原图：橙框内是真实狗目标','模型实际看到的裁剪输入'])):
    sx.append(f'<image href="{uri}" x="{20+i*345}" y="58" width="320" height="320"/>')
    sx.append(f'<text x="{180+i*345}" y="35" text-anchor="middle" font-family="Noto Sans SC" font-size="20">{label}</text>')
for j,text in enumerate(['真实类别：狗','同一架构三份训练结果','分别认成：人 / 车 / 人','AI状态描述：背向','这个例子三次都认错']):
    sx.append(f'<text x="715" y="{100+j*52}" font-family="Noto Sans SC" font-size="21" fill="#263442">{text}</text>')
sx.append('<text x="20" y="420" font-family="Noto Sans SC" font-size="21" fill="#627180">先看模型的实际输入；原图中的清晰上下文，并不都进入分类器。</text></svg>')
(FIG/'12_real_case.svg').write_text(''.join(sx),encoding='utf-8')
cairosvg.svg2png(url=str(FIG/'12_real_case.svg'),write_to=str(FIG/'12_real_case.png'),output_width=1800)

def block(label,text):
    heading(label)
    p(text)

def topic(title,question,stem,caption,done,read,result,say):
    page(title,question)
    block('我具体做了什么',done)
    picture(stem,caption)
    block('这张图怎么读',read)
    block('结果说明什么',result)
    block('向老师可以这样讲',say)

page('我的工作量：逐项看图讲清楚', '从老师提出分层建议开始回顾；按解释顺序整理，工作之间有交叉进行。')
p('上次老师让我分层。我完成分层后，继续换图片核验、补做比较，并把低分细化成具体的薄弱情况。',13,True)
picture('04_workflow','图1  整体工作路线：不是只做一次分组，而是沿着结果继续检查。')
table(['已完成工作','具体工作量或产出'],[
    ['整理目标、描述状态、分成五组','狗、猫、马、羊、人、车六类；按五项条件分组。'],
    ['沿用模型，换三批图片核验','8种模型架构×3份训练结果；后续核验5400个目标。'],
    ['随机分组、去掉一项规则再试','每批随机分组10000次；依次去掉五项规则。'],
    ['固定数量，改变困难图片比例','每次120个目标；21种比例，每种抽样1000次。'],
    ['只看大小对照、核验具体条件','比较简单方案；换来源、在大小相近组内检查条件。'],
    ['案例、论文、计算材料整理','加入实际模型输入；核对结果，更新论文与图表。'],
], [7,10.2],10.5)
p('下面每项工作都单独解释：为什么继续做、怎样操作、怎样看图、现在能得出什么结论。',10,color='627180')

topic('1  我怎样把图片分成五组？', '先回答：老师说的“分层”，我实际怎样落实？',
      '05_split_example','图2  用COCO的一类400个目标解释实际分组数量；这不是虚构实验结果。',
      '看五项条件：大小（占画面多大）、可见部分（身体露出多少）、遮挡（被其他东西挡住）、画面切掉（超出图片边缘）、朝向（正面、背面等）。按事先规则打分，在每类内部排序、均分五组。真实类别来自数据集；状态由两套AI和程序提供。',
      '从上往下看：某一类400个目标，每组80个。六类分别这样分，因此每组有480个目标，而且类别比例相同。第一组预计较容易，第五组预计较困难。',
      '分组完成后，能比较条件不同的图片，而不是只报所有图片的平均分。分组时没有用模型的答错记录，因此不是先挑错误再说它难。',
      '“我在每个类别内部划分五个等数量组，再分别测试，减少类别比例不同带来的影响。”')
p('注意：这只是规则预计的相对难度，不是人工认定每张图片的难度；AI状态描述也没有完整人工核对。',10,color='627180')

topic('2  分组后，模型五组成绩分别怎样？', '分组完成后，下一步是计算每一组的表现。',
      '06_all_five_groups','图3  三批核验图片的五组结果；本地AI描述下，24份已训练模型的平均F1。',
      '沿用8种模型架构、每种3份训练结果，共24份模型。分别计算五组的分类表现；后续比较没有重新训练这些模型。',
      '横轴1到5表示从规则预计较容易到较困难。纵轴是F1综合评分，换算成满分100，越高越好。看整条线的变化，不只看两个端点。',
      '第五组都比第一组低。但中间几组不一定严格一步一步下降，例如VOC2007的第二组到第三组有局部回升。所以可以说两端有明显差别，不能说每一层都必然更难。',
      '“五组分别测完后，最容易组和最困难组有明显分差；中间层的表现有波动，我也把它保留下来。”')
p('F1是综合评分，不是“100张图片里认对多少张”的正确率。两套AI负责写状态描述；24份模型负责分类，这是两种不同任务。',10,color='627180')

topic('3  我为什么又换了三批图片？', '原样本有差别，还要检查换图片、换来源能否再看到。',
      '01_three_sources','图4  三批核验目标，共5400个；图中为本地AI描述结果，在线描述下也出现相同方向。',
      '保留原来的模型和分组规则，建立COCO新样本2400个，再加入VOC2007的600个和Objects365的2400个目标。三批都包含相同六类。',
      '每个来源有两根柱：蓝色是第一组，橙色是第五组。三批的蓝柱都高于橙柱。COCO是新样本核验；后两批还增加了不同来源。',
      '本地描述下，分差分别为25.50、16.57、41.80分；在线描述下为21.30、14.44、41.25分。差别在已测来源重复出现，范围约14到42分。',
      '“我不是只测原来那一批；换了三批、共5400个目标后，较困难组平均分仍更低。”')
p('COCO、VOC2007、Objects365是三个公开图片数据集，相当于三套来源的考题。5400是目标样本数，不一定是5400张不同原图。',10,color='627180')

topic('4  随便分五组，也会有这么大差别吗？', '这是检查偶然性：规则真的提供了信息，还是随便分都有差别？',
      '07_random_check','图5  每批随机分组10000次；灰条显示其中中间95%的范围，橙点是按五项规则得到的分差。',
      '在每个类别内随机打乱目标，再等分五组。组大小、类别比例和模型保持不变；每批做10000次，计算第一组与第五组的分差。',
      '先找竖虚线：零表示两组分数一样。灰条在零附近，表示随机分组通常差别不大；橙点远在右边，表示按规则分组时第一组明显更高。',
      '实际规则分差明显大于随机波动。10000次随机结果都没有达到实际分差的绝对值。这支持规则与模型表现有关，但还不能说规则已经最优。',
      '“为了排除随便分组的偶然差别，我补做了随机分组对照。真实规则分组的差别明显更大。”')
p('这些是对已有目标反复分组，不是重新采集10000批新图片；图上灰条也不是所有随机结果的最大范围。',10,color='627180')

topic('5  我又把分组规则一项一项去掉', '检查：差别是不是必须依赖某一项规则才能出现？',
      '08_remove_one_rule','图6  本地AI描述下的结果；横条越长表示第一组比第五组高得越多。在线结果也保留相同方向。',
      '先保留完整五项作参照，然后分别去掉大小、可见部分、遮挡、被画面切掉、朝向。每次重新打分、重新划五组，再比较分差。',
      '每个小图第一行是五项全用。往下每行只去掉一项。所有横条仍大于零，表示第一组平均分仍高于第五组。',
      '去掉任意一项后，主要方向还在。Objects365去掉大小时，分差从41.8分缩到32.4分。但删一项也会改变分组成员，不能直接把变化当成该项的单独贡献。',
      '“我逐项去掉规则再测，检查结果是否依赖某一项。方向还在，但目前不能确定各项独立起多大作用。”')
p('也试过把不同项的分值范围拉齐、等权相加。结果方向保留；正式统计比较还没有确认每项差异，因此不据此宣称某项最重要。',10,color='627180')

topic('6  数量不变，换困难图片的比例会怎样？', '避免把“图片更多”和“图片更难”混在一起。',
      '09_fixed_120_curve','图7  每次120个目标，六类各20个；21个实际比例点，每点抽样1000次，展示本地描述下平均F1。',
      '固定每次测试总量和类别比例，只增加第五组目标的比例：从0%到100%，每隔5%测一次。模型保持不变。',
      '横轴表示120个目标中有多少来自第五组，往右困难目标更多。纵轴是综合分。三条线对应三批来源，整体向下。',
      '本地描述下，从起点到终点，三批F1分别降低20.78、15.92、34.77分。说明同一个模型，换一套测试条件，总分就会变化；不是模型又训练后变差了。',
      '“我把数量固定为120个，只改变困难目标占比。总分仍明显改变，所以评价模型还要交代考题的构成。”')
p('0%从第一至第四组混合选，不是只选第一组。抽样1000次是反复选已有目标，也不等于新增1000批图片。',10,color='627180')

topic('7  最新：只看目标大小，会不会也够用？', '做一个更简单的方法作对照，才能知道复杂规则新增了什么。',
      '10_size_only_comparison','图8  两种方案均按每类五等分；这里展示本地描述比较，完整核验还包括在线结果。',
      '不使用AI状态，只根据数据集标注框计算目标占原图的面积，按大小排序分五组。然后与原来的五项规则比较。',
      '蓝柱是五项规则，橙柱是只看大小；纵轴是第一组与第五组的分差，不是模型总分。Objects365的橙柱甚至比蓝柱更高。',
      '只看大小也能产生很大分差。五项规则没有在所有设置下更强；统计比较也没有确认其正向优势。因此“分差大”还不足以证明复杂规则一定更有价值。',
      '“我新增了只按大小分组的简单对照，发现它也很强。后面重点就转到状态能否提供更具体的测试问题。”')
p('完整比较中，COCO在线描述的五项分差也比只看大小更小；不能只展示有利于五项规则的结果。',10,color='627180')

topic('8  我怎样把“分数低”说得更具体？', '老师通常还会问：那到底该重点测什么？',
      '03_back_facing_dog',f'图9  COCO中ResNet-50的狗识别结果；背向狗子组有{case_n}个目标，数值是三份训练结果的平均。',
      '在最初的图片中整理表现较差的“模型＋具体条件”清单，先保留175项，再把同一清单放到新图片中核验。图中的ResNet-50是其中一种图片分类模型。',
      '左柱是全部真实狗样本中认成狗的比例，81.92%；右柱是其中背向狗的比例，57.25%。这里是Recall，可以理解为真实狗被认对的比例。',
      '总体结果会掩盖具体情况。这个例子提示背向狗值得补测，但不代表每个模型都一样，也没有证明背向就是错误原因；有些原来薄弱的条件换图片后没有复现。',
      '“我进一步把总分拆到具体条件。例如这个模型对全部狗识别率约82%，背向狗只有约57%，就能提出重点测试背向狗。”')
p('175项是模型与条件的组合，条件和目标会重叠，不等于175个独立缺陷。VOC很多条件样本太少，不能把它们说成已稳定复现。',10,color='627180')

topic('9  这些情况会不会只是因为目标太小？', '既然大小本身很强，就继续检查朝向、遮挡等信息是否还有用。',
      '11_after_size_grouping','图10  两套AI描述的核验项合并计数。蓝色仍较差；橙色相同或更好；灰色样本不足。',
      '从原来保留的清单中排除纯大小条件，剩141项。在每个类别内按大小分成五档，在相近大小档中比较某状态与非该状态。',
      '三行都从同一份141项清单出发。COCO有119项能比较，其中109项仍较差；Objects365有122项能比较，111项仍较差。VOC只有4项能比较，其余主要是灰色。',
      '在COCO和Objects365中，很多状态的较差表现，在粗略调整大小后仍存在。VOC证据不足。“样本不足”表示暂时无法判断，不表示没有问题。',
      '“为了排除大小的影响，我在大小相近的组内再比其他状态。两批中多数可比较项仍较差，另一批样本太少，需要保留不确定性。”')
p('五个大小档并不等于把大小完全配成一样，其他条件也可能同时出现；这些是关联证据，计数不能当作独立错误数。',10,color='627180')

topic('10  我还补了真实图片，解释模型看到了什么', '只有统计图还不够：具体图片可以帮助理解测试流程。',
      '12_real_case','图11  复用既有真实案例：原图、实际模型输入、真实类别和三份模型预测。状态由AI自动描述。',
      '从已有测试结果中整理真实案例，把原图里的目标框、送给模型的裁剪输入和预测并排放出来。案例没有重新生成，也没有人工改状态标签。',
      '先看左图橙框，真实类别是狗。中间是模型实际看到的裁剪输入，已经很模糊。右边显示同一架构三份训练结果分别认成了人、车、人。',
      '原图的上下文和模型实际输入不同。小目标经放大、裁剪后，可能丢失信息，所以不能把错误全部归到朝向或模型结构上。这个案例解释现象，不估计所有图片的错误率。',
      '“我补了真实输入案例，帮助说明模型究竟看到了什么。除了状态，目标裁剪和预处理也需要纳入结果解释。”')
p('这一例是在符合既有条件的错误案例中按固定顺序选出；它不是人工核对，也不是随机抽样，不能代表总体失败比例。',10,color='627180')

page('11  最后整理了什么？怎样向老师汇报？', '把前面的实验变成可查看、可解释、可重新计算的材料。')
block('已经完成的材料工作','我更新了中英文论文，把新对照、具体条件、真实图片和局限写进去；同时整理图表、补充结果与计算材料。利用保存的预测核对计算，8张新增结果表和60项历史比较与原结果一致。这是结果复算，不是从零重新训练。')
heading('约3分钟汇报稿')
for text in [
    '老师，上次您建议我对测试图片做分层。我先按照目标大小、可见部分、遮挡、画面切掉和朝向，把六类目标分别划成五个等数量组，用原来训练好的模型分别测试。',
    '有了初步差别以后，我保留规则又换了三批图片核验，共5400个目标。三批里较困难组的平均分都低于较容易组，两套AI描述下，差值约14到42分；但中间组不是每一步都下降。',
    '我又补了三种检查：每批随机分组10000次；分别去掉一项规则重新分组；每次固定120个目标，逐步增加困难目标比例。这些检查帮助说明差别不是随便分组就能出现，同一个模型的总分也需要结合测试图片构成来解释。',
    '最新我补了只看目标大小的对照。它本身也很强，所以还不能说五项规则一定更好。我接着找更具体的测试问题，例如这个模型对全部狗的识别率约82%，背向狗约57%。再换图片、在大小相近的组内比较后，两批里很多可比较条件仍较差；另一批样本少，暂时不好判断。',
    '最后我加入了真实输入案例，更新了论文和计算材料。现在的主要成果，是把总体成绩展开成具体测试条件，帮助决定接下来重点测什么。状态是AI自动描述，找到的还是相关薄弱情况，没有证明错误原因，也没有完成模型修复。',
]: p(text,11)
heading('讲图时，用这四句就够了')
p('① 我为什么做这一步？\n② 图的横轴、纵轴或颜色是什么意思？\n③ 最应该看哪两个位置，它们有什么差别？\n④ 这个差别让我知道了什么，还有什么不能确定？',11)
p('数值与核查来源：random_stratification_ablation_20261006、controlled_risk_distribution_20261005、scale_only_confirmation_20261007、submission_closeout_20261007及docs/论文_最新.md。实验目录位于当前工作目录artifacts下；逐图数据与生成脚本在随附图表素材包中。',8,color='627180')

doc.core_properties.title='工作量：逐图解释版'
doc.core_properties.subject='用通俗文字和逐项图表解释老师分层建议之后的已完成工作'
doc.core_properties.author='QE研究工作整理'
doc.save(OUT/f'{NAME}.docx')
(OUT/f'{NAME}.md').write_text('\n'.join(md),encoding='utf-8')
for ext in ['docx','md']: shutil.copy2(OUT/f'{NAME}.{ext}',ROOT/'docs'/f'{NAME}.{ext}')
shutil.copy2(Path(__file__),OUT/Path(__file__).name)
used=[s for s in ['04_workflow','05_split_example','06_all_five_groups','01_three_sources','07_random_check','08_remove_one_rule','09_fixed_120_curve','10_size_only_comparison','03_back_facing_dog','11_after_size_grouping','12_real_case']]
for ext in ['svg','png']:
    (FIG/f'02_same_number.{ext}').unlink(missing_ok=True)
assert len(doc.inline_shapes)==11
assert len(doc.element.xpath('//w:br[@w:type="page"]'))==11
(OUT/'audit.json').write_text(json.dumps({'used_figures':used,'planned_pages':12,'figures':11,
 'input_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
 'condition_counts':condition_counts,'new_experiments_run':False,'render_verified':False},ensure_ascii=False,indent=2))
source_dir=OUT/'source_data'
for src in inputs:
    target=source_dir/src.relative_to(ROOT)
    target.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(src,target)
target=source_dir/'scripts/assets/NotoSansSC-Regular.ttf'
target.parent.mkdir(parents=True,exist_ok=True)
shutil.copy2(ROOT/'scripts/assets/NotoSansSC-Regular.ttf',target)
shutil.copy2(Path(__file__),source_dir/'scripts/build_illustrated_workload_word.py')
(OUT/'图表素材说明.md').write_text('''# 工作量逐图解释版：图表素材

正文使用11张图，文件名见audit.json的used_figures清单。

- figures：文字可编辑的SVG、Word使用的PNG；两个流程图另有Draw.io源文件。
- source_data：原始结果表、五组评分数组、论文案例数值、真实案例图、字体和生成脚本。
- audit.json：源文件哈希与核查记录。

实际图片来自已有实验，没有生成新图片或新增实验。数据图以本地AI描述结果举例，正文也说明在线核验和不利于复杂规则的结果。

复现时，使用装有numpy、matplotlib、python-docx、cairosvg、Pillow的Python运行source_data/scripts/build_illustrated_workload_word.py。输出写入source_data/artifacts/workload_illustrated_20261008及source_data/docs。
''',encoding='utf-8')
with zipfile.ZipFile(ROOT/'docs/工作量_逐图解释版_图表素材_20261008.zip','w',zipfile.ZIP_DEFLATED) as z:
    for folder in [FIG,source_dir]:
        for path in folder.rglob('*'):
            if path.is_file():z.write(path,path.relative_to(OUT))
    for name in ['audit.json','图表素材说明.md','build_illustrated_workload_word.py']:
        z.write(OUT/name,name)
print(json.dumps({'word':str(ROOT/'docs'/f'{NAME}.docx'),'figures':len(doc.inline_shapes),'planned_pages':12},ensure_ascii=False))
