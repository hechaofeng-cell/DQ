#!/usr/bin/env python3
"""Independent audit of existing-feature scoring and downstream diagnostic exports."""
from __future__ import annotations
import csv
import hashlib
import json
from collections import Counter, defaultdict
from pathlib import Path
import numpy as np
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data/QE_VisualState_CoverRisk_v1_20261013'
FULL=ROOT/'artifacts/QE_ThreeStrategy_Full_20261013'
BASE=ROOT/'artifacts/QE_VisualState_CoverRisk_v1_20261013'
OUT=ROOT/'artifacts/QE_FeatureRisk_Diagnostics_20261013'
CLASSES=('dog','cat','horse','sheep','person','car')
DIMS=('scale','visibility','occlusion','truncation','orientation')
BUDGETS=(20,40,100,200,400,600,800,1000,1200,1300)
MAIN=dict(scale='target_scale',visibility='body_visibility',occlusion='occlusion_level',truncation='truncation_level',orientation='body_orientation')
MAP=dict(body_visibility='body_visibility',head_visibility='part_visibility',leg_visibility='leg_visibility',body_orientation='body_orientation',head_orientation='head_orientation',occlusion_level='occlusion_level',truncation_level='truncation_level')

def read(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()

def close(a,b,label):
    assert np.isclose(float(a),float(b),atol=1e-8,rtol=0),(label,a,b)

def scales(area):return 'tiny' if area<.01 else 'small' if area<.03 else 'medium' if area<.1 else 'large'

def mapped(value,mapping):
    if value in ('unknown','not_observable'):return (0.,.5,1.,0.)
    score=mapping[value];return (score,score,score,1.)

def independent_dimensions(area,features,c,config,parts=False):
    s=sum(area<edge for edge in config['scale_edges'])/3
    result={'scale':(s,s,s,1.)}
    for d in DIMS[1:]:result[d]=mapped(features[MAIN[d]],config['maps'][MAIN[d]])
    if parts and c!='car':
        for d,components in [('visibility',config['noncar_visibility_components']),('orientation',config['noncar_orientation_components'])]:
            result[d]=tuple(sum(w*mapped(features[field],config['maps'][MAP[field]])[j] for field,w in components.items()) for j in range(4))
    return result

def statcheck(row,prefix,values,label):
    values=np.asarray(values,dtype=float)
    for suffix,value in [('mean',values.mean()),('selection_p025',np.quantile(values,.025)),('selection_p975',np.quantile(values,.975))]:
        close(row[prefix+'_'+suffix],value,(label,prefix,suffix))

def main():
    freeze_counts={}
    for name in ('input_freeze.sha256.json','result_freeze.sha256.json','selection_analysis_freeze.sha256.json','inventory_freeze.sha256.json','report_freeze.sha256.json'):
        mapping=json.loads((OUT/name).read_text())
        for path,expected in mapping.items():assert digest(path)==expected,(name,path,'freeze changed')
        freeze_counts[name]=len(mapping)
    config=json.loads((DATA/'risk_protocol.json').read_text())
    states=[tuple(s) for s in json.loads((DATA/'coverage_dictionary.json').read_text())['states']]
    assert len(states)==len(set(states))==71
    manifest=read(DATA/'manifest.csv');master={r['sample_id']:r for r in manifest}
    assert len(master)==7800
    native={}; packed={}
    for row in manifest:
        ann=json.loads((DATA/row['annotation_path']).read_text());sid=row['sample_id']
        assert ann['sample_id']==sid and ann['class_name']==row['class_name']
        packed[sid]=ann
        for b in ('local','online'):native[b,sid]=ann[b]['native']
    wide=read(OUT/'target_scores.csv');lookup={(r['annotation_backend'],r['sample_id']):r for r in wide}
    assert len(wide)==len(lookup)==15600
    dims={};field_counts=Counter()
    for row in wide:
        b,sid,c=row['annotation_backend'],row['sample_id'],row['class_name']
        m=master[sid];f=native[b,sid]['features'];area=float(m['bbox_area_ratio'])
        assert c==m['class_name'] and row['source_pool']==m['source_pool'] and row['official_label']==m['label']
        assert int(row['native_field_count'])==len(f)==(18 if c=='car' else 19)
        main_dim=independent_dimensions(area,f,c,config);part_dim=independent_dimensions(area,f,c,config,True);dims[b,sid]=main_dim
        for d in DIMS:
            label=scales(area) if d=='scale' else f[MAIN[d]]
            assert row[d+'_label']==label
            close(row[d+'_known_weight'],main_dim[d][3],(b,sid,d,'known weight'))
            for j,bound in enumerate(('lower','midpoint','upper')):
                close(row[d+'_risk_'+bound+'_0_100'],100*main_dim[d][j],(b,sid,d,bound,'risk'))
                close(row[d+'_contribution_'+bound+'_0_20'],20*main_dim[d][j],(b,sid,d,bound,'contribution'))
                close(row['sensitivity_parts_'+d+'_risk_'+bound+'_0_100'],100*part_dim[d][j],(b,sid,d,bound,'parts dimension'))
        for j,bound in enumerate(('lower','midpoint','upper')):
            score=sum(20*main_dim[d][j] for d in DIMS);part_score=sum(20*part_dim[d][j] for d in DIMS)
            close(row['global_risk_'+bound],score,(b,sid,bound,'global'))
            close(row['sensitivity_equal_parts_'+bound],part_score,(b,sid,bound,'parts global'))
        assert row['fully_known']==str(all(v[3]==1 for v in main_dim.values()))
        assert row['evaluable']==str(all(v[3]>=.5-1e-12 for v in main_dim.values()))
        assert row['sensitivity_equal_parts_fully_known']==str(all(v[3]==1 for v in part_dim.values()))
        assert row['sensitivity_equal_parts_evaluable']==str(all(v[3]>=.5-1e-12 for v in part_dim.values()))
        close(row['sensitivity_parts_minus_global_midpoint'],sum(20*(part_dim[d][1]-main_dim[d][1]) for d in DIMS),(b,sid,'parts difference'))
        assert int(row['unknown_dimensions'])==sum(v[3]==0 for v in main_dim.values())
        field_counts[b,sid]=len(f)+1
    print('Main and existing-part scores independently verified: 15,600 targets',flush=True)
    seen=defaultdict(set);feature_rows=0;state_counts=Counter()
    with (OUT/'feature_scores_long.csv').open(encoding='utf-8-sig',newline='') as stream:
        for row in csv.DictReader(stream):
            b,sid,field=row['annotation_backend'],row['sample_id'],row['feature_id'];m=master[sid];f=native[b,sid]['features'];area=float(m['bbox_area_ratio'])
            value=scales(area) if field=='target_scale' else f[field]
            expected=json.dumps(value,ensure_ascii=False) if isinstance(value,list) else value
            assert row['feature_value']==expected and field not in seen[b,sid]
            seen[b,sid].add(field);feature_rows+=1
            if field!='target_scale':assert row['evidence']==native[b,sid]['evidence'][field]
            primary=field in MAIN.values();assert row['participates_equal_global']==str(primary)
            scored=field=='target_scale' or field in MAP
            if scored:
                if field=='target_scale':values=dims[b,sid]['scale']
                else:values=mapped(value,config['maps'][MAP[field]])
                for j,bound in enumerate(('lower','midpoint','upper')):close(row['feature_score_'+bound+'_0_100'],values[j]*100,(b,sid,field,bound))
                assert row['feature_score_known']==str(values[3]==1)
            else:
                assert all(row['feature_score_'+bound+'_0_100']=='' for bound in ('lower','midpoint','upper'))
            for state in value if isinstance(value,list) else [value]:state_counts[b,m['class_name'],field,state]+=1
    assert feature_rows==309400 and all(len(seen[k])==v for k,v in field_counts.items())
    assert len(seen)==15600
    for row in read(OUT/'feature_state_counts.csv'):
        key=row['annotation_backend'],row['class_name'],row['feature_id'],row['state'];count=state_counts[key]
        assert int(row['count'])==count
        close(row['proportion'],count/1300,('field state',key))
        fixed=(key[2],key[3]) in states
        assert row['in_fixed_coverage_denominator']==str(fixed)
        if fixed:close(row['support_completion'],min(count/20,1),('field state support',key))
    pred=dict(np.load(FULL/'merged_predictions.npz',allow_pickle=False));global_pos={str(s):i for i,s in enumerate(pred['sample_ids'])}
    arrays={};tier_by_sid={}
    for b in ('local','online'):
        arrays[b]={}
        for c in CLASSES:
            root=BASE/'prepared' if b=='local' else BASE/'online_analysis/prepared'
            z=dict(np.load(root/f'{c}.npz',allow_pickle=False));arrays[b][c]=z
            order=sorted(range(1300),key=lambda i:(float(z['risk_midpoint'][i]),str(z['sample_ids'][i])))
            for rank,index in enumerate(order):
                sid=str(z['sample_ids'][index]);tier=rank//260+1;tier_by_sid[b,sid]=tier
                assert int(lookup[b,sid]['risk_tier_integer'])==tier and lookup[b,sid]['risk_tier']=='Q'+str(tier)
    for row in read(OUT/'target_risk_tier_assignments.csv'):
        assert int(row['risk_layer'])==tier_by_sid[row['backend'],row['sample_id']]
    audit_pool_summaries(lookup,dims,tier_by_sid)
    for row in read(OUT/'risk_tier_model_run_metrics.csv'):
        b,c,tier=row['backend'],row['category'],int(row['risk_layer']);ci=CLASSES.index(c)
        ids=[str(s) for s in arrays[b][c]['sample_ids'] if tier_by_sid[b,str(s)]==tier]
        run=next(i for i in range(24) if str(pred['models'][i])==row['model'] and int(pred['seeds'][i])==int(row['seed']))
        errors=int(np.count_nonzero(pred['predictions'][run,[global_pos[s] for s in ids]]!=ci))
        assert int(row['targets'])==260 and int(row['errors'])==errors
        close(row['error_rate'],errors/260,('tier model error',b,c,tier,run))
    errors=read(OUT/'risk_tier_error_relationship.csv')
    for row in errors:
        b,c,tier=row['backend'],row['category'],int(row['risk_layer'])
        categories=CLASSES if c=='ALL_CLASSES' else (c,)
        rates=[]
        for category in categories:
            ci=CLASSES.index(category);ids=[str(s) for s in arrays[b][category]['sample_ids'] if tier_by_sid[b,str(s)]==tier]
            rates.append((pred['predictions'][:,[global_pos[s] for s in ids]]!=ci).mean(axis=1))
        rates=np.mean(rates,axis=0);close(row['mean_error_rate'],rates.mean(),('tier mean error',b,c,tier))
        close(row['min_run_error_rate'],rates.min(),('tier min error',b,c,tier));close(row['max_run_error_rate'],rates.max(),('tier max error',b,c,tier))
    print('Feature values, evidence, ordinal scores, blank nominal scores, tiers and observed errors verified',flush=True)
    audit_selection_stats(arrays,dims,tier_by_sid,master)
    audit_manifests(master,native,lookup,arrays,tier_by_sid)
    audit_report(wide,lookup,dims,errors)
    inventory=json.loads((OUT/'inventory_audit.json').read_text())
    assert inventory['fresh_fine_targets']==60 and inventory['early_verified_target_matches']==0
    preserved=read(OUT/'fine_scores_preserved.csv');original=read(ROOT/'artifacts/fresh_vlm_feature_scoring_20261012/scores.csv')
    assert preserved==original
    result=dict(status='pass',targets=7800,main_and_existing_parts_scores_recomputed=15600,
        per_feature_values_evidence_score_rows_verified=309400,nominal_scores_blank=True,
        fixed_states_per_source=426,selected_stats_settings_verified=60,
        full_pool_three_methods_tier_state_dimension_stats_identical=True,
        source_correct_stratified_manifests=10,source_correct_selected_manifests=6,
        fresh_score_rows_exactly_preserved=60,early_dog_annotations_mapped=0,
        frozen_files_verified=freeze_counts,semantic_accuracy_of_vlm_not_assumed=True,
        new_vlm_inference=False,score_is_failure_probability=False)
    (OUT/'independent_diagnostics_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    paths=[Path(__file__).resolve(),OUT/'independent_diagnostics_audit.json',OUT/'report_freeze.sha256.json',OUT/'selection_analysis_freeze.sha256.json',OUT/'result_freeze.sha256.json',OUT/'inventory_freeze.sha256.json']
    (OUT/'independent_diagnostics_freeze.sha256.json').write_text(json.dumps({str(p):digest(p) for p in paths},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2),flush=True)


def audit_pool_summaries(lookup,dims,tiers):
    for row in read(OUT/'risk_strata_summary.csv'):
        b,c,tier=row['annotation_backend'],row['class_name'],int(row['risk_tier'][1:])
        rows=[r for (bb,s),r in lookup.items() if bb==b and r['class_name']==c and tiers[b,s]==tier]
        assert len(rows)==int(row['targets'])==260
        mid=np.array([float(r['global_risk_midpoint']) for r in rows])
        for key,value in [('min',mid.min()),('max',mid.max()),('mean',mid.mean()),('distinct_scores',len(np.unique(mid)))]:close(row[key],value,(b,c,tier,key))
        for bound in ('lower','upper'):close(row[bound+'_mean'],np.mean([float(r['global_risk_'+bound]) for r in rows]),(b,c,tier,bound))
        unknown=sum(r['fully_known']!='True' for r in rows)
        assert int(row['unknown_targets'])==unknown and int(row['fully_known_targets'])==260-unknown
        close(row['unknown_fraction'],unknown/260,(b,c,tier,'unknown'))
        if tier<5:
            following=[float(r['global_risk_midpoint']) for (bb,s),r in lookup.items() if bb==b and r['class_name']==c and tiers[b,s]==tier+1]
            assert row['upper_boundary_tied_next_tier']==str(float(mid.max())==min(following))
    for row in read(OUT/'dimension_score_summary.csv'):
        b,c,d=row['annotation_backend'],row['class_name'],row['dimension']
        values=np.array([v[d] for (bb,s),v in dims.items() if bb==b and lookup[bb,s]['class_name']==c]);assert len(values)==1300
        mid=values[:,1]*100
        for key,value in [('min',mid.min()),('max',mid.max()),('mean',mid.mean()),('distinct_scores',len(np.unique(mid))),('weighted_contribution_mean_0_20',values[:,1].mean()*20),('risk_lower_mean_0_100',values[:,0].mean()*100),('risk_upper_mean_0_100',values[:,2].mean()*100)]:close(row[key],value,(b,c,d,key))
        assert int(row['unknown_targets'])==np.count_nonzero(values[:,3]<1)
    for row in read(OUT/'equal_parts_sensitivity_summary.csv'):
        b,c=row['annotation_backend'],row['class_name'];rows=[r for (bb,s),r in lookup.items() if bb==b and r['class_name']==c]
        global_scores=np.array([float(r['global_risk_midpoint']) for r in rows]);parts=np.array([float(r['sensitivity_equal_parts_midpoint']) for r in rows])
        close(row['global_mean'],global_scores.mean(),(b,c,'parts global mean'));close(row['equal_parts_mean'],parts.mean(),(b,c,'parts mean'));close(row['mean_delta'],np.mean(parts-global_scores),(b,c,'parts delta'))
        assert int(row['changed_targets'])==np.count_nonzero(np.abs(parts-global_scores)>1e-12)


def audit_selection_stats(arrays,dims,tiers,master):
    families={}
    for name,ids in [('selection_risk_layers.csv',('category','risk_layer')),('selection_dimension_scores.csv',('category','dimension')),('selection_feature_states.csv',('category','field','state')),('selection_source_composition.csv',('category','source_pool')),
                     ('selection_risk_layer_totals.csv',('risk_layer',)),('selection_dimension_score_totals.csv',('dimension',)),('selection_source_composition_totals.csv',('source_pool',))]:
        families[name]={(r['backend'],int(r['budget_per_class']),r['method'])+tuple(r[k] for k in ids):r for r in read(OUT/name)}
    for b in ('local','online'):
        for budget in BUDGETS:
            chosen=dict(np.load(FULL/'selection'/f'{b}_n{6*budget}.npz'))
            for method,selections in chosen.items():
                if selections.ndim==2:selections=selections[None]
                base=(b,budget,method);reps=len(selections)
                totaltiers=np.zeros((reps,5));totaldim={d:{k:np.zeros(reps) for k in ('lower','midpoint','upper','unknown')} for d in DIMS}
                pools=sorted({m['source_pool'] for m in master.values()});totalpool={p:np.zeros(reps) for p in pools}
                for ci,c in enumerate(CLASSES):
                    z=arrays[b][c];ids=selections[:,ci];sids=[str(s) for s in z['sample_ids']]
                    tier=np.array([tiers[b,s] for s in sids]);source=np.array([master[s]['source_pool'] for s in sids])
                    for t in range(1,6):
                        counts=np.count_nonzero(tier[ids]==t,axis=1);totaltiers[:,t-1]+=counts
                        row=families['selection_risk_layers.csv'][base+(c,str(t))]
                        statcheck(row,'targets',counts,(base,c,t));statcheck(row,'fraction',counts/budget,(base,c,t))
                    for d in DIMS:
                        values=np.asarray([dims[b,s][d] for s in sids]);row=families['selection_dimension_scores.csv'][base+(c,d)]
                        for j,key in enumerate(('lower','midpoint','upper')):
                            observed=20*values[ids,j].mean(axis=1);totaldim[d][key]+=observed/6;statcheck(row,key+'_points',observed,(base,c,d))
                        unknown=(values[ids,3]<1-1e-9).mean(axis=1);totaldim[d]['unknown']+=unknown/6;statcheck(row,'unknown_fraction',unknown,(base,c,d))
                    counts=z['matrix'][ids].sum(axis=1)
                    states=[tuple(s) for s in json.loads((DATA/'coverage_dictionary.json').read_text())['states']]
                    for j,(field,state) in enumerate(states):
                        row=families['selection_feature_states.csv'][base+(c,field,state)];v=counts[:,j]
                        statcheck(row,'targets',v,(base,c,field,state));statcheck(row,'fraction',v/budget,(base,c,field,state));statcheck(row,'support_completion',np.minimum(v/20,1),(base,c,field,state))
                        close(row['presence_coverage_probability'],np.mean(v>0),(base,c,field,state,'presence'))
                        close(row['zero_support_probability'],np.mean(v==0),(base,c,field,state,'zero'))
                        close(row['fully_supported_probability'],np.mean(v>=20),(base,c,field,state,'support20'))
                    for p in pools:
                        observed=np.count_nonzero(source[ids]==p,axis=1);totalpool[p]+=observed
                        row=families['selection_source_composition.csv'][base+(c,p)];statcheck(row,'targets',observed,(base,c,p));statcheck(row,'fraction',observed/budget,(base,c,p))
                for t in range(1,6):
                    row=families['selection_risk_layer_totals.csv'][base+(str(t),)];statcheck(row,'targets',totaltiers[:,t-1],(base,t));statcheck(row,'fraction',totaltiers[:,t-1]/(6*budget),(base,t))
                for d in DIMS:
                    row=families['selection_dimension_score_totals.csv'][base+(d,)]
                    for key in ('lower','midpoint','upper'):statcheck(row,key+'_points',totaldim[d][key],(base,d))
                    statcheck(row,'unknown_fraction',totaldim[d]['unknown'],(base,d))
                for p in pools:
                    row=families['selection_source_composition_totals.csv'][base+(p,)];statcheck(row,'targets',totalpool[p],(base,p));statcheck(row,'fraction',totalpool[p]/(6*budget),(base,p))
            print(f'Selected-state and risk-dimension diagnostics verified: {b}, {budget}/class',flush=True)


def audit_manifests(master,native,lookup,arrays,tiers):
    index=read(OUT/'manifest_index.csv');assert len(index)==16
    for item in index:
        b=item['backend'];rows=read(item['path']);ids=[r['sample_id'] for r in rows]
        assert len(ids)==len(set(ids))==int(item['targets'])
        if item['kind']=='whole_pool_risk_layer':
            tier=int(item['risk_layer']);expected={sid for sid in master if tiers[b,sid]==tier}
            assert set(ids)==expected and Counter(r['class_name'] for r in rows)==Counter(dict.fromkeys(CLASSES,260))
        else:
            selection=dict(np.load(FULL/'selection'/f'{b}_n1200.npz'))[item['method']]
            if selection.ndim==3:selection=selection[0];assert item['random_repetition']=='0'
            expected=[str(arrays[b][c]['sample_ids'][i]) for ci,c in enumerate(CLASSES) for i in selection[ci]]
            assert ids==expected
        for row in rows:
            sid=row['sample_id'];score=lookup[b,sid];original=master[sid]
            for key,value in score.items():assert row[key]==value,('source score manifest',item['path'],sid,key)
            for key in ('lower','midpoint','upper'):close(row['score_'+key],score['global_risk_'+key],('legacy score alias',b,sid,key))
            assert int(row['high_risk_q5'])==int(tiers[b,sid]==5)
            for key in ('image_path','crop_path','marked_image_path','feature_crop_path','annotation_path'):
                path=DATA/original[key];assert row[key]==str(path) and path.is_file(),(b,sid,key,'path')
            for field,value in native[b,sid]['features'].items():assert row['feature_'+field]==(json.dumps(value,ensure_ascii=False) if isinstance(value,list) else value)
    print('All 16 source-specific image/label/score manifests independently verified',flush=True)


def audit_report(wide,lookup,dims,errors):
    text=(OUT/'report.md').read_text();summary=read(FULL/'summary.csv');names={'coverage':'覆盖选样','risk':'风险分层随机','risk_coverage':'风险＋覆盖'}
    for b in ('local','online'):
        row=next(r for r in summary if r['source']==b and r['method']=='coverage' and r['total_targets']=='7800')
        selected=[r for r in wide if r['annotation_backend']==b];means=[np.mean([dims[b,r['sample_id']][d][1]*100 for r in selected]) for d in DIMS]
        known=sum(r['fully_known']=='True' for r in selected)
        cells=[b]+[f'{v:.2f}' for v in means]+[f"{float(row['risk_midpoint']):.2f}",f"[{float(row['risk_lower']):.2f}, {float(row['risk_upper']):.2f}]",f'{known}/7800',f"{float(row['presence_coverage']):.2%}",f"{float(row['support_completion']):.2%}"]
        assert '| '+' | '.join(cells)+' |' in text
        for method in names:
            row=next(r for r in summary if r['source']==b and r['method']==method and r['total_targets']=='1200')
            cells=[b,names[method],f"{float(row['presence_coverage']):.2%}",f"{float(row['support_completion']):.2%}",f"{float(row['risk_midpoint']):.2f}",f"[{float(row['risk_lower']):.2f}, {float(row['risk_upper']):.2f}]",f"{float(row['q5_fraction']):.2%}",f"{float(row['risk_unknown_fraction']):.2%}",f"{float(row['macro_f1']):.2%}"]
            assert '| '+' | '.join(cells)+' |' in text
        cells=[b]+[f"{float(next(r for r in errors if r['backend']==b and r['category']=='ALL_CLASSES' and int(r['risk_layer'])==q)['mean_error_rate']):.2%}" for q in range(1,6)]
        assert '| '+' | '.join(cells)+' |' in text
    strata=read(OUT/'risk_strata_summary.csv');ties=sum(r['upper_boundary_tied_next_tier']=='True' for r in strata if r['risk_tier']!='Q5')
    assert f'48条相邻层边界中有{ties}条' in text and ties==47
    assert '新增丰富细项的独立标注只有60目标' in text and '不是模型失败概率' in text
    doc=Document(OUT/'report.docx');assert len(doc.tables)>=6
    for table in doc.tables:assert len(table.rows)>=2
    for row in read(OUT/'equal_parts_sensitivity_summary.csv'):
        cells=[row['annotation_backend'],row['class_name'],f"{float(row['global_mean']):.2f}",f"{float(row['equal_parts_mean']):.2f}",row['changed_targets']]
        assert '| '+' | '.join(cells)+' |' in text
    contributions=read(OUT/'selection_dimension_score_totals.csv')
    for b in ('local','online'):
        for method in names:
            cells=[b,names[method]]+[f"{float(next(r for r in contributions if r['backend']==b and r['total_targets']=='1200' and r['method']==method and r['dimension']==d)['midpoint_points_mean']):.3f}" for d in DIMS]
            assert '| '+' | '.join(cells)+' |' in text

if __name__=='__main__':main()
