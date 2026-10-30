#!/usr/bin/env python3
"""Add evidence-bounded conclusions to the main-body teacher report."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import zipfile

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/工作量_正文整合版_20261008.docx'
NAME = '工作量_含最终结论与投稿依据_20261008'
OUT = ROOT / 'artifacts/workload_conclusions_20261008'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document(SOURCE)
body = doc.element.body

def txt(e):
    return ''.join(e.xpath('.//w:t/text()'))

def replace(e, value):
    nodes = e.xpath('.//w:t')
    nodes[0].text = value
    for n in nodes[1:]:
        n.text = ''

for p in doc.paragraphs:
    if p.text.startswith('上次老师让我分层。'):
        p.text = ('核心结论：总体成绩会掩盖具体条件下的薄弱表现。分层与换来源核验能把这些问题展开；'
                  '目前不能据此认定多维评分一定优于只看大小，也不能保证论文发表。')
    elif p.text.startswith('阅读顺序：'):
        p.text = ('阅读顺序：分层与模型降幅（第3—7页）→换来源核验（第8—10页）→随机与去项检查（第11—12页）'
                  '→固定数量实验（第13—15页）→后续对照（第16—19页）→最终结论与投稿依据（第20—21页）→汇报稿（第22页）。')

closing = next(e for e in body if txt(e).startswith('11  最后整理'))
replace(closing, '13  最后整理了什么？怎样向老师汇报？')
# The original page break before the closing speech now starts the conclusion.
before = list(body)

def para(text, size=11.5, bold=False, color=None, after=7, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.16
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = 'Noto Sans CJK SC'
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p

def title(text):
    para(text, size=18, bold=True, color='17365D', after=12, keep=True)

def sub(text):
    para(text, size=12.5, bold=True, color='17365D', after=5, keep=True)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for c, value in zip(t.rows[0].cells, headers):
        c.text = value
    for row in rows:
        for c, value in zip(t.add_row().cells, row):
            c.text = value
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs:
                    r.font.size = Pt(10.5)
                    r.bold = i == 0
                    r.font.name = 'Noto Sans CJK SC'
                    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    para('', after=2)
    return t

title('11  最终结论：这些实验究竟说明了什么？')
para('一句话：同一个模型的总体分数，不能完整说明它在哪些条件下可靠；测试报告需要同时写清图片构成和具体薄弱条件。', bold=True, color='17365D')
sub('结论一：分层能把总体成绩里的差别展开')
para('在六类目标、24个固定分类器和三个确认池中，两套状态描述分出的困难组，平均F1均低于容易组，差约14.44—41.80个百分点。随机分组的波动小得多。这说明当前规则确实能区分表现；中间五层并非总是逐层下降。')
sub('结论二：测试图片构成会改变测得的成绩')
para('每次固定120个目标、六类比例相同，只增加困难组比例，平均F1仍降低约12.29—34.77个百分点。模型权重没有变化：变化的是这次测什么。因此，比较模型成绩时，需要交代困难条件占多少。这里的起点是第1—4组混合，与“第1组对第5组”不是同一对照。')
sub('结论三：状态分析能给出更具体的后续测试对象')
para('例如ResNet-50对全部狗的识别率为81.92%，对背向狗为57.25%，低24.67个百分点。冻结已发现的条件后，再换图片核验；在目标大小分组相近的比较中，COCO和Objects365仍有很多条件表现较差。它们可作为后续重点测试对象，VOC2007多数条件样本不足。')
sub('最终判断的边界')
table(['已经支持', '目前不能推出'], [
    ('总体指标之外，报告具体条件表现有用。', 'F1降低说明模型被改坏，或模型已经得到改进。'),
    ('当前分层差别超过随机分组波动。', '多维评分一定优于只按目标大小分组。'),
    ('部分薄弱条件能在新来源中重复看到。', '某个状态就是错误原因，或AI状态标签全部准确。'),
])
para('应用意义：报告应回答“测试了哪些条件、每种条件有多少样本、模型分别表现如何”，让后续测试有具体方向。', bold=True)
doc.add_page_break()

title('12  为什么值得投稿？还缺哪部分说服力？')
para('我的判断：这些结果具备组织成测试领域实证论文的基础，但还不能承诺发表。仅说“难图片分数低”说服力有限；论文需要证明这套测试组织和诊断报告提供了什么额外信息。', bold=True, color='17365D')
sub('论文可以围绕这三项贡献组织')
table(['贡献', '这份工作的具体内容'], [
    ('一套可核查的测试流程', '先按预定状态规则分层，不读取分类器对错；再用固定预测评价。分组依据、样本数量、分类结果分别报告，便于检查与复算。'),
    ('一组有控制的实验证据', '固定模型与类别比例，补做随机分组、去掉字段、只看大小、固定120个目标和换来源核验，说明测试构成如何影响成绩。'),
    ('具体的条件诊断与适用边界', '把一个总分展开为背向、模糊等条件报告；先冻结条件，再在新图片中核验，区分可重复的问题与样本不足的条件。'),
])
sub('目前最需要讲清楚的三个问题')
para('① 比现有工作多了什么？已有可解释薄弱子集研究，不能把“用AI描述失败图片”本身当成全部新意。需要明确本文的问题、信息使用方式与验证步骤有哪些实际区别。')
para('② 多维状态比便宜的大小信息多提供什么？只看大小的端点差也很强；完整评分没有得到稳定优于它的统计证据。当前额外价值更适合从具体条件诊断说明，不能靠F1降幅更大来论证。')
para('③ 自动状态描述可信到什么程度？两套AI标签不等于人工真值，现有一致性与外部核验仍有局限。应给出可信度证据，并把尚不确定的判断保留为边界。')
sub('向老师回答“为什么能写成论文”')
para('“我的研究问题是：只报一个总分会不会漏掉模型的薄弱条件，以及怎样组织测试才能把这些问题说明白。我给出了固定模型下的分层、构成变化和新来源核验结果，并能报告具体条件。当前成果可以按测试实证研究来写；是否达到发表要求，还要看相对已有工作的新增价值和标签可信度。”')
para('当前定位：有证据支撑的投稿候选；工作量、图多、分数下降本身都不是发表保证。', bold=True, color='17365D')
doc.add_page_break()

# Move the new body elements before the original closing-page heading, keeping
# all original image relationships and portrait/landscape sections intact.
added = [e for e in list(body) if e not in before and e.tag != qn('w:sectPr')]
for e in added:
    body.remove(e)
    closing.addprevious(e)

doc.core_properties.title = '工作量：正文图表、最终结论与投稿依据'
doc.core_properties.subject = '明确已支持结论、论文贡献、投稿价值及证据边界'
target = ROOT / 'docs' / f'{NAME}.docx'
doc.save(target)

def media(path):
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}

assert media(SOURCE) == media(target)
assert len(doc.inline_shapes) == 19
assert len(doc.sections) == 7
pages = [[]]
for e in body:
    if e.tag == qn('w:sectPr'):
        continue
    if e.xpath('.//w:br[@w:type="page"]') or e.xpath('./w:pPr/w:sectPr'):
        pages.append([])
    else:
        pages[-1].append(e)
assert len(pages) == 22
md = []
for i, elements in enumerate(pages, 1):
    md.append(f'\n<!-- 第{i}页 -->\n')
    for e in elements:
        if e.tag == qn('w:tbl'):
            for row in e.xpath('./w:tr'):
                md.append('| ' + ' | '.join(txt(c) for c in row.xpath('./w:tc')) + ' |\n')
        elif txt(e).strip():
            md.append(txt(e) + '\n\n')
(ROOT / 'docs' / f'{NAME}.md').write_text(''.join(md), encoding='utf-8')
(OUT / 'audit.json').write_text(json.dumps({
    'source': str(SOURCE), 'word': str(target), 'expected_pages': 22,
    'figures': 19, 'image_bytes_preserved': True, 'conclusion_pages': [20, 21],
    'source_of_claims': 'docs/论文_最新.md, sections 1, 5, 7, 8',
    'render_verified': False,
}, ensure_ascii=False, indent=2), encoding='utf-8')
print(target)
