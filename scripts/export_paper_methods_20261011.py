#!/usr/bin/env python3
"""Export a standalone Methods chapter with native editable Word equations."""
from pathlib import Path
import argparse
import hashlib
import json
import re
import sys
import zipfile
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts/methods_section_20261011'
sys.path.insert(0, str(ROOT / 'artifacts/teacher_methodology_20261011/vendor'))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image as PILImage, ImageDraw
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Spacer, Table, TableStyle, KeepTogether
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
import pymupdf

NAME = '论文方法章节_视觉状态驱动模型测试_20261011'
SOURCE = ROOT / 'docs' / (NAME + '.md')
FONT = ROOT / 'scripts/assets/NotoSansSC-Regular.ttf'
FIGURE = ROOT / 'output/imagegen/qe_method_final_stage_image2_20261010.png'


def mr(value, plain=False):
    node = OxmlElement('m:r')
    if plain:
        props = OxmlElement('m:rPr')
        sty = OxmlElement('m:sty')
        sty.set(qn('m:val'), 'p')
        props.append(sty)
        node.append(props)
    properties = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Cambria Math')
    fonts.set(qn('w:hAnsi'), 'Cambria Math')
    properties.append(fonts)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '22')
    properties.append(size)
    node.append(properties)
    text = OxmlElement('m:t')
    text.set(qn('xml:space'), 'preserve')
    text.text = str(value)
    node.append(text)
    return node


def flat(*values):
    result = []
    for value in values:
        if isinstance(value, (list, tuple)):
            result.extend(flat(*value))
        elif isinstance(value, str):
            result.append(mr(value))
        elif value is not None:
            result.append(value)
    return result


def part(tag, nodes):
    value = OxmlElement('m:' + tag)
    value.extend(flat(nodes))
    return value


def script(base, lower=None, upper=None):
    tag = 'sSubSup' if lower is not None and upper is not None else 'sSub' if lower is not None else 'sSup'
    value = OxmlElement('m:' + tag)
    value.append(part('e', base))
    if lower is not None:
        value.append(part('sub', lower))
    if upper is not None:
        value.append(part('sup', upper))
    return value


def frac(numerator, denominator):
    value = OxmlElement('m:f')
    value.append(part('num', numerator))
    value.append(part('den', denominator))
    return value


def delim(body, begin='(', end=')'):
    value = OxmlElement('m:d')
    properties = OxmlElement('m:dPr')
    for name, symbol in [('begChr', begin), ('endChr', end)]:
        setting = OxmlElement('m:' + name)
        setting.set(qn('m:val'), symbol)
        properties.append(setting)
    value.append(properties)
    value.append(part('e', body))
    return value


def nary(lower, upper, body, symbol='∑'):
    value = OxmlElement('m:nary')
    properties = OxmlElement('m:naryPr')
    for name, setting in [('chr', symbol), ('limLoc', 'undOvr'),
                          ('subHide', '0' if lower is not None else '1'),
                          ('supHide', '0' if upper is not None else '1')]:
        item = OxmlElement('m:' + name)
        item.set(qn('m:val'), setting)
        properties.append(item)
    value.append(properties)
    value.append(part('sub', lower))
    value.append(part('sup', upper))
    value.append(part('e', body))
    return value


def hat(body):
    value = OxmlElement('m:acc')
    properties = OxmlElement('m:accPr')
    char = OxmlElement('m:chr')
    char.set(qn('m:val'), '̂')
    properties.append(char)
    value.append(properties)
    value.append(part('e', body))
    return value


def a():
    return script('A', 'i', '(r)')


def support(set_name='S'):
    return flat(script('n', 'c,s', '(r)'), delim(set_name))


def rec(state=True):
    return script(mr('Rec', True), 'c,s' if state else 'c', '(r)' if state else None)


def indicator(condition):
    return flat(mr('1', True), delim(condition, '[', ']'))


def binom(upper, lower):
    value = frac(upper, lower)
    props = OxmlElement('m:fPr')
    kind = OxmlElement('m:type')
    kind.set(qn('m:val'), 'noBar')
    props.append(kind)
    value.insert(0, props)
    return delim(value)


def display_math(number):
    if number == 1:
        return flat(script('z', 'i', '(r)'), '=', script('φ', 'r'), delim(flat(script('I', 'i'), ',', script('M', 'i'), ',', script('V', 'i'))), '.')
    if number == 2:
        body = flat(indicator(flat(script('y', 'i'), '=c')), indicator(flat('s∈', a())))
        return flat(support(), '=', nary('i∈S', None, body), '.')
    if number == 3:
        c = script('C', 'c', '(r)')
        summand = flat(mr('min', True), delim(flat(frac(support(), 'τ'), ',1')))
        left = flat(c, delim('S'), '=', frac('1', flat('|', script('U', 'c'), '|')), nary(flat('s∈', script('U', 'c')), None, summand))
        right = flat(script('C', None, '(r)'), delim('S'), '=', frac('1', 'K'), nary('c=1', 'K', flat(script('C', 'c', '(r)'), delim('S'))))
        return flat(left, ',   ', right, '.')
    if number == 4:
        lower = flat('s∈', a(), ', ', support(), '<τ')
        body = flat(mr('log', True), delim(flat(support(), '+2')), '−', mr('log', True), delim(flat(support(), '+1')))
        return flat(script('g', 'c', '(r)'), delim('i|S'), '=', nary(lower, None, delim(body, '[', ']')), '.')
    if number == 5:
        return flat(script('H', 'i', '(r)'), '=', nary('d∈D', None, flat(script('h', 'd'), delim('i,r'))), '.')
    if number == 6:
        low_union = nary('ℓ=1', '4', script('Q', 'c,ℓ', '(r)'), '⋃')
        return flat(script('T', 'c', '(r)'), delim('ρ'), '=', mr('Sample', True),
                    delim(flat(script('Q', 'c,5', '(r)'), ',ρm')), '∪', mr('Sample', True),
                    delim(flat(low_union, ',', delim('1−ρ'), 'm')), '.')
    if number == 7:
        numerator = nary('i∈T', None, flat(indicator(flat(script('y', 'i'), '=c')),
                          indicator(flat('s∈', a())), indicator(flat(script(hat('y'), 'i'), '=', script('y', 'i')))))
        return flat(rec(), delim('T'), '=', frac(numerator, support('T')), '.')
    if number == 8:
        return flat(script('Δ', 'c,s', '(r)'), delim('T'), '=', rec(), delim('T'), '−', rec(False), delim('T'), '.')
    if number == 9:
        return flat('𝔼', delim(flat(mr('1', True), delim(mr('type observed', True))), '[', ']'),
                    '=1−', frac(binom('N−e', 'q'), binom('N', 'q')), '.')
    raise ValueError(number)


def inline_math(tex):
    simple = {'P': 'P', 'K': 'K', 'K=6': 'K=6', 'i': 'i', 'r': 'r', 'c': 'c', 's': 's',
              'S': 'S', 'D': 'D', 'd': 'd', 'm': 'm', 'q': 'q', 'N': 'N', 'e': 'e',
              'Km': 'Km', 'n': 'n', 'v': 'v', 't': 't', r'\tau': 'τ', r'\rho': 'ρ'}
    if tex in simple:
        return flat(simple[tex])
    if tex in ('O(tnv)', 'O(nv+n+t)'):
        return flat('O', delim(tex[2:-1]))
    if tex == r'O(n\log n)':
        return flat('O', delim(flat('n', mr('log', True), ' n')))
    if tex in ('I_i', 'b_i', 'y_i', 'x_i', 'M_i', 'V_i', 'a_i'):
        return flat(script(tex[0], 'i'))
    if tex == 'U_c':
        return flat(script('U', 'c'))
    if tex == 'h_d':
        return flat(script('h', 'd'))
    if tex == r'f_\theta':
        return flat(script('f', 'θ'))
    if tex == r'\phi_r':
        return flat(script('φ', 'r'))
    if tex == 'A_i^{(r)}':
        return flat(a())
    if tex == r'\mathbf{1}[\cdot]':
        return indicator('·')
    if tex == r'r\in\{\mathrm{local},\mathrm{online}\}':
        return flat('r∈', delim(flat(mr('local', True), ',', mr('online', True)), '{', '}'))
    if tex == r'\rho m':
        return flat('ρm')
    if tex == r'T^{(r)}(\rho)':
        return flat(script('T', None, '(r)'), delim('ρ'))
    if tex == r'\hat y_i=f_\theta(x_i)':
        return flat(script(hat('y'), 'i'), '=', script('f', 'θ'), delim(script('x', 'i')))
    raise ValueError('Unmapped inline formula: ' + tex)


def parse():
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    blocks, index, equations = [], 0, 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == '$$':
            latex = []
            index += 1
            while lines[index].strip() != '$$':
                latex.append(lines[index])
                index += 1
            equations += 1
            blocks.append(('equation', '\n'.join(latex), equations))
        elif line == '```algorithm':
            algorithm = []
            index += 1
            while index < len(lines) and lines[index].strip() != '```':
                algorithm.append(lines[index])
                index += 1
            blocks.append(('algorithm', algorithm, None))
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                cells = [v.strip() for v in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(':?-+:?', c) for c in cells):
                    rows.append(cells)
                index += 1
            blocks.append(('table', rows, None))
            continue
        elif line.startswith('!['):
            target = re.fullmatch(r'!\[[^]]*\]\(([^)]+)\)', line)[1]
            blocks.append(('image', str((SOURCE.parent / target).resolve()), None))
        elif line.startswith('# '):
            blocks.append(('title', line[2:], None))
        elif line.startswith('## '):
            blocks.append(('heading', line[3:], None))
        elif line.startswith('### '):
            blocks.append(('subheading', line[4:], None))
        elif re.match(r'^(?:图|表|算法)\d+(?:[ \u3000]|$)', line):
            blocks.append(('caption', line, None))
        else:
            blocks.append(('body', line, None))
        index += 1
    assert equations == 9
    starts = ('z_i', 'n_{c,s}', 'C_c', 'g_c', 'H_i', 'T_c', r'\operatorname{Rec}', r'\Delta', r'\mathbb{E}')
    for block, value, number in blocks:
        if block == 'equation':
            assert value.startswith(starts[number - 1])
    return blocks


def font(run, size=11, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体' if not bold else '黑体')


def word_inline(p, text, size=11):
    for chunk in re.split(r'(\$[^$]+\$)', text):
        if chunk.startswith('$') and chunk.endswith('$'):
            math = OxmlElement('m:oMath')
            math.extend(inline_math(chunk[1:-1]))
            p._p.append(math)
        else:
            font(p.add_run(chunk), size)


def word_export(blocks):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [('Normal', 11), ('Title', 17), ('Heading 1', 13), ('Heading 2', 11.5)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('000000')
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体' if name == 'Normal' else '黑体')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for index, (kind, text, number) in enumerate(blocks):
        if kind == 'equation':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(7)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            formula = OxmlElement('m:oMathPara')
            math = OxmlElement('m:oMath')
            math.extend(display_math(number))
            formula.append(math)
            p._p.append(formula)
        elif kind == 'image':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(text, width=Cm(17))
        elif kind == 'table':
            table = doc.add_table(rows=len(text), cols=2)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            table.columns[0].width, table.columns[1].width = Cm(3.0), Cm(14.0)
            for row_number, (row, cells) in enumerate(zip(table.rows, text)):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell, value in zip(row.cells, cells):
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    run = p.add_run(value)
                    font(run, 9.5, row_number == 0)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            if kind in ('title', 'heading', 'subheading'):
                style = {'title': 'Title', 'heading': 'Heading 1', 'subheading': 'Heading 2'}[kind]
                p = doc.add_paragraph(text, style=style)
                p.paragraph_format.space_before = Pt(9 if kind != 'title' else 0)
                p.paragraph_format.space_after = Pt(5)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.keep_together = True
                if index + 1 < len(blocks) and blocks[index + 1][0] == 'equation':
                    p.paragraph_format.keep_with_next = True
                if kind == 'caption' and text.startswith('算法1'):
                    p.paragraph_format.keep_with_next = True
                if kind == 'body':
                    p.paragraph_format.first_line_indent = Cm(.75)
                word_inline(p, text, 9 if kind == 'caption' else 11)
    doc.core_properties.title = '研究方法：视觉状态驱动的模型测试'
    doc.core_properties.author = 'QE项目'
    path = SOURCE.with_suffix('.docx')
    doc.save(path)
    return path


def math_image(tex, size=12):
    tex = tex.strip().rstrip('.')
    if r'\substack' in tex:
        tex = tex.replace(r'\substack{s\in A_i^{(r)}\\n_{c,s}^{(r)}(S)<\tau}',
                          r's\in A_i^{(r)},\,n_{c,s}^{(r)}(S)<\tau')
    digest = hashlib.sha256((tex + str(size)).encode()).hexdigest()[:16]
    path = OUT / 'equations' / (digest + '.png')
    if not path.exists():
        fig = plt.figure(figsize=(9, 1), dpi=210)
        fig.text(.02, .45, '$' + tex + '$', size=size)
        fig.savefig(path, bbox_inches='tight', pad_inches=.015, transparent=False)
        plt.close(fig)
    with PILImage.open(path) as im:
        return path, im.width * 72 / 210, im.height * 72 / 210


def pdf_inline(text):
    parts = []
    for chunk in re.split(r'(\$[^$]+\$)', text):
        if chunk.startswith('$') and chunk.endswith('$'):
            path, w, h = math_image(chunk[1:-1], 10.5)
            parts.append(f'<img src="{path}" width="{w:.3f}" height="{h:.3f}" valign="middle"/>')
        else:
            parts.append(escape(chunk))
    return ''.join(parts)


def pdf_paragraph(text, style):
    paragraph = Paragraph(pdf_inline(text), style)
    for fragment in paragraph.frags:
        callback = getattr(fragment, 'cbDefn', None)
        if not fragment.text and getattr(callback, 'kind', None) == 'img':
            # Give CJK line breaking an object character for inline math images.
            # The image callback still controls rendering and occupies its own width.
            fragment.text = '\uFFFC'
    return paragraph


def pdf_export(blocks):
    pdfmetrics.registerFont(TTFont('NotoSC', str(FONT)))
    pdfmetrics.registerFontFamily('NotoSC', normal='NotoSC', bold='NotoSC')
    spec = {'title': (17, 25), 'heading': (13, 21), 'subheading': (11.5, 18),
            'body': (11, 17.5), 'caption': (9, 14), 'cell': (9.5, 15),
            'algorithm': (9.5, 16)}
    styles = {name: ParagraphStyle(name, fontName='NotoSC', fontSize=size, leading=leading,
                     wordWrap='CJK', spaceAfter=6, firstLineIndent=22 if name == 'body' else 0,
                     spaceBefore=8 if name in ('heading', 'subheading') else 0,
                     keepWithNext=False)
              for name, (size, leading) in spec.items()}
    path = SOURCE.with_suffix('.pdf')
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=56.7, rightMargin=56.7,
                           topMargin=53, bottomMargin=48, title='研究方法：视觉状态驱动的模型测试', author='QE项目')
    width = A4[0] - 113.4
    story, pending_headings, i = [], [], 0
    def equation_flowables(tex):
        image, w, h = math_image(tex)
        factor = min(1, width / w)
        return [Spacer(1, 3), Image(str(image), width=w * factor, height=h * factor), Spacer(1, 9)]
    while i < len(blocks):
        kind, text, number = blocks[i]
        if kind in ('title', 'heading', 'subheading') or (kind == 'caption' and text.startswith(('算法1', '表1'))):
            pending_headings.append(pdf_paragraph(text, styles[kind]))
            i += 1
            continue
        if kind == 'equation':
            flowables = equation_flowables(text)
        elif kind == 'image':
            with PILImage.open(text) as im:
                h = width * im.height / im.width
            figure = [Image(text, width=width, height=h), Spacer(1, 5)]
            if i + 1 < len(blocks) and blocks[i + 1][0] == 'caption':
                figure.append(pdf_paragraph(blocks[i + 1][1], styles['caption']))
                i += 1
            flowables = figure
        elif kind == 'table':
            rows = [[pdf_paragraph(c, styles['cell']) for c in row] for row in text]
            first_width = 174 if text[0][0] == '符号' else 85
            table = Table(rows, colWidths=[first_width, width - first_width], repeatRows=1)
            table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .45, colors.HexColor('#999999')),
                                      ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F1F1')),
                                      ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                                      ('TOPPADDING', (0, 0), (-1, -1), 5),
                                      ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
            flowables = [table, Spacer(1, 6)]
        elif kind == 'algorithm':
            flowables = []
            for line in text:
                number_match = re.match(r'^(\d+) ( +)', line)
                indent = len(number_match[2]) if number_match else 0
                algorithm_style = ParagraphStyle('algorithm_line', parent=styles['algorithm'],
                                                 leftIndent=indent * 3, spaceAfter=2)
                flowables.append(pdf_paragraph(line.strip(), algorithm_style))
            flowables.append(Spacer(1, 5))
        else:
            paragraph = pdf_paragraph(text, styles[kind])
            flowables = [paragraph]
            if kind == 'body':
                if i + 1 < len(blocks) and blocks[i + 1][0] == 'equation':
                    flowables.extend(equation_flowables(blocks[i + 1][1]))
                    i += 1
        story.append(KeepTogether([*pending_headings, *flowables]))
        pending_headings = []
        i += 1
    def footer(canvas, document):
        canvas.setFont('NotoSC', 8.5)
        canvas.setFillColor(colors.HexColor('#555555'))
        canvas.drawCentredString(A4[0] / 2, 25, str(document.page))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def audit_and_preview(docx_path, pdf_path, blocks, figure_bytes=None):
    namespaces = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    from lxml import etree
    with zipfile.ZipFile(docx_path) as package:
        xml = package.read('word/document.xml')
        tree = etree.fromstring(xml)
        display_count = len(tree.findall('.//m:oMathPara', namespaces))
        assert display_count == 9
        assert b'\\frac' not in xml and b'\\tau' not in xml and b'/home/hcf' not in xml
        media = [p for p in package.namelist() if p.startswith('word/media/')]
        expected_figure = FIGURE.read_bytes() if figure_bytes is None else figure_bytes
        assert len(media) == 1 and package.read(media[0]) == expected_figure
    pdf = pymupdf.open(pdf_path)
    previews = OUT / 'preview'
    previews.mkdir(exist_ok=True)
    thumbnails = []
    for index, page in enumerate(pdf, 1):
        assert len(page.get_text().strip()) > 40, ('Possible blank page', index)
        path = previews / f'page_{index:02d}.png'
        page.get_pixmap(matrix=pymupdf.Matrix(1.65, 1.65), alpha=False).save(path)
        im = PILImage.open(path).convert('RGB')
        im.thumbnail((480, 680))
        tile = PILImage.new('RGB', (510, 720), '#E6EAED')
        tile.paste(im, ((510 - im.width) // 2, 25))
        ImageDraw.Draw(tile).text((12, 7), f'Page {index}', fill='black')
        thumbnails.append(tile)
    montage = PILImage.new('RGB', (1020, 720 * ((len(pdf) + 1) // 2)), 'white')
    for index, tile in enumerate(thumbnails):
        montage.paste(tile, ((index % 2) * 510, (index // 2) * 720))
    montage.save(previews / 'contact_sheet.png')
    audit = {'source': str(SOURCE.relative_to(ROOT)),
             'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
             'word_display_equations': display_count,
             'word_figure_embedded': True, 'pdf_pages': len(pdf), 'new_experiments': False,
             'pdf_generation': 'Generated from the same Markdown source; not a render of Word.',
             'basis': {
                 'problem_and_states': ['docs/论文_最新.md', 'scripts/analyze_dual_vlm_study.py'],
                 'coverage_and_quota_selection': ['scripts/freeze_independent_sixclass_selection.py'],
                 'risk_scoring': ['scripts/analyze_sixclass_difficulty.py'],
                 'risk_mixture': ['scripts/run_controlled_risk_distribution.py'],
                 'equal_error_breadth': ['scripts/run_class_stratified_error_controls.py', 'scripts/run_coverage_incremental_value.py'],
                 'weak_condition_discovery': ['scripts/prepare_independent_sixclass_pool.py'],
                 'condition_confirmation': ['scripts/analyze_independent_sixclass.py'],
                 'state_agreement_and_sensitivity': ['scripts/analyze_dual_vlm_study.py', 'scripts/analyze_label_disagreement_sensitivity.py'],
                 'risk_ablations_and_controls': ['scripts/run_random_stratification_ablation.py', 'scripts/run_scale_only_confirmation.py'],
                 'error_cluster_definition': ['scripts/run_voc2012_confirmation_benchmark.py', 'scripts/run_controlled_risk_distribution.py'],
             }}
    (OUT / 'writing_basis.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'word': str(docx_path), 'pdf': str(pdf_path), 'markdown': str(SOURCE),
                      'pdf_pages': len(pdf), 'editable_display_equations': display_count}, ensure_ascii=False, indent=2))


def main():
    global SOURCE, OUT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE,
                        help='Markdown source; Word and PDF are written beside it.')
    parser.add_argument('--output-dir', type=Path, default=OUT,
                        help='Directory for equation images, previews and the writing audit.')
    args = parser.parse_args()
    SOURCE = args.source.resolve()
    OUT = args.output_dir.resolve()
    (OUT / 'equations').mkdir(parents=True, exist_ok=True)
    blocks = parse()
    word = word_export(blocks)
    pdf = pdf_export(blocks)
    audit_and_preview(word, pdf, blocks)


if __name__ == '__main__':
    main()
