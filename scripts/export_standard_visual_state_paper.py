#!/usr/bin/env python3
"""Create a conventional manuscript with numbered citations, three-line tables and OMML."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

from export_empirical_paper_docx import add_page_number, set_run_font
from export_sixclass_stratified_paper import run as mr, sub, frac, sum_term

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/visual_state_testing_paper.md'
ORIGINAL=ROOT/'docs/visual_state_testing_empirical_paper_updated_20261006.md'
OUTPUT=SOURCE.with_suffix('.docx')
AUDIT=ROOT/'artifacts/visual_state_paper_updated_20261006/standard_format_audit.json'


def superscript(base,exponent):
    node=OxmlElement('m:sSup');body=OxmlElement('m:e');tail=OxmlElement('m:sup')
    body.extend(base);tail.extend(exponent);node.extend([body,tail]);return node


def inline_math(expression):
    math=OxmlElement('m:oMath')
    symbols={'U_c':sub('U','c'),'n_s(S)':sub('n','s'),'[0,h_{max}]':sub('h','max')}
    if expression=='U_c':parts=[symbols[expression]]
    elif expression=='n_s(S)':parts=[symbols[expression],mr('(S)')]
    elif expression=='[0,h_{max}]':parts=[mr('[0, '),symbols[expression],mr(']')]
    elif expression=='10^{-4}':parts=[superscript([mr('10')],[mr('−4')])]
    elif expression=='2/2^8=0.0078125':parts=[frac([mr('2')],[superscript([mr('2')],[mr('8')])]),mr(' = 0.0078125')]
    elif expression=='F1_c=2P_cR_c/(P_c+R_c)':parts=[sub('F1','c'),mr(' = '),frac([mr('2'),sub('P','c'),sub('R','c')],[sub('P','c'),mr(' + '),sub('R','c')])]
    elif expression=='1-\\sum_k p_k(x)^2':
        parts=[mr('1 − '),sum_term([mr('k')],[superscript([sub('p','k'),mr('(x)')],[mr('2')])])]
    elif expression=='\\max(2,\\min(8,\\operatorname{round}(\\sqrt m)))':
        radical=OxmlElement('m:rad');pr=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');pr.append(hide)
        degree=OxmlElement('m:deg');body=OxmlElement('m:e');body.append(mr('m'));radical.extend([pr,degree,body])
        parts=[mr('max(2, min(8, round('),radical,mr(')))')]
    else:
        value=expression.replace('\\geq',' ≥ ').replace('\\leq',' ≤ ')
        assert '\\' not in value and '_' not in value and '^' not in value,expression
        parts=[mr(value)]
    math.extend(parts)
    return math


def compact_citation(raw):
    values=[int(v) for v in raw[1:-1].split(',')]
    result=[];start=previous=values[0]
    for value in values[1:]+[None]:
        if value is not None and value==previous+1:previous=value;continue
        result.append(str(start) if start==previous else f'{start}–{previous}')
        start=previous=value
    return '['+','.join(result)+']'


def add_inline(paragraph,value,size=10.5,bold=False,citations=True,east='宋体'):
    # All math tokens are real Word objects, including equations inside table cells.
    pattern=r'(\$[^$]+\$|\*\*[^*]+\*\*|`[^`]+`|\[\d+(?:,\d+)*\])'
    for token in re.split(pattern,value):
        if not token:continue
        if token.startswith('$'):
            paragraph._p.append(inline_math(token[1:-1]));continue
        if citations and re.fullmatch(r'\[\d+(?:,\d+)*\]',token):
            node=OxmlElement('w:hyperlink');node.set(qn('w:anchor'),'ref_'+token[1:-1].split(',')[0])
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
            pos=OxmlElement('w:vertAlign');pos.set(qn('w:val'),'superscript');pr.append(pos)
            font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');pr.append(font)
            fontsize=OxmlElement('w:sz');fontsize.set(qn('w:val'),'18');pr.append(fontsize)
            r.append(pr);t=OxmlElement('w:t');t.text=compact_citation(token);r.append(t);node.append(r);paragraph._p.append(node)
            continue
        strong=token.startswith('**');code=token.startswith('`')
        token=token[2:-2] if strong else token[1:-1] if code else token
        r=paragraph.add_run(token);set_run_font(r,'Times New Roman',east,size,bold or strong)


def border(parent,side,style,size=None):
    node=OxmlElement('w:'+side);node.set(qn('w:val'),style)
    if size is not None:node.set(qn('w:sz'),str(size));node.set(qn('w:color'),'000000')
    parent.append(node)


def three_line_table(document,lines):
    rows=[[x.strip() for x in line.strip('|').split('|')] for line in lines]
    assert all(re.fullmatch(r':?-{3,}:?',x.replace(' ','')) for x in rows[1])
    rows.pop(1);assert len({len(row) for row in rows})==1
    table=document.add_table(rows=len(rows),cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    props=table._tbl.tblPr
    edges=OxmlElement('w:tblBorders')
    for side in ('left','right','insideH','insideV'):border(edges,side,'nil')
    border(edges,'top','single',8);border(edges,'bottom','single',8);props.append(edges)
    width=16/len(rows[0])
    if len(rows[0])==2:widths=[4.,12.]
    elif len(rows[0])==3:widths=[3.4,6.3,6.3]
    else:widths=[width]*len(rows[0])
    for column,cm in zip(table.columns,widths):column.width=Cm(cm)
    for i,values in enumerate(rows):
        pr=table.rows[i]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=table.cell(i,j);cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            # Explicit nil cell borders prevent table-style grid inheritance.
            cb=OxmlElement('w:tcBorders')
            for side in ('top','left','right'):border(cb,side,'nil')
            border(cb,'bottom','single' if i==0 else 'nil',4 if i==0 else None);cp.append(cb)
            margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'65');edge.set(qn('w:type'),'dxa');margins.append(edge)
            cp.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.1
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or j>=2 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p,value,8.5,i==0)
    gap=document.add_paragraph();gap.paragraph_format.space_after=Pt(2);gap.paragraph_format.line_spacing=.5


def bookmark(paragraph,number):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),f'ref_{number}')
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    paragraph._p.append(start);return end


def build():
    document=Document();section=document.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2.5);section.bottom_margin=Cm(2.5);section.left_margin=Cm(2.5);section.right_margin=Cm(2.5)
    add_page_number(section.footer.paragraphs[0])
    for name,size in [('Normal',10.5),('Title',18),('Subtitle',12),('Heading 1',12),('Heading 2',11),('Heading 3',10.5),('Caption',9)]:
        style=document.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=None
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if name.startswith('Heading') or name=='Title' else '宋体')
        style.paragraph_format.space_after=Pt(5)
        if name.startswith('Heading'):
            style.font.bold=True;style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(9)
    lines=SOURCE.read_text().splitlines();index=0;reference_section=False;abstract=False
    existing=Document(ORIGINAL.with_suffix('.docx'))
    equations=[deepcopy(x) for x in existing._element.findall('.//'+qn('m:oMath'))]
    assert len(equations)==3
    displays=0
    while index<len(lines):
        line=lines[index].strip();index+=1
        if not line:continue
        if line=='$$':
            while index<len(lines) and lines[index].strip()!='$$':index+=1
            assert index<len(lines);index+=1
            p=document.add_paragraph();p.paragraph_format.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6)
            p.add_run('\t');p._p.append(equations[displays]);displays+=1;p.add_run(f'\t({displays})')
            continue
        if line.startswith('|'):
            group=[line]
            while index<len(lines) and lines[index].startswith('|'):group.append(lines[index]);index+=1
            three_line_table(document,group);continue
        picture=re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)',line)
        if picture:
            file=Path(picture[2]);width_px,height_px=Image.open(file).size
            width_cm=min(15.8,17*width_px/height_px)
            p=document.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(file),width=Cm(width_cm))
            p=document.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_inline(p,picture[1],9)
            continue
        heading=re.match(r'^(#{1,4})\s+(.+)',line)
        if heading:
            level=len(heading[1]);title=heading[2]
            if level==1:
                p=document.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_inline(p,title,18,True,east='黑体')
            else:
                p=document.add_paragraph(style='Heading '+str(level-1));add_inline(p,title,12 if level==2 else 11 if level==3 else 10.5,True,east='黑体')
            reference_section=title=='参考文献' if level==2 else reference_section
            abstract=title in ('摘要','Abstract') if level==2 else abstract
            continue
        if line.startswith('Visual-State Coverage and Risk-Stratified Testing'):
            p=document.add_paragraph(style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_inline(p,line,12);continue
        if re.match(r'^\*\*表(?:\d+|A\d+)\s',line):
            p=document.add_paragraph(style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            add_inline(p,line,9,True);continue
        ref=re.match(r'^\[(\d+)\]\s+(.+)',line) if reference_section else None
        p=document.add_paragraph();p.paragraph_format.line_spacing=1.35;p.paragraph_format.space_after=Pt(4)
        if ref:
            number=int(ref[1]);end=bookmark(p,number)
            p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(-.75)
            add_inline(p,line,9,citations=False);p._p.append(end)
        else:
            p.paragraph_format.first_line_indent=Pt(0) if abstract or line.startswith('**关键词') or re.match(r'^\d+\. ',line) else Cm(.74)
            add_inline(p,line)
    assert displays==3
    document.core_properties.title='基于视觉状态覆盖与风险分层的多类别图像分类模型测试：一项实证研究'
    document.core_properties.subject='多类别视觉状态测试实证研究'
    document.core_properties.author=''
    document.save(OUTPUT)


def audit():
    original=ORIGINAL.read_text();clean=SOURCE.read_text()
    # All 21 result tables retain exact cell content; only captions/references are renumbered.
    for n in range(1,22):
        a=original.split(f'**表{n} ')[1].split('\n\n',1)[1].split('\n\n',1)[0]
        b=clean.split(f'**表{n+4} ')[1].split('\n\n',1)[1].split('\n\n',1)[0]
        assert a==b,(n,'result changed')
    assert re.findall(r'^\*\*表(\d+) ',clean,re.M)==[str(n) for n in range(1,26)]
    bibliography=clean.split('## 参考文献')[1].split('## 附录A')[0]
    assert re.findall(r'^\[(\d+)\]',bibliography,re.M)==[str(n) for n in range(1,14)]
    body=clean.split('## 参考文献')[0]
    cited={int(n) for token in re.findall(r'\[(\d+(?:,\d+)*)\]',body) for n in token.split(',')}
    assert cited==set(range(1,14))
    assert not any(token in clean for token in ('中文论文整合稿','旧稿','本版本','本地审阅链接'))
    with ZipFile(OUTPUT) as z:
        e=ET.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math',
            'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        tables=e.findall('.//w:tbl',ns);assert len(tables)==26
        for table in tables:
            edges=table.find('./w:tblPr/w:tblBorders',ns)
            for side in ('left','right','insideV','insideH'):assert edges.find('w:'+side,ns).get(qn('w:val'))=='nil'
            for side in ('top','bottom'):assert edges.find('w:'+side,ns).get(qn('w:val'))=='single'
            rows=table.findall('./w:tr',ns)
            assert rows[0].find('./w:trPr/w:tblHeader',ns) is not None
            for ri,row in enumerate(rows):
                for cell in row.findall('./w:tc',ns):
                    cb=cell.find('./w:tcPr/w:tcBorders',ns)
                    assert cb.find('w:bottom',ns).get(qn('w:val'))==('single' if ri==0 else 'nil')
                    assert cell.find('./w:tcPr/w:shd',ns) is None
        texts=[x.text or '' for x in e.findall('.//w:t',ns)]
        assert not any('$' in t or '\\sum' in t or '\\sqrt' in t for t in texts)
        assert not any('/home/hcf' in t or 'scripts/' in t for t in texts)
        bookmarks={x.get(qn('w:name')) for x in e.findall('.//w:bookmarkStart',ns)}
        assert {f'ref_{n}' for n in range(1,14)}<=bookmarks
        citations=e.findall('.//w:hyperlink',ns);assert citations
        assert all(x.get(qn('w:anchor')) in bookmarks for x in citations)
        assert all(x.find('.//w:vertAlign',ns).get(qn('w:val'))=='superscript' for x in citations)
        equations=len(e.findall('.//m:oMath',ns))
        inline_count=len(re.findall(r'(?<!\$)\$(?!\$)[^$]+\$(?!\$)',clean))
        assert equations==inline_count+3,(equations,inline_count)
        assert len(e.findall('.//a:blip',ns))==15
    result=dict(status='pass',markdown=str(SOURCE),word=str(OUTPUT),three_line_tables=26,result_tables_unchanged=21,
                bibliography_entries=13,cited_reference_ids=sorted(cited),superscript_citation_groups=len(citations),
                editable_equations=equations,display_equations_numbered=3,inline_equations=inline_count,figures=15,
                version_notes_and_local_paths_removed=True,source_sha256=hashlib.sha256(clean.encode()).hexdigest(),
                scope='A4 single-column manuscript; structural OOXML audit, not a venue-specific template')
    AUDIT.parent.mkdir(parents=True,exist_ok=True);AUDIT.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':
    build();audit()
