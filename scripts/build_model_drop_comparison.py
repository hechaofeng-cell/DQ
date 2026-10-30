#!/usr/bin/env python3
"""Compare Q1-to-Q5 performance drops, pairing the same model and seed."""
from pathlib import Path
import csv
import json
import hashlib
import shutil
import zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/model_drop_comparison_20261008'
FIG=OUT/'figures';FIG.mkdir(parents=True,exist_ok=True)
NAME='八种模型_分层降幅对比_20261008'
INPUT=ROOT/'artifacts/model_comparison_figures_20261008/model_tier_scores.csv'
RANDOM=ROOT/'artifacts/random_stratification_ablation_20261006/tables/random_summary.csv'
MODELS=['ResNet-18','ResNet-50','MobileNetV3-Small','MobileNetV3-Large','DenseNet-121','EfficientNet-B0','ConvNeXt-Tiny','Swin-T']
POOLS={'B':'COCO','C':'VOC2007','D':'Objects365'}
SOURCES={'local':'本地AI描述分组','online':'在线AI描述分组'}
def read(p):
    with p.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
rows=read(INPUT)
lookup={(r['pool'],r['label_source'],r['model'],r['seed'],int(r['tier'])):float(r['macro_f1_100']) for r in rows}
assert len(lookup)==len(rows)==720
seeds={m:sorted({r['seed'] for r in rows if r['model']==m}) for m in MODELS}
assert all(len(v)==3 for v in seeds.values())
drops={}
run_rows=[]
summary=[]
for source in SOURCES:
    for pool in POOLS:
        for m in MODELS:
            vals=[]
            for seed in seeds[m]:
                low=lookup[pool,source,m,seed,1];high=lookup[pool,source,m,seed,5]
                delta=low-high;assert delta>0
                vals.append(delta)
                run_rows.append([pool,POOLS[pool],source,m,seed,low,high,delta])
            a=np.array(vals);drops[pool,source,m]=a
            summary.append([pool,POOLS[pool],source,m,a.mean(),a.min(),a.max()])
    for pool in POOLS:
        observed=-100*float(next(r for r in read(RANDOM) if r['pool']==pool and r['source']==source)['observed_delta'])
        assert abs(np.mean([drops[pool,source,m].mean() for m in MODELS])-observed)<1e-9

for filename,headers,records in [
    ('paired_model_drops.csv',['pool','dataset','label_source','model','seed','Q1_f1_100','Q5_f1_100','drop_percentage_points'],run_rows),
    ('model_drop_summary.csv',['pool','dataset','label_source','model','mean_drop_pp','minimum_of_3_runs','maximum_of_3_runs'],summary)]:
    with (OUT/filename).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(records)
font_manager.fontManager.addfont(str(ROOT/'scripts/assets/NotoSansSC-Regular.ttf'))
plt.rcParams.update({'font.family':'Noto Sans SC','font.size':14,'svg.fonttype':'none','pdf.fonttype':42,
    'axes.spines.top':False,'axes.spines.right':False,'axes.linewidth':.7,'xtick.major.size':0,'ytick.major.size':0})
limit=float(np.ceil((max(a.max() for a in drops.values())+5)/10)*10)
figures=[]
for source,label in SOURCES.items():
    fig,axes=plt.subplots(1,3,figsize=(15,6.2),layout='constrained',sharey=True)
    for ax,(pool,name) in zip(axes,POOLS.items()):
        mean=np.array([drops[pool,source,m].mean() for m in MODELS])
        lo=np.array([drops[pool,source,m].min() for m in MODELS])
        hi=np.array([drops[pool,source,m].max() for m in MODELS])
        ax.barh(range(8),mean,height=.55,color='#3574A8')
        ax.errorbar(mean,range(8),xerr=[mean-lo,hi-mean],fmt='none',ecolor='#244761',capsize=3,lw=1)
        for i,v in enumerate(mean):ax.text(hi[i]+.8,i,f'{v:.2f}',va='center',fontsize=13)
        ax.set(title=name,yticks=range(8),yticklabels=MODELS,ylim=(7.65,-.65),xlim=(0,limit),
               xticks=np.arange(0,limit+1,10),xlabel='第1组 → 第5组：降幅（百分点）')
    stem=f'01_drop_{source}'
    for ext in ['png','svg','pdf']:fig.savefig(FIG/f'{stem}.{ext}',dpi=250,bbox_inches='tight',facecolor='white')
    plt.close(fig);figures.append((stem,label))

doc=Document();sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=Cm(29.7),Cm(21)
sec.left_margin=sec.right_margin=Cm(1.2);sec.top_margin=sec.bottom_margin=Cm(1.1)
normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(10)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
normal.paragraph_format.space_after=Pt(4)
for i,(stem,label) in enumerate(figures):
    if i:doc.add_page_break()
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
    run=p.add_run(f'8种模型的分层降幅对比（{label}）')
    run.font.size=Pt(17);run.font.bold=True;run.font.color.rgb=RGBColor.from_string('24577E')
    p=doc.add_paragraph('降幅 = 第1组Macro-F1 − 第5组Macro-F1；单位：百分点。')
    p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/f'{stem}.png'),width=Cm(26.9))
    doc.add_paragraph('横条为同一模型3份训练结果的配对降幅平均值；黑色短线为3份配对降幅的最小—最大范围。模型保持不变，比较的是规则预计较容易组与较困难组的表现差异。')
doc.core_properties.title='8种模型的分层性能降幅对比'
doc.save(OUT/f'{NAME}.docx');shutil.copy2(OUT/f'{NAME}.docx',ROOT/'docs'/f'{NAME}.docx')
shutil.copy2(INPUT,OUT/INPUT.name)
shutil.copy2(Path(__file__),OUT/Path(__file__).name)
audit={'baseline':'Q1_macro_f1','comparison':'Q5_macro_f1','drop':'Q1 minus Q5','unit':'percentage points',
       'paired_runs':144,'mean_values':48,'aggregate_checks':6,'render_verified':False,
       'input_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [INPUT,RANDOM]}}
(OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
(OUT/'说明.txt').write_text('图显示的是层间降幅，而非绝对F1。基准固定为第1组，对比第5组；不能称为重新训练前后的下降，也不能把较大降幅当作模型更强。每个降幅先在同一个模型、同一个训练种子下计算，再对3份结果取平均；单位是百分点，不是相对下降百分比。没有新增训练或推理。\n',encoding='utf-8')
def package():
    with zipfile.ZipFile(ROOT/'docs/八种模型_分层降幅对比图与数据_20261008.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in FIG.glob('*'):z.write(p,p.relative_to(OUT))
        for p in OUT.glob('*'):
            if p.is_file() and p.suffix not in ['.docx']:z.write(p,p.name)
package()
print(json.dumps({'word':str(ROOT/'docs'/f'{NAME}.docx'),'figures':2,
     'COCO_ResNet50_local_drop_pp':float(drops['B','local','ResNet-50'].mean())},ensure_ascii=False))
