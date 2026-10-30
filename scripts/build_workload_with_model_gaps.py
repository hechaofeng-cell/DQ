#!/usr/bin/env python3
"""Compare Q1-to-Q5 performance drops, pairing the same model and seed."""
from pathlib import Path
import csv
import json
import hashlib
import shutil
import zipfile
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts/workload_model_gaps_20261008'
FIG=OUT/'figures';FIG.mkdir(parents=True,exist_ok=True)
NAME='工作量_含模型降幅多图版_20261008'
INPUT=ROOT/'artifacts/model_comparison_figures_20261008/model_tier_scores.csv'
RANDOM=ROOT/'artifacts/random_stratification_ablation_20261006/tables/random_summary.csv'
MODELS=['ResNet-18','ResNet-50','MobileNetV3-Small','MobileNetV3-Large','DenseNet-121','EfficientNet-B0','ConvNeXt-Tiny','Swin-T']
POOLS={'B':'COCO','C':'VOC2007','D':'Objects365'}
SOURCES={'local':'本地AI描述分组','online':'在线AI描述分组'}
def read(p):
    with p.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
rows=read(INPUT)
lookup={(r['pool'],r['label_source'],r['model'],r['seed'],int(r['tier'])):float(r['macro_f1_100']) for r in rows}
assert len(lookup)==len(rows)==720
seeds={m:sorted({r['seed'] for r in rows if r['model']==m}) for m in MODELS}
assert all(len(v)==3 for v in seeds.values())
drops={}
run_rows=[]
summary=[]
for source in SOURCES:
    for pool in POOLS:
        for m in MODELS:
            vals=[]
            for seed in seeds[m]:
                low=lookup[pool,source,m,seed,1];high=lookup[pool,source,m,seed,5]
                delta=low-high;assert delta>0
                vals.append(delta)
                run_rows.append([pool,POOLS[pool],source,m,seed,low,high,delta])
            a=np.array(vals);drops[pool,source,m]=a
            summary.append([pool,POOLS[pool],source,m,a.mean(),a.min(),a.max()])
    for pool in POOLS:
        observed=-100*float(next(r for r in read(RANDOM) if r['pool']==pool and r['source']==source)['observed_delta'])
        assert abs(np.mean([drops[pool,source,m].mean() for m in MODELS])-observed)<1e-9

for filename,headers,records in [
    ('paired_model_drops.csv',['pool','dataset','label_source','model','seed','Q1_f1_100','Q5_f1_100','drop_percentage_points'],run_rows),
    ('model_drop_summary.csv',['pool','dataset','label_source','model','mean_drop_pp','minimum_of_3_runs','maximum_of_3_runs'],summary)]:
    with (OUT/filename).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(records)

from docx.enum.section import WD_SECTION_START
from docxcompose.composer import Composer
from docx.oxml import OxmlElement
OLD=ROOT/'docs/工作量_逐图解释版_20261008.docx'
MIX=ROOT/'artifacts/controlled_risk_distribution_20261005/tables/model_summary.csv'
CONTRAST=ROOT/'artifacts/controlled_risk_distribution_20261005/tables/distribution_contrasts.csv'
MODEL_IDS={r['model']:r['model_id'] for r in rows}
mix_rows=[r for r in read(MIX) if r['total']=='120' and r['budget_role']=='common']
contrast_rows=[r for r in read(CONTRAST) if r['total']=='120' and r['budget_role']=='common']
mixed={}
mixed_export=[]
for source in SOURCES:
    for pool in POOLS:
        for model in MODELS:
            rr=sorted([r for r in mix_rows if r['pool']==pool and r['source']==source and r['model']==MODEL_IDS[model]],key=lambda r:float(r['high_risk_percent']))
            assert len(rr)==21
            xx=np.array([float(r['high_risk_percent']) for r in rr])
            yy=np.array([100*float(r['macro_f1_mean']) for r in rr])
            assert xx.tolist()==list(range(0,101,5))
            loss=yy[0]-yy
            mixed[pool,source,model]=(xx,loss)
            contrast=next(r for r in contrast_rows if r['pool']==pool and r['source']==source and r['model']==MODEL_IDS[model])
            assert abs(loss[-1]+100*float(contrast['macro_f1_delta']))<1e-8
            for percent,score,delta in zip(xx,yy,loss):mixed_export.append([pool,source,model,percent,score,delta])
with (OUT/'fixed_120_model_drops.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(['pool','label_source','model','high_risk_percent','macro_f1_100','drop_vs_0_percent_pp']);w.writerows(mixed_export)
font_manager.fontManager.addfont(str(ROOT/'scripts/assets/NotoSansSC-Regular.ttf'))
plt.rcParams.update({'font.family':'Noto Sans SC','font.size':14,'svg.fonttype':'none','pdf.fonttype':42,
    'axes.spines.top':False,'axes.spines.right':False,'axes.linewidth':.7,'xtick.major.size':0,'ytick.major.size':0})
COLORS=['#1B6CA8','#E38D26','#409773','#BD5858','#8763AB','#87674E','#CD70AA','#4C8791']
MARKERS=['o','s','^','D','v','P','X','>']
figures=[]
limit=float(np.ceil((max(a.max() for a in drops.values())+5)/10)*10)
def save(fig,stem,title,caption):
    for ext in ['png','svg','pdf']:fig.savefig(FIG/f'{stem}.{ext}',dpi=250,bbox_inches='tight',facecolor='white')
    plt.close(fig);figures.append({'stem':stem,'title':title,'caption':caption})

for source,label in SOURCES.items():
    fig,axes=plt.subplots(1,3,figsize=(15,6.2),layout='constrained',sharey=True)
    for ax,(pool,name) in zip(axes,POOLS.items()):
        mean=np.array([drops[pool,source,m].mean() for m in MODELS])
        lo=np.array([drops[pool,source,m].min() for m in MODELS]);hi=np.array([drops[pool,source,m].max() for m in MODELS])
        ax.barh(range(8),mean,height=.55,color='#3574A8')
        ax.errorbar(mean,range(8),xerr=[mean-lo,hi-mean],fmt='none',ecolor='#244761',capsize=3,lw=1)
        for i,v in enumerate(mean):ax.text(hi[i]+.8,i,f'{v:.2f}',va='center',fontsize=13)
        ax.set(title=name,yticks=range(8),yticklabels=MODELS,ylim=(7.65,-.65),xlim=(0,limit),
               xticks=np.arange(0,limit+1,10),xlabel='第1组 → 第5组：降幅（百分点）')
    save(fig,f'01_endpoint_{source}',f'各模型从第1组到第5组降低多少（{label}）',
         '降幅＝第1组F1－第5组F1。每条横线对应一种模型；横条是3份训练结果的配对降幅均值，黑色短线是这3份配对降幅的最小—最大范围。')

for source,label in SOURCES.items():
    fig,axes=plt.subplots(1,3,figsize=(15,6.3),layout='constrained',sharey=True)
    for ax,(pool,name) in zip(axes,POOLS.items()):
        for i,m in enumerate(MODELS):
            losses=[np.mean([lookup[pool,source,m,s,1]-lookup[pool,source,m,s,t] for s in seeds[m]]) for t in range(1,6)]
            assert abs(losses[-1]-drops[pool,source,m].mean())<1e-9
            ax.plot(range(1,6),losses,color=COLORS[i],marker=MARKERS[i],lw=1.8,markersize=5,label=m)
        ax.axhline(0,color='#A9B3BC',ls='--',lw=.8)
        ax.set(title=name,xticks=range(1,6),xlim=(.85,5.15),ylim=(-10,60),yticks=[-10,0,10,20,30,40,50,60],xlabel='测试组别：1 → 5')
    axes[0].set_ylabel('相对第1组降低多少（百分点）')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='outside lower center',ncol=4,frameon=False,fontsize=13)
    save(fig,f'02_tier_drop_{source}',f'降幅在五个组中怎样变化（{label}）',
         '每条线是一种模型。以该模型第1组为零点，显示“第1组F1－当前组F1”的3份训练结果平均值。往上表示下降更多；负值表示该组比第1组高，中间组不一定逐步变差。')

matrix=np.array([[drops[p,s,m].mean() for p in POOLS for s in SOURCES] for m in MODELS])
fig,ax=plt.subplots(figsize=(13.8,6),layout='constrained')
ax.pcolormesh(np.arange(7)-.5,np.arange(9)-.5,matrix,cmap='Blues',vmin=0,vmax=60,shading='flat',rasterized=False)
ax.set_xlim(-.5,5.5);ax.set_ylim(7.5,-.5)
ax.set(xticks=range(6),xticklabels=[f'{POOLS[p]}\n'+('本地' if s=='local' else '在线') for p in POOLS for s in SOURCES],
       yticks=range(8),yticklabels=MODELS)
ax.tick_params(pad=9)
for i in range(8):
    for j in range(6):ax.text(j,i,f'{matrix[i,j]:.2f}',ha='center',va='center',fontsize=16,color='white' if matrix[i,j]>32 else '#233D53')
for spine in ax.spines.values():spine.set_visible(False)
save(fig,'03_drop_overview','一张图看齐8种模型、3个来源、2套分组的降幅',
     '每行是一种模型；每个来源分别列本地和在线分组。数字均为“第1组F1－第5组F1”（百分点），颜色越深表示降幅越大。这里共48个模型平均降幅，全部大于零。')

fig,axes=plt.subplots(1,3,figsize=(15,6.1),layout='constrained',sharey=True,sharex=True)
for ax,(pool,name) in zip(axes,POOLS.items()):
    ax.plot([0,60],[0,60],ls='--',lw=1,color='#A9B3BC')
    for i,m in enumerate(MODELS):
        ax.scatter(drops[pool,'local',m].mean(),drops[pool,'online',m].mean(),color=COLORS[i],marker=MARKERS[i],s=80,label=m)
    ax.set(title=name,xlim=(0,60),ylim=(0,60),xticks=[0,15,30,45,60],yticks=[0,15,30,45,60],
           xlabel='本地分组降幅（百分点）',aspect='equal')
axes[0].set_ylabel('在线分组降幅（百分点）')
fig.legend(*axes[0].get_legend_handles_labels(),loc='outside lower center',ncol=4,frameon=False,fontsize=13)
save(fig,'04_source_comparison','同一模型，两套AI分组的降幅是否接近？',
     '每个点是一种模型的3份训练结果平均降幅。越靠近虚线，两套分组测得的降幅越接近；在线降幅更大时点位于线上方。两套AI负责状态描述，不是被比较的8种分类模型；接近不等于标签一定正确。')

for source,label in SOURCES.items():
    fig,axes=plt.subplots(1,3,figsize=(15,6.3),layout='constrained',sharey=True)
    for ax,(pool,name) in zip(axes,POOLS.items()):
        for i,m in enumerate(MODELS):
            xx,loss=mixed[pool,source,m]
            ax.plot(xx,loss,color=COLORS[i],marker=MARKERS[i],markevery=4,markersize=5,lw=1.8,label=m)
        ax.axhline(0,color='#A9B3BC',ls='--',lw=.8)
        ax.set(title=name,xlim=(-2,102),ylim=(-5,50),xticks=[0,25,50,75,100],yticks=[0,10,20,30,40,50],
               xlabel='第5组目标占比（%）')
    axes[0].set_ylabel('相对0%起点降低多少（百分点）')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='outside lower center',ncol=4,frameon=False,fontsize=13)
    save(fig,f'05_fixed120_{source}',f'固定120个目标，困难目标越多时降幅怎样变化（{label}）',
         '每次120个目标、六类各20个；21种占比，每种抽样1000次。每条线是一种模型的平均F1降幅。此处基准是0%第5组（从第1—4组混合取样），与前面单独第1组的基准不同；连线连接已有比例点，没有新增观测。')

supp=Document();sec=supp.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width,sec.page_height=Cm(29.7),Cm(21)
sec.left_margin=sec.right_margin=Cm(1.2);sec.top_margin=sec.bottom_margin=Cm(1.1)
normal=supp.styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(10)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
normal.paragraph_format.space_after=Pt(4)
for i,record in enumerate(figures):
    if i:supp.add_page_break()
    p=supp.add_paragraph();p.paragraph_format.keep_with_next=True
    run=p.add_run(f'新增图D{i+1}  {record["title"]}');run.font.size=Pt(16);run.font.bold=True
    run.font.color.rgb=RGBColor.from_string('24577E')
    p=supp.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/f'{record["stem"]}.png'),width=Cm(26.6))
    supp.add_paragraph(record['caption'])
supp.save(OUT/'模型降幅图表补充.docx')

doc=Document(OLD)
locator=doc.add_paragraph('新增：第13—20页集中展示8种模型的降幅对比、五组降幅变化、两套分组对照和固定数量实验；主要口径为“第1组F1－第5组F1”，单位是百分点。')
for run in locator.runs:run.font.size=Pt(10);run.font.bold=True;run.font.color.rgb=RGBColor.from_string('24577E')
target=next(p for p in doc.paragraphs if p.text.startswith('下面每项工作都单独解释'))
target._p.addnext(locator._p)
section=doc.add_section(WD_SECTION_START.NEW_PAGE)
section.orientation=WD_ORIENT.LANDSCAPE;section.page_width,section.page_height=Cm(29.7),Cm(21)
section.left_margin=section.right_margin=Cm(1.2);section.top_margin=section.bottom_margin=Cm(1.1)
composer=Composer(doc);composer.append(supp)
composer.doc.core_properties.title='工作量：含8种模型降幅多图对比'
composer.doc.core_properties.subject='原工作量逐图解释＋8张模型降幅对比图'
composer.save(OUT/f'{NAME}.docx');shutil.copy2(OUT/f'{NAME}.docx',ROOT/'docs'/f'{NAME}.docx')
assert len(composer.doc.inline_shapes)==19
assert len(composer.doc.sections)==2
assert composer.doc.sections[0].orientation != composer.doc.sections[1].orientation
md=OLD.with_suffix('.md').read_text()+'\n\n# 新增：8种模型的降幅对比\n'
for i,rec in enumerate(figures):md+=f'\n## 图D{i+1} {rec["title"]}\n\n![图D{i+1}]({FIG / (rec["stem"]+".png")})\n\n{rec["caption"]}\n'
(ROOT/'docs'/f'{NAME}.md').write_text(md,encoding='utf-8')
shutil.copy2(Path(__file__),OUT/Path(__file__).name)
input_paths=[INPUT,RANDOM,MIX,CONTRAST,OLD]
audit={'old_figures':11,'new_figures':figures,'total_figures':19,'planned_pages':20,
    'pairwise_endpoint_runs':len(run_rows),'endpoint_means':len(summary),'fixed120_endpoints_verified':48,
    'models':MODELS,'primary_baseline':'Q1','fixed120_baseline':'0% Q5 (Q1-Q4 mixture)',
    'input_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in input_paths},
    'new_training_or_inference':False,'render_verified':False}
(OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
for p in [MIX,CONTRAST]:shutil.copy2(p,OUT/p.name)
shutil.copy2(INPUT,OUT/INPUT.name)
with zipfile.ZipFile(ROOT/'docs/工作量_模型降幅多图素材_20261008.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in FIG.glob('*'):z.write(p,p.relative_to(OUT))
    for p in OUT.glob('*'):
        if p.is_file() and p.suffix in ['.py','.json','.csv']:z.write(p,p.name)
    for p in input_paths:z.write(p,Path('source_data')/p.relative_to(ROOT))
    z.write(OLD.with_suffix('.md'),Path('source_data')/OLD.with_suffix('.md').relative_to(ROOT))
    z.write(Path(__file__),'source_data/scripts/'+Path(__file__).name)
    z.write(ROOT/'scripts/assets/NotoSansSC-Regular.ttf','source_data/scripts/assets/NotoSansSC-Regular.ttf')
print(json.dumps({'word':str(ROOT/'docs'/f'{NAME}.docx'),'new_figures':8,'total_figures':19,'expected_pages':20},ensure_ascii=False))
